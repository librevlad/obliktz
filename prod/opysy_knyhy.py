# -*- coding: utf-8 -*-
r"""Чернетки інвентаризаційних описів запасів продовольчої служби з книг обліку («Облік СП», «Облік ОП»: сухі
пайки, тара, одноразовий посуд, миючі засоби, серветки) за типовою формою наказу Мінфіну від 17.06.2015 № 572 —
лише читання книг.

Опис = місце зберігання на кінець дня перед датою «станом на» за робочими таблицями книг (прихід, переміщення, списання;
старі назви підрозділів — через `tblRename`); він може складатися з кількох місць книги (`knyha`). Для кожного опису й коду
дві сторони. A — що за обліком, на весь опис: документи, які ФЕС провів (без списань з номером із конфігу «znyshcheno» і без
документів у станах `_pending`, яких ФЕС ще не провів); переміщення між місцями опису для нього нічого не міняють. B — що
фактично є, по місцях книги (усі документи; ΣB+ — сума додатних). Звичайні рядки разом дають min(max(ΣA, 0), ΣB+): власний
облік місця, затиснутий у [0, B]; зайве знімається з останніх місць, бракуюче додається першим місцям із запасом B. Фактично
більше — рядок «фактично є, за обліком нема» на кожне таке місце, з приміткою місця й документів; обліку більше — один рядок
опису «за обліком є, фактично нуль» з приміткою, які документи за ним стоять. Лише «znyshcheno» без таких документів — рядки
за місцем і датою, як у iprod. Кожен документ, якого ФЕС не провів, і кожне знищене, що торкаються опису, названо в примітці
якогось рядка опису; неназвані дописано в примітку першого окремого рядка, а коли такого рядка нема, їх називає розділ
«Довідки». Мінус B місця (книга) і мінус ΣA опису (облік) — у «Мінусах» з причиною: місце з мінусом книги рядків не має, але
його облік лишається в ΣA; опис з мінусом обліку йде з фактичними рядками й обліком 0. Графи «Фактична наявність» порожні:
кількість вносить комісія (сума — формулою).

Значення частини — у JSON-конфігу поряд з описами: наказ, комісія, книги, категорії (порядок і субрахунок), місця
(місця книги → опис, «де зберігаються», МВО), номенклатурні номери ФЕС, позиції, описані деінде («poza_opysom»).
МВО і «перевірив» із ключем `dobovka` — з аркуша «Підписанти» найновішої Добовки на дату «станом на».

    build(cfg, out, dry=False, dobovka=<Добовка>)       — «Звіти › Описи 572» програми; книги `knyhy` «Облік СП» /
                                                         «Облік ОП» / «Облік РП» — вивантаження книг програми

Від сервера програми в `cfg` (межі перенесення, окрім `build` і Добовки): `_knyhy_names` — імена книг для «Довідки»
(«Облік СП 6.xlsx (з програми)»; без них — імена файлів, а не шляхи тимчасових вивантажень), `_pending` — стани
документів, яких ФЕС ще не провів («роздрукувати», «на підписі», «їде на ФЕС»); типово порожньо — як у iprod. Ще одна
відмінність: `load_config` читає й файл «UTF-8 з BOM» (Блокнот).
"""
import collections
import datetime
import io
import json
import os
import re
import shutil
import tempfile
import warnings

import excel_names
from prod.opysy_forma import (FIRST_ROW, MONEY_FMT, UA_MONTHS_GEN, WIDTHS, Sheet, Signers, lines_of, money_words,
                              num, page_setup, parse_date, person_of, sign_row, ua_date, ua_date_text, words)

INT_FMT = '#,##0'
QTY3_FMT = '#,##0.000'
EPS = 1e-6
# qty і summa — за даними обліку (графи 9–11); lost — окремий рядок, де облік і фактична наявність розходяться; fact —
# очікувана фактична кількість: у звичайного рядка вона дорівнює qty, у «за обліком є, фактично нуль» — 0, у «фактично є,
# за обліком нема» — кількість різниці, а qty і summa — 0.
Line = collections.namedtuple('Line', 'cat code fes name unit price qty summa note lost fact')


# ---------------------------------------------------------------- конфіг

def load_config(path):
    with io.open(path, encoding='utf-8-sig') as f:
        cfg = json.load(f)
    for key in ('stan_na', 'dok', 'ustanova', 'vch', 'edrpou', 'sluzhba', 'komisia', 'knyhy', 'kategorii', 'mista'):
        if not cfg.get(key):
            raise SystemExit('конфіг %s: нема «%s»' % (path, key))
    cfg['_dir'] = os.path.dirname(os.path.abspath(path))
    cfg['_path'] = os.path.abspath(path)
    return cfg


# ---------------------------------------------------------------- книги

def _code(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return v


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_book(path):
    """Робочі таблиці книги (копія — книгу може тримати відкритою Excel): позиції, перейменування, події."""
    import openpyxl
    from openpyxl.utils import range_boundaries
    if not os.path.exists(path):
        raise SystemExit('книги нема: %s' % path)
    tmp = os.path.join(tempfile.mkdtemp(prefix='opysy_'), 'book.xlsx')
    shutil.copy2(path, tmp)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(tmp, data_only=True)

    def table(name):
        for ws in wb.worksheets:
            if name in ws.tables:
                c1, r1, c2, r2 = range_boundaries(ws.tables[name].ref)
                hdr = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
                return [dict(zip(hdr, (ws.cell(r, c).value for c in range(c1, c2 + 1)))) for r in range(r1 + 1, r2 + 1)]
        raise SystemExit('%s: нема таблиці %s' % (os.path.basename(path), name))

    try:
        items = {}
        for r in table('tblItems'):
            c = _code(r.get('код'))
            if c in (None, ''):
                continue
            items[c] = dict(name=_txt(r.get('найменування')), cat=_txt(r.get('категорія')), unit=_txt(r.get('од.вим.')),
                            price=float(r.get('ціна') or 0))
        rename = {_txt(r.get('стара назва')): _txt(r.get('нова назва')) for r in table('tblRename')
                  if r.get('стара назва') and r.get('нова назва')}
        ev = []
        for name, kind, frm, to in (('tblIncoming', 'in', None, 'кому'), ('tblMovement', 'mv', 'від_кого', 'кому'),
                                    ('tblWriteoffs', 'wo', 'підрозділ', None)):
            for r in table(name):
                d, c = _date(r.get('док_дата')), _code(r.get('код'))
                if d is None or c in (None, ''):              # порожній рядок вставки таблиці — не подія
                    continue
                ev.append(dict(kind=kind, date=d, num=_txt(r.get('док_номер')), code=c,
                               qty=float(r.get('кількість') or 0),
                               frm=_txt(r.get(frm)) if frm else '', to=_txt(r.get(to)) if to else '',
                               status=_txt(r.get('статус'))))
    finally:
        wb.close()
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    return dict(path=path, items=items, rename=rename, events=ev)


def ledger(book, cut, lost_num, pending, place_of):
    """Залишки на кінець дня `cut` по місцях і кодах: {(місце, код): [A, B]} — A за обліком (документи, які ФЕС провів), B
    фактично (усі документи); місця — у нових назвах. Другий результат — документи, яких ФЕС не провів: у станах `pending` і
    списання з номером `lost_num` (поле `zn`); переміщення між двома місцями одного опису (`place_of`) нічого не міняє для
    опису, тож за проведене. Документ — рядок книги з ключами kind, date, num, code, qty, frm, to, status."""
    ren = book['rename']
    sides = collections.defaultdict(lambda: [0.0, 0.0])
    waits = []
    for e in book['events']:
        if e['date'] > cut:
            continue
        frm, to = ren.get(e['frm'], e['frm']), ren.get(e['to'], e['to'])
        if e['kind'] == 'in':
            effect = [(to, e['qty'])]
        elif e['kind'] == 'mv':
            effect = [(frm, -e['qty']), (to, e['qty'])] if frm != to else []
        else:
            effect = [(frm, -e['qty'])]
        destroyed = e['kind'] == 'wo' and bool(lost_num) and e['num'].upper() == lost_num.upper()
        inner = e['kind'] == 'mv' and place_of.get(frm) is not None and place_of.get(frm) is place_of.get(to)
        counted = not destroyed and (e.get('status') not in pending or inner)
        for m, d in effect:
            sides[(m, e['code'])][1] += d
            if counted:
                sides[(m, e['code'])][0] += d
        if effect and not counted:
            waits.append(dict(e, frm=frm, to=to, zn=destroyed))
    return sides, waits


def qty_text(q):
    """Кількість у примітці: 1000 -> «1000», 12.5 -> «12,5»."""
    return '%d' % round(q) if abs(q - round(q)) < EPS else ('%.3f' % q).rstrip('0').replace('.', ',')


def doc_text(w, mine):
    """Документ, якого ФЕС не провів, для примітки: «прихід №12 від 21.09.2026, 1000, на підписі». Переміщення називає й друге
    місце: «до «…»», коли воно йде зі `mine` (місця опису), і «з «…»», коли сюди."""
    name = {'in': 'прихід', 'mv': 'переміщення', 'wo': 'списання'}[w['kind']]
    text = '%s №%s від %s' % (name, w['num'] or 'б/н', w['date'].strftime('%d.%m.%Y'))
    if w['kind'] == 'mv':
        if w['frm'] in mine:
            text += ' до «%s»' % w['to']
        else:
            text += ' з «%s»' % w['frm']
    return '%s, %s, %s' % (text, qty_text(w['qty']), w['status'])


def touches(w, mine):
    """Чи міняє документ залишок опису (його місця `mine`): прихід — сюди, списання — звідси, переміщення — одним кінцем."""
    if w['kind'] == 'in':
        return w['to'] in mine
    if w['kind'] == 'wo':
        return w['frm'] in mine
    return (w['frm'] in mine) != (w['to'] in mine)


def at(w, m):
    """Чи документ зачіпає саме місце `m`: прихід — сюди, списання — звідси, переміщення — обидва кінці."""
    if w['kind'] == 'in':
        return w['to'] == m
    if w['kind'] == 'wo':
        return w['frm'] == m
    return m in (w['frm'], w['to'])


def by_date(docs):
    return sorted(docs, key=lambda w: (w['date'], w['kind'], w['num']))


def make_line(it, code, fes, qty, acc, fact, note, lost):
    """Рядок опису: `acc` — кількість стоїть у графах 9–11, `fact` — комісія очікує її й фактично (рядок «за обліком є,
    фактично нуль» — лише acc, «фактично є, за обліком нема» — лише fact)."""
    qty = round(qty, 6)
    return Line(it['cat'], code, fes.get(str(code), ''), it['name'], it['unit'], it['price'], qty if acc else 0.0,
                round(qty * it['price'] + 1e-9, 2) if acc else 0.0, note, lost, qty if fact else 0.0)


def allocate(places, total):
    """Розкласти `total` звичайних одиниць по місцях опису `places` = [(місце, A, B)] у порядку конфігу (B > 0): кожне місце бере
    власний облік, затиснутий у [0, B]; зайве над `total` знімається з останніх місць, нестача додається першим місцям, у яких ще
    є запас фактичного залишку. -> {місце: звичайна кількість}."""
    got = {m: round(min(max(a, 0.0), b), 6) for m, a, b in places}
    extra = round(sum(got.values()) - total, 6)
    if extra >= EPS:
        for m, _a, _b in reversed(places):
            take = min(got[m], extra)
            got[m] = round(got[m] - take, 6)
            extra = round(extra - take, 6)
    elif extra <= -EPS:
        lack = -extra
        for m, _a, b in places:
            add = min(round(b - got[m], 6), lack)
            got[m] = round(got[m] + add, 6)
            lack = round(lack - add, 6)
    return got


def zn_parts(killed, destroyed_note):
    """Знищене, ще не списане у ФЕС, для примітки: по датах, з кількістю — «знищено 25.09.2026, у ФЕС не списано: 4»."""
    per = collections.defaultdict(float)
    for w in killed:
        per[w['date']] += w['qty']
    return ['%s: %s' % (destroyed_note(d), qty_text(q)) for d, q in sorted(per.items())]


def why_text(docs, mine, killed, destroyed_note):
    """Документи за різницею: «ФЕС ще не провів: прихід №12 від 21.09.2026, 1000, на підписі; …», далі знищене."""
    head = ['ФЕС ще не провів: ' + '; '.join(doc_text(w, mine) for w in docs)] if docs else []
    return '; '.join(head + zn_parts(killed, destroyed_note))


def minus_text(minus_books):
    """Мінус книги в примітці рядка, за яким нема документа: «мінус у книзі «склад ППД»: −20, див. «Мінуси в книгах»»."""
    return '; '.join('мінус у книзі «%s»: −%s, див. «Мінуси в книгах»' % (m, qty_text(-b)) for m, b in minus_books)


def opys_lines(p, it, code, members, waits, fes, destroyed_note, info):
    """Рядки однієї позиції опису `p`; `members` — [(місце, A, B)] його місць книги (і з нулями з обох боків), `waits` — документи,
    яких ФЕС не провів. Облік — на весь опис (ΣA), фактично — по місцях (B < 0 — мінус книги, його вже названо; решта ΣB+; облік
    місця з мінусом лишається в ΣA). Звичайні рядки разом min(max(ΣA, 0), ΣB+), по місцях — `allocate`. Решта фактичного залишку
    місця — його рядок «фактично є, за обліком нема» з приміткою місця й документів, що його торкаються (немає таких — документи
    за різницею опису, а немає й їх — мінус книги, який лишив залишок без документа). ΣA − ΣB+ більше нуля — один рядок опису «за
    обліком є, фактично нуль»; лише «znyshcheno» — рядки за місцем і датою, як у iprod. ΣA менше нуля — мінус обліку
    (`info['neg']`), фактичні рядки йдуть з обліком 0. Кожен документ, якого ФЕС не провів, і кожне знищене, що торкаються опису,
    названо в примітці якогось рядка: неназвані дописано в примітку першого окремого рядка, а нема окремого рядка (різниця нульова
    чи фактичного залишку нема) — вони йдуть у `info['quiet']`, розділ «Довідки»."""
    mine = set(p['knyha'])
    notes = p.get('prymitky') or {}
    order = {m: i for i, m in enumerate(p['knyha'])}
    members = sorted(members, key=lambda x: order[x[0]])
    sa = round(sum(a for _m, a, _b in members), 6)                        # облік опису: переміщення між його місцями зникають
    present = [(m, a, b) for m, a, b in members if b >= EPS]              # місця з фактичним залишком
    sb = round(sum(b for _m, _a, b in present), 6)
    minus_books = [(m, b) for m, _a, b in members if b < -EPS]            # мінус книги місця названо в `collect`
    if sa < -EPS:
        info['neg'].append((p['nazva'], code, it['name'], sa, 'облік'))
    behind = by_date(w for w in waits if w['code'] == code and not w['zn'] and touches(w, mine))
    killed = [w for w in waits if w['code'] == code and w['zn'] and w['frm'] in mine]
    got = allocate(present, min(max(sa, 0.0), sb))
    made = [make_line(it, code, fes, got[m], True, True, notes.get(m) or '', False)
            for m, _a, _b in present if got[m] >= EPS]
    rests = []                                                            # фактично є, за обліком нема — своє на кожне місце
    for m, _a, b in present:
        rest = round(b - got[m], 6)
        if rest >= EPS:
            own = [w for w in behind if at(w, m)]
            own_zn = [w for w in killed if w['frm'] == m]
            if not own and not own_zn:                                    # документ за різницею стоїть деінде в описі
                own, own_zn = behind, killed
            rests.append([m, rest, own, own_zn])
    named = {id(w) for _m, _r, own, own_zn in rests for w in own + own_zn}
    diff = round(sa - sb, 6)
    acc = None                                                            # за обліком є, фактично нуль — один рядок опису
    if diff >= EPS:
        if not behind and killed and abs(sum(w['qty'] for w in killed) - diff) < EPS:
            per = collections.defaultdict(float)                          # лише «znyshcheno»: за місцем і датою, як у iprod
            for w in killed:
                per[(w['frm'], w['date'])] += w['qty']
            for (m, d), q in sorted(per.items(), key=lambda kv: (order[kv[0][0]], kv[0][1])):
                note = '; '.join(x for x in (notes.get(m), destroyed_note(d)) if x)
                made.append(make_line(it, code, fes, q, True, False, note, True))
        else:
            acc = diff
        named |= {id(w) for w in behind + killed}
    left = [w for w in behind if id(w) not in named]                      # жоден рядок опису їх не назвав
    left_zn = [w for w in killed if id(w) not in named]
    if (left or left_zn) and rests:                                       # окремий рядок є: до його примітки
        rests[0][2], rests[0][3] = by_date(rests[0][2] + left), rests[0][3] + left_zn
    for m, rest, own, own_zn in rests:
        why = why_text(own, mine, own_zn, destroyed_note) or minus_text(minus_books)
        made.append(make_line(it, code, fes, rest, False, True, '; '.join(x for x in (notes.get(m), why) if x), True))
    if acc is not None:
        why = why_text(behind, mine, killed, destroyed_note) or minus_text(minus_books)
        made.append(make_line(it, code, fes, acc, True, False, why, True))
    if (left or left_zn) and not rests:                                   # окремого рядка нема: до «Довідки»
        info['quiet'].append((p['nazva'], code, it['name'], why_text(left, mine, left_zn, destroyed_note)))
    return made


# ---------------------------------------------------------------- описи

def collect(cfg, books):
    stan = parse_date(cfg['stan_na'])
    cut = stan - datetime.timedelta(days=1)
    cats = [k['kod'] for k in cfg['kategorii']]
    subr = {k['kod']: k['subrahunok'] for k in cfg['kategorii']}
    fes = {str(k): str(v) for k, v in (cfg.get('fes_kody') or {}).items()}
    skip = {str(x['kod']): x.get('prychyna', '') for x in cfg.get('poza_opysom') or []}
    zn = cfg.get('znyshcheno') or {}
    pending = tuple(cfg.get('_pending') or ())
    names = cfg.get('_knyhy_names') or []
    place_of = {}
    for p in cfg['mista']:
        for m in p['knyha']:
            place_of[m] = p
    lines = collections.defaultdict(list)            # nazva опису -> [Line]
    info = dict(cut=cut, skipped=[], neg=[], quiet=[], unknown=collections.defaultdict(list), nofes=set(), books=[])

    def destroyed_note(d):
        return (zn.get('prymitka') or 'не списано').format(data=d.strftime('%d.%m.%Y'))
    for i, book in enumerate(books):
        info['books'].append(names[i] if i < len(names) else os.path.basename(book['path']))
        sides, waits = ledger(book, cut, zn.get('nomer'), pending, place_of)
        groups = collections.defaultdict(list)                     # (опис, код) -> [(місце, A, B)]
        outside = {}                                               # описані деінде, але місце в описі: (опис, код) -> [опис, позиція, ΣA]
        for (m, c), (a, b) in sorted(sides.items(), key=lambda kv: (str(kv[0][0]), str(kv[0][1]))):
            a, b = round(a, 6), round(b, 6)
            it = book['items'].get(c) or dict(name='код %s' % c, cat='', unit='', price=0.0)
            p = place_of.get(m)
            if str(c) in skip or p is None or it['cat'] not in cats:   # поза описами чи поза конфігом
                judged = str(c) in skip and p is not None          # місце в описі: його облік судять по опису, а не по місцю
                if judged:
                    outside.setdefault((id(p), c), [p, it, 0.0])[2] += a
                if b < -EPS or (a < -EPS and not judged):          # мінус книги — по місцю, як у iprod
                    info['neg'].append((m, c, it['name'], b, 'книга') if b < -EPS else (m, c, it['name'], a, 'облік'))
                    continue
                if abs(a) < EPS and abs(b) < EPS:
                    continue
                if str(c) in skip:                                 # описана деінде: залишок і кожен документ ФЕС окремо
                    if b >= EPS:
                        info['skipped'].append((m, c, it['name'], it['unit'], b, skip[str(c)]))
                    for w in by_date(w for w in waits if w['code'] == c and at(w, m)):
                        why = destroyed_note(w['date']) if w['zn'] else 'ФЕС ще не провів: ' + doc_text(w, {m})
                        info['skipped'].append((m, c, it['name'], it['unit'], w['qty'], '%s; %s' % (skip[str(c)], why)))
                else:
                    info['unknown'][(m, it['cat'])].append('%s %s' % (c, it['name']))
                continue
            if b < -EPS:                                           # мінус книги: місце в рядки не йде, його облік лишається в ΣA
                info['neg'].append((m, c, it['name'], b, 'книга'))
            groups[(id(p), c)].append((m, a, b))                   # й нулі з обох боків: документи за ними теж називають
        for (_pid, c), (p, it, sa) in outside.items():
            if sa < -EPS:
                info['neg'].append((p['nazva'], c, it['name'], round(sa, 6), 'облік'))
        for (_pid, c), members in groups.items():
            p = place_of[members[0][0]]
            it = book['items'][c]
            made = opys_lines(p, it, c, members, waits, fes, destroyed_note, info)
            if made and str(c) not in fes:
                info['nofes'].add((c, it['name']))
            lines[p['nazva']].extend(made)
    if info['unknown']:
        raise SystemExit('залишки поза конфігом (місце книги · категорія; враховано й документи, яких ФЕС ще не провів): %s'
                         % '; '.join(
            '%s · %s: %s' % (m, cat or '—', ', '.join(v[:5])) for (m, cat), v in sorted(info['unknown'].items())))
    info['neg'].sort(key=lambda t: (str(t[0]), str(t[1]), t[4]))
    out = []
    for p in cfg['mista']:
        ls = sorted(lines.get(p['nazva'], []), key=lambda ln: (cats.index(ln.cat), ln.name.lower(), ln.price, ln.lost,
                                                               ln.note, ln.qty, ln.fact))
        if not ls:
            continue
        subs = []
        for ln in ls:
            if subr[ln.cat] not in subs:
                subs.append(subr[ln.cat])
        out.append(dict(place=p, nazva=p['nazva'], lines=ls, subs=subs, subr=subr))
    return out, info


def qty_fmt(q):
    return INT_FMT if abs(q - round(q)) < EPS else QTY3_FMT


def sum_fmt(lines):
    """Формат «Разом» графи 6: за рядками, які вона підсумовує, — комісія вносить у кожен max(qty, fact), тож дробовий рядок
    робить дробовим і підсумок."""
    return QTY3_FMT if any(qty_fmt(max(ln.qty, ln.fact)) == QTY3_FMT for ln in lines) else INT_FMT


def write_opys(ws, cfg, op, mvo, perev, buh):
    sh = Sheet(ws)
    stan = parse_date(cfg['stan_na'])
    year = stan.year
    de = '%s військової частини %s' % (op['place']['de'], cfg['vch'])
    sh.put('I1:L1', 'ЗАТВЕРДЖЕНО')
    sh.put('I2:L2', 'Наказ Міністерства фінансів України')
    sh.put('A3:D3', cfg['ustanova'], size=11, h='center')
    sh.put('I3:L3', '17.06.2015  № 572')
    sh.put('A4:D4', '(установа)', size=7, h='center')
    sh.put('A5:B5', 'Ідентифікаційний код за ЄДРПОУ', size=8, border='box')
    sh.put('C5:D5', '  '.join(cfg['edrpou']), size=9, h='center', border='box')
    for r, hh in ((1, 15.2), (2, 15.2), (3, 16.52), (4, 15), (5, 23.12)):
        sh.height(r, hh)
    sh.put('A7:L7', 'ІНВЕНТАРИЗАЦІЙНИЙ ОПИС', size=12, bold=True, h='center')
    head = 'запасів %s %s' % (cfg['sluzhba'], de)
    sh.put('A8:L8', head, size=11, bold=True, h='center')
    sh.height(8, 15.5 * lines_of(head, 153, 11) + 3)
    sh.put('A9:L9', '(запаси)', size=7, h='center')
    sh.put('A10:L10', ua_date_text(cfg.get('data_skladannia') or cfg['stan_na'], year), h='center')
    sh.put('A11:L11', '(дата складання)', size=7, h='center')
    basis = ('На підставі розпорядчого документа %s виконано зняття фактичних залишків запасів, які обліковуються '
             'на субрахунку(ах) %s та зберігаються: %s станом на «%d» %s %d р.'
             % (cfg['dok'], ', '.join(op['subs']), de, stan.day, UA_MONTHS_GEN[stan.month - 1], stan.year))
    sh.put('A13:L13', basis, h='justify', v='top')
    sh.height(13, 13.6 * lines_of(basis, 153, 10) + 4)
    sh.put('A15:L15', 'Розписка', bold=True, h='center')
    rozp = ('      До початку проведення інвентаризації всі видаткові та прибуткові документи на матеріальні '
            'цінності здано в бухгалтерську службу і всі матеріальні цінності, що надійшли на мою '
            'відповідальність, оприбутковано, а ті, що вибули, списано.')
    sh.put('A16:L16', rozp, h='justify', v='top')
    sh.height(16, 28.4)
    sh.put('A18:L18', 'Матеріально відповідальна особа:')
    sign_row(sh, 19, mvo[0], mvo[1], h=41.6)
    sh.put('A22:L22', 'Інвентаризація: розпочата %s' % ua_date_text(cfg.get('pochatok'), year))
    sh.put('A23:L23', '                          закінчена %s' % ua_date_text(cfg.get('kinets'), year))
    sh.put('A25:L25', 'При інвентаризації встановлено таке:')
    sh.put('A26:A27', '№\nз/п', size=8, h='center', border='box')
    sh.put('B26:B27', 'Рахунок, субрахунок', size=8, h='center', border='box')
    sh.put('C26:D26', 'Матеріальні цінності', size=9, h='center', border='box')
    sh.put('E26:E27', 'Одиниця виміру', size=8, h='center', border='box')
    sh.put('F26:H26', 'Фактична наявність', size=9, h='center', border='box')
    sh.put('I26:K26', 'За даними бухгалтерського обліку', size=9, h='center', border='box')
    sh.put('L26:L27', 'Інші відомості', size=8, h='center', border='box')
    for col, t in zip('CDFGHIJK', ('найменування, вид, сорт, група', 'номенклатурний номер (за наявності)',
                                   'кількість', 'вартість', 'сума', 'кількість', 'вартість', 'сума')):
        sh.put('%s27' % col, t, size=8, h='center', border='box')
    for i, col in enumerate('ABCDEFGHIJKL'):
        sh.put('%s28' % col, str(i + 1), size=8, bold=True, h='center', border='box')
    sh.height(26, 26)
    sh.height(27, 40)
    sh.height(28, 15)
    lines = op['lines']
    r = FIRST_ROW
    for n, ln in enumerate(lines, 1):
        sh.put('A%d' % r, n, size=9, h='center', border='box')
        sh.put('B%d' % r, op['subr'][ln.cat].split()[0], size=9, h='center', border='box')
        sh.put('C%d' % r, ln.name, size=9, border='box')
        sh.put('D%d' % r, ln.fes or None, size=9, h='center', border='box')
        sh.put('E%d' % r, ln.unit, size=9, h='center', border='box')
        sh.put('F%d' % r, None, size=9, h='center', border='box', fmt=qty_fmt(max(ln.qty, ln.fact)))
        sh.put('G%d' % r, ln.price, size=9, h='center', border='box', fmt=MONEY_FMT)
        sh.put('H%d' % r, '=IF(F{0}="","",ROUND(F{0}*G{0},2))'.format(r), size=9, h='center', border='box',
               fmt=MONEY_FMT)
        sh.put('I%d' % r, ln.qty, size=9, h='center', border='box', fmt=qty_fmt(ln.qty))
        sh.put('J%d' % r, ln.price, size=9, h='center', border='box', fmt=MONEY_FMT)
        sh.put('K%d' % r, ln.summa, size=9, h='center', border='box', fmt=MONEY_FMT)
        sh.put('L%d' % r, ln.note or None, size=8, h='center', border='box')
        sh.height(r, max(15.0, 12.0 * max(lines_of(ln.name, WIDTHS['C'], 9), lines_of(ln.note, WIDTHS['L'], 8)) + 3))
        r += 1
    last = r - 1
    rng = lambda col: '%s%d:%s%d' % (col, FIRST_ROW, col, last)                 # noqa: E731
    tot_q = round(sum(ln.qty for ln in lines), 6)
    tot_s = round(sum(ln.summa for ln in lines), 2)
    sh.put('A%d' % r, 'Разом', size=9, bold=True, h='center', border='box')
    for col in 'BCDEGJL':
        sh.put('%s%d' % (col, r), 'X', size=9, bold=True, h='center', border='box')
    sh.put('F%d' % r, '=IF(COUNT({0})=0,"",SUM({0}))'.format(rng('F')), size=9, bold=True, h='center',
           border='box', fmt=sum_fmt(lines))
    sh.put('H%d' % r, '=IF(COUNT({0})=0,"",SUM({0}))'.format(rng('H')), size=9, bold=True, h='center',
           border='box', fmt=MONEY_FMT)
    sh.put('I%d' % r, tot_q, size=9, bold=True, h='center', border='box', fmt=qty_fmt(tot_q))
    sh.put('K%d' % r, tot_s, size=9, bold=True, h='center', border='box', fmt=MONEY_FMT)
    sh.height(r, 25.76)
    r += 2
    n = len(lines)
    nd = 0 if abs(tot_q - round(tot_q)) < EPS else 3
    g = '%s (%s)' % (num(tot_q, nd), words(int(round(tot_q, 3))))
    for label, val, hh in (('Разом за описом: а) кількість порядкових номерів', words(n), 15.2),
                           ('б) загальна кількість одиниць (фактично)', '', 15.2),
                           ('в) вартість фактична', '', 15.2),
                           ('г) загальна кількість одиниць за даними бухгалтерського обліку', g, 28.4),
                           ('ґ) вартість за даними бухгалтерського обліку', money_words(tot_s), 15.2)):
        sh.put('A%d:C%d' % (r, r), label)
        sh.put('D%d:I%d' % (r, r), val, h='center', border='under')
        sh.height(r, max(hh, 13.6 * lines_of(val, 64, 10) + 2))
        sh.put('D%d:I%d' % (r + 1, r + 1), '(прописом)', size=7, h='center')
        sh.height(r + 1, 10)
        r += 2
    r += 1
    kom = cfg['komisia']
    sh.put('A%d:C%d' % (r, r), 'Голова комісії')
    r = sign_row(sh, r, kom['holova'][0], kom['holova'][1])
    for i, (rank, pib) in enumerate(kom.get('chleny', [])):
        if i == 0:
            sh.put('A%d:C%d' % (r, r), 'Члени комісії')
        r = sign_row(sh, r, rank, pib)
    r += 1
    stmt = ('Усі цінності, пронумеровані в цьому інвентаризаційному описі з №1 до №%d, перевірено комісією в '
            'натурі в моїй присутності та внесено в опис. У зв’язку з цим претензій до інвентаризаційної комісії '
            'не маю. Цінності, перелічені в описі, знаходяться на моєму відповідальному зберіганні.' % n)
    sh.put('A%d:L%d' % (r, r), stmt, h='justify', v='top')
    sh.height(r, 41.6)
    r += 2
    sign_date = ua_date_text(cfg.get('data_pidpysu'), year)
    sh.put('A%d:L%d' % (r, r), 'Матеріально відповідальна особа:')
    r = sign_row(sh, r + 1, mvo[0], mvo[1], date=sign_date, h=41.6) + 1
    sh.put('A%d:L%d' % (r, r), 'Командир (начальник) підрозділу:')
    r = sign_row(sh, r + 1, '', '') + 1
    sh.put('A%d:C%d' % (r, r), 'Інформацію за даними бухгалтерського обліку вніс:')
    r = sign_row(sh, r, buh[0], buh[1], h=54.8) + 1
    sh.put('A%d:L%d' % (r, r), 'Вказані в цьому описі дані перевірив:')
    r = sign_row(sh, r + 1, perev[0], perev[1], date=sign_date, h=41.6) + 1
    sh.put('A%d:L%d' % (r, r), 'Графи 9–11 заповнюються бухгалтерською службою.', size=7)
    page_setup(ws, r)
    return dict(n=n, qty=tot_q, summa=tot_s, lost=round(sum(ln.summa for ln in lines if ln.lost), 2),
                fresh=round(sum(ln.fact * ln.price for ln in lines if ln.lost and ln.qty == 0), 2))


CAUSES = {'книга': 'у книзі за всіма документами',
          'облік': 'за обліком, без документів, яких ФЕС ще не провів'}              # причини мінуса: сторона B чи A


def states_of(pending):
    """«роздрукувати», «на підписі» чи «їде на ФЕС»."""
    quoted = ['«%s»' % s for s in pending]
    return ', '.join(quoted[:-1]) + (' чи ' if len(quoted) > 1 else '') + quoted[-1]


def rule_text(zn, pending):
    """Правило граф 9–11: що лишається на боці обліку, хоч у книзі це вже зроблено. Без станів від сервера — лише списання (iprod)."""
    if not pending:
        return 'усі списання — як передані в бухгалтерію' + (', крім позначених «%s»' % zn if zn else '')
    kept = (['позначених «%s»' % zn] if zn else []) + ['тих, яких ФЕС ще не провів (%s)' % states_of(pending)]
    return ('усі документи — як проведені у ФЕС, крім %s: списання лишаються за обліком, прихід є лише фактично, '
            'переміщення — за обліком там, звідки, і фактично там, куди; різниця за позицією опису йде окремими рядками, '
            'а в їхніх примітках названо ці документи' % ' і '.join(kept))


def write_dovidka(ws, cfg, ops, stats, info, sources):
    from openpyxl.styles import Alignment, Font
    bold = Font(name='Times New Roman', size=10, bold=True)
    norm = Font(name='Times New Roman', size=10)
    r = 1
    zn = (cfg.get('znyshcheno') or {}).get('nomer', '')
    pending = list(cfg.get('_pending') or ())

    def put(col, text, b=False, fmt=None):
        c = ws.cell(r, col, text)
        c.font = bold if b else norm
        if fmt:
            c.number_format = fmt

    put(1, 'Чернетки інвентаризаційних описів запасів з книг обліку станом на %s' % ua_date(parse_date(cfg['stan_na'])),
        True)
    r += 1
    put(1, 'Графи 9–11 — залишок у книгах на кінець %s (%s): %s'
        % (info['cut'].strftime('%d.%m.%Y'), rule_text(zn, pending), '; '.join(info['books'])))
    r += 1
    put(1, 'Складено %s програмою («Звіти › Описи 572»), конфіг «%s»'
        % (datetime.datetime.now().strftime('%d.%m.%Y %H:%M'), os.path.basename(cfg['_path'])))
    r += 2
    put(1, 'Описи', True)
    r += 1
    for j, t in enumerate(('Аркуш', 'Місце', 'МВО (джерело)', 'Рядків', 'Кількість', 'Сума, грн',
                           'з них за обліком є, фактично нуль, грн', 'фактично є, за обліком нема, грн')):
        put(j + 1, t, True)
    r += 1
    for op, st in zip(ops, stats):
        mvo = '%s — %s' % (op['mvo'][1], op['mvo'][2]) if op['mvo'][1] else 'не вказано — заповнити від руки'
        for j, v in enumerate((op['sheet'], op['place']['de'], mvo, st['n'], st['qty'], st['summa'], st['lost'],
                               st['fresh'])):
            put(j + 1, v, fmt=MONEY_FMT if j in (5, 6, 7) else None)
        r += 1
    for j, v in enumerate(('Разом', '', '', sum(s['n'] for s in stats), round(sum(s['qty'] for s in stats), 6),
                           round(sum(s['summa'] for s in stats), 2), round(sum(s['lost'] for s in stats), 2),
                           round(sum(s['fresh'] for s in stats), 2))):
        put(j + 1, v, True, fmt=MONEY_FMT if j in (5, 6, 7) else None)
    r += 2
    lost = [(op['sheet'], ln) for op in ops for ln in op['lines'] if ln.lost and ln.fact == 0]      # облік є, фактично нуль
    fresh = [(op['sheet'], ln) for op in ops for ln in op['lines'] if ln.lost and ln.qty == 0]      # фактично є, обліку нема
    if lost:
        why = (['у книгах «%s»' % zn] if zn else []) + (['документи, яких ФЕС ще не провів'] if pending else [])
        put(1, 'За обліком є, фактично нуль (%s)' % ' чи '.join(why), True)
        r += 1
        for sheet, ln in lost:
            put(2, '%s · %s %s: %s %s × %s = %s грн; %s' % (sheet, ln.code, ln.name, num(ln.qty, 3), ln.unit,
                                                           num(ln.price, 2), num(ln.summa, 2), ln.note))
            r += 1
        r += 1
    if fresh:
        put(1, 'Фактично є, за обліком нема (документи, яких ФЕС ще не провів)', True)
        r += 1
        for sheet, ln in fresh:
            put(2, '%s · %s %s: %s %s × %s = %s грн; %s' % (sheet, ln.code, ln.name, num(ln.fact, 3), ln.unit,
                                                           num(ln.price, 2), num(ln.fact * ln.price, 2), ln.note))
            r += 1
        r += 1
    if info['quiet']:
        put(1, 'Документи, яких ФЕС ще не провів, і знищене без окремого рядка в описі (різниця нульова чи фактичного залишку нема)',
            True)
        r += 1
        for nazva, c, name, text in info['quiet']:
            put(2, '%s · %s %s: %s' % (nazva, c, name, text))
            r += 1
        r += 1
    if info['skipped']:
        put(1, 'Поза цими описами (конфіг «poza_opysom»)', True)
        r += 1
        for m, c, name, unit, q, why in info['skipped']:
            put(2, '%s · %s %s: %s %s — %s' % (m, c, name, num(q, 3), unit, why))
            r += 1
        r += 1
    if info['nofes']:
        put(1, 'Без номенклатурного номера ФЕС (конфіг «fes_kody»)', True)
        r += 1
        for c, name in sorted(info['nofes'], key=str):
            put(2, '%s %s' % (c, name))
            r += 1
        r += 1
    if info['neg']:
        put(1, 'Мінуси в книгах — виправити облік (місце з мінусом у книзі рядків не має, його облік лишається в графах 9–11; '
               'опис з мінусом обліку йде з обліком 0)', True)
        r += 1
        for m, c, name, q, cause in info['neg']:
            put(2, '%s · %s %s: %s (%s)' % (m, c, name, num(q, 3), CAUSES[cause]))
            r += 1
        r += 1
    put(1, 'Підписанти', True)
    r += 1
    for s in sources:
        put(2, s)
        r += 1
    for col, w in zip('ABCDEFGH', (22, 44, 46, 8, 14, 16, 18, 18)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=False)


def sheet_title(op):
    return re.sub(r'[\[\]:*?/\\]', ' ', op['nazva'])[:31]


def build(cfg, out_path, dry=False, log=print, dobovka=None):
    """Описи з книг `cfg['knyhy']`; `dobovka` — Добовка для МВО й «перевірив» з ключем `dobovka`."""
    books = [read_book(p) for p in cfg['knyhy']]
    ops, info = collect(cfg, books)
    stan = parse_date(cfg['stan_na'])
    signers, dob_file = None, None
    need_dob = any((p.get('mvo') or {}).get('dobovka') for p in cfg['mista']) or \
        (cfg.get('pereviryv') or {}).get('dobovka')
    if need_dob:
        dob_file = dobovka
        if not dob_file:
            raise SystemExit('Добовки нема: вкажіть теку Добовки в «Налаштування» → «Теки»')
        signers = Signers(dob_file, stan)
    sources = ['Добовка: %s (рядок на %s)' % (dob_file, stan.strftime('%d.%m.%Y'))] if dob_file else []
    perev = person_of(cfg.get('pereviryv'), signers, 'перевірив')
    buh = cfg.get('buhgalter') or ['', '']
    sources.append('Перевірив: %s, %s — %s' % perev)
    for op in ops:
        op['mvo'] = person_of(op['place'].get('mvo'), signers, 'МВО %s' % op['nazva'])
        op['sheet'] = sheet_title(op)
        sources.append('МВО %s: %s' % (op['nazva'], '%s, %s — %s' % (op['mvo'][1], op['mvo'][0], op['mvo'][2])
                                        if op['mvo'][1] else 'не вказано'))
        log('%-10s рядків %3d  %14s грн  (за обліком є, фактично нуль %s; фактично є, за обліком нема %s)  МВО %s'
            % (op['sheet'], len(op['lines']), num(sum(ln.summa for ln in op['lines']), 2),
               num(sum(ln.summa for ln in op['lines'] if ln.lost), 2),
               num(sum(ln.fact * ln.price for ln in op['lines'] if ln.lost and ln.qty == 0), 2), op['mvo'][1] or '—'))
    for m, c, name, unit, q, why in info['skipped']:
        log('поза описами: %s · %s %s %s %s — %s' % (m, c, name, num(q, 3), unit, why))
    for m, c, name, q, cause in info['neg']:
        log('мінус: %s · %s %s: %s (%s)' % (m, c, name, num(q, 3), CAUSES[cause]))
    for nazva, c, name, text in info['quiet']:
        log('без окремого рядка: %s · %s %s: %s' % (nazva, c, name, text))
    if dry:
        return ops, info
    import openpyxl
    wb = openpyxl.Workbook()
    dov = wb.active
    dov.title = 'Довідка'
    stats = [write_opys(wb.create_sheet(op['sheet']), cfg, op, op['mvo'], perev[:2], buh) for op in ops]
    write_dovidka(dov, cfg, ops, stats, info, sources)
    wb.calculation.fullCalcOnLoad = True
    folder = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(folder, exist_ok=True)
    tmp = os.path.join(folder, '_tmp_' + os.path.basename(out_path))
    excel_names.save_book(wb, tmp)
    try:
        os.replace(tmp, out_path)
    except PermissionError:
        raise SystemExit('файл відкрито в Excel — закрийте і повторіть: %s (нова версія — %s)' % (out_path, tmp))
    log('записано: %s' % out_path)
    return ops, info
