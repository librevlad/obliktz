# -*- coding: utf-8 -*-
"""Спільне бланка інвентаризаційного опису запасів за типовою формою наказу Мінфіну від 17.06.2015 № 572 — з
iprod `tools/inv/opysy_prod.py` (коміт 4e143c6) дослівно, цілими блоками: константи бланка, дати й суми словами,
підписанти з аркуша «Підписанти» Добовки, аркуш опису й друк. Ним складає описи з книг обліку `prod.opysy_knyhy`;
опис продуктів і води з бази ПС (решта файла iprod) — у випуску 2.0.0. Єдина відмінність від iprod — слова відмови
`person_of` без Добовки: вони називають «Налаштування» → «Теки», а не `settings.ini`."""
import datetime
import math
import os


UA_MONTHS_GEN = ('січня', 'лютого', 'березня', 'квітня', 'травня', 'червня', 'липня', 'серпня', 'вересня',
                 'жовтня', 'листопада', 'грудня')
ALL = '*'                             # «vch» місця: усі склади частини одним описом
MIN_QTY = {'шт': 1.0}                 # менше — залишок округлення ПС (не в опис)
MIN_QTY_DEFAULT = 0.05                # кг: хвости розбивки рядків ПС по договорах — десятки грамів
FIRST_ROW = 29                        # перший рядок таблиці опису
TNR = 'Times New Roman'
WIDTHS = dict(A=5, B=10, C=31, D=12, E=8, F=10, G=11, H=13, I=10, J=11, K=13, L=19)
TITLE_ROWS = '26:28'                  # шапка таблиці на кожній сторінці
QTY_FMT = '#,##0.000'
MONEY_FMT = '#,##0.00'
KIT_FMT = '#,##0.00000'
BLANK_DATE = '«___» _______________ %dр.'


# ---------------------------------------------------------------- дрібниці

def parse_date(s):
    return datetime.datetime.strptime(s.strip(), '%d.%m.%Y').date()


def ua_date(d):
    """«1» жовтня 2026р."""
    return '«%d» %s %dр.' % (d.day, UA_MONTHS_GEN[d.month - 1], d.year)


def ua_date_text(s, year):
    """Дата з конфігу (дд.мм.рррр) — «1» жовтня 2026р.; порожня — рядок під руку."""
    return ua_date(parse_date(s)) if (s or '').strip() else BLANK_DATE % year


def num(x, nd):
    """1234.5 -> '1 234,50' (нерозривний пробіл між тисячами)."""
    return '{:,.{nd}f}'.format(x, nd=nd).replace(',', '\xa0').replace('.', ',')


_U1 = ('', 'один', 'два', 'три', 'чотири', "п'ять", 'шість', 'сім', 'вісім', "дев'ять")
_U1F = ('', 'одна', 'дві', 'три', 'чотири', "п'ять", 'шість', 'сім', 'вісім', "дев'ять")
_U10 = ('десять', 'одинадцять', 'дванадцять', 'тринадцять', 'чотирнадцять', "п'ятнадцять", 'шістнадцять',
        'сімнадцять', 'вісімнадцять', "дев'ятнадцять")
_UD = ('', '', 'двадцять', 'тридцять', 'сорок', "п'ятдесят", 'шістдесят', 'сімдесят', 'вісімдесят',
       "дев'яносто")
_UH = ('', 'сто', 'двісті', 'триста', 'чотириста', "п'ятсот", 'шістсот', 'сімсот', 'вісімсот', "дев'ятсот")


def _triple(n, fem):
    w = []
    if n >= 100:
        w.append(_UH[n // 100])
        n %= 100
    if 10 <= n <= 19:
        w.append(_U10[n - 10])
        n = 0
    elif n >= 20:
        w.append(_UD[n // 10])
        n %= 10
    if n:
        w.append((_U1F if fem else _U1)[n])
    return w


def plural(n, forms):
    if 11 <= n % 100 <= 14:
        return forms[2]
    if n % 10 == 1:
        return forms[0]
    if 2 <= n % 10 <= 4:
        return forms[1]
    return forms[2]


def words(n, fem=False):
    """Ціле число прописом (fem — рід одиниць: гривня — жіночий)."""
    n = int(n)
    if n == 0:
        return 'нуль'
    w = []
    mln, rest = divmod(n, 1000000)
    th, units = divmod(rest, 1000)
    if mln:
        w += _triple(mln, False) + [plural(mln, ('мільйон', 'мільйони', 'мільйонів'))]
    if th:
        w += _triple(th, True) + [plural(th, ('тисяча', 'тисячі', 'тисяч'))]
    w += _triple(units, fem)
    return ' '.join(w)


def money_words(x):
    """12345.6 -> 'дванадцять тисяч триста сорок п'ять грн. 60 коп.'"""
    kop = int(round(x * 100))
    return '%s грн. %02d коп.' % (words(kop // 100, fem=True), kop % 100)


def lines_of(text, width, size):
    """Скільки рядків займе текст у клітинці ширини `width` (одиниці Excel) шрифтом `size`."""
    cpl = max(1, int(width * 10.0 / size))
    return sum(max(1, math.ceil(len(part) / cpl)) for part in str(text or '').split('\n'))


class Signers:
    """Аркуш «Підписанти» Добовки: (підрозділ, роль) -> посада, звання, ПІБ на дату (останній непорожній
    рядок не пізніше дати). Шапка — три рядки: підрозділ / роль / поле, дані з 4-го рядка."""

    def __init__(self, path, date):
        import openpyxl
        self.path, self.date = path, date
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb['Підписанти']
        rows = ws.iter_rows(values_only=True)
        h1, h2, h3 = next(rows), next(rows), next(rows)
        cols, unit, role = {}, None, None
        for i, (a, b, c) in enumerate(zip(h1, h2, h3)):
            if a:
                unit, role = str(a).strip().lower(), None
            if b:
                role = str(b).strip().lower()
            if unit and role and c:
                cols[(unit, role, str(c).strip().lower())] = i
        self.cols, self.found = cols, {}
        for r in rows:
            d = r[1] if len(r) > 1 else None
            if not isinstance(d, datetime.datetime) or d.date() > date:
                continue
            for k, i in cols.items():
                v = r[i] if i < len(r) else None
                if v not in (None, ''):
                    self.found[k] = str(v).strip()
        wb.close()

    def person(self, unit, role):
        unit, role = unit.strip().lower(), role.strip().lower()
        get = lambda f: self.found.get((unit, role, f), '')               # noqa: E731
        if not any((unit, role, f) in self.cols for f in ('посада', 'звання', 'піб')):
            raise SystemExit('Добовка %s: на аркуші «Підписанти» нема ролі «%s · %s»'
                             % (os.path.basename(self.path), unit, role))
        return get('посада'), get('звання'), get('піб')


def person_of(spec, signers, what):
    """Підписант із конфігу: {"dobovka": [підрозділ, роль]} або {"posada": …, "pib": …}
    -> (посада з званням, «Ім'я ПРІЗВИЩЕ», джерело)."""
    spec = spec or {}
    if spec.get('dobovka'):
        if signers is None:
            raise SystemExit('%s: потрібна Добовка, вкажіть теку Добовки в «Налаштування» → «Теки»' % what)
        posada, zv, pib = signers.person(*spec['dobovka'])
        text = ', '.join(x for x in (posada, zv) if x)
        return text, pib, 'Добовка «Підписанти» %s · %s' % tuple(spec['dobovka'])
    return spec.get('posada', ''), spec.get('pib', ''), 'конфіг'


class Sheet:
    """Запис у аркуш із стилями форми (Times New Roman, тонкі рамки, об'єднання)."""

    def __init__(self, ws):
        from openpyxl.styles import Alignment, Border, Font, Side
        self.ws, self.Alignment, self.Font = ws, Alignment, Font
        thin = Side(style='thin')
        self.box = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.under = Border(bottom=thin)
        for col, w in WIDTHS.items():
            ws.column_dimensions[col].width = w

    def put(self, rng, value=None, size=10, bold=False, h='left', v='center', border=None, fmt=None,
            wrap=True):
        ws = self.ws
        if ':' in rng:
            ws.merge_cells(rng)
        first = ws[rng.split(':')[0]]
        if value is not None:
            first.value = value
        cells = ws[rng] if ':' in rng else ((first,),)
        for row in cells:
            for c in row:
                c.font = self.Font(name=TNR, size=size, bold=bold)
                c.alignment = self.Alignment(horizontal=h, vertical=v, wrap_text=wrap)
                if border == 'box':
                    c.border = self.box
                elif border == 'under':
                    c.border = self.under
                if fmt:
                    c.number_format = fmt
        return first

    def height(self, r, h):
        self.ws.row_dimensions[r].height = h


def sign_row(sh, r, posada, pib, date=None, h=15.2):
    """Рядок підпису: [дата A:C] посада D:G | підпис H:I | ініціали, прізвище J:L + підписи під рядком."""
    if date is not None:
        sh.put('A%d:C%d' % (r, r), date, h='center', border='under')
    sh.put('D%d:G%d' % (r, r), posada, h='center', v='bottom', border='under')
    sh.put('H%d:I%d' % (r, r), None, h='center', border='under')
    sh.put('J%d:L%d' % (r, r), pib, h='center', v='bottom', border='under')
    sh.height(r, max(h, 13.6 * max(lines_of(posada, 41, 10), lines_of(pib, 40, 10)) + 2))
    sh.put('D%d:G%d' % (r + 1, r + 1), '(посада)', size=7, h='center')
    sh.put('H%d:I%d' % (r + 1, r + 1), '(підпис)', size=7, h='center')
    sh.put('J%d:L%d' % (r + 1, r + 1), '(ініціали, прізвище)', size=7, h='center')
    sh.height(r + 1, 10)
    return r + 2


def page_setup(ws, last_row):
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    m = ws.page_margins
    m.left, m.right, m.top, m.bottom, m.header, m.footer = 0.4, 0.4, 0.5, 0.75, 0.5, 0.3
    ws.oddFooter.center.text = '&P із &N'
    # імена _xlnm.* українського Excel не впізнає — після запису їх переписує excel_names.localize_file
    ws.print_area = 'A1:L%d' % last_row
    ws.print_title_rows = TITLE_ROWS
    ws.row_breaks.append(Break(id=24))
    ws.sheet_view.showGridLines = False
