# -*- coding: utf-8 -*-
"""Книга «Облік ОП» — запасний шлях обліку посуду, миючих засобів і серветок.

Програма вивантажує в книгу все, що знає про книгу ОП, і приймає її назад: коли програма
недоступна, книгу ведуть руками, а потім програма — навіть чиста — приймає її й має ту саму
базу. Будова — як у «Облік ОП 6.xlsx», якою служба вела облік до програми: ті самі аркуші,
таблиці з A1, ті самі колонки в тому самому порядку (за ними працюють скрипти служби);
колонки, яких у старій книзі не було, — лише після наявних. Формули книги лишаються
формулами (обчислювані колонки таблиць: рядок, вписаний руками, отримує їх сам), але кожна
несе й значення на момент вивантаження: скрипти читають значення.
"""
import contextlib
import datetime
import hashlib
import json
import math
import os
import re
import sqlite3
import traceback
import zipfile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

import op_norms
import stable_hash

BOOK = "ОП"
# Відбитки книги, яку програма сама записала чи прийняла, — байтів і вмісту (як у книги «Закупівлі»): така книга на диску —
# не руками правлена, вивантаження перезаписує її без звірки з базою. app_setting, серед state_db.SERVER_SETTINGS.
SHA_SETTING = "op_book_sha"
KEEP_SHAS = 10
GROUP_OF_TYPE = {"одноразовий посуд": "ОП.1", "серветки": "ОП.2", "миючі засоби": "ОП.3",
                 "господарчі товари": "ОП.4"}
TYPE_OF_GROUP = {v: k for k, v in GROUP_OF_TYPE.items()}
CAT_OF_GROUP = {"ОП.1": "ОП", "ОП.2": "ПС", "ОП.3": "МЗ", "ОП.4": "ГТ"}
GROUP_OF_CAT = {v: k for k, v in CAT_OF_GROUP.items()}
# Назви груп — ті самі, що в довіднику бази (nomen_group.name, міграція 026): слова відмови «рядок норми не з групи позиції» при
# прийомі книги збігаються зі словами сервера (state_db._save_items).
GROUP_NAME = {g: t[:1].upper() + t[1:] for g, t in TYPE_OF_GROUP.items()}
STATUSES = ["немає витяга", "на підписі", "їде на ФЕС", "проведено", "переробка"]
ARCHIVE = "архів"
UNSURE = "уточнити"

# «рядок норми» — рядок норми наказу 390, вибраний у картці позиції (назва рядка з op_norms.PER_RATION): за ним акт ділить
# межу між кодами; книга 1.12.0 і старіша колонки не має — тоді рядків прийом не чіпає (`lineCol`).
LINE_COL = "рядок норми"
ITEM_COLS = ["код", "найменування", "категорія", "од.вим.", "ціна", "поз", "тип", "дободач", "примітка",
             "рядок 2/Прод", "множник", "перевірено", LINE_COL]
IN_COLS = ["док_тип", "док_дата", "док_номер", "від_кого", "кому", "найменування", "код", "найм", "од_вим",
           "ціна", "кількість", "категорія", "сума", "примітка", "скан", "статус", "рапорт", "наказ",
           "графа 2/прод"]
MV_COLS = ["док_тип", "док_дата", "док_номер", "від_кого", "кому", "найменування", "код", "од_вим", "ціна",
           "кількість", "категорія", "зал_відпр", "зал_одерж", "нм", "одвим", "примітка", "статус", "скан"]
WR_COLS = ["док_тип", "док_дата", "док_номер", "підрозділ", "найменування", "код", "транз_найменування",
           "кількість", "категорія", "статус", "примітка", "рапорт", "наказ", "скан", "кому", "графа 2/прод", "норма_дд"]
SUB_COLS = ["назва", "підпорядкований", "вид", "закритий", "псевдоніми"]
SHEETS = {"incoming": ("Прибуток", "tblIncoming", IN_COLS), "movement": ("Переміщення", "tblMovement", MV_COLS),
          "writeoffs": ("Списання", "tblWriteoffs", WR_COLS)}
# Акти списання за нормами: рядок «Акти» знаходить свій документ у «Списання» за датою, номером і підрозділом;
# підписанти — знімком, яким акт склали; дободачі — широка таблиця, як у Добовці; пункти — tblPoints на «Довідка».
ACT_COLS = ["док_дата", "док_номер", "підрозділ", "серія", "вид", "пункт", "період_з", "період_по", "дд",
            "накладна_номер", "накладна_дата", "рапорт_дата", "звірка_номер", "звірка_дата", "розрахунок_дата", "примітка"]
SIGNER_COLS = ["док_дата", "док_номер", "підрозділ", "роль", "порядок", "посада", "звання", "ім'я_прізвище"]
RATION_HEAD = ["дата", "витяг", "джерело"]
POINT_COLS = ["пункт", "підрозділ", "їдальня", "назва_у_звірці", "закрито"]
KIND_WORDS = {"op390": "за нормами 390"}
PAPER_CELLS = {"накладна": ("накладна_номер", "накладна_дата"), "рапорт": (None, "рапорт_дата"),
               "звірка": ("звірка_номер", "звірка_дата"), "розрахунок": (None, "розрахунок_дата")}


def _me(table, col):
    return f"{table}[[#This Row],[{col}]]"


def _rest_formula(side):
    """Залишок місця з колонки side («від_кого» чи «кому») на дату рядка — «зал_відпр»/«зал_одерж»."""
    code, place, day = _me("tblMovement", "код"), _me("tblMovement", side), _me("tblMovement", "док_дата")
    return (f'SUMIFS(tblIncoming[кількість],tblIncoming[код],{code},tblIncoming[кому],{place},'
            f'tblIncoming[док_дата],"<="&{day})'
            f'+SUMIFS(tblMovement[кількість],tblMovement[код],{code},tblMovement[кому],{place},'
            f'tblMovement[док_дата],"<="&{day})'
            f'-SUMIFS(tblMovement[кількість],tblMovement[код],{code},tblMovement[від_кого],{place},'
            f'tblMovement[док_дата],"<="&{day})'
            f'-SUMIFS(tblWriteoffs[кількість],tblWriteoffs[код],{code},tblWriteoffs[підрозділ],{place},'
            f'tblWriteoffs[док_дата],"<="&{day})')


def _lookup(table, n):
    return f"VLOOKUP({_me(table, 'код')},Номенклатура!$A:$E,{n},FALSE)"


# Обчислювані колонки: та сама формула в описі колонки таблиці й у клітинках.
CALC = {
    "Номенклатура": {"поз": f'{_me("tblItems", "код")}&" "&{_me("tblItems", "найменування")}&" "&'
                            f'{_me("tblItems", "ціна")}'},
    "Прибуток": {"код": f"INT(TRIM(LEFT({_me('tblIncoming', 'найменування')},5)))",
                 "од_вим": _lookup("tblIncoming", 4), "ціна": _lookup("tblIncoming", 5),
                 "категорія": _lookup("tblIncoming", 3),
                 "сума": f"{_me('tblIncoming', 'ціна')}*{_me('tblIncoming', 'кількість')}"},
    "Переміщення": {"код": f"INT(TRIM(LEFT({_me('tblMovement', 'найменування')},5)))",
                    "од_вим": _lookup("tblMovement", 4), "ціна": _lookup("tblMovement", 5),
                    "категорія": _lookup("tblMovement", 3),
                    "зал_відпр": _rest_formula("від_кого"), "зал_одерж": _rest_formula("кому")},
    "Списання": {"код": f"INT(TRIM(LEFT({_me('tblWriteoffs', 'найменування')},5)))",
                 "категорія": _lookup("tblWriteoffs", 3)},
}


def _num(v):
    """Ціле — цілим: у книзі 9000, а не 9000.0."""
    v = float(v or 0)
    return int(v) if v.is_integer() else round(v, 6)


def _code_cell(code):
    return int(code) if str(code).isdigit() else str(code)


def _pos(it):
    """«поз» книги: «код найменування ціна» — ціна з комою, як її показує Excel українською."""
    price = ("%g" % float(it["price"] or 0)).replace(".", ",")
    return f"{it['code']} {it['name']} {price}"


def _checked_text(it):
    """«перевірено»: хто й коли, «уточнити» — прив'язку ще не перевірено, порожньо — прив'язки немає."""
    if it.get("checked"):
        return it["checked"]
    return UNSURE if it.get("row2") else None


def model_from_db(con):
    """Усе, що база знає про книгу ОП, — у вигляді книги."""
    import state_db                                              # noqa: PLC0415 — модуль сервера
    items = []
    for r in con.execute("""
            SELECT n.id, n.code, n.name, g.code AS grp, u.code AS uom, n.per_ration, n.norm_line, n.note, n.archived_at,
                   COALESCE(p.price_kop, n.app_price_kop) AS price_kop
              FROM nomen n JOIN nomen_group g ON g.id = n.group_id JOIN uom u ON u.id = n.uom_id
              LEFT JOIN nomen_last_price p ON p.nomen_id = n.id
             WHERE g.book = 'ОП' ORDER BY n.code"""):
        b = con.execute("""SELECT rl.row_no, m.factor, m.checked_by, m.checked_on FROM nomen_report_line m
                             JOIN report_line rl ON rl.id = m.report_line_id JOIN report_form f ON f.id = rl.form_id
                            WHERE m.nomen_id = ? AND f.code = '2/Прод' AND rl.row_no IS NOT NULL""",
                        (r["id"],)).fetchone()
        skip = con.execute("""SELECT s.reason FROM nomen_form_skip s JOIN report_form f ON f.id = s.form_id
                               WHERE s.nomen_id = ? AND f.code = '2/Прод'""", (r["id"],)).fetchone()
        checked = (" ".join(x for x in (b["checked_by"], b["checked_on"]) if x) or None) if b else None
        items.append({"code": r["code"], "name": r["name"], "cat": CAT_OF_GROUP[r["grp"]], "uom": r["uom"],
                      "price": (r["price_kop"] or 0) / 100, "type": TYPE_OF_GROUP[r["grp"]],
                      "archived": bool(r["archived_at"]), "per_ration": r["per_ration"], "note": r["note"] or "",
                      "row2": b["row_no"] if b else None, "factor": b["factor"] if b else 1.0,
                      "checked": checked, "skip2": skip["reason"] if skip else None, "line": r["norm_line"] or None})
    codes = {i["code"] for i in items}
    subs = [{"name": r["name"], "parent": r["parent"] or "", "kind": r["kind"], "closed": not r["is_active"],
             "aliases": [a[0] for a in con.execute("SELECT name FROM subdivision_alias WHERE subdivision_id = ? "
                                                   "AND source = 'книга ОП' ORDER BY id", (r["id"],))]}
            for r in con.execute("""SELECT s.id, s.name, p.name AS parent, k.code AS kind, s.is_active
                                      FROM subdivision s JOIN subdivision_kind k ON k.id = s.kind_id
                                      LEFT JOIN subdivision p ON p.id = s.parent_id ORDER BY s.sort, s.id""")]
    parties = [{"name": r["name"], "kind": r["kind"]} for r in con.execute(
        "SELECT c.name, k.code AS kind FROM counterparty c JOIN counterparty_kind k ON k.id = c.kind_id ORDER BY c.name")]
    fes, meta = state_db._load_doc_fes(con), state_db._load_doc_meta(con)
    norms = {(r[0], r[1]): r[2] for r in con.execute("SELECT i.document_id, n.code, i.norm FROM act_item i "
                                                     "JOIN nomen n ON n.id = i.nomen_id")}
    docs, at = [], {}
    for journal, rows in state_db._load_docs(con).items():
        for r in rows:
            date, paper, no, src, dst, code, qty, note = r[:8]
            doc_id = r[10] if len(r) > 10 else None
            if str(code) not in codes:
                continue
            key = (journal, doc_id) if doc_id else (journal, date, no, src, dst)
            d = at.get(key)
            if d is None:
                f, m = fes.get(str(doc_id), {}), meta.get(str(doc_id), {})
                d = at[key] = {"journal": journal, "type": paper or "", "date": date, "no": str(no), "src": src or "",
                               "dst": dst or "", "status": f.get("status", ""), "report": m.get("report", ""),
                               "order": m.get("order", ""), "scan": m.get("scan", ""), "col": m.get("col"),
                               "lines": []}
                docs.append(d)
            same = next((x for x in d["lines"] if x["code"] == str(code) and x["note"] == (note or "")), None)
            if same:                                              # рядки однієї позиції по партіях — один рядок книги
                same["qty"] = _num(float(same["qty"]) + float(qty))
            else:                                                 # норма акта — одна на код, у злитті не важить
                line = {"code": str(code), "qty": _num(qty), "note": note or ""}
                if journal == "writeoffs":
                    line["norm"] = norms.get((doc_id, str(code)))
                d["lines"].append(line)
    renames = [[a[0], a[1]] for a in con.execute(
        "SELECT a.name, s.name FROM subdivision_alias a JOIN subdivision s ON s.id = a.subdivision_id "
        "WHERE a.source = 'книга ОП' ORDER BY a.id")]
    model = {"items": items, "subs": subs, "parties": parties, "renames": renames, "docs": docs}
    model.update(_acts_model(con))
    return model


def _acts_model(con):
    """Пункти харчування, дободачі, акти й підписанти — частина книги ОП (повна альтернатива програмі).
    Акт і підписант названі ключем свого документа: дата, номер (як у «Списання» після читання) і підрозділ."""
    points = [{"name": r["name"], "sub": r["sub"], "hall": r["hall"] or "", "title": r["title"] or "",
               "closed": r["closed_at"] or ""}
              for r in con.execute("""SELECT p.*, s.name AS sub FROM food_point p
                                        JOIN subdivision s ON s.id = p.subdivision_id ORDER BY p.id""")]
    days = {}

    def day(on):
        return days.setdefault(on, {"date": on, "order": "", "src": "", "values": {}})

    for r in con.execute("""SELECT p.name, d.on_date, d.rations FROM ration_day d
                              JOIN food_point p ON p.id = d.point_id ORDER BY d.on_date, p.id"""):
        day(r["on_date"])["values"][r["name"]] = r["rations"]
    for r in con.execute("SELECT * FROM ration_date ORDER BY on_date"):
        day(r["on_date"]).update(order=r["order_no"] or "", src=r["source"] or "")
    papers, people = {}, {}
    for r in con.execute("SELECT * FROM act_paper ORDER BY document_id, role"):
        papers.setdefault(r["document_id"], {})[r["role"]] = {"no": r["number"] or "", "date": r["doc_date"] or ""}
    for r in con.execute("SELECT * FROM act_signer ORDER BY document_id, role, sort"):
        people.setdefault(r["document_id"], []).append(r)
    acts, signers = [], []
    for r in con.execute("""SELECT a.*, d.number, d.doc_date, s.name AS sub, p.name AS point FROM act a
                              JOIN document d ON d.id = a.document_id JOIN subdivision s ON s.id = d.from_subdivision_id
                              LEFT JOIN food_point p ON p.id = a.point_id ORDER BY d.doc_date, d.number, d.id"""):
        head = {"date": r["doc_date"], "no": " ".join(str(r["number"]).split()), "sub": r["sub"]}
        acts.append({**head, "series": r["series"], "kind": r["kind"], "point": r["point"] or "", "from": r["period_from"],
                     "to": r["period_to"], "rations": r["rations"], "papers": papers.get(r["document_id"], {}),
                     "note": r["note"] or ""})
        signers += [{**head, "role": x["role"], "sort": x["sort"], "pos": x["position"] or "", "rank": x["rank"] or "",
                     "name": x["name"] or ""} for x in people.get(r["document_id"], [])]
    return {"points": points, "rations": {"points": [p["name"] for p in points], "days": [days[k] for k in sorted(days)]},
            "acts": acts, "signers": signers}


def _balances(model, as_of):
    """{(код, місце): к-сть} на кінець дня as_of — як SUMIFS аркуша «Залишки»."""
    out = {}
    for d in model["docs"]:
        if as_of and d["date"] > as_of:
            continue
        for ln in d["lines"]:
            q = float(ln["qty"] or 0)
            if d["journal"] == "incoming":
                out[(ln["code"], d["dst"])] = out.get((ln["code"], d["dst"]), 0) + q
            elif d["journal"] == "movement":
                out[(ln["code"], d["src"])] = out.get((ln["code"], d["src"]), 0) - q
                out[(ln["code"], d["dst"])] = out.get((ln["code"], d["dst"]), 0) + q
            else:
                out[(ln["code"], d["src"])] = out.get((ln["code"], d["src"]), 0) - q
    return out


def _table(ws, name, cols, rows_n, calc=None):
    """Таблиця з A1; колонки з `calc` — обчислювані: Excel сам продовжує їх на новий рядок."""
    ref = f"A1:{get_column_letter(len(cols))}{max(rows_n + 1, 2)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    t._initialise_columns()
    for col, title in zip(t.tableColumns, cols):
        col.name = title
        if calc and title in calc:
            col.calculatedColumnFormula = TableFormula(attr_text=calc[title])
    ws.add_table(t)


def _doc_order(d):
    return (d["date"], d["no"], d["src"], d["dst"])


def write_book(model, path, as_of=None):
    as_of = as_of or datetime.date.today().isoformat()
    items = sorted(model["items"], key=lambda i: (i["archived"], int(i["code"]) if str(i["code"]).isdigit() else 0,
                                                  str(i["code"])))
    by_code = {i["code"]: i for i in items}
    places = [s["name"] for s in model["subs"] if not s["closed"]]
    wb = Workbook()
    cached = {}                                                   # аркуш -> {клітинка: значення}

    def put(ws, ref, formula, value):
        ws[ref] = formula
        cached.setdefault(ws.title, {})[ref] = value

    # Довідка: місця, типи, статуси; перейменування; контрагенти.
    ws = wb.active
    ws.title = "Довідка"
    ws.append(["Підрозділи", "типи", "статус"])
    types = list(GROUP_OF_TYPE) + [ARCHIVE]
    n_ref = max(len(places), len(types), len(STATUSES))
    for n in range(n_ref):
        ws.append([places[n] if n < len(places) else None, types[n] if n < len(types) else None,
                   STATUSES[n] if n < len(STATUSES) else None])
    _table(ws, "Таблиця4", ["Підрозділи", "типи", "статус"], n_ref)
    ws["H1"], ws["I1"] = "стара назва", "нова назва"
    for n, (old, new) in enumerate(model["renames"], 2):
        ws.cell(n, 8, old)
        ws.cell(n, 9, new)
    ws.add_table(Table(displayName="tblRename", ref=f"H1:I{max(len(model['renames']) + 1, 2)}"))
    ws["K1"], ws["L1"] = "контрагент", "вид"
    for n, p in enumerate(model["parties"], 2):
        ws.cell(n, 11, p["name"])
        ws.cell(n, 12, p["kind"])
    ws.add_table(Table(displayName="tblParties", ref=f"K1:L{max(len(model['parties']) + 1, 2)}"))
    if model.get("points") is not None:
        for i, title in enumerate(POINT_COLS):
            ws.cell(1, 14 + i, title)
        for n, p in enumerate(model["points"], 2):
            for i, v in enumerate((p["name"], p["sub"], p["hall"] or None, p["title"] or None, _day_cell(p["closed"]))):
                ws.cell(n, 14 + i, v)
            ws.cell(n, 18).number_format = "dd.mm.yyyy"
        ws.add_table(Table(displayName="tblPoints", ref=f"N1:R{max(len(model['points']) + 1, 2)}"))

    # Номенклатура.
    ws = wb.create_sheet("Номенклатура")
    ws.append(ITEM_COLS)
    for r, it in enumerate(items, 2):
        line = op_norms.LINES.get(it.get("line") or "")
        ws.append([_code_cell(it["code"]), it["name"], it["cat"], it["uom"], _num(it["price"]), None,
                   ARCHIVE if it["archived"] else it["type"], it["per_ration"], it["note"] or None,
                   it["skip2"] or it["row2"], _num(it["factor"]) if it["row2"] else None, _checked_text(it),
                   line["name"] if line else None])
        put(ws, f"F{r}", "=" + CALC["Номенклатура"]["поз"], _pos(it))
    _table(ws, "tblItems", ITEM_COLS, len(items), CALC["Номенклатура"])

    # Журнали.
    rest_of = {}
    for journal, (title, table, cols) in SHEETS.items():
        ws = wb.create_sheet(title)
        ws.append(cols)
        f = {k: "=" + v for k, v in CALC[title].items()}
        r = 1
        for d in sorted((d for d in model["docs"] if d["journal"] == journal), key=_doc_order):
            for ln in d["lines"]:
                r += 1
                it = by_code[ln["code"]]
                date = datetime.date.fromisoformat(d["date"])
                code = _code_cell(it["code"])
                if journal == "incoming":
                    ws.append([d["type"], date, d["no"], d["src"], d["dst"], _pos(it), None, _pos(it), None, None,
                               _num(ln["qty"]), None, None, ln["note"] or None, d["scan"] or None,
                               d["status"] or None, d["report"] or None, d["order"] or None, d["col"]])
                    put(ws, f"G{r}", f["код"], code)
                    put(ws, f"I{r}", f["од_вим"], it["uom"])
                    put(ws, f"J{r}", f["ціна"], _num(it["price"]))
                    put(ws, f"L{r}", f["категорія"], it["cat"])
                    put(ws, f"M{r}", f["сума"], round(float(it["price"]) * float(ln["qty"]), 2))
                elif journal == "movement":
                    ws.append([d["type"], date, d["no"], d["src"], d["dst"], _pos(it), None, None, None,
                               _num(ln["qty"]), None, None, None, _pos(it), it["uom"], ln["note"] or None,
                               d["status"] or None, d["scan"] or None])
                    put(ws, f"G{r}", f["код"], code)
                    put(ws, f"H{r}", f["од_вим"], it["uom"])
                    put(ws, f"I{r}", f["ціна"], _num(it["price"]))
                    put(ws, f"K{r}", f["категорія"], it["cat"])
                    bal = rest_of.setdefault(d["date"], _balances(model, d["date"]))
                    put(ws, f"L{r}", f["зал_відпр"], _num(bal.get((ln["code"], d["src"]), 0)))
                    put(ws, f"M{r}", f["зал_одерж"], _num(bal.get((ln["code"], d["dst"]), 0)))
                else:
                    ws.append([d["type"], date, d["no"], d["src"], _pos(it), None, None, _num(ln["qty"]), None,
                               d["status"] or None, ln["note"] or None, d["report"] or None, d["order"] or None,
                               d["scan"] or None, d["dst"] or None, d["col"], ln.get("norm")])
                    put(ws, f"F{r}", f["код"], code)
                    put(ws, f"I{r}", f["категорія"], it["cat"])
                ws.cell(r, 2).number_format = "dd.mm.yyyy"
        _table(ws, table, cols, r - 1, CALC[title])
        _validations(ws, journal, r)
    _act_sheets(wb, model)

    _rests_sheet(wb, model, items, places, as_of, put)
    _moves_sheet(wb, model, items, as_of)

    ws = wb.create_sheet("Підрозділи")
    ws.append(SUB_COLS)
    for s in model["subs"]:
        ws.append([s["name"], s["parent"] or None, s["kind"], "так" if s["closed"] else None,
                   ", ".join(s["aliases"]) or None])
    _table(ws, "tblSubs", SUB_COLS, len(model["subs"]))

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    _put_cached(path, cached)


def _validations(ws, journal, last):
    """Списки, як у старій книзі: позиції, місця, статуси."""
    def lst(formula):
        return DataValidation(type="list", formula1=formula, allow_blank=True)
    pos = lst("Номенклатура!$F$2:$F$2000")
    place = lst("Довідка!$A$2:$A$200")
    status = lst("Довідка!$C$2:$C$6")
    for dv in (pos, place, status):
        ws.add_data_validation(dv)
    end = max(last, 1) + 500
    name_col = "E" if journal == "writeoffs" else "F"
    pos.add(f"{name_col}2:{name_col}{end}")
    if journal == "incoming":
        place.add(f"E2:E{end}")
        status.add(f"P2:P{end}")
    elif journal == "movement":
        place.add(f"D2:E{end}")
        status.add(f"Q2:Q{end}")
    else:
        place.add(f"D2:D{end}")
        status.add(f"J2:J{end}")


def _day_cell(iso):
    """Дата бази («рррр-мм-дд») — датою клітинки; порожньо — порожня клітинка, інше — як записано."""
    if not iso:
        return None
    try:
        return datetime.date.fromisoformat(iso)
    except ValueError:
        return iso


def _act_sheets(wb, model):
    """«Акти», «Підписанти актів», «Дободачі» — після журналів; розділу немає в моделі (None) — немає й аркуша."""
    if model.get("acts") is not None:
        ws = wb.create_sheet("Акти")
        ws.append(ACT_COLS)
        for a in model["acts"]:
            p = {role: a["papers"].get(role) or {} for role in PAPER_CELLS}
            ws.append([_day_cell(a["date"]), a["no"], a["sub"], a["series"], KIND_WORDS.get(a["kind"], a["kind"]),
                       a["point"] or None, _day_cell(a["from"]), _day_cell(a["to"]), a["rations"],
                       p["накладна"].get("no") or None, _day_cell(p["накладна"].get("date")),
                       _day_cell(p["рапорт"].get("date")), p["звірка"].get("no") or None,
                       _day_cell(p["звірка"].get("date")), _day_cell(p["розрахунок"].get("date")), a["note"] or None])
            for c in (1, 7, 8, 11, 12, 14, 15):
                ws.cell(ws.max_row, c).number_format = "dd.mm.yyyy"
        _table(ws, "tblActs", ACT_COLS, len(model["acts"]))
    if model.get("signers") is not None:
        ws = wb.create_sheet("Підписанти актів")
        ws.append(SIGNER_COLS)
        for s in model["signers"]:
            ws.append([_day_cell(s["date"]), s["no"], s["sub"], s["role"], s["sort"], s["pos"] or None,
                       s["rank"] or None, s["name"] or None])
            ws.cell(ws.max_row, 1).number_format = "dd.mm.yyyy"
        _table(ws, "tblActSigners", SIGNER_COLS, len(model["signers"]))
    if model.get("rations") is not None:
        ws = wb.create_sheet("Дободачі")
        names = model["rations"]["points"]
        ws.append(RATION_HEAD + names)
        for d in model["rations"]["days"]:
            ws.append([_day_cell(d["date"]), d["order"] or None, d["src"] or None] + [d["values"].get(p) for p in names])
            ws.cell(ws.max_row, 1).number_format = "dd.mm.yyyy"
        _table(ws, "tblRations", RATION_HEAD + names, len(model["rations"]["days"]))


def _rest_cell(col, r):
    """Залишок позиції рядка r у місці з шапки колонки col на дату D2 — як аркуш «Залишки» старої книги."""
    return (f"=SUMIFS(Прибуток!$K:$K,Прибуток!$G:$G,$A{r},Прибуток!$E:$E,{col}$4,Прибуток!$B:$B,\"<=\"&$D$2)"
            f"+SUMIFS(Переміщення!$J:$J,Переміщення!$G:$G,$A{r},Переміщення!$E:$E,{col}$4,Переміщення!$B:$B,\"<=\"&$D$2)"
            f"-SUMIFS(Переміщення!$J:$J,Переміщення!$G:$G,$A{r},Переміщення!$D:$D,{col}$4,Переміщення!$B:$B,\"<=\"&$D$2)"
            f"-SUMIFS(Списання!$H:$H,Списання!$F:$F,$A{r},Списання!$D:$D,{col}$4,Списання!$B:$B,\"<=\"&$D$2)")


def _rests_sheet(wb, model, items, places, as_of, put):
    """«Залишки»: D2 — дата, шапка в рядку 4 (за нею працюють скрипти служби й ручна картка)."""
    ws = wb.create_sheet("Залишки")
    ws["D1"] = "Залишки на дату:"
    ws["D2"] = datetime.date.fromisoformat(as_of)
    ws["D2"].number_format = "dd.mm.yyyy"
    for c, title in enumerate(["код", "кат", "найм", "разом"] + places, 1):
        ws.cell(4, c, title)
    bal = _balances(model, as_of)
    last = get_column_letter(4 + len(places))
    for r, it in enumerate([i for i in items if not i["archived"]], 5):
        ws.cell(r, 1, _code_cell(it["code"]))
        ws.cell(r, 2, it["cat"])
        ws.cell(r, 3, _pos(it))
        tot = 0.0
        for c, place in enumerate(places, 5):
            col = get_column_letter(c)
            v = bal.get((it["code"], place), 0)
            tot += v
            put(ws, f"{col}{r}", _rest_cell(col, r), _num(v))
        put(ws, f"D{r}", f"=SUM(E{r}:{last}{r})" if places else "=0", _num(tot))


def _moves_sheet(wb, model, items, as_of, place="склад"):
    """«Рух ТМЦ»: місце й період — початковий, надійшло, вибуло, кінцевий (як у старій книзі).
    Значення без формул: ручна картка бере звідси лише підсумки."""
    ws = wb.create_sheet("Рух ТМЦ")
    start = as_of[:4] + "-01-01"
    ws["C3"], ws["D3"], ws["E3"], ws["F3"] = "період", datetime.date.fromisoformat(start), \
        datetime.date.fromisoformat(as_of), place
    ws["D3"].number_format = ws["E3"].number_format = "dd.mm.yyyy"
    for c, title in enumerate(["код", "кат", "найм", "початковий", "надійшло", "вибуло", "кінцевий"], 1):
        ws.cell(4, c, title)
    before = (datetime.date.fromisoformat(start) - datetime.timedelta(days=1)).isoformat()
    a, b = _balances(model, before), _balances(model, as_of)
    got, gone = {}, {}
    for d in model["docs"]:
        if not start <= d["date"] <= as_of:
            continue
        for ln in d["lines"]:
            q = float(ln["qty"] or 0)
            if d["journal"] in ("incoming", "movement") and d["dst"] == place:
                got[ln["code"]] = got.get(ln["code"], 0) + q
            if d["journal"] in ("movement", "writeoffs") and d["src"] == place:
                gone[ln["code"]] = gone.get(ln["code"], 0) + q
    for r, it in enumerate([i for i in items if not i["archived"]], 5):
        c = it["code"]
        for col, v in enumerate([_code_cell(c), it["cat"], it["name"], _num(a.get((c, place), 0)),
                                 _num(got.get(c, 0)), _num(gone.get(c, 0)), _num(b.get((c, place), 0))], 1):
            ws.cell(r, col, v)


_CELL = re.compile(r'<c r="(?P<ref>[A-Z]+\d+)"(?P<attrs>[^>]*)><f>(?P<f>.*?)</f><v></v></c>', re.S)


def _put_cached(path, cached):
    """openpyxl пише формули без значень (`<f>…</f><v></v>`): дописуємо значення, які порахувала
    програма, — книга, ще не відкрита в Excel, дає скриптам числа й текст, а не порожнечу."""
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheets = _sheet_paths(files)
    for title, values in cached.items():
        part = sheets[title]
        xml = files[part].decode("utf-8")

        def fill(m, values=values):
            v = values.get(m.group("ref"))
            if v is None:
                return m.group(0)
            if isinstance(v, str):
                return f'<c r="{m.group("ref")}"{m.group("attrs")} t="str"><f>{m.group("f")}</f><v>{escape(v)}</v></c>'
            return f'<c r="{m.group("ref")}"{m.group("attrs")}><f>{m.group("f")}</f><v>{v}</v></c>'

        files[part] = _CELL.sub(fill, xml).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)


def _sheet_paths(files):
    """Назва аркуша → шлях його XML у книзі (атрибути тегу <sheet> — у будь-якому порядку)."""
    wbx = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    target = {}
    for tag in re.findall(r"<Relationship\b[^>]*>", rels):
        rid, tgt = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
        if rid and tgt:
            target[rid.group(1)] = tgt.group(1)
    out = {}
    for tag in re.findall(r"<sheet\b[^>]*>", wbx):
        name, rid = re.search(r'\bname="([^"]+)"', tag), re.search(r'\br:id="([^"]+)"', tag)
        if name and rid:
            t = target[rid.group(1)].lstrip("/")
            out[_unescape(name.group(1))] = t if t.startswith("xl/") else "xl/" + t
    return out


def _unescape(s):
    return s.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'").replace("&amp;", "&")


def export_book(con, path, as_of=None):
    """Книга з бази — цілою або ніяк: пишеться поруч під одним і тим самим тимчасовим ім'ям («Облік ОП.xlsx.part» —
    той, що лишився після обриву, наступний запис перезаписує) і стає на місце одним кроком. Збій посеред запису лишає
    попередню книгу як була; книга, відкрита в Excel, заміні не піддається — PermissionError, як і раніше.
    Готова книга запам'ятовується своєю (`remember_sha`): відбиток байтів і вмісту (вона ж, прочитана назад); не
    читається назад — лише байтів, а у відповіді `readBack` — що не так (до 50 рядків), як у книги «Закупівлі»: книга
    на місці, але «Прийняти книгу…» її не прийме, поки ці рядки не виправлено."""
    m = model_from_db(con)
    path = os.fspath(path)
    part = path + ".part"
    problems = None
    try:
        write_book(m, part, as_of)
        sha = stable_hash.stable_sha(part)
        with open(part, "rb") as f:
            data = f.read()
        try:
            mark = fingerprint(read_book(data))
        except BookError as e:                    # через дані самої бази назад не читається: впізнаємо лише за байтами
            mark, problems = None, e.problems[:50]
        except Exception:                         # noqa: BLE001 — несподіване при читанні назад: книга однаково стає
            traceback.print_exc()                 # людині — речення, назва винятку й текст помилки — лише в консоль
            mark, problems = None, ["Книгу назад не прочитано."]
        os.replace(part, path)
    except BaseException:
        with contextlib.suppress(OSError):
            os.remove(part)
        raise
    with contextlib.suppress(sqlite3.Error):      # не запам'ятали — наступного разу книгу звірять із базою
        remember_sha(con, sha, mark)
    out = {"path": path, "items": len(m["items"]), "docs": len(m["docs"]),
           "lines": sum(len(d["lines"]) for d in m["docs"])}
    if problems:
        out["readBack"] = problems
    return out


def _strip(model):
    """Модель книги без службового: номерів рядків аркушів (`row` — лише в записах актів і підписантів) і звіту
    перенесення. Глибше `row` — дані: пункт харчування може зватися «row», і його дні входять у відбиток."""
    out = {k: v for k, v in model.items() if k != "report"}
    for part in ("acts", "signers"):
        if isinstance(out.get(part), list):
            out[part] = [{k: v for k, v in r.items() if k != "row"} if isinstance(r, dict) else r for r in out[part]]
    return out


def fingerprint(model):
    """Відбиток вмісту книги: sha256 канонічного JSON її моделі (без номерів рядків, ключі відсортовані). Той самий вміст —
    той самий відбиток, хоч би байти файла були інші: книгу відкрили в Excel і зберегли."""
    text = json.dumps(_strip(model), sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _own(con):
    """Що програма пам'ятає про свої книги ОП: `sha` — відбитки байтів (швидкий шлях), `print` — відбитки вмісту."""
    row = con.execute("SELECT value FROM app_setting WHERE key = ?", (SHA_SETTING,)).fetchone()
    try:
        got = json.loads(row[0]) if row and row[0] else {}
    except ValueError:
        got = {}
    got = got if isinstance(got, dict) else {}
    return {key: [x for x in got[key] if isinstance(x, str) and x] if isinstance(got.get(key), list) else []
            for key in ("sha", "print")}


def own_shas(con):
    """Відбитки байтів книг, які програма сама записала чи прийняла (новіші спереду; не більше KEEP_SHAS)."""
    return _own(con)["sha"]


def own_prints(con):
    """Відбитки вмісту таких книг (новіші спереду; не більше KEEP_SHAS)."""
    return _own(con)["print"]


def remember_sha(con, sha=None, fingerprint=None):
    """Запам'ятати книгу, яку програма записала чи прийняла: відбиток байтів і/або вмісту, не більше KEEP_SHAS кожного."""
    own = _own(con)
    for key, value in (("sha", sha), ("print", fingerprint)):
        if value:
            own[key] = [value] + [x for x in own[key] if x != value]
        own[key] = own[key][:KEEP_SHAS]
    con.execute("INSERT INTO app_setting(key, value) VALUES(?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (SHA_SETTING, json.dumps(own)))


def own_book(con, path):
    """(своя, модель): чи файл на диску — книга, яку програма сама записала чи прийняла, і книга, прочитана для відповіді
    (None, коли читати не довелось). Спершу відбиток байтів (збіг за будь-яким із двох — stable_hash.shas); далі вміст:
    книгу відкрили в Excel і зберегли — байти інші, а дані ті самі. Незчитна книга — BookError (вона не своя)."""
    if set(stable_hash.shas(path)) & set(own_shas(con)):
        return True, None
    with open(path, "rb") as f:
        model = read_book(f.read())
    return fingerprint(model) in set(own_prints(con)), model


def is_own(con, path):
    try:
        return own_book(con, path)[0]
    except BookError:
        return False


# ---------------------------------------------------------------- читання й прийом

class BookError(ValueError):
    """Книга не приймається: перелік причин з аркушем і рядком («Прибуток, рядок 12: …»)."""

    def __init__(self, problems):
        super().__init__("; ".join(problems[:5]))
        self.problems = list(problems)


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v or "").strip()
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    return ""


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _code_of(name, fallback=None):
    """Код — з початку тексту «найменування» (як INT(LEFT(…,5)) у книзі), інакше з колонки «код»."""
    m = re.match(r"\s*(\d+)", str(name or ""))
    if m:
        return m.group(1)
    return _text(fallback) if fallback not in (None, "") else ""


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [_text(h).lower() for h in next(it, [])]
    for n, row in enumerate(it, 2):
        yield n, {h: row[i] if i < len(row) else None for i, h in enumerate(head) if h}


def _number(v):
    """Число клітинки; порожньо, не число чи не скінченне («nan», «inf») — None: далі воно — відмова з рядком, а не 500."""
    try:
        x = float(str(v).replace(",", ".")) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None
    return x if x is None or math.isfinite(x) else None


def _price(v):
    """Ціна позиції: облік веде ціни в копійках — частки копійки книга не несе."""
    x = _number(v)
    return None if x is None else x


HEAD_NAMES = {"type": "док_тип", "status": "статус", "report": "рапорт", "order": "наказ", "scan": "скан",
              "col": "графа 2/прод"}


def _comma(v, places=None):
    """Число з комою, як його пише людина: 0,07474; з places — рівно стільки знаків."""
    text = f"{v:.{places}f}" if places is not None else ("%.10g" % v)
    return text.replace(".", ",")


def _real_day(iso):
    """Дата за календарем із роком у межах обліку — те саме правило, що й для дня дободач (state_db._ration_day)."""
    import state_db                                              # noqa: PLC0415 — модуль сервера
    try:
        state_db._ration_day(iso)
    except state_db.SaveError:
        return False
    return True


def _book_day(v, where, what, problems):
    """Дата з клітинки нових аркушів: «рррр-мм-дд»; порожньо — ""; не дата чи рік з опискою — у перелік відмов."""
    if not _text(v):
        return ""
    iso = _date(v)
    if not iso or not _real_day(iso):
        problems.append(f"{where}: {what} «{_text(v)}» — не дата або рік з опискою")
        return ""
    return iso


def _norm_cell(v, where, problems):
    """«норма_дд» рядка «Списання»: число, більше за нуль; порожньо — рядок не акта (None)."""
    if not _text(v):
        return None
    x = _number(v)
    if x is None or not math.isfinite(x) or x <= 0:
        problems.append(f"{where}: «норма_дд» — число, більше за нуль")
        return None
    return x


def _whole_cell(v, where, what, top, problems):
    """Ціле від 0 до `top` з клітинки (порожньо — 0); інше — у перелік відмов."""
    import state_db                                              # noqa: PLC0415
    if not _text(v):
        return 0
    n = state_db._whole_number(v)
    if n is None or not 0 <= n <= top:
        problems.append(f"{where}: {what} — ціле число від 0 до {format(top, ',').replace(',', ' ')}")
        return 0
    return n


def _read_acts(ws, problems):
    """Аркуш «Акти»: рядок на акт. Вид — словами книги (KIND_WORDS) чи кодом; порожньо — «за нормами 390».
    Чи є документ акта в «Списання» і пункт у «Довідка» — перевіряє прийом (`_apply_acts`)."""
    import state_db                                              # noqa: PLC0415
    out, seen = [], {}
    for n, r in _rows(ws):
        if not any(_text(v) for v in r.values()):
            continue
        where = f"Акти, рядок {n}"
        word = _text(r.get("вид"))
        a = {"date": _book_day(r.get("док_дата"), where, "док_дата", problems), "no": _text(r.get("док_номер")),
             "sub": _text(r.get("підрозділ")), "series": _text(r.get("серія")),
             "kind": next((k for k, w in KIND_WORDS.items() if word in (k, w)), None) if word else "op390",
             "point": _text(r.get("пункт")), "from": _book_day(r.get("період_з"), where, "період_з", problems),
             "to": _book_day(r.get("період_по"), where, "період_по", problems), "note": _text(r.get("примітка")),
             "papers": {}, "row": where}
        if not (a["date"] and a["no"] and a["sub"] and a["from"] and a["to"]):
            problems.append(f"{where}: потрібні дата, номер, підрозділ і період")
        elif a["from"] > a["to"]:
            problems.append(f"{where}: початок періоду пізніше за кінець")
        if a["series"] not in ("ОП", "МЗ"):
            problems.append(f"{where}: серія — ОП чи МЗ")
        if a["kind"] is None:
            problems.append(f"{where}: вид акта — «{KIND_WORDS['op390']}»")
        days = 1
        if a["from"] and a["to"] and a["from"] <= a["to"]:
            days += (datetime.date.fromisoformat(a["to"]) - datetime.date.fromisoformat(a["from"])).days
        a["rations"] = _whole_cell(r.get("дд"), where, "«дд»", state_db.RATIONS_MAX * days, problems)
        for role, (no, day) in PAPER_CELLS.items():
            rec = {"no": _text(r.get(no)) if no else "", "date": _book_day(r.get(day), where, day, problems)}
            if rec["no"] or rec["date"]:
                a["papers"][role] = rec
        key = (a["date"], a["no"], a["sub"])
        if key in seen:
            problems.append(f"{where}: акт №{a['no']} від {a['date']} ({a['sub']}) уже є в рядку {seen[key]}")
        seen.setdefault(key, n)
        out.append(a)
    return out


def _read_signers(ws, problems):
    """Аркуш «Підписанти актів»: знімок підписантів. Чи є такий акт — перевіряє прийом (`_apply_acts`)."""
    out, seen = [], {}
    for n, r in _rows(ws):
        if not any(_text(v) for v in r.values()):
            continue
        where = f"Підписанти актів, рядок {n}"
        s = {"date": _book_day(r.get("док_дата"), where, "док_дата", problems), "no": _text(r.get("док_номер")),
             "sub": _text(r.get("підрозділ")), "role": _text(r.get("роль")),
             "sort": _whole_cell(r.get("порядок"), where, "«порядок»", 99, problems), "pos": _text(r.get("посада")),
             "rank": _text(r.get("звання")), "name": _text(r.get("ім'я_прізвище")), "row": where}
        if not (s["date"] and s["no"] and s["sub"]):
            problems.append(f"{where}: потрібні дата, номер і підрозділ акта")
        if not s["role"]:
            problems.append(f"{where}: потрібна роль")
        key = (s["date"], s["no"], s["sub"], s["role"], s["sort"])
        if key in seen:
            problems.append(f"{where}: «{s['role']}» з порядком {s['sort']} акта №{s['no']} уже є в рядку {seen[key]}")
        seen.setdefault(key, n)
        out.append(s)
    return out


def _read_rations(ws, problems):
    """Аркуш «Дободачі»: дата, витяг, джерело й колонка на пункт. Правила ті самі, що й у записі вікна
    (state_db._save_rations): дата за календарем у межах обліку, д/д — ціле від 0 до RATIONS_MAX."""
    import state_db                                              # noqa: PLC0415
    it = ws.iter_rows(values_only=True)
    head = list(next(it, None) or [])
    cols, names = [], {}
    for i in range(len(RATION_HEAD), len(head)):
        name = _text(head[i])
        if not name:
            continue
        if name.lower() in names:
            problems.append(f"Дободачі: пункт «{name}» у шапці двічі")
        else:
            names[name.lower()] = i
            cols.append((i, name))
    titled = {i for i in range(len(RATION_HEAD), len(head)) if _text(head[i])}
    out, seen = {"points": [name for _, name in cols], "days": []}, {}
    for n, row in enumerate(it, 2):
        row = list(row or [])
        if not any(_text(v) for v in row):
            continue
        where = f"Дободачі, рядок {n}"
        day = _date(row[0])
        if not _text(row[0]):
            problems.append(f"{where}: немає дати")
            continue
        if not day or not _real_day(day):
            problems.append(f"{where}: «{_text(row[0])}» — не дата або рік з опискою")
            continue
        if day in seen:
            problems.append(f"{where}: дата {day} уже є в рядку {seen[day]}")
            continue
        seen[day] = n
        values = {}
        for i, name in cols:
            if i < len(row) and _text(row[i]):
                values[name] = _whole_cell(row[i], f"{where}, «{name}»", "д/д", state_db.RATIONS_MAX, problems)
        if any(_text(v) for i, v in enumerate(row) if i >= len(RATION_HEAD) and i not in titled):
            problems.append(f"{where}: число в колонці без назви пункту")
        out["days"].append({"date": day, "order": _text(row[1]) if len(row) > 1 else "",
                            "src": _text(row[2]) if len(row) > 2 else "", "values": values})
    return out


def _points_table(ws, problems):
    """Таблиця tblPoints аркуша «Довідка»: шапка «пункт» у рядку 1; немає — None (книга старої версії).
    Поруч на аркуші інші таблиці різної довжини: рядок, у якому колонки пунктів порожні, пропускається."""
    rows = ws.iter_rows(values_only=True)
    head = [_text(h).lower() for h in next(rows, None) or []]
    if "пункт" not in head:
        return None
    col = {k: head.index(k) if k in head else None for k in POINT_COLS}
    out, seen = [], {}
    for n, row in enumerate(rows, 2):
        def get(k, row=row):
            return row[col[k]] if col[k] is not None and col[k] < len(row) else None
        if not any(_text(get(k)) for k in POINT_COLS):
            continue
        where = f"Довідка, рядок {n}"
        name, sub = _text(get("пункт")), _text(get("підрозділ"))
        if not name:
            problems.append(f"{where}: пункт харчування без назви")
            continue
        if name.lower() in seen:
            problems.append(f"{where}: пункт «{name}» уже є в рядку {seen[name.lower()]}")
            continue
        seen[name.lower()] = n
        if _fold(name) in RATION_HEAD:               # шапка «Дободачі» мала б дві колонки з однією назвою
            problems.append(f"{where}: пункт не може зватися «{name}»: так зветься колонка аркуша «Дободачі»")
        if not sub:
            problems.append(f"{where}: пункт «{name}» без підрозділу")
        out.append({"name": name, "sub": sub, "hall": _text(get("їдальня")), "title": _text(get("назва_у_звірці")),
                    "closed": _book_day(get("закрито"), where, "закрито", problems)})
    return out


def _split_checked(text):
    """«перевірено» книги → (хто, коли): «Влад 2026-10-08» → («Влад», «2026-10-08»); «уточнити» — ніхто."""
    t = _text(text)
    if not t or t.lower() == UNSURE:
        return None, None
    m = re.match(r"^(.*?)\s*(\d{4}-\d{2}-\d{2})$", t)
    if m and m.group(1):
        return m.group(1).strip(), m.group(2)
    return t, None


def read_book(data, legacy=False):
    """Модель з книги «Облік ОП» (значення клітинок, не формули). Помилки — з аркушем і рядком.

    `legacy` — перше перенесення старої книги: розбіжні поля документа зводяться до першого
    непорожнього, статус з «примітки» переходить у «статус», «транз_найменування» — у
    «примітку», ціна з частками копійки округлюється; кожна така правка — у model["report"]."""
    import io                                                    # noqa: PLC0415
    from openpyxl import load_workbook                           # noqa: PLC0415
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:                                       # noqa: BLE001 — будь-який зіпсований файл
        raise BookError([f"файл не читається як книга Excel: {e}"]) from e
    problems, report = [], []
    missing = [s for s in ("Номенклатура", "Прибуток", "Переміщення", "Списання") if s not in wb.sheetnames]
    if missing:
        raise BookError([f"немає аркуша «{s}»" for s in missing])
    renames, parties, subs = [], [], []
    if "Довідка" in wb.sheetnames:
        renames, parties = _reference(wb["Довідка"])
    alias = {old: new for old, new in renames if old and new} if legacy else {}
    items = []
    # Книга 1.12.0 і старіша колонки «рядок норми» не має: рядків норми позицій вона не несе й прийом їх не чіпає.
    line_col = LINE_COL in [_text(h).lower() for h in next(wb["Номенклатура"].iter_rows(max_row=1, values_only=True), ())]
    for n, r in _rows(wb["Номенклатура"]):
        code = _code_of(r.get("код"))
        if not code:
            continue
        where = f"Номенклатура, рядок {n}"
        kind = _text(r.get("тип")).lower()
        cat = _text(r.get("категорія")).upper()
        archived = kind == ARCHIVE
        grp = GROUP_OF_TYPE.get(kind) or GROUP_OF_CAT.get(cat)
        if grp is None:
            problems.append(f"{where}: невідомий тип «{_text(r.get('тип'))}» і категорія «{_text(r.get('категорія'))}»")
            continue
        price = _price(r.get("ціна"))
        if price is None:
            if _text(r.get("ціна")):                            # «nan», «inf», слово — не ціна, а не нуль мовчки
                problems.append(f"{where}: ціна «{_text(r.get('ціна'))}» — не число")
            price = 0.0
        line = None
        if line_col:
            try:
                line = op_norms.line_key(r.get(LINE_COL))
            except ValueError:
                problems.append(f"{where}: «{LINE_COL}» «{_text(r.get(LINE_COL))}» — такого рядка норми наказу 390 немає")
            if line and op_norms.LINES[line]["group"] != grp:       # рядок норми — рядок групи позиції (група — з цього ж рядка книги)
                problems.append(f"{where}: Позиція {code}: рядок норми «{op_norms.LINES[line]['name']}» "
                                f"не з групи «{GROUP_NAME[grp]}».")
        if abs(price * 100 - round(price * 100)) > 1e-6:
            if legacy:
                report.append(f"{where}: ціна {_comma(price)} — до копійок {_comma(round(price, 2), 2)} "
                              "(облік веде ціни в копійках)")
                price = round(price, 2)
            else:
                problems.append(f"{where}: ціна {_comma(price)} з частками копійки — облік веде ціни в копійках")
        row2 = r.get("рядок 2/прод")
        skip = None
        if isinstance(row2, str) and not row2.strip().isdigit():
            skip, row2 = row2.strip() or None, None
        per = _number(r.get("дободач"))
        items.append({"code": code, "name": _text(r.get("найменування")), "cat": CAT_OF_GROUP[grp],
                      "uom": _text(r.get("од.вим.")) or "шт", "price": price,
                      "type": TYPE_OF_GROUP[grp], "archived": archived, "per_ration": per,
                      "note": _text(r.get("примітка")),
                      "row2": int(row2) if row2 not in (None, "") else None,
                      "factor": _number(r.get("множник")) or 1.0,
                      "checked": None if _text(r.get("перевірено")).lower() == UNSURE else (_text(r.get("перевірено")) or None),
                      "skip2": skip, "line": line})
    codes = {i["code"] for i in items}
    docs, at = [], {}
    # Книга 1.11.0 колонки «норма_дд» не має: тоді норм рядків вона не несе й прийом їх не чіпає.
    norm_col = "норма_дд" in [_text(h).lower() for h in next(wb["Списання"].iter_rows(max_row=1, values_only=True), ())]
    norm_at = {}
    for journal, (title, _table_name, _cols) in SHEETS.items():
        for n, r in _rows(wb[title]):
            name = r.get("найменування")
            if not _text(name) and not r.get("док_дата") and not _text(r.get("док_номер")):
                continue
            where = f"{title}, рядок {n}"
            code = _code_of(name, r.get("код"))
            date = _date(r.get("док_дата"))
            no = _text(r.get("док_номер"))
            qty = _number(r.get("кількість")) or 0
            if not date:
                problems.append(f"{where}: немає дати")
            if not no:
                problems.append(f"{where}: немає номера документа")
            if qty <= 0:
                problems.append(f"{where}: кількість має бути більшою за нуль")
            elif not legacy and abs(qty * 1000 - round(qty * 1000)) > 1e-6:
                problems.append(f"{where}: кількість {_comma(qty)} — облік веде кількість до тисячних")
            if code not in codes:
                problems.append(f"{where}: позиції «{_text(name)}» немає в «Номенклатура»")
            status = _text(r.get("статус"))
            note = _text(r.get("примітка"))
            if legacy and not status and note in STATUSES:
                report.append(f"{where}: статус «{note}» стояв у «примітка» — перенесено в «статус»")
                status, note = note, ""
            if status and status not in STATUSES:
                problems.append(f"{where}: невідомий статус «{status}»")
            alt = _text(r.get("транз_найменування"))
            if alt:
                if legacy:
                    report.append(f"{where}: «транз_найменування» «{alt}» перенесено в «примітка»")
                    note = (note + " " + alt).strip()
                else:
                    problems.append(f"{where}: колонку «транз_найменування» програма не веде — перенесіть текст у «примітка»")
            if journal == "incoming":
                src, dst = _text(r.get("від_кого")), _text(r.get("кому")) or "склад"
            elif journal == "movement":
                src, dst = _text(r.get("від_кого")), _text(r.get("кому"))
            else:
                src, dst = _text(r.get("підрозділ")), _text(r.get("кому"))
            if journal != "incoming" and src in alias:
                src = alias[src]
            if journal != "writeoffs" and dst in alias:
                dst = alias[dst]
            col = r.get("графа 2/прод")
            col = int(col) if str(col or "").strip().isdigit() else None
            head = {"type": _text(r.get("док_тип")), "status": status, "report": _text(r.get("рапорт")),
                    "order": _text(r.get("наказ")), "scan": _text(r.get("скан")), "col": col}
            key = (journal, date, no, src, dst)
            d = at.get(key)
            if d is None:
                d = at[key] = {"journal": journal, "date": date, "no": no, "src": src, "dst": dst, **head, "lines": []}
                docs.append(d)
            elif any(d[k] != v for k, v in head.items()):
                differ = [HEAD_NAMES[k] for k, v in head.items() if d[k] != v]
                if legacy:
                    for k, v in head.items():
                        if not d[k] and v:
                            d[k] = v
                    report.append(f"{where}: у рядках документа №{no} різні {', '.join(differ)} — узято перше непорожнє")
                else:
                    problems.append(f"{where}: у рядках одного документа різні {', '.join(differ)}")
            line = {"code": code, "qty": _num(qty), "note": note}
            if journal == "writeoffs" and norm_col:               # норма акта — одна на код документа
                line["norm"] = _norm_cell(r.get("норма_дд"), where, problems)
                first = norm_at.setdefault((key, code), (line["norm"], n))
                if first[0] != line["norm"]:
                    problems.append(f"{where}: «норма_дд» коду {code} не та, що в рядку {first[1]} цього документа")
            d["lines"].append(line)
    if "Підрозділи" in wb.sheetnames:
        for n, r in _rows(wb["Підрозділи"]):
            name = _text(r.get("назва"))
            if name:
                subs.append({"name": name, "parent": _text(r.get("підпорядкований")),
                             "kind": _text(r.get("вид")) or "склад", "closed": bool(_text(r.get("закритий"))),
                             "aliases": [a.strip() for a in _text(r.get("псевдоніми")).split(",") if a.strip()]})
    # Аркуші актів — не обов'язкові: книга 1.11.0 їх не має, і тоді ці дані в базі не змінюються (None).
    acts = _read_acts(wb["Акти"], problems) if "Акти" in wb.sheetnames else None
    signers = _read_signers(wb["Підписанти актів"], problems) if "Підписанти актів" in wb.sheetnames else None
    rations = _read_rations(wb["Дободачі"], problems) if "Дободачі" in wb.sheetnames else None
    points = _points_table(wb["Довідка"], problems) if "Довідка" in wb.sheetnames else None
    if problems:
        raise BookError(problems)
    model = {"items": items, "subs": subs, "parties": parties, "renames": renames, "docs": docs, "acts": acts,
             "signers": signers, "rations": rations, "points": points, "normCol": norm_col, "lineCol": line_col}
    if legacy:
        for old, new in sorted(alias.items()):
            report.append(f"місце «{old}» → «{new}» за tblRename")
        model["report"] = report
    return model


def _reference(ws):
    """tblRename (стара назва → нова назва) і tblParties (контрагент → вид) з аркуша «Довідка»."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    head = grid[0] if grid else []
    pos = {_text(h).lower(): i for i, h in enumerate(head) if h}
    renames, parties = [], []

    def cell(row, key):
        i = pos.get(key)
        return _text(row[i]) if i is not None and i < len(row) else ""

    for row in grid[1:]:
        if cell(row, "стара назва"):
            renames.append([cell(row, "стара назва"), cell(row, "нова назва")])
        if cell(row, "контрагент"):
            parties.append({"name": cell(row, "контрагент"), "kind": cell(row, "вид") or "інше"})
    return renames, parties


def _doc_label(d):
    return f"{SHEETS[d['journal']][0]} · {d['date']} · №{d['no']} · {d['src'] or '—'} → {d['dst'] or '—'}"


def _doc_sig(d):
    """Підпис документа для порівняння; норма рядка акта — теж (немає норми — 0: норма буває лише більшою за нуль)."""
    return (d["type"], d["status"], d["report"], d["order"], d["scan"], d["col"],
            tuple(sorted((ln["code"], float(ln["qty"]), ln["note"], float(ln.get("norm") or 0)) for ln in d["lines"])))


def _doc_key(d):
    return (d["journal"], d["date"], d["no"], d["src"], d["dst"])


def _names(con):
    """Назва підрозділу за назвою чи псевдонімом."""
    out = {r[0]: r[0] for r in con.execute("SELECT name FROM subdivision")}
    for alias, name in con.execute("SELECT a.name, s.name FROM subdivision_alias a "
                                   "JOIN subdivision s ON s.id = a.subdivision_id"):
        out.setdefault(alias, name)
    return out


def _resolved(con, model):
    """Модель з назвами місць, як їх знає база (псевдонім → назва підрозділу): у документах, актах, підписантах і пунктах."""
    names = _names(con)

    def fix(d):
        d = dict(d)
        if d["journal"] != "incoming":
            d["src"] = names.get(d["src"], d["src"])
        if d["journal"] != "writeoffs":
            d["dst"] = names.get(d["dst"], d["dst"])
        return d
    out = dict(model, docs=[fix(d) for d in model["docs"]])
    for part in ("acts", "signers", "points"):
        if model.get(part) is not None:
            out[part] = [dict(x, sub=names.get(x["sub"], x["sub"])) for x in model[part]]
    return out


def _fold(name):
    """Назва пункту для порівняння — як на сторінці (pointKey): без зайвих пробілів і без різниці регістру."""
    return " ".join(str(name or "").split()).lower()


def _act_key(a):
    return (a["date"], a["no"], a["sub"])


def _act_label(key):
    return f"Акти · {key[0]} · №{key[1]} · {key[2]}"


def _signers_by_act(m):
    out = {}
    for s in m.get("signers") or []:
        out.setdefault(_act_key(s), []).append((s["role"], s["sort"], s["pos"], s["rank"], s["name"]))
    return {k: sorted(v) for k, v in out.items()}


def _act_sig(a, signers=None):
    """Підпис акта для порівняння: дані акта (пункт — як його знаходить прийом, `_fold`), папери й (коли книга несе аркуш
    підписантів) підписанти."""
    return (a["series"], a["kind"], _fold(a["point"]), a["from"], a["to"], a["rations"],
            sorted((k, v.get("no", ""), v.get("date", "")) for k, v in a["papers"].items()), a["note"],
            None if signers is None else signers.get(_act_key(a), []))


def _ration_records(r):
    """Дободачі як записи: (дата, пункт) → д/д і (дата, "") → (витяг, джерело); пункт — як його знаходить прийом (`_fold`)."""
    out = {}
    for d in r["days"]:
        for p, v in d["values"].items():
            out[(d["date"], _fold(p))] = v
        if d["order"] or d["src"]:
            out[(d["date"], "")] = (d["order"], d["src"])
    return out


def diff(con, model):
    """Що зміниться в базі, коли прийняти книгу: документи, позиції, нові місця й контрагенти, акти, дободачі й пункти
    харчування. Аркуша нових розділів у книзі немає (None) — ці дані не порівнюються: прийом їх не змінює."""
    have = model_from_db(con)
    model = _resolved(con, model)
    if not model.get("normCol", True):           # книга без «норма_дд» норм рядків не змінює
        have = dict(have, docs=[dict(d, lines=[dict(ln, norm=None) for ln in d["lines"]]) for d in have["docs"]])
    if not model.get("lineCol", True):           # книга без «рядок норми» рядків норми позицій не змінює
        have = dict(have, items=[dict(i, line=None) for i in have["items"]])
    mine, theirs = {_doc_key(d): d for d in model["docs"]}, {_doc_key(d): d for d in have["docs"]}
    out = {"docs": {"new": [_doc_label(d) for k, d in mine.items() if k not in theirs],
                    "changed": [_doc_label(d) for k, d in mine.items() if k in theirs and _doc_sig(d) != _doc_sig(theirs[k])],
                    "removed": [_doc_label(d) for k, d in theirs.items() if k not in mine]}}
    hi = {i["code"]: i for i in have["items"]}
    out["items"] = {"new": [i["code"] for i in model["items"] if i["code"] not in hi],
                    "changed": [i["code"] for i in model["items"] if i["code"] in hi and i != hi[i["code"]]],
                    "removed": [c for c in hi if c not in {i["code"] for i in model["items"]}]}
    names = _names(con)
    used = {d["dst"] for d in model["docs"] if d["journal"] != "writeoffs"} | {
        d["src"] for d in model["docs"] if d["journal"] != "incoming"}
    out["subs"] = sorted(n for n in used if n and n not in names)
    hp = {p["name"]: p["kind"] for p in have["parties"]}
    out["parties"] = sorted(p["name"] for p in model["parties"] if hp.get(p["name"]) != p["kind"])
    out.update(_diff_acts(have, model))
    out["same"] = not any(out["docs"][k] for k in out["docs"]) and not any(
        out["items"][k] for k in out["items"]) and not out["subs"] and not out["parties"] and not any(
        out["acts"][k] for k in out["acts"]) and not out["rations"]["changed"] and not any(
        out["points"][k] for k in out["points"])
    return out


def _diff_acts(have, model):
    """Акти (за ключем документа), дободачі (кількість різних записів і самі записи — «Прийняти» звіряє показане з тим,
    що є тепер; `fromBook` — записів, яких у базі немає такими: їх сторожить вивантаження) і пункти харчування (за назвою
    без різниці регістру)."""
    out = {"acts": {"new": [], "changed": [], "removed": []}, "rations": {"changed": 0, "fromBook": 0, "keys": []},
           "points": {"new": [], "changed": [], "removed": []}}
    hs, ms = _signers_by_act(have), (_signers_by_act(model) if model.get("signers") is not None else None)
    if model.get("acts") is not None:
        old, new = {_act_key(a): a for a in have["acts"]}, {_act_key(a): a for a in model["acts"]}
        out["acts"] = {"new": [_act_label(k) for k in sorted(new.keys() - old.keys())],
                       "changed": [_act_label(k) for k in sorted(new.keys() & old.keys())
                                   if _act_sig(new[k], ms) != _act_sig(old[k], None if ms is None else hs)],
                       "removed": [_act_label(k) for k in sorted(old.keys() - new.keys())]}
    elif ms is not None:                         # аркуша «Акти» немає, підписанти є: змінені підписанти — змінений акт
        out["acts"]["changed"] = [_act_label(k) for k in sorted(_act_key(a) for a in have["acts"])
                                  if ms.get(k, []) != hs.get(k, [])]
    if model.get("rations") is not None:
        mine, base = _ration_records(model["rations"]), _ration_records(have["rations"])
        keys = sorted(k for k in mine.keys() | base.keys() if mine.get(k) != base.get(k))
        out["rations"] = {"changed": len(keys), "fromBook": sum(1 for k, v in mine.items() if base.get(k) != v),
                          "keys": [f"{day} · {point or 'витяг і джерело'}" for day, point in keys]}
    if model.get("points") is not None:
        def sig(p):
            return (p["name"], p["sub"], p["hall"], p["title"], p["closed"])
        old, new = {_fold(p["name"]): p for p in have["points"]}, {_fold(p["name"]): p for p in model["points"]}
        out["points"] = {"new": [new[k]["name"] for k in sorted(new.keys() - old.keys())],
                         "changed": [new[k]["name"] for k in sorted(new.keys() & old.keys()) if sig(new[k]) != sig(old[k])],
                         "removed": [old[k]["name"] for k in sorted(old.keys() - new.keys())]}
    return out


def apply(con, model):
    """Книга ОП у базі стає точно такою, як у файлі. Облік ТЗ не змінюється; місця й контрагенти
    лише додаються. Усе або нічого: відмова лишає базу як була."""
    import state_db                                              # noqa: PLC0415
    plan = diff(con, model)
    con.execute("BEGIN IMMEDIATE")
    try:
        _ensure_subs(con, model)
        model = _resolved(con, model)
        _ensure_items(con, model)
        state_db._save_parties(con, {p["name"]: p["kind"] for p in model["parties"]})
        held = _held_acts(con)
        before = state_db.negative_days(con)
        ours = {i["code"] for i in model["items"]} | state_db.op_codes(con)
        docs = {j: [r for r in rows if str(r[5]) not in ours] for j, rows in state_db._load_docs(con).items()}
        price = {i["code"]: i["price"] for i in model["items"]}
        for d in model["docs"]:
            for ln in d["lines"]:
                docs[d["journal"]].append([d["date"], d["type"], d["no"], d["src"], d["dst"], ln["code"], ln["qty"],
                                           ln["note"], price[ln["code"]] if d["journal"] == "incoming" else 0])
        try:
            state_db._save_docs(con, docs, seed_drop_limit=False)
        except state_db.SaveError as e:
            raise BookError([str(e)]) from e
        new = {k: v for k, v in state_db.negative_days(con).items() if k not in before}
        if new:
            raise BookError([state_db._negative_words(con, new)])
        _apply_doc_fields(con, model)
        _apply_acts(con, model, held)
        edits = {}
        for i in model["items"]:
            by, on = _split_checked(i["checked"])
            edits[i["code"]] = {"row": i["row2"], "factor": i["factor"], "checked": by, "checkedOn": on,
                                "skip": i["skip2"]}
        state_db._save_form2_map(con, edits)
        _drop_unused_items(con, {i["code"] for i in model["items"]})
        con.execute("COMMIT")
    except BaseException:
        con.execute("ROLLBACK")
        raise
    return plan


def _ensure_subs(con, model):
    """Місця з аркуша «Підрозділи» (у чистій програмі — усе дерево) і назви з документів."""
    kinds = dict(con.execute("SELECT code, id FROM subdivision_kind"))
    have = dict(con.execute("SELECT name, id FROM subdivision"))
    aliases = {r[0] for r in con.execute("SELECT name FROM subdivision_alias")}
    sort = con.execute("SELECT COALESCE(MAX(sort), 0) FROM subdivision").fetchone()[0] or 0
    for s in model["subs"]:
        if s["name"] in have:
            continue
        if s["kind"] not in kinds:
            kinds[s["kind"]] = con.execute("INSERT INTO subdivision_kind(code, name) VALUES(?, ?)",
                                           (s["kind"], s["kind"])).lastrowid
        sort += 1
        have[s["name"]] = con.execute(
            "INSERT INTO subdivision(name, kind_id, parent_id, sort, is_active) VALUES(?, ?, ?, ?, ?)",
            (s["name"], kinds[s["kind"]], have.get(s["parent"]), sort, 0 if s["closed"] else 1)).lastrowid
        for a in s["aliases"]:
            if a not in aliases and a not in have:
                con.execute("INSERT INTO subdivision_alias(subdivision_id, name, source) VALUES(?, ?, 'книга ОП')",
                            (have[s["name"]], a))
                aliases.add(a)
    for old, new in model["renames"]:
        if new in have and old not in have and old not in aliases:
            con.execute("INSERT INTO subdivision_alias(subdivision_id, name, source) VALUES(?, ?, 'книга ОП')",
                        (have[new], old))
            aliases.add(old)
    # Назва з документа, якої немає ні в довіднику, ні серед псевдонімів, — помилка, а не новий підрозділ.
    unknown = set()
    for d in model["docs"]:
        names = (d["dst"],) if d["journal"] == "incoming" else (d["src"], d["dst"]) if d["journal"] == "movement" \
            else (d["src"],)
        unknown |= {n for n in names if n and n not in have and n not in aliases}
    if unknown:
        # У старій книзі аркуша «Підрозділи» немає: місце заводять у програмі чи зіставляють у вікні перенесення.
        how = ("додайте його в аркуш «Підрозділи» книги чи заведіть у програмі («Підрозділи» → «+ Підрозділ»)"
               if model["subs"] else "заведіть такий підрозділ («Підрозділи» → «+ Підрозділ») або зіставте його з наявним")
        raise BookError([f"невідоме місце «{n}»: {how}" for n in sorted(unknown)])


def _ensure_items(con, model):
    groups = dict(con.execute("SELECT code, id FROM nomen_group"))
    uoms = dict(con.execute("SELECT code, id FROM uom"))
    for it in model["items"]:
        if it["uom"] not in uoms:
            uoms[it["uom"]] = con.execute("INSERT INTO uom(code, name) VALUES(?, ?)", (it["uom"], it["uom"])).lastrowid
        grp = groups[GROUP_OF_TYPE[it["type"]]]
        row = con.execute("SELECT id, archived_at FROM nomen WHERE code = ?", (it["code"],)).fetchone()
        archived = ((row["archived_at"] if row and row["archived_at"] else datetime.date.today().isoformat())
                    if it["archived"] else None)
        vals = (it["name"], grp, uoms[it["uom"]], it["per_ration"], it["note"] or None, archived,
                int(round(float(it["price"]) * 100)) or None)
        if row is None:
            con.execute("INSERT INTO nomen(code, name, group_id, uom_id, per_ration, note, archived_at, app_price_kop, "
                        "source) VALUES(?, ?, ?, ?, ?, ?, ?, ?, 'program')", (it["code"], *vals))
        else:
            con.execute("UPDATE nomen SET name = ?, group_id = ?, uom_id = ?, per_ration = ?, note = ?, archived_at = ?, "
                        "app_price_kop = ? WHERE id = ?", (*vals, row["id"]))
        if model.get("lineCol", True):                           # книга без «рядок норми» рядків норми не змінює
            con.execute("UPDATE nomen SET norm_line = ? WHERE code = ?", (it.get("line"), it["code"]))


def _doc_id(con, d):
    """id документа бази за ключем книги (журнал, дата, номер, сторони)."""
    import state_db                                              # noqa: PLC0415
    for journal, date, no, src, dst, doc in state_db.doc_ids(con):
        if (journal, date, str(no), src, dst) == _doc_key(d):
            return doc
    return None


def _op_doc_ids(con):
    return {r[0] for r in con.execute("""SELECT DISTINCT l.document_id FROM document_line l
                                           JOIN nomen n ON n.id = l.nomen_id JOIN nomen_group g ON g.id = n.group_id
                                          WHERE g.book = 'ОП'""")}


def _apply_doc_fields(con, model):
    """Статус ФЕС і поля документів книги ОП — як у книзі; документів ТЗ не чіпає."""
    import state_db                                              # noqa: PLC0415
    ops = _op_doc_ids(con)
    fes = {k: v for k, v in state_db._load_doc_fes(con).items() if int(k) not in ops}
    meta = {k: v for k, v in state_db._load_doc_meta(con).items() if int(k) not in ops}
    ids = {(j, dt, str(no), s, t): doc for j, dt, no, s, t, doc in state_db.doc_ids(con)}
    for d in model["docs"]:
        doc = ids.get(_doc_key(d))
        if doc is None:
            continue
        if d["status"]:
            fes[str(doc)] = {"status": d["status"]}
        if d["report"] or d["order"] or d["scan"] or d["col"]:
            meta[str(doc)] = {"report": d["report"], "order": d["order"], "scan": d["scan"], "col": d["col"]}
    state_db._save_doc_fes(con, fes)
    state_db._save_doc_meta(con, meta)


def _held_acts(con):
    """Акти бази за ключем документа (дата, номер, підрозділ) з паперами, підписантами й нормами — до того, як прийом
    перепише документи: документ, рядки якого в книзі змінено, записується наново, і дані акта йшли б разом зі старим."""
    held = {}
    for r in con.execute("""SELECT a.*, d.doc_date, d.number, s.name AS sub FROM act a JOIN document d ON d.id = a.document_id
                              JOIN subdivision s ON s.id = d.from_subdivision_id"""):
        doc = r["document_id"]
        held[(r["doc_date"], " ".join(str(r["number"]).split()), r["sub"])] = {
            "doc": doc, "act": {k: r[k] for k in ("kind", "series", "point_id", "period_from", "period_to", "rations",
                                                  "off_norm", "note")},
            "papers": con.execute("SELECT role, number, doc_date FROM act_paper WHERE document_id = ?", (doc,)).fetchall(),
            "signers": con.execute("SELECT role, sort, position, rank, name FROM act_signer WHERE document_id = ?",
                                   (doc,)).fetchall(),
            "norms": dict(con.execute("SELECT n.code, i.norm FROM act_item i JOIN nomen n ON n.id = i.nomen_id "
                                      "WHERE i.document_id = ?", (doc,)).fetchall())}
    return held


def _off_norm(con, doc, series, point, d1, d2, day):
    """Кількість документа акта не за нормою: розрахунок на дату акта (op_acts.compute) дає інші рядки
    або не робиться зовсім (немає пункту, дат)."""
    import op_acts                                               # noqa: PLC0415
    import state_db                                              # noqa: PLC0415
    try:
        prop = op_acts.compute(con, {"series": series, "point": point, "from": d1, "to": d2, "date": day, "docId": doc})
    except state_db.SaveError:
        return True
    have = dict(con.execute("SELECT n.code, SUM(l.qty_milli) FROM document_line l JOIN nomen n ON n.id = l.nomen_id "
                            "WHERE l.document_id = ? GROUP BY n.code", (doc,)).fetchall())
    return {ln["code"]: ln["qty_milli"] for ln in prop["lines"]} != have


def _sync_points(con, points, problems):
    """Пункти з tblPoints: наявний (та сама назва за `_fold`) оновлюється, новий заводиться. Пункт із невідомим підрозділом —
    одна відмова, що його називає; сам пункт лишається, яким був. Повертає пункти бази, яких у книзі немає: вони йдуть геть
    після актів (за якими акти — відмова)."""
    import uuid                                                  # noqa: PLC0415
    subs = dict(con.execute("SELECT name, id FROM subdivision"))
    have = {_fold(name): (pid, name) for pid, name in con.execute("SELECT id, name FROM food_point").fetchall()}
    kept = set()
    for p in points:
        hit = have.get(_fold(p["name"]))
        if hit:
            kept.add(hit[0])
        sub = subs.get(p["sub"])
        if sub is None:
            problems.append(f"Довідка, пункт «{p['name']}»: немає підрозділу «{p['sub']}»")
            continue
        vals = (p["name"], sub, p["hall"] or None, p["title"] or None, p["closed"] or None)
        if hit:
            con.execute("UPDATE food_point SET name = ?, subdivision_id = ?, hall = ?, title = ?, closed_at = ? WHERE id = ?",
                        (*vals, hit[0]))
        else:
            con.execute("INSERT INTO food_point(ext_id, name, subdivision_id, hall, title, closed_at) VALUES(?, ?, ?, ?, ?, ?)",
                        ("fp-" + uuid.uuid4().hex[:10], *vals))
    return [x for x in have.values() if x[0] not in kept]


def _apply_acts(con, model, held):
    """Пункти, дободачі, акти, підписанти й норми рядків — з книги, у транзакції прийому. Аркуша (колонки «норма_дд»)
    немає — ці дані лишаються, як були; акт, документ якого прийом записав наново, стає до документа з тим самим ключем.
    Дободачі пишуться за правилами вікна (state_db._save_rations, _save_ration_dates). «Не за нормою» — розрахунок
    на дату акта дає інше; акт, якого книга не змінила (і його документа теж), лишає свою позначку."""
    import op_acts                                               # noqa: PLC0415
    import state_db                                              # noqa: PLC0415
    problems = []
    book_points = model.get("points")
    gone = _sync_points(con, book_points, problems) if book_points is not None else []
    # Пункт, який називають «Дободачі» й «Акти», — з tblPoints книги (є таблиця) чи з бази (немає), за `_fold`. Пункт, що
    # його книга прибирає, — невідомий; пункт книги з невідомим підрозділом уже має свою відмову, другої за ним немає.
    found = {_fold(name): (i, e) for i, e, name in con.execute("SELECT id, ext_id, name FROM food_point").fetchall()}
    known = {_fold(p["name"]) for p in book_points} if book_points is not None else set(found)

    def point_of(name):
        return found.get(_fold(name), (None, None)) if _fold(name) in known else (None, None)
    if model.get("rations") is not None:
        r = model["rations"]
        problems += [f"Дободачі: пункту «{p}» немає в «Довідка»" for p in r["points"] if _fold(p) not in known]
        try:
            state_db._save_rations(con, {f"{point_of(p)[1]}|{d['date']}": v for d in r["days"] for p, v in d["values"].items()
                                         if point_of(p)[1]})
            state_db._save_ration_dates(con, {d["date"]: {"order": d["order"], "src": d["src"]} for d in r["days"]})
        except state_db.SaveError as e:
            problems.append(str(e))
    ops = _op_doc_ids(con)
    ids = {(d, " ".join(str(no).split()), src): i for j, d, no, src, dst, i in state_db.doc_ids(con)
           if j == "writeoffs" and not dst and i in ops}
    alive = {r[0] for r in con.execute("SELECT document_id FROM act")}
    acts = {}                                                    # документ акта → ключ
    if model.get("acts") is not None:
        for a in model["acts"]:
            key = _act_key(a)
            doc = ids.get(key)
            if doc is None:
                problems.append(f"{a['row']}: немає документа вибуття №{a['no']} від {a['date']} ({a['sub']}) в «Списання»")
                continue
            # Пункту немає — відмова, а акт однаково пишеться (без пункту): його підписанти й норми не відмовляють удруге.
            if a["point"] and _fold(a["point"]) not in known:
                problems.append(f"{a['row']}: пункту «{a['point']}» немає в «Довідка»")
            pid, ext = point_of(a["point"]) if a["point"] else (None, None)
            acts[doc] = key
            old = held.get(key)
            same = old is not None and old["doc"] == doc and doc in alive and (
                old["act"]["series"], old["act"]["point_id"], old["act"]["period_from"], old["act"]["period_to"],
                old["act"]["rations"]) == (a["series"], pid, a["from"], a["to"], a["rations"])
            off = bool(old["act"]["off_norm"]) if same else _off_norm(con, doc, a["series"], ext or "", a["from"], a["to"],
                                                                      a["date"])
            try:
                op_acts._write_act(con, doc, {"series": a["series"], "point": ext, "from": a["from"], "to": a["to"],
                                              "rations": a["rations"], "lines": []},
                                   {"papers": a["papers"], "note": a["note"], "signers": []}, off_norm=off)
            except state_db.SaveError as e:
                problems.append(f"{a['row']}: {e}")
        for (doc,) in con.execute("SELECT document_id FROM act").fetchall():
            if doc not in acts:
                con.execute("DELETE FROM act WHERE document_id = ?", (doc,))
    else:
        for key, old in held.items():
            doc = ids.get(key)
            if doc is None:                                      # документ акта прибрано з книги — акт іде з ним
                continue
            acts[doc] = key
            if old["doc"] == doc and doc in alive:
                continue
            o = old["act"]
            point = con.execute("SELECT ext_id FROM food_point WHERE id = ?", (o["point_id"],)).fetchone()
            off = _off_norm(con, doc, o["series"], point[0] if point else "", o["period_from"], o["period_to"], key[0])
            con.execute("""INSERT INTO act(document_id, kind, series, point_id, period_from, period_to, rations, off_norm, note)
                           VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (doc, o["kind"], o["series"], o["point_id"], o["period_from"], o["period_to"], o["rations"],
                         int(off), o["note"]))
            con.executemany("INSERT INTO act_paper(document_id, role, number, doc_date) VALUES(?, ?, ?, ?)",
                            [(doc, *p) for p in old["papers"]])
    _apply_signers(con, model, held, ids, acts, problems)
    _apply_norms(con, model, held, ids, acts, problems)
    for p, name in gone:
        if con.execute("SELECT 1 FROM act WHERE point_id = ?", (p,)).fetchone():
            problems.append(f"Довідка: пункт «{name}» прибрано, а за ним складено акти. Пункт закривають («закрито»), "
                            "а не прибирають")
        else:
            con.execute("DELETE FROM food_point WHERE id = ?", (p,))
    if problems:
        raise BookError(problems)


def _apply_signers(con, model, held, ids, acts, problems):
    """Підписанти актів — з аркуша; аркуша немає — ті, що були в акта."""
    rows = {doc: [tuple(s)[:5] for s in (held.get(key) or {}).get("signers", [])] for doc, key in acts.items()}
    if model.get("signers") is not None:
        rows = {doc: [] for doc in acts}
        for s in model["signers"]:
            doc = ids.get(_act_key(s))
            if doc not in acts:
                problems.append(f"{s['row']}: немає акта №{s['no']} від {s['date']} ({s['sub']})")
                continue
            rows[doc].append((s["role"], s["sort"], s["pos"] or None, s["rank"] or None, s["name"] or None))
    for doc, people in rows.items():
        con.execute("DELETE FROM act_signer WHERE document_id = ?", (doc,))
        con.executemany("INSERT INTO act_signer(document_id, role, sort, position, rank, name) VALUES(?, ?, ?, ?, ?, ?)",
                        [(doc, *x) for x in people])


def _apply_norms(con, model, held, ids, acts, problems):
    """Норма на д/д рядків акта — з «норма_дд»; колонки немає — ті, що були. Норма рядка не акта — відмова."""
    lines = {doc: dict(con.execute("SELECT n.code, NULL FROM document_line l JOIN nomen n ON n.id = l.nomen_id "
                                   "WHERE l.document_id = ?", (doc,)).fetchall()) for doc in acts}
    if model.get("normCol", True):
        for d in model["docs"]:
            if d["journal"] != "writeoffs":
                continue
            norms = {ln["code"]: ln["norm"] for ln in d["lines"] if ln.get("norm") is not None}
            doc = ids.get((d["date"], " ".join(str(d["no"]).split()), d["src"])) if not d["dst"] else None
            if doc in lines:
                lines[doc] = {c: norms.get(c) for c in lines[doc]}
            elif norms:
                problems.append(f"{_doc_label(d)}: «норма_дд» буває лише в рядках акта. Додайте акт на аркуш «Акти» "
                                "або очистіть «норма_дд»")
    else:
        for doc, key in acts.items():
            kept = (held.get(key) or {}).get("norms", {})
            lines[doc] = {c: kept.get(c) for c in lines[doc]}
    for doc, norms in lines.items():
        con.execute("DELETE FROM act_item WHERE document_id = ?", (doc,))
        con.executemany("INSERT INTO act_item(document_id, nomen_id, norm) VALUES(?, (SELECT id FROM nomen WHERE code = ?), ?)",
                        [(doc, c, n) for c, n in norms.items() if n is not None])


def _drop_unused_items(con, keep):
    """Позиції книги ОП, яких у книзі немає й на які не посилається жоден документ, — геть."""
    for nid, code in con.execute("""SELECT n.id, n.code FROM nomen n JOIN nomen_group g ON g.id = n.group_id
                                     WHERE g.book = 'ОП'""").fetchall():
        if code in keep or con.execute("SELECT 1 FROM document_line WHERE nomen_id = ? LIMIT 1", (nid,)).fetchone():
            continue
        con.execute("DELETE FROM nomen_report_line WHERE nomen_id = ?", (nid,))
        con.execute("DELETE FROM nomen_form_skip WHERE nomen_id = ?", (nid,))
        con.execute("DELETE FROM fes_item_map WHERE nomen_id = ?", (nid,))
        con.execute("DELETE FROM nomen WHERE id = ?", (nid,))


def canonical(con):
    """Порівнювана проєкція книги ОП у базі: модель і залишки по кодах і місцях на кожну дату руху.
    Партії книги ОП програма розкладає сама за порядком надходження — у порівняння вони не входять."""
    m = model_from_db(con)
    dates = sorted({d["date"] for d in m["docs"]})
    rests = {dt: sorted((k, round(v, 3)) for k, v in _balances(m, dt).items() if abs(v) > 1e-9) for dt in dates}
    docs = sorted((_doc_key(d), _doc_sig(d)) for d in m["docs"])
    return {"items": sorted(m["items"], key=lambda i: i["code"]), "docs": docs, "rests": rests,
            "parties": m["parties"], "subs": [(s["name"], s["parent"], s["kind"], s["closed"]) for s in m["subs"]],
            "acts": sorted(_act_key(a) + _act_sig(a)[:-1] for a in m["acts"]),
            "signers": sorted(_act_key(s) + (s["role"], s["sort"], s["pos"], s["rank"], s["name"]) for s in m["signers"]),
            "rations": m["rations"],
            "points": sorted((p["name"], p["sub"], p["hall"], p["title"], p["closed"]) for p in m["points"])}


# ---------------------------------------------------------------- перше перенесення

def legacy_places(model):
    """Назви місць, які називають документи старої книги (після tblRename), за абеткою."""
    used = set()
    for d in model["docs"]:
        if d["journal"] == "incoming":
            used.add(d["dst"])
        elif d["journal"] == "movement":
            used.update((d["src"], d["dst"]))
        else:
            used.add(d["src"])
    return sorted(n for n in used if n)


def _group_by_code(code):
    """Група архівної позиції — за діапазоном коду книги: 9xxx — серветки, 80xx/81xx — миючі,
    82xx — господарчі, решта — одноразовий посуд (категорія архіву в старій книзі ненадійна)."""
    c = str(code)
    if c.startswith("9"):
        return "ОП.2"
    if c.startswith(("80", "81")):
        return "ОП.3"
    if c.startswith("82"):
        return "ОП.4"
    return "ОП.1"


def _round_quantities(model, report):
    """Облік веде кількість до тисячних, а стара книга мала й п'ять знаків (19,00313 кг). Рядок
    округлюється до тисячних; вибуття, яке в книзі забирало все до нуля, забирає рівно стільки,
    скільки після округлення лишилося (інакше залишок ішов би в мінус на 0,001)."""
    order = {"incoming": 0, "movement": 1, "writeoffs": 2}
    exact, rounded = {}, {}
    for d in sorted(model["docs"], key=lambda d: (d["date"], order[d["journal"]])):
        for ln in d["lines"]:
            q = float(ln["qty"])
            r = round(q + 1e-12, 3)
            src = None if d["journal"] == "incoming" else (ln["code"], d["src"])
            dst = (ln["code"], d["dst"]) if d["journal"] != "writeoffs" else None
            if src is not None:
                have_r, have_e = rounded.get(src, 0.0), exact.get(src, 0.0)
                if r > have_r + 1e-9 and q <= have_e + 1e-9 and r - have_r <= 0.0015:
                    r = round(have_r, 3)
            if abs(r - q) > 1e-9:
                report.append(f"{SHEETS[d['journal']][0]} №{d['no']} від {d['date']}, код {ln['code']}: "
                              f"кількість {_comma(q)} → {_comma(r)} (облік веде тисячні)")
                ln["qty"] = _num(r)
            if src is not None:
                rounded[src] = rounded.get(src, 0.0) - r
                exact[src] = exact.get(src, 0.0) - q
            if dst is not None:
                rounded[dst] = rounded.get(dst, 0.0) + r
                exact[dst] = exact.get(dst, 0.0) + q


def normalize_legacy(model, mapping):
    """Стара книга → модель програми: місця — за таблицею, яку погодив власник (стара назва →
    підрозділ програми; порожнє — лишити як є); архівні позиції — у групу за діапазоном коду.
    Кожна правка — у model["report"]."""
    report = list(model.get("report", []))
    mapping = {k: v for k, v in (mapping or {}).items() if v and k != v}
    for d in model["docs"]:
        for side in ("src", "dst"):
            if d[side] in mapping and not (d["journal"] == "incoming" and side == "src") \
                    and not (d["journal"] == "writeoffs" and side == "dst"):
                d[side] = mapping[d[side]]
    for old, new in sorted(mapping.items()):
        report.append(f"місце «{old}» → «{new}»")
    for it in model["items"]:
        if it["archived"]:
            grp = _group_by_code(it["code"])
            if TYPE_OF_GROUP[grp] != it["type"]:
                report.append(f"позиція {it['code']}: група «{TYPE_OF_GROUP[grp]}» замість «{it['type']}» (архів)")
                it["type"], it["cat"] = TYPE_OF_GROUP[grp], CAT_OF_GROUP[grp]
    _round_quantities(model, report)
    model["report"] = report
    return model
