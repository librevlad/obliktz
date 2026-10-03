# -*- coding: utf-8 -*-
"""Відомість закуплених (отриманих) матеріально-технічних засобів — відомість МТЗ.

Форму щомісяця збирає вищий штаб із кожної частини: чотирнадцять граф — служба
забезпечення, група, найменування, одиниця виміру, кількість цілим числом, ціна
й сума в гривнях, частина, належність, КПКВ, КЕКВ, джерело надходжень і
примітка. Служба подає свої рядки виконавцеві частини в тій самій формі, тож
їх можна перенести у відомість частини як є.

Тут — запис рядків служби у цій формі (`save_mtz`) і читання відомості частини
(`read_mtz`), з якою служба звіряє свої рядки.
"""
import io
import math
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

FONT = "Times New Roman"
F_BODY = Font(name=FONT, size=12)
F_BOLD = Font(name=FONT, size=12, bold=True)
F_SIGN = Font(name=FONT, size=14)
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
FMT_QTY = "#,##0"
FMT_QTY_PART = "#,##0.###"      # дробову кількість відомість не приймає, але й ховати її не можна
FMT_MONEY = "#,##0.00"

# Графи бланка: (заголовок, ширина). Порядок і назви — як у зразку відомості. Ширини
# свої: зразок розрахований на відомість усієї частини й утричі ширший за аркуш, а
# рядки однієї служби мають друкуватися так, щоб їх можна було прочитати.
COLUMNS = [
    ("№\n з/п", 5.5), ("Служба забезпечення", 16), ("Група ОВТ, МТЗ", 17),
    ("Найменування предмету закупівлі", 46), ("Одиниця виміру", 10.5), ("Кількість", 11.5),
    ("Вартість за одиницю (грн)", 15), ("Сума\n(грн)", 16.5),
    ("Найменування військової частини (підрозділу)", 16), ("Належність", 13.5), ("КПКВ ", 12.5),
    ("КЕКВ", 8), ("Джерело\nнадходжень", 14), ("Примітка", 20),
]
FIRST = 7                       # перший рядок майна — як у відомості частини
LINE_PT = 15.6                  # рядок тексту кеглем 12

# Ширина знаків Times New Roman у тисячних кегля. Рядок із перенесеним текстом сам
# не росте: книгу пише не Excel, і висоту треба вказати. Рахуємо її так, як переносить
# клітинка, — за словами.
_WIDTHS = {
    250: " .,'", 278: "іїijlt;:/\\", 333: "ІЇfrI!’ʼ`()[]-", 400: 'ґзгsJ°"', 444: "єэтаесёacez?“”",
    510: "ҐЗьявклорухчбдbdghknopquvxy0123456789«»–*#_", 535: "ийнпцъ", 580: "РБЬГFPS+=<>×",
    611: "ЕТЁELTZ", 680: "мфыЧЄЭВКСЯЛДBCR", 722: "жЪУАИЙНОПХЦwADGHKNOQUVXY", 790: "юшщФm&",
    950: "ЫМЖMW%@№", 1030: "ШЩЮ—",
}
_EM = {ch: w / 1000 for w, chars in _WIDTHS.items() for ch in chars}
PX_EM = 16                      # кегль 12 на екрані
BOLD = 1.09                     # жирний ширший за звичайний
SPARE = 1.04                    # друк переносить трохи інакше, ніж екран: краще зайвий рядок


def _width(text, bold=False):
    """Ширина тексту кеглем 12, у пікселях."""
    return sum(_EM.get(ch, 0.6) for ch in text) * PX_EM * (BOLD if bold else 1.0)


def _lines(text, width, bold=False):
    """Скільки рядків займе текст у графі такої ширини (у знаках Excel)."""
    room = max(20.0, width * 7 - 5) / SPARE
    total = 0
    for part in str(text if text is not None else "").split("\n"):
        used, lines = 0.0, 1
        for word in part.split():
            w = _width(word, bold)
            gap = _width(" ", bold) if used else 0.0
            if used and used + gap + w > room:
                lines, used, gap = lines + 1, 0.0, 0.0
            if w > room:                              # слово, довше за рядок, рветься посередині
                extra = math.ceil(w / room) - 1
                lines += extra
                w -= extra * room
            used += gap + w
        total += lines
    return total


def _shown(v, fmt):
    """Число так, як його покаже клітинка: «1 258 236,63»."""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v if v is not None else "")
    digits = 2 if fmt == FMT_MONEY else 0
    return f"{v:,.{digits}f}".replace(",", " ").replace(".", ",")


def build_mtz(spec: dict) -> Workbook:
    """Книга з відомістю. spec: {title: [рядки заголовка], rows: [[служба, група,
    найменування, одиниця, кількість, ціна, частина, належність, КПКВ, КЕКВ, джерело,
    примітка]], total: підпис підсумку, signs: [{pos, rank, name}]}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Відомість"
    last_col = get_column_letter(len(COLUMNS))
    widths = [width for _, width in COLUMNS]
    rows = [(list(row) + [""] * 12)[:12] for row in spec.get("rows") or []]

    # Число, яке не вміщається в графу, Excel показує ґратками: під найбільшу ціну
    # й під підсумок графа ширшає.
    def number(v):
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0

    prices = [number(row[5]) for row in rows]
    sums = [number(row[4]) * number(row[5]) for row in rows]
    for col, shown, bold in ((7, max(prices, default=0), False), (8, sum(sums), True)):
        need = (_width(_shown(shown, FMT_MONEY), bold) * SPARE + 9) / 7
        widths[col - 1] = max(widths[col - 1], round(need, 1))
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for r, text in enumerate((spec.get("title") or [])[:3], 2):
        ws.merge_cells(f"A{r}:{last_col}{r}")
        cell = ws.cell(r, 1, text)
        cell.font, cell.alignment = (F_BODY if r == 3 else F_BOLD), CENTER
        ws.row_dimensions[r].height = LINE_PT * _lines(text, sum(widths), r != 3)
    for i, (label, _) in enumerate(COLUMNS, 1):
        head = ws.cell(5, i, label)
        head.font, head.alignment, head.border = F_BOLD, CENTER, BOX
        num = ws.cell(6, i, i)
        num.font, num.alignment, num.border = F_BODY, CENTER, BOX
    ws.row_dimensions[5].height = LINE_PT * max(_lines(label, widths[i], True) for i, (label, _) in enumerate(COLUMNS))
    ws.row_dimensions[6].height = LINE_PT

    for n, row in enumerate(rows, 1):
        r = FIRST + n - 1
        service, group, name, unit, qty, price, part, belongs, kpkv, kekv, src, note = row
        kekv = int(kekv) if str(kekv).strip().isdigit() else kekv
        values = [n, service, group, name, unit, qty, price, f"=F{r}*G{r}", part, belongs, kpkv, kekv, src, note or None]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.alignment, cell.border = F_BODY, CENTER, BOX
        whole = not isinstance(qty, float) or qty.is_integer()
        ws.cell(r, 6).number_format = FMT_QTY if whole else FMT_QTY_PART
        ws.cell(r, 7).number_format = FMT_MONEY
        ws.cell(r, 8).number_format = FMT_MONEY
        ws.cell(r, 12).number_format = "0"
        tall = max(_lines(v, widths[c]) for c, v in enumerate(values) if isinstance(v, str) and not v.startswith("="))
        ws.row_dimensions[r].height = LINE_PT * max(2, tall)

    end = FIRST + len(rows)                   # рядок підсумку
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(end, c)
        cell.font, cell.alignment, cell.border = F_BOLD, CENTER, BOX
    # У зразку підпис підсумку стоїть у графі служби, широкій на всю назву; тут графа
    # вузька, тож підпис займає графи 2–5.
    ws.cell(end, 2, spec.get("total") or "Всього:")
    ws.merge_cells(start_row=end, start_column=2, end_row=end, end_column=5)
    if rows:
        ws.cell(end, 6, f"=SUM(F{FIRST}:F{end - 1})").number_format = FMT_QTY
        ws.cell(end, 8, f"=SUM(H{FIRST}:H{end - 1})").number_format = FMT_MONEY
    ws.row_dimensions[end].height = LINE_PT * max(2, _lines(ws.cell(end, 2).value, sum(widths[1:5]), True))

    r = end + 2
    for sign in spec.get("signs") or []:
        ws.cell(r, 1, sign.get("pos") or "").font = F_SIGN
        ws.cell(r + 1, 1, sign.get("rank") or "").font = F_SIGN
        ws.cell(r + 1, 8, sign.get("name") or "").font = F_SIGN
        r += 4

    ws.freeze_panes = f"A{FIRST}"
    ws.print_title_rows = "5:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddFooter.center.text = "&P з &N"
    pr = wb.properties
    pr.creator = pr.lastModifiedBy = "Облік ТЗ ПС"
    pr.created = pr.modified = datetime.now()
    return wb


def save_mtz(spec: dict, folder: str) -> str:
    """Записує відомість у теку вивантажень і повертає шлях до файла."""
    from excel_export import safe_name                              # noqa: PLC0415
    os.makedirs(folder, exist_ok=True)
    name = safe_name(spec.get("file") or "Відомість МТЗ")
    path = os.path.join(folder, name + ".xlsx")
    wb = build_mtz(spec)
    try:
        wb.save(path)
    except PermissionError:                  # файл із таким іменем відкрито в Excel
        path = os.path.join(folder, f"{name} {datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path


# ------------------------------------------------------- відомість частини
# Графа відомості за початком її заголовка. У файлах частин таблиця буває
# зсунута на стовпець чи рядок, тож графи шукаємо за шапкою, а не за місцем.
HEADS = [("no", "№"), ("service", "служба забезпечення"), ("group", "група"),
         ("name", "найменування предмет"), ("unit", "одиниця"), ("qty", "кількість"),
         ("price", "вартість"), ("sum", "сума"), ("part", "найменування військової"),
         ("belongs", "належність"), ("kpkv", "кпкв"), ("kekv", "кекв"), ("src", "джерело"),
         ("note", "примітка")]
MAX_ROWS = 50000


class NotMtz(ValueError):
    """Файл не схожий на відомість МТЗ."""


def _norm(v):
    return " ".join(str(v if v is not None else "").split()).lower()


def _number(v):
    """Число з клітинки: 1 250,50 і «1250.5» — те саме; не число — None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def read_mtz(data: bytes) -> dict:
    """Рядки відомості МТЗ із файла Excel: {sheet, rows: [{no, service, group, name,
    unit, qty, price, part, belongs, kpkv, kekv, src, note}]}. Сума не читається:
    у клітинці формула, і вона дорівнює кількості на ціну."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as e:                                        # noqa: BLE001
        raise NotMtz("файл не відкривається як книга Excel (.xlsx)") from e
    try:
        for ws in wb.worksheets:
            grid = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                grid.append(list(row))
                if i >= MAX_ROWS:
                    break
            for hr, row in enumerate(grid[:40]):
                cols = {}
                for c, v in enumerate(row):
                    text = _norm(v)
                    for key, start in HEADS:
                        if key not in cols and text.startswith(start):
                            cols[key] = c
                            break
                if not {"service", "name", "qty", "price"} <= set(cols):
                    continue
                return {"sheet": ws.title, "rows": _rows(grid[hr + 1:], cols)}
    finally:
        wb.close()
    raise NotMtz("це не відомість МТЗ: у файлі немає шапки з графами «Служба забезпечення», "
                 "«Найменування предмету закупівлі», «Кількість» і «Вартість за одиницю»")


def _rows(grid, cols):
    out = []

    def at(row, key):
        c = cols.get(key)
        return row[c] if c is not None and c < len(row) else None

    for row in grid:
        service = " ".join(str(at(row, "service") or "").split())
        name = " ".join(str(at(row, "name") or "").split())
        if service.lower().startswith("всього"):
            break
        qty, price = _number(at(row, "qty")), _number(at(row, "price"))
        if not service or not name or qty is None:
            continue                          # рядок із номерами граф, порожній рядок
        if _norm(at(row, "service")) == "2" and _norm(at(row, "name")) == "4":
            continue
        no = _number(at(row, "no"))
        kekv = at(row, "kekv")
        out.append({
            "no": int(no) if no is not None and float(no).is_integer() else (no if no is not None else ""),
            "service": service, "group": " ".join(str(at(row, "group") or "").split()), "name": name,
            "unit": " ".join(str(at(row, "unit") or "").split()), "qty": qty, "price": price if price is not None else 0,
            "part": " ".join(str(at(row, "part") or "").split()),
            "belongs": " ".join(str(at(row, "belongs") or "").split()),
            "kpkv": " ".join(str(at(row, "kpkv") or "").split()),
            "kekv": str(int(kekv)) if isinstance(kekv, (int, float)) and not isinstance(kekv, bool)
            else " ".join(str(kekv or "").split()),
            "src": " ".join(str(at(row, "src") or "").split()),
            "note": " ".join(str(at(row, "note") or "").split()),
        })
    return out
