# -*- coding: utf-8 -*-
"""Вивантаження таблиць програми в Excel — для друку.

Друк із браузера залежить від браузера: поля, масштаб, переноси, шапка на
кожній сторінці виходять кожного разу по-різному. Excel друкує передбачувано, і
в ньому документ можна ще поправити перед підписом. Тому програма не друкує
сама, а збирає .xlsx, уже розмічений під аркуш А4, і відкриває його в Excel.

Опис вивантаження приходить від сторінки у вигляді простого словника, щоб
застосунок не знав нічого про openpyxl, а цей модуль — нічого про облік:

    {"file": "Накладна №17 від 2026-03-02",
     "sheets": [{
        "name": "Накладна", "orientation": "portrait",
        "top": ["Військова частина А0000", "продовольча служба"],
        "title": "НАКЛАДНА № 17", "subtitle": "від 02.03.2026",
        "lines": ["Від кого: склад", "Кому: їдальня"],
        "head": [["№", "Найменування", {"t": "Рух", "span": 3}], [...]],
        "widths": [6, 42, 10],
        "rows": [[1, "Термос ТВН-12", 2]],
        "num": [2], "money": [], "total": ["Разом", "", 2],
        "after": ["Усього найменувань: 1."],
        "signs": ["Відпустив", "Прийняв"]}]}
"""
import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Times New Roman"
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=13, bold=True)
F_SMALL = Font(name=FONT, size=9)
A_WRAP = Alignment(wrap_text=True, vertical="top")
A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_RIGHT = Alignment(horizontal="right", vertical="top")
A_TOP = Alignment(horizontal="center", vertical="top")
FMT_QTY = "#,##0.###"
# Ціле — окремим форматом: «#,##0.###» в українському Excel показує 1 як «1,».
FMT_QTY_INT = "#,##0"
FMT_MONEY = "#,##0.00"
FMT_DATE = "dd.mm.yyyy"
ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")

MAX_SHEET_NAME = 31
BAD_NAME = re.compile(r'[\\/:*?"<>|\[\]]')


def safe_name(name: str, limit: int = 120) -> str:
    """Ім'я файла без символів, яких не приймає Windows; дати — як у документах."""
    out = re.sub(r"\b(\d{4})-(\d{2})-(\d{2})\b", r"\3.\2.\1", str(name or "вивантаження"))
    out = BAD_NAME.sub("-", out).strip(" .")
    return (out or "вивантаження")[:limit]


def qty_format(v) -> str:
    """Формат кількості: ціле без коми, дробове — до тисячних."""
    try:
        return FMT_QTY_INT if float(v).is_integer() else FMT_QTY
    except (TypeError, ValueError):
        return FMT_QTY


def as_date(v):
    """Рядок «2026-09-11» -> дата для клітинки; решта — як є."""
    m = ISO_DATE.match(v) if isinstance(v, str) else None
    if not m:
        return None
    try:
        return datetime(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def clean_workbook(wb):
    """Прибирає зі шаблона все, чого програма не передавала: посилання на
    надбудови з диска автора (Excel питав би про оновлення зв'язків), автора,
    дату останнього друку й решту властивостей чужого файла."""
    wb._external_links = []                                   # noqa: SLF001
    pr = wb.properties
    pr.creator = pr.lastModifiedBy = "Облік ТЗ ПС"
    pr.title = pr.description = pr.category = pr.keywords = None
    pr.lastPrinted = None
    pr.created = pr.modified = datetime.now()
    return wb


def _sheet_title(name: str, used: set) -> str:
    base = BAD_NAME.sub("-", str(name or "Аркуш"))[:MAX_SHEET_NAME] or "Аркуш"
    title, n = base, 2
    while title in used:
        suffix = f" ({n})"
        title = base[:MAX_SHEET_NAME - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def _cell(spec):
    """Клітинка шапки: рядок або {"t": текст, "span": ширина}."""
    if isinstance(spec, dict):
        return str(spec.get("t", "")), max(1, int(spec.get("span", 1)))
    return str(spec if spec is not None else ""), 1


def _fill_sheet(ws, sh: dict) -> None:
    widths = sh.get("widths") or []
    head = sh.get("head") or []
    ncols = max([len(widths)] + [sum(_cell(c)[1] for c in row) for row in head]
                + [len(r) for r in sh.get("rows") or []] + [1])
    num = set(sh.get("num") or [])
    money = set(sh.get("money") or [])
    r = 1

    def full_width(text, font, align, height=None):
        nonlocal r
        ws.cell(row=r, column=1, value=text).font = font
        ws.cell(row=r, column=1).alignment = align
        if ncols > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        if height:
            ws.row_dimensions[r].height = height
        r += 1

    for line in sh.get("top") or []:
        full_width(line, F_SMALL, Alignment(horizontal="right"))
    if sh.get("title"):
        r += 0 if not sh.get("top") else 1
        full_width(sh["title"], F_TITLE, Alignment(horizontal="center"), 20)
    if sh.get("subtitle"):
        full_width(sh["subtitle"], F_BODY, Alignment(horizontal="center"))
    if sh.get("title") or sh.get("subtitle"):
        r += 1
    for line in sh.get("lines") or []:
        full_width(line, F_BODY, Alignment(horizontal="left", wrap_text=True))
    if sh.get("lines"):
        r += 1

    # Шапка таблиці: кілька рядків, клітинки з об'єднанням по ширині.
    head_first = r
    for row in head:
        c = 1
        for spec in row:
            text, span = _cell(spec)
            cell = ws.cell(row=r, column=c, value=text)
            cell.font, cell.alignment = F_HEAD, A_CENTER
            for k in range(c, c + span):
                ws.cell(row=r, column=k).border = BOX
            if span > 1:
                ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
            c += span
        for k in range(c, ncols + 1):
            ws.cell(row=r, column=k).border = BOX
        r += 1
    head_last = r - 1

    def put(cell, c, v, font):
        when = as_date(v)
        cell.value = when or v
        # Назва, номер чи примітка, що починається з «=», — текст, а не формула.
        if isinstance(v, str) and v.startswith("="):
            cell.data_type = "s"
        cell.font, cell.border = font, BOX
        if when:
            cell.number_format, cell.alignment = FMT_DATE, A_TOP
        elif (c - 1) in money:
            cell.number_format, cell.alignment = FMT_MONEY, A_RIGHT
        elif (c - 1) in num:
            cell.number_format, cell.alignment = qty_format(v), A_RIGHT
        else:
            cell.alignment = A_WRAP

    for data in sh.get("rows") or []:
        for c, v in enumerate(data, 1):
            put(ws.cell(row=r, column=c), c, v, F_BODY)
        for c in range(len(data) + 1, ncols + 1):
            ws.cell(row=r, column=c).border = BOX
        r += 1

    if sh.get("total"):
        for c, v in enumerate(sh["total"], 1):
            put(ws.cell(row=r, column=c), c, v, F_HEAD)
        r += 1

    if sh.get("after"):
        r += 1
        for line in sh["after"]:
            full_width(line, F_BODY, Alignment(horizontal="left", wrap_text=True))

    signs = sh.get("signs") or []
    if signs:
        r += 2
        # Підписи — рискою на всю половину аркуша, підпис дрібно під нею.
        half = max(1, ncols // max(1, len(signs)))
        for i, text in enumerate(signs):
            c1 = 1 + i * half
            c2 = ncols if i == len(signs) - 1 else c1 + half - 1
            for k in range(c1, c2 + 1):
                ws.cell(row=r, column=k).border = Border(top=THIN)
            cell = ws.cell(row=r, column=c1, value=text)
            cell.font, cell.alignment = F_SMALL, Alignment(horizontal="center", wrap_text=True)
            if c2 > c1:
                ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        r += 1

    for i in range(ncols):
        w = widths[i] if i < len(widths) and widths[i] else 12
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    # Друк: А4, по ширині аркуша, шапка таблиці повторюється на кожній сторінці.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ("landscape" if sh.get("orientation") == "landscape"
                                 else "portrait")
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    if head:
        ws.print_title_rows = f"{head_first}:{head_last}"
        ws.freeze_panes = f"A{head_last + 1}"           # координатою: комірка може бути в об'єднанні


def build(spec: dict) -> Workbook:
    wb = clean_workbook(Workbook())
    wb.remove(wb.active)
    used = set()
    sheets = spec.get("sheets") or []
    if not sheets:
        raise ValueError("немає жодного аркуша для вивантаження")
    for sh in sheets:
        ws = wb.create_sheet(_sheet_title(sh.get("name"), used))
        _fill_sheet(ws, sh)
    return wb


def save(spec: dict, folder: str) -> str:
    """Записує книгу в теку вивантажень і повертає шлях до файла.

    Якщо файл із таким іменем уже відкритий в Excel, записати поверх нього не
    вийде — тоді до імені додається час, а старий файл лишається як був.
    """
    os.makedirs(folder, exist_ok=True)
    name = safe_name(spec.get("file") or "вивантаження")
    path = os.path.join(folder, name + ".xlsx")
    wb = build(spec)
    try:
        wb.save(path)
    except PermissionError:
        path = os.path.join(folder, f"{name} {datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path


# ---------------------------------------------------------------- накладна
# Бланк — власна накладна служби «Накладна (вимога)» за Додатком 25 до
# Інструкції з обліку військового майна (п. 24 розд. IV). Файл узято як є й
# лише заповнюємо: шрифти, рамки, об'єднання й поля друку лишаються такими,
# якими їх звикла бачити служба.
TEMPLATE = "nakladna-vymoga.xlsx"
SHEET = "Накладна-вимога"
FIRST_LINE = 23          # перший рядок таблиці
LINE_SLOTS = 3           # скільки рядків у бланку заготовлено
TOTAL_ROW = 26           # рядок «Всього» в бланку


def template_path(name: str = TEMPLATE) -> str:
    """Бланк поруч із модулем; у зібраному .exe — у розпакованому каталозі."""
    import sys
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, "templates", name)


def _shift_down(ws, from_row: int, by: int) -> None:
    """Зсуває все від `from_row` униз на `by` рядків разом з об'єднаннями й
    висотами. openpyxl.insert_rows цього не робить: об'єднання лишаються на
    місці й налазять на нові рядки, тому зсув зроблено вручну."""
    if by <= 0:
        return
    moved = [m for m in list(ws.merged_cells.ranges) if m.min_row >= from_row]
    for m in moved:
        ws.unmerge_cells(str(m))
    ws.move_range(f"A{from_row}:{get_column_letter(ws.max_column)}{ws.max_row}",
                  rows=by, cols=0, translate=False)
    for m in moved:
        ws.merge_cells(start_row=m.min_row + by, start_column=m.min_col,
                       end_row=m.max_row + by, end_column=m.max_col)
    heights = {r: d.height for r, d in ws.row_dimensions.items()
               if r >= from_row and d.height}
    for r in sorted(heights, reverse=True):
        ws.row_dimensions[r + by].height = heights[r]
        ws.row_dimensions[r].height = None


def _shift_up(ws, from_row: int, by: int) -> None:
    """Прибирає `by` рядків над `from_row`: усе від `from_row` піднімається.

    Об'єднання в прибраних рядках зникають разом із ними, нижчі — зсуваються;
    висоти рядків ідуть слідом, як і в `_shift_down`."""
    if by <= 0:
        return
    gone = from_row - by
    moved = []
    for m in list(ws.merged_cells.ranges):
        if m.max_row >= gone:
            ws.unmerge_cells(str(m))
            if m.min_row >= from_row:
                moved.append(m)
    last_row, last_col = ws.max_row, get_column_letter(ws.max_column)
    ws.move_range(f"A{from_row}:{last_col}{last_row}", rows=-by, cols=0, translate=False)
    for m in moved:
        ws.merge_cells(start_row=m.min_row - by, start_column=m.min_col,
                       end_row=m.max_row - by, end_column=m.max_col)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if r >= gone}
    for r in range(gone, last_row + 1):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        if r >= from_row and h:
            ws.row_dimensions[r - by].height = h


def _copy_style(src, dst) -> None:
    from copy import copy                                    # noqa: PLC0415
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.fill = copy(src.fill)


def _person(p) -> str:
    """{pos, rank, name} -> «посада, звання Ім'я ПРІЗВИЩЕ», як у бланку служби."""
    if not p or not p.get("name"):
        return ""
    return ", ".join(x for x in (p.get("pos"), " ".join(
        x for x in (p.get("rank"), p.get("name")) if x)) if x)


def fill_invoice(spec: dict):
    """Заповнює бланк накладної. Повертає книгу openpyxl.

    spec: no, date (YYYY-MM-DD), operation, basis, sender, receiver, lines
    (name, code, uom, cat, price, qty, note), qty_words, sum_words;
    place — місце складання на дату накладної; підписанти {pos, rank, name}:
    sender_person, receiver_person (МВО), chief (начальник служби), accountant
    (хто відобразив в обліку), fes (начальник ФЕС). Кого немає — рядок лишається
    для підпису від руки.
    """
    from openpyxl import load_workbook
    wb = load_workbook(template_path())
    ws = wb[SHEET]
    # Допоміжний аркуш бланка рахує «прописом» через надбудову Excel
    # (sumpropua.xla). Її може не бути на іншому комп'ютері, тому слова
    # пишуться готовими, а допоміжний аркуш прибирається.
    for name in list(wb.sheetnames):
        if name != SHEET:
            del wb[name]

    clean_workbook(wb)

    lines = spec.get("lines") or []
    extra = max(0, len(lines) - LINE_SLOTS)
    if extra:
        _shift_down(ws, TOTAL_ROW, extra)
        for r in range(FIRST_LINE + LINE_SLOTS, FIRST_LINE + LINE_SLOTS + extra):
            ws.row_dimensions[r].height = ws.row_dimensions[FIRST_LINE].height
            for c in range(1, 11):
                _copy_style(ws.cell(row=FIRST_LINE, column=c), ws.cell(row=r, column=c))
    last = FIRST_LINE + max(len(lines), LINE_SLOTS) - 1
    total = TOTAL_ROW + extra
    # У бланку служби рядки 24 і 25 приховані — накладна на дві чи три позиції
    # друкувалася без другої й третьої, хоч у підсумку вони враховані.
    for r in range(FIRST_LINE, total + 1):
        ws.row_dimensions[r].hidden = False

    date = spec.get("date") or ""
    try:
        when = datetime.strptime(date, "%Y-%m-%d")
    except ValueError:
        when = None

    snd = spec.get("sender_person") or {}
    rcv = spec.get("receiver_person") or {}
    # Реквізити частини — з довідника програми: бланк лишається бланком, а
    # назва частини, ЄДРПОУ, назва служби й строк дії живуть в обліку.
    unit = spec.get("unit") or {}
    # Реквізити частини — завжди з налаштувань: чого немає, те лишається
    # порожнім, а не береться зі зразка бланка.
    ws["A5"] = unit.get("legalName") or None
    # Код за ЄДРПОУ — вісім знаків, а не число: записаний числом, він губив нуль попереду.
    code = str(unit.get("edrpou") or "").strip()
    ws["B7"] = code or None
    if code:
        ws["B7"].number_format = "@"
    if unit.get("serviceFull"):
        ws["E15"] = f"Служба забезпечення: {unit['serviceFull']}"
    if unit.get("validDays") is not None:
        days = int(unit["validDays"])
        ws["B9"] = f"=B15+{days}" if days else "=B15"
    ws["A10"] = f"Накладна (вимога) №{spec.get('no', '')}"
    ws["E11"] = spec.get("place") or None
    ws["B15"] = when or date
    ws["A16"] = f"Вид операції: {spec.get('operation') or 'Розподіл (видача)'}"
    # Підстави в документі немає — рядок лишається для руки, а не з вигаданою метою.
    ws["E16"] = f"Підстава (мета): {spec.get('basis') or ''}"
    # Одержувач і той, хто передає, — люди (МВО на дату накладної), а не
    # підрозділи; підрозділ-одержувач — у «Приймає», як у бланку служби.
    ws["A17"] = f"Відповідальний одержувач: {_person(rcv) or spec.get('receiver', '')}"
    ws["A18"] = f"Передає: {_person(snd) or spec.get('sender', '')}"
    ws["E18"] = f"Приймає: {spec.get('receiver', '')}"

    # Значення ставляться через `.value =`: openpyxl.cell(..., value=None)
    # нічого не записує, і в порожніх рядках лишалися б дані зі зразка.
    def put(r, c, v):
        ws.cell(row=r, column=c).value = v

    for i in range(max(len(lines), LINE_SLOTS)):
        r = FIRST_LINE + i
        ln = lines[i] if i < len(lines) else None
        put(r, 1, i + 1 if ln else None)
        put(r, 2, ln.get("name") if ln else None)
        put(r, 3, ln.get("code") if ln else None)
        put(r, 4, ln.get("uom") if ln else None)
        put(r, 5, (ln.get("cat") or None) if ln else None)
        put(r, 6, (ln.get("price") or None) if ln else None)
        put(r, 7, ln.get("qty") if ln else None)
        put(r, 8, f"=G{r}" if ln else None)
        put(r, 9, f"=F{r}*H{r}" if ln else None)
        put(r, 10, (ln.get("note") or None) if ln else None)
        ws.cell(row=r, column=6).number_format = FMT_MONEY
        ws.cell(row=r, column=9).number_format = FMT_MONEY

    put(total, 7, f"=SUM(G{FIRST_LINE}:G{last})")
    put(total, 8, f"=G{total}")
    put(total, 9, f"=SUM(I{FIRST_LINE}:I{last})")
    ws.cell(row=total, column=9).number_format = FMT_MONEY

    # «Всього передано» й «на суму» — текстом замість формул надбудови.
    put(32 + extra, 2, spec.get("qty_words") or "")
    put(34 + extra, 2, spec.get("sum_words") or "")
    # Підписанти — з довідників на дату накладної. Кого там немає, того рядок
    # лишається порожнім для підпису від руки, а не з чужим прізвищем зі зразка.
    chief = spec.get("chief") or {}
    if chief.get("pos"):
        put(30 + extra, 1, chief["pos"])
    put(30 + extra, 7, chief.get("name") or None)
    put(38 + extra, 2, snd.get("pos") or spec.get("sender", ""))
    put(38 + extra, 7, snd.get("name") or None)
    put(40 + extra, 2, rcv.get("pos") or spec.get("receiver", ""))
    put(40 + extra, 7, rcv.get("name") or None)
    buh = spec.get("accountant") or {}
    put(53 + extra, 2, f"{buh.get('pos') or '_' * 34}    ______________    {buh['name']}"
        if buh.get("name") else "_" * 34 + "    ______________   " + "_" * 29)
    fes = spec.get("fes") or {}
    if fes.get("pos"):
        # У бланку посада начфіна стоїть двома рядками; ділимо по тире, як там.
        head, sep, tail = fes["pos"].partition("–")
        if not sep:
            head, sep, tail = fes["pos"].partition("-")
        put(57 + extra, 1, (head.strip() + " " + sep).strip() if sep else fes["pos"])
        put(58 + extra, 1, tail.strip() or None)
    put(59 + extra, 2, " " * 72 + "______________" + " " * 25 + (fes.get("name") or "_" * 22))

    ws.print_area = f"A1:J{ws.max_row}"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


def save_invoice(spec: dict, folder: str) -> str:
    os.makedirs(folder, exist_ok=True)
    name = safe_name(spec.get("file") or f"Накладна {spec.get('no', '')}")
    path = os.path.join(folder, name + ".xlsx")
    wb = fill_invoice(spec)
    try:
        wb.save(path)
    except PermissionError:
        path = os.path.join(folder, f"{name} {datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path


# ------------------------------------------------ узагальнююча відомість
# Бланк — відомість, яку служба підписує з підрозділом під час звірки (Додаток 1
# до Інструкції з обліку військового майна, п. 7 розд. II). Узято файл служби як
# є: у ньому 14 рядків таблиці (12–25), рядок «Всього» (26) і підписи нижче. Для
# іншої кількості рядків таблиця розсувається чи стискається разом з
# об'єднаннями, а все, що під нею, їде слідом.
RECON_TEMPLATE = "uzag-vidomist.xlsx"
RECON_FIRST = 12
RECON_SLOTS = 14
RECON_TOTAL = 26
FMT_INT = "0_ "


def _date(iso: str):
    try:
        return datetime.strptime(iso or "", "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


def _dmy(iso: str) -> str:
    d = _date(iso)
    return d.strftime("%d.%m.%Y") if d else (iso or "")


def _qty_cell(cell, v):
    """Кількість: ціле — як у бланку, дробове — з тисячними. Порожнє лишається
    порожнім: «за фінансовим обліком» без числа — це «не звірено з ФЕС», не нуль."""
    cell.value = v
    if v is not None:
        cell.number_format = FMT_INT if float(v).is_integer() else FMT_QTY


def fill_recon(ws, st: dict) -> None:
    """Заповнює аркуш бланка однією відомістю.

    st: no, date, from, to (YYYY-MM-DD), title (підрозділ у родовому відмінку),
    lines [{name, uom, price, fin, acc, fact, note}], signer_pos, signer_name,
    chief_pos, chief_name.
    """
    lines = st.get("lines") or []
    n = max(1, len(lines))
    d = n - RECON_SLOTS
    if d > 0:
        _shift_down(ws, RECON_TOTAL, d)
        for r in range(RECON_FIRST + RECON_SLOTS, RECON_FIRST + n):
            for c in range(1, 11):
                _copy_style(ws.cell(row=RECON_FIRST, column=c), ws.cell(row=r, column=c))
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    elif d < 0:
        _shift_up(ws, RECON_FIRST + RECON_SLOTS, -d)
    last = RECON_FIRST + n - 1
    total = RECON_TOTAL + d

    def put(r, c, v):
        ws.cell(row=r, column=c).value = v

    put(3, 1, f"Узагальнююча відомість  № {st.get('no') or ''}".rstrip())
    put(5, 1, st.get("title") or "")
    # Частина — з реквізитів, а не зі зразка бланка.
    unit = (st.get("unit") or "").strip()
    put(6, 1, f" військової частини {unit}".rstrip())
    put(9, 1, f" {st['legal_name']}" if st.get("legal_name") else None)
    put(8, 5, f"{_dmy(st.get('from'))} по {_dmy(st.get('to'))}")
    when = _date(st.get("date"))
    put(8, 8, when or st.get("date") or "")
    ws.cell(row=8, column=8).number_format = "dd.mm.yyyy"

    for i in range(n):
        r = RECON_FIRST + i
        ln = lines[i] if i < len(lines) else {}
        # Висоту під довгу назву підбирає Excel: висоти зі зразка під інші назви.
        ws.row_dimensions[r].height = None
        put(r, 1, i + 1 if ln else None)
        put(r, 2, ln.get("name"))
        put(r, 5, ln.get("uom"))
        put(r, 6, ln.get("price") or None)
        ws.cell(row=r, column=6).number_format = FMT_MONEY
        _qty_cell(ws.cell(row=r, column=7), ln.get("fin"))
        _qty_cell(ws.cell(row=r, column=8), ln.get("acc"))
        _qty_cell(ws.cell(row=r, column=9), ln.get("fact"))
        put(r, 10, ln.get("note") or None)
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="center",
                                                       wrap_text=True)
        ws.cell(row=r, column=10).alignment = Alignment(vertical="center", wrap_text=True)

    put(total, 1, "Всього")
    put(total, 6, None)
    for c, col in ((7, "G"), (8, "H"), (9, "I")):
        put(total, c, f"=SUM({col}{RECON_FIRST}:{col}{last})")
        # Сума формулою — ціла чи ні, наперед невідомо; «Загальний» не дає «5,».
        ws.cell(row=total, column=c).number_format = "General"

    year = (st.get("to") or "")[:4] or str(datetime.now().year)
    put(28 + d, 1, "Реквізити книги обліку військового майна №14, реєстраційний номер "
                   f"______ за номенклатурою на {year} рік.")
    signer = (st.get("signer_pos") or "").rstrip(", ")
    chief = (st.get("chief_pos") or "").rstrip(", ")
    put(30 + d, 1, signer + ("," if signer else ""))
    put(30 + d, 10, st.get("signer_name") or None)
    put(33 + d, 1, chief + ("," if chief else ""))
    put(33 + d, 10, st.get("chief_name") or None)

    ws.print_area = f"A1:J{35 + d}"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_recon(spec: dict):
    """Книга з відомостями: по аркушу на кожну, у тому порядку, як прийшли."""
    from openpyxl import load_workbook
    wb = clean_workbook(load_workbook(template_path(RECON_TEMPLATE)))
    base = wb.worksheets[0]
    sts = spec.get("statements") or []
    if not sts:
        raise ValueError("немає жодної відомості для вивантаження")
    # Копії робляться з чистого бланка — до того, як заповнено перший аркуш.
    sheets = [base] + [wb.copy_worksheet(base) for _ in sts[1:]]
    used = set()
    for ws, st in zip(sheets, sts):
        ws.title = _sheet_title(st.get("sheet") or st.get("sub") or "Відомість", used)
        fill_recon(ws, st)
    return wb


def _save_wb(wb, folder: str, name: str) -> str:
    os.makedirs(folder, exist_ok=True)
    name = safe_name(name)
    path = os.path.join(folder, name + ".xlsx")
    try:
        wb.save(path)
    except PermissionError:
        path = os.path.join(folder, f"{name} {datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path


def save_recon(spec: dict, folder: str) -> str:
    return _save_wb(build_recon(spec), folder, spec.get("file") or "Узагальнююча відомість")


# ------------------------------------------------ журнал результатів звірки
# Бланк — «Журнал результатів звірки обліку військового майна» (Додаток 9) з
# комплекту електронних форм: титульна сторінка, таблиця з 7 граф і пояснення.
JOURNAL_TEMPLATE = "zhurnal-zvirky.xlsx"
JOURNAL_FIRST = 6


def build_recon_journal(spec: dict):
    """spec: service, unit, rows [[дата, хто звіряє, підрозділ, з ким (посада,
    ПІБ), результати, рішення начальника, примітки]]."""
    from openpyxl import load_workbook
    wb = clean_workbook(load_workbook(template_path(JOURNAL_TEMPLATE)))
    cover = wb["Титульна сторінка"]
    cover["B15"] = spec.get("service") or "продовольча служба"
    cover["B18"] = f"військова частина {spec.get('unit') or ''}".strip()
    ws = wb["Аркуш1"]
    font = Font(name=FONT, size=12)
    rows = spec.get("rows") or []
    for i, row in enumerate(rows):
        r = JOURNAL_FIRST + i
        for c in range(1, 8):
            v = row[c - 1] if c - 1 < len(row) else None
            cell = ws.cell(row=r, column=c)
            if c == 1 and _date(v):
                cell.value = _date(v)
                cell.number_format = "dd.mm.yyyy"
            else:
                cell.value = v if v not in ("", None) else None
            cell.font, cell.border = font, BOX
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if c == 1 else "left")
    last = JOURNAL_FIRST + max(len(rows), 1) - 1
    ws.column_dimensions["E"].width = 30
    ws.print_area = f"A1:G{last}"
    ws.print_title_rows = "4:5"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = wb.index(ws)
    # Бланк зберігся з вибраною титульною сторінкою: разом з активним журналом Excel
    # відкривав їх групою, де правка лягає на обидва аркуші. Вибраний аркуш — один.
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet is ws
    return wb


def save_recon_journal(spec: dict, folder: str) -> str:
    return _save_wb(build_recon_journal(spec), folder,
                    spec.get("file") or "Журнал результатів звірки")


# ------------------------------------------------ інвентаризація
def save_inventory(spec: dict, folder: str) -> str:
    """Інвентаризаційні описи (форма наказу Мінфіну №572) і акт — одна книга."""
    from inventory_export import build_inventory                      # noqa: PLC0415
    return _save_wb(build_inventory(spec), folder, spec.get("file") or "Інвентаризація")
