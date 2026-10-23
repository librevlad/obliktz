# -*- coding: utf-8 -*-
"""Сухпай за фактичними днями: акт списання книги СП → дні й особи за Добовкою (аркуш «Сухпай») →
квартал звіту 1/прод рядка («1прод»). Перенесено з iprod (`src/suhpay_days.py`, `tools/prod1/sp_book_1prod.py`).

Акт видачі покриває кілька днів (видача наперед), і дні можуть лягти в різні квартали: рядок акта
(пайки, вода, підігрівачі) ділиться на дні в частках осіб цього акта за Добовкою. Аркуш «Сухпай»: рядок
блоків (підрозділи), рядок підблоків, рядок полів «осіб / рапорт / акт», далі рядки днів з датою в колонці
«дата». Підблок «усього» — сума інших, не рахується. У клітинці «акт» буває кілька актів
(«СП2026-116,СП2026-104») — особи такого дня діляться між ними за добовими видачами актів. Номер акта з
«/…» (вода до акта, «СП2026-103/1») бере дні акта без «/…». Акт сухпаю програми ділиться за своїми днями
рапорту (sp_acts.day_shares), вода — за кілограмами днів. Акт без днів у Добовці й акти, датовані
раніше за `days_from` (квартали, уже подані за датою акта), лишаються на даті акта; «ЗНИЩЕНО» — «не в
1прод (втрати)»."""
import collections
import datetime
import os
import re
import warnings

from prod import ps_calendar as C

SHEET = "Сухпай"
LOSS_MARK = "ЗНИЩЕНО"
LOSS_TEXT = "не в 1прод (втрати)"
EPS = 1e-9


def _txt(v):
    return "" if v is None else str(v).strip()


def act_base(num):
    """Номер акта без «/…» і пробілів: «СП2026-103/1» -> «СП2026-103»."""
    return act_key(num).split("/")[0]


def act_key(num):
    """Номер акта без пробілів, з «/…»: ключ часток акта сухпаю програми — у нього свої дні рапорту."""
    return re.sub(r"\s+", "", _txt(num))


def split_acts(cell):
    """Клітинка «акт» Добовки -> номери актів (без «/…»): «СП2026-116,СП2026-104» -> (…116, …104)."""
    parts = [act_base(x) for x in re.split(r"[,;\n]+", _txt(cell)) if _txt(x)]
    return tuple(dict.fromkeys(p for p in parts if p))


def newest_dobovka(folder):
    """Найновіша «Добовка дд.мм.xlsx» у теці (за часом збереження) або None."""
    import dobovka  # noqa: PLC0415
    return dobovka.newest(folder)


def dobovka_days(path):
    """Аркуш «Сухпай» Добовки -> (groups {(акти,…): {дата: осіб}}, no_act [(дата, блок, підблок, осіб)]). Аркуш читає
    один читач Добовки програми (`dobovka._sukhpai`)."""
    import dobovka                                               # noqa: PLC0415
    import openpyxl                                              # noqa: PLC0415
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise ValueError(f"у Добовці немає аркуша «{SHEET}»: {os.path.basename(path)}")
        rows = [list(r) for r in wb[SHEET].iter_rows(values_only=True)]
    finally:
        wb.close()
    got = {"sukhpai": {}}
    if not dobovka._sukhpai(rows, got):
        raise ValueError(f"Добовка, аркуш «{SHEET}»: немає рядка полів «осіб / акт»")
    groups = collections.defaultdict(lambda: collections.defaultdict(float))
    no_act = []
    for unit in got["sukhpai"].values():
        for iso, (n, _rep, act) in unit["days"].items():
            if not n:
                continue
            d = datetime.date.fromisoformat(iso)
            acts = split_acts(act)
            if acts:
                groups[acts][d] += float(n)
            else:
                no_act.append((d, unit["block"], unit["sub"], float(n)))
    return {k: dict(v) for k, v in groups.items()}, no_act


def rows_of(model, items=None):
    """Рядки списання книги СП -> [dict(key, date, num, base, unit, code, note, qty, f, loss)]; key — ключ
    документа книги, f — «дободач» позиції (0 — не пайок: вода, підігрівач, тара)."""
    per = {i["code"]: float(i["per_ration"] or 0) for i in (items or model["items"])}
    out = []
    for d in model["docs"]:
        if d["journal"] != "writeoffs":
            continue
        for ln in d["lines"]:
            out.append({"key": (d["journal"], d["date"], d["no"], d["src"], d["dst"]),
                        "date": datetime.date.fromisoformat(d["date"]), "num": d["no"], "base": act_base(d["no"]),
                        "unit": d["src"], "code": ln["code"], "note": ln["note"], "qty": float(ln["qty"]),
                        "f": per.get(ln["code"], 0.0), "loss": LOSS_MARK in d["no"].upper()})
    return out


def act_shares(rows, groups, days_from=None, own=()):
    """{акт без «/…»: {дата: частка}} — частки днів акта (сума 1) за Добовкою. Особи клітинки з кількома
    актами діляться між ними за добовими видачами актів. Без днів у Добовці й датовані раніше за
    `days_from` — не входять (лишаються на даті акта). `own` — номери без «/…» актів сухпаю програми: вони діляться
    за своїми днями рапорту (sp_acts.day_shares), тож слів про дні Добовки для них немає. -> (shares, notes)."""
    dv = collections.defaultdict(float)
    first_date = {}
    for r in rows:
        if r["loss"]:
            continue
        dv[r["base"]] += r["qty"] * r["f"]
        first_date[r["base"]] = min(first_date.get(r["base"], r["date"]), r["date"])
    days = collections.defaultdict(lambda: collections.defaultdict(float))
    notes = []
    own = set(own)
    for acts, by_day in groups.items():
        known = [a for a in acts if dv.get(a, 0.0) > EPS]
        if not known:
            if not own & set(acts):
                notes.append(f"Добовка: акт {', '.join(acts)} — у книзі немає пайків з таким номером "
                             f"({_fmt(sum(by_day.values()))} д/в)")
            continue
        tot = sum(dv[a] for a in known)
        for a in known:
            w = dv[a] / tot
            for d, n in by_day.items():
                days[a][d] += n * w
    shares = {}
    for a, by_day in days.items():
        if days_from and first_date.get(a) and first_date[a] < days_from:
            continue
        s = sum(by_day.values())
        if s > EPS:
            shares[a] = {d: n / s for d, n in sorted(by_day.items())}
            if abs(s - dv.get(a, 0.0)) > 0.01 and a not in own:
                notes.append(f"акт {a}: у книзі {_fmt(dv.get(a, 0.0))} д/в, у Добовці {_fmt(s)} — дні беруться з "
                             "Добовки, кількість — з книги")
    for a in sorted(dv):
        if dv[a] > EPS and a not in shares and a not in own \
                and not (days_from and first_date.get(a) and first_date[a] < days_from):
            notes.append(f"акт {a} ({first_date[a].strftime('%d.%m.%Y')}): днів у Добовці немає — за датою акта")
    return shares, notes


def split_row(r, shares):
    """[(дата, частка)] рядка: акт сухпаю програми — за своїми днями рапорту (рядок води — за кілограмами днів: ключ
    номер і код), акт книги — за днями Добовки акта без «/…», інакше — дата акта цілком."""
    key = act_key(r.get("num") or r["base"])
    sh = shares.get((key, r.get("code"))) or shares.get(key) or shares.get(r["base"])
    if not sh or r["loss"]:
        return [(r["date"], 1.0)]
    return list(sh.items())


def row_label(r, shares):
    """«1прод» рядка: «III кв 2026» або «III кв 2026: 40; IV кв 2026: 20» (одиниці рядка)."""
    if r["loss"]:
        return LOSS_TEXT
    per = collections.OrderedDict()
    for d, part in split_row(r, shares):
        k = C.report_quarter(d)
        per[k] = per.get(k, 0.0) + r["qty"] * part
    if len(per) == 1:
        return C.quarter_label(next(iter(per)))
    return "; ".join(f"{C.quarter_label(k)}: {_fmt(v)}" for k, v in sorted(per.items()))


def _fmt(x):
    s = ("%.3f" % x).rstrip("0").rstrip(".")
    return s.replace(".", ",")


def fill_q1prod(con, dobovka_path=None, days_from=None):
    """«1прод» рядків списання книги СП, яких не задала людина й не принесла книга, — за днями з Добовки
    (без Добовки — за датою акта). -> {"changed": n, "notes": […], "no_act": n}.

    Читання стану і запис — одна транзакція BEGIN IMMEDIATE … COMMIT (ROLLBACK на збої), як у книги й у збереження
    стану: список рядків, який записується цілим, не губить чужої зміни, а збій посеред запису не лишає часткового
    оновлення. Пишеться, лише коли змінилась хоч одна мітка. Викликати без відкритої транзакції: функція відкриває
    свою, тож у сеансі, де транзакція вже відкрита, BEGIN відмовить (SAVEPOINT тут не відкривається)."""
    import book_xlsx as B                                        # noqa: PLC0415
    import state_db                                              # noqa: PLC0415
    groups, no_act = dobovka_days(dobovka_path) if dobovka_path else ({}, [])
    con.execute("BEGIN IMMEDIATE")
    try:
        model = B.model_from_db(con, B.SP)
        rows = rows_of(model)
        import sp_acts                                           # noqa: PLC0415
        program = sp_acts.day_shares(con, days_from)             # акти сухпаю програми — за своїми днями рапорту
        shares, notes = act_shares(rows, groups, days_from,
                                   {act_base(k if isinstance(k, str) else k[0]) for k in program})
        shares.update(program)
        ids = {(j, dt, str(no), s, t): doc for j, dt, no, s, t, doc in state_db.doc_ids(con)}
        extra = state_db._load_line_extra(con)
        changed = 0
        for r in rows:
            doc = ids.get(r["key"])
            if doc is None:
                continue
            rec = extra.setdefault(str(doc), {}).setdefault(f"{r['code']}|{r['note']}", {"best": "", "q1": "", "fixed": False})
            if rec.get("fixed"):
                continue
            label = row_label(r, shares)
            if rec.get("q1") != label:
                rec["q1"] = label
                changed += 1
        if changed:
            state_db._save_line_extra(con, extra)
        con.execute("COMMIT")
    except BaseException:
        con.execute("ROLLBACK")
        raise
    return {"changed": changed, "notes": notes, "no_act": len(no_act)}
