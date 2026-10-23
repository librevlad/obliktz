# -*- coding: utf-8 -*-
"""Акти сухпаю (вид «sp», `sp_acts`) у книзі «Облік СП»: книга несе їх туди й назад, і книга, прийнята в чисту програму,
дає ту саму базу.

Чотири аркуші після «Списання», таблиці з A1: «Акти» — рядок на акт (одержувач, папери, примітка); «Рядки актів» — рядок
на позицію акта (роль: пайок, підігрівач, вода, інше; у пайка — людино-дні й «з підігрівачем»; норма); «Дні актів» —
день рапорту й осіб; «Підписанти актів» — знімок підписантів. Рядок кожного аркуша називає свій акт ключем документа
(«док_дата», «док_номер», «підрозділ» — місце, з якого списано): документ акта — вибуття в «Списання» з тими самими
датою, номером і підрозділом, без «кому». «Списання» не міняється: скрипти служби читають його за позицією колонок.

Похідного аркуші не несуть: вид («sp»), серію («СП»), пункт (немає), період і людино-дні акта прийом рахує з «Дні актів»
так само, як `sp_acts.compute`, а «не за нормою» — 0. Колонка з числом, яке програма однаково перерахує, або мовчки б
ігнорувалась, або сперечалась би з днями.

Прийом пише документи через `state_db._save_docs`: документ зі зміненими рядками видаляється й вставляється знову, і дані
акта йшли б за ним каскадом. Мовчки акт не губиться: книга з аркушами — акт збирається з них на документ з тим самим
ключем; книга без аркушів (стара «Облік СП 6», складена руками) — акт бази лишається за своїм документом (`held`), а
документ, якого книга не несе, забирає свій акт (перегляд каже це: «Прибрані акти»). Кількість акта не за правилами
прийом бере з книги, а «Зауваження» перегляду її називають (`off_rules`; рішення власника 23.10.2026).
"""
import datetime
import json
import math

from openpyxl.utils import get_column_letter

import book_xlsx as B
import op_acts
import sp_acts
import state_db

SHEETS = ("Акти", "Рядки актів", "Дні актів", "Підписанти актів")
TABLES = ("tblActs", "tblActLines", "tblActDays", "tblActSigners")
KEY_COLS = ("док_дата", "док_номер", "підрозділ")
ACT_COLS = KEY_COLS + ("одержувач",) + tuple(f"{p}_{w}" for p in sp_acts.PAPER_ROLES for w in ("номер", "дата")) + (
    "примітка",)
LINE_COLS = KEY_COLS + ("код", "роль", "людино-дні", "з підігрівачем", "норма")
DAY_COLS = KEY_COLS + ("дата", "осіб")
SIGNER_COLS = KEY_COLS + ("роль", "порядок", "посада", "звання", "ім'я_прізвище")
ROLES = ("пайок", "підігрівач", "вода", "інше")
RATION = ROLES[0]
DATE_FORMAT = "dd.mm.yyyy"
NO_SHEETS = ("У книзі немає аркушів актів сухпаю («Акти», «Рядки актів», «Дні актів», «Підписанти актів»): її документи "
             "вибуття приходять без даних актів, а акти, що вже є в програмі, лишаються за своїми документами.")
# Поля рядка `act`, які тримає `held`: акт, документ якого прийом записав наново, стає до нього з тими самими полями.
ACT_FIELDS = ("kind", "series", "point_id", "period_from", "period_to", "rations", "off_norm", "note", "party")


# ---------------------------------------------------------------- база → модель

def model(con):
    """Акти сухпаю бази — частини моделі книги (`book_xlsx.ACT_PARTS`): акти за (дата, номер, місце), рядки акта в
    порядку запису, дні за датою, підписанти за роллю й порядком; запит на таблицю, не на акт. Кожен запис названо
    ключем документа: дата, номер без зайвих пробілів, місце."""
    keys, acts = {}, []
    for r in con.execute("""SELECT a.document_id, a.party, a.note, d.doc_date, d.number, s.name AS sub FROM act a
                              JOIN document d ON d.id = a.document_id JOIN subdivision s ON s.id = d.from_subdivision_id
                             WHERE a.kind = ?""", (sp_acts.KIND,)):
        keys[r["document_id"]] = {"date": r["doc_date"], "no": " ".join(str(r["number"]).split()), "sub": r["sub"]}
        acts.append((r["document_id"], {"party": r["party"] or "", "note": r["note"] or ""}))
    acts.sort(key=lambda a: (keys[a[0]]["date"], keys[a[0]]["no"], keys[a[0]]["sub"], a[0]))
    papers, lines, days, people = {}, {}, {}, {}
    for r in con.execute("SELECT document_id, role, number, doc_date FROM act_paper ORDER BY document_id, role"):
        if r["document_id"] in keys and r["role"] in sp_acts.PAPER_ROLES:
            papers.setdefault(r["document_id"], {})[r["role"]] = {"no": r["number"] or "", "date": r["doc_date"] or ""}
    for r in con.execute("""SELECT i.document_id, n.code, i.role, i.people_days, i.heater, i.norm FROM act_item i
                              JOIN nomen n ON n.id = i.nomen_id ORDER BY i.rowid"""):
        if r["document_id"] in keys:
            ration = r["role"] == RATION
            lines.setdefault(r["document_id"], []).append(
                {"code": str(r["code"]), "role": r["role"] or "", "pdays": r["people_days"],
                 "heater": bool(r["heater"]) if ration else None, "norm": r["norm"]})
    for r in con.execute("SELECT document_id, on_date, persons FROM act_day ORDER BY document_id, on_date"):
        if r["document_id"] in keys:
            days.setdefault(r["document_id"], []).append({"day": r["on_date"], "persons": r["persons"]})
    for r in con.execute("""SELECT document_id, role, sort, position, rank, name FROM act_signer
                             ORDER BY document_id, role, sort"""):
        if r["document_id"] in keys:
            people.setdefault(r["document_id"], []).append(
                {"role": r["role"], "sort": r["sort"], "pos": r["position"] or "", "rank": r["rank"] or "",
                 "name": r["name"] or ""})
    out = {"acts": [], "act_lines": [], "act_days": [], "act_signers": []}
    for doc, act in acts:
        key = keys[doc]
        out["acts"].append({**key, "party": act["party"], "papers": papers.get(doc, {}), "note": act["note"]})
        out["act_lines"] += [{**key, **x} for x in lines.get(doc, [])]
        out["act_days"] += [{**key, **x} for x in days.get(doc, [])]
        out["act_signers"] += [{**key, **x} for x in people.get(doc, [])]
    return out


# ---------------------------------------------------------------- модель → книга

def _day(iso):
    """Дата бази («рррр-мм-дд») — датою клітинки; порожньо — порожня клітинка, інше — як записано (читання назве рядок)."""
    if not iso:
        return None
    try:
        return datetime.date.fromisoformat(iso)
    except ValueError:
        return iso


def _norm_cell(ws, ref, v, put):
    """«норма» рядка акта — як «дободач» на «Номенклатура», але без округлення: 1/3 — формулою «=1/3» зі значенням,
    ціле — цілим, решта — числом як є (округлення зробило б щойно вивантажену книгу «зміненою»)."""
    if v is None:
        return
    v = float(v)
    if abs(v - 1 / 3) < 1e-9:
        put(ws, ref, "=1/3", 1 / 3)
    else:
        ws[ref] = int(v) if v.is_integer() else v


def write(wb, model, put):
    """Чотири аркуші актів сухпаю після журналів (порядок `SHEETS`); частин актів у моделі немає — немає й аркушів.
    `put(ws, ref, формула, значення)` — формула зі значенням для скриптів (`book_xlsx.write_book`)."""
    if model.get("acts") is None:
        return

    def key(r):
        return [_day(r["date"]), r["no"], r["sub"]]

    parts = (
        [key(a) + [a["party"] or None]
         + [cell for p in sp_acts.PAPER_ROLES for cell in ((a["papers"].get(p) or {}).get("no") or None,
                                                          _day((a["papers"].get(p) or {}).get("date")))]
         + [a["note"] or None] for a in model["acts"]],
        [key(r) + [B._code_cell(r["code"]), r["role"] or None, r["pdays"],
                   ("так" if r["heater"] else "ні") if r["role"] == RATION else None, None]
         for r in model["act_lines"]],
        [key(r) + [_day(r["day"]), r["persons"]] for r in model["act_days"]],
        [key(r) + [r["role"], r["sort"], r["pos"] or None, r["rank"] or None, r["name"] or None]
         for r in model["act_signers"]])
    for sheet, table, cols, rows in zip(SHEETS, TABLES, (ACT_COLS, LINE_COLS, DAY_COLS, SIGNER_COLS), parts):
        ws = wb.create_sheet(sheet)
        ws.append(list(cols))
        dates = [n for n, c in enumerate(cols, 1) if c == "дата" or c.endswith("_дата")]
        for n, row in enumerate(rows, 2):
            ws.append(row)
            for c in dates:
                ws.cell(n, c).number_format = DATE_FORMAT
        if sheet == "Рядки актів":
            col = get_column_letter(cols.index("норма") + 1)
            for n, r in enumerate(model["act_lines"], 2):
                _norm_cell(ws, f"{col}{n}", r["norm"], put)
        B._table(ws, table, cols, len(rows))


# ---------------------------------------------------------------- книга → модель

def _book_day(v, where, what, problems):
    """Дата з клітинки аркушів актів: «рррр-мм-дд»; порожньо — ""; не дата чи рік з опискою — у перелік відмов (правило
    дня дободач `state_db._ration_day`, як у книги ОП)."""
    if not B._text(v):
        return ""
    iso = B._date(v)
    try:
        state_db._ration_day(iso)
    except state_db.SaveError:
        problems.append(f"{where}: {what} «{B._text(v)}» — не дата або рік з опискою")
        return ""
    return iso


def _key(r):
    return (r["date"], r["no"], r["sub"])


def _words(key):
    return f"№{key[1]} від {key[0]} ({key[2]})"


def label(key):
    """Акт у переліках вікна прийому: «Акти · <дата> · №<номер> · <місце>» — як акти книги ОП."""
    return f"Акти · {key[0]} · №{key[1]} · {key[2]}"


def _rows(ws, sheet):
    """Непорожні рядки аркуша: (де — «<аркуш>, рядок N», рядок за шапкою)."""
    for n, r in B._rows(ws):
        if any(B._text(v) for v in r.values()):
            yield n, f"{sheet}, рядок {n}", r


def _head(r, where, alias, problems, act=False):
    """Ключ акта з рядка: дата, номер, підрозділ (колишня назва місця з tblRename старої книги — `alias`, як «підрозділ»
    рядків «Списання»)."""
    sub = B._text(r.get("підрозділ"))
    key = {"date": _book_day(r.get("док_дата"), where, "док_дата", problems), "no": B._text(r.get("док_номер")),
           "sub": alias.get(sub, sub)}
    if not all(_key(key)):
        problems.append(f"{where}: потрібні дата, номер і підрозділ" + ("" if act else " акта"))
    return key


def _whole(v, top):
    """Ціле від 1 до `top` з клітинки чи None."""
    n = state_db._whole_number(v)
    return n if n is not None and 0 < n <= top else None


def _read_acts(ws, alias, problems):
    out, seen = [], {}
    for n, where, r in _rows(ws, "Акти"):
        a = {**_head(r, where, alias, problems, act=True), "party": B._text(r.get("одержувач")), "papers": {},
             "note": B._text(r.get("примітка")), "row": where}
        for p in sp_acts.PAPER_ROLES:
            rec = {"no": B._text(r.get(f"{p}_номер")), "date": _book_day(r.get(f"{p}_дата"), where, f"{p}_дата", problems)}
            if rec["no"] or rec["date"]:
                a["papers"][p] = rec
        if _key(a) in seen:
            problems.append(f"{where}: акт {_words(_key(a))} уже є в рядку {seen[_key(a)]}")
        seen.setdefault(_key(a), n)
        out.append(a)
    return out


def _read_lines(ws, alias, problems):
    out, seen = [], {}
    for n, where, r in _rows(ws, "Рядки актів"):
        x = {**_head(r, where, alias, problems), "code": B._text(r.get("код")), "role": B._text(r.get("роль")).lower(),
             "pdays": None, "heater": None, "norm": None, "row": where}
        if not x["code"]:
            problems.append(f"{where}: потрібен код")
        pdays, heater = r.get("людино-дні"), B._text(r.get("з підігрівачем")).lower()
        if x["role"] not in ROLES:
            problems.append(f"{where}: роль — пайок, підігрівач, вода чи інше")
        elif x["role"] == RATION:
            x["pdays"] = _whole(pdays, state_db.RATIONS_MAX)
            if x["pdays"] is None:
                problems.append(f"{where}: людино-дні пайка — ціле число більше нуля")
            if heater not in ("так", "ні", ""):
                problems.append(f"{where}: «з підігрівачем» — «так» чи «ні»")
            x["heater"] = heater == "так"
        elif B._text(pdays) or heater:
            problems.append(f"{where}: людино-дні й «з підігрівачем» — лише в пайка")
        if B._text(r.get("норма")):
            x["norm"] = B._number(r.get("норма"))
            if x["norm"] is None or not math.isfinite(x["norm"]) or x["norm"] <= 0:
                problems.append(f"{where}: норма — число більше нуля")
        key = (*_key(x), x["code"])
        if x["code"] and key in seen:
            problems.append(f"{where}: код {x['code']} акта №{x['no']} уже є в рядку {seen[key]}")
        seen.setdefault(key, n)
        out.append(x)
    return out


def _read_days(ws, alias, problems):
    out, seen = [], {}
    for n, where, r in _rows(ws, "Дні актів"):
        x = {**_head(r, where, alias, problems), "day": "", "persons": _whole(r.get("осіб"), state_db.RATIONS_MAX),
             "row": where}
        if not B._text(r.get("дата")):
            problems.append(f"{where}: немає дати")
        else:
            x["day"] = _book_day(r.get("дата"), where, "дата", problems)
        if x["persons"] is None:
            problems.append(f"{where}: осіб — ціле число більше нуля")
        key = (*_key(x), x["day"])
        if x["day"] and key in seen:
            problems.append(f"{where}: день {x['day']} акта №{x['no']} уже є в рядку {seen[key]}")
        seen.setdefault(key, n)
        out.append(x)
    return out


def _read_signers(ws, alias, problems):
    """Як «Підписанти актів» книги ОП (`op_book._read_signers`); підписант без посади, звання й імені — відмова, а не
    рядок, який запис акта мовчки відкинув би (`op_acts._clean_signers`)."""
    out, seen = [], {}
    for n, where, r in _rows(ws, "Підписанти актів"):
        sort = state_db._whole_number(r.get("порядок")) if B._text(r.get("порядок")) else 0
        if sort is None or not 0 <= sort <= 99:
            problems.append(f"{where}: «порядок» — ціле число від 0 до 99")
            sort = 0
        s = {**_head(r, where, alias, problems), "role": B._text(r.get("роль")), "sort": sort,
             "pos": B._text(r.get("посада")), "rank": B._text(r.get("звання")), "name": B._text(r.get("ім'я_прізвище")),
             "row": where}
        if not s["role"]:
            problems.append(f"{where}: потрібна роль")
        if not (s["pos"] or s["rank"] or s["name"]):
            problems.append(f"{where}: потрібні посада, звання чи ім'я")
        key = (*_key(s), s["role"], s["sort"])
        if key in seen:
            problems.append(f"{where}: «{s['role']}» з порядком {s['sort']} акта №{s['no']} уже є в рядку {seen[key]}")
        seen.setdefault(key, n)
        out.append(s)
    return out


def _writeoffs(docs):
    """Документи вибуття без «кому» за ключем акта (номер — без зайвих пробілів): {ключ: {код: кількість}} у порядку
    рядків «Списання»."""
    out = {}
    for d in docs:
        if d["journal"] == "writeoffs" and not d["dst"]:
            q = out.setdefault((d["date"], " ".join(str(d["no"]).split()), d["src"]), {})
            for ln in d["lines"]:
                q[ln["code"]] = q.get(ln["code"], 0) + float(ln["qty"] or 0)
    return out


def _check(acts, lines, days, signers, docs, problems):
    """Аркуші між собою й зі «Списання» книги: рядок — свого акта, акт — свого документа й днів; кожен код документа акта
    має свій рядок на «Рядки актів» і навпаки (вікно акта збирає акт за ролями рядків); людино-дні пайків — сума осіб."""
    docs = _writeoffs(docs)
    by_act = {}
    for a in acts:
        by_act.setdefault(_key(a), a)
    for part in (lines, days, signers):
        for x in part:
            if all(_key(x)) and _key(x) not in by_act:
                problems.append(f"{x['row']}: немає акта {_words(_key(x))} на аркуші «Акти»")
    mine, dates = {}, {}
    for x in lines:
        mine.setdefault(_key(x), []).append(x)
    for x in days:
        dates.setdefault(_key(x), []).append(x)
    for k, a in by_act.items():
        if not all(k):
            continue
        codes = docs.get(k)
        if codes is None:
            problems.append(f"{a['row']}: немає документа вибуття {_words(k)} на аркуші «Списання»")
        else:
            for x in mine.get(k, []):
                if x["code"] and x["code"] not in codes:
                    problems.append(f"{x['row']}: коду {x['code']} немає в документі акта {_words(k)} на аркуші «Списання»")
            have = {x["code"] for x in mine.get(k, [])}
            problems += [f"Рядки актів: в акта {_words(k)} немає рядка коду {c} зі «Списання»"
                         for c in codes if c not in have]
        if not dates.get(k):
            problems.append(f"{a['row']}: немає днів рапорту на аркуші «Дні актів»")
            continue
        rations = [x for x in mine.get(k, []) if x["role"] == RATION]
        if rations and all(x["pdays"] for x in rations) and all(x["persons"] for x in dates[k]):
            got, want = sum(x["pdays"] for x in rations), sum(x["persons"] for x in dates[k])
            if got != want:
                problems.append(f"Рядки актів: людино-дні пайків акта {_words(k)} — {got}, а в «Дні актів» — {want}")


def read(wb, docs, alias, problems):
    """Чотири частини актів з книги (`book_xlsx.ACT_PARTS`) — лише дані, без номерів рядків; аркушів актів немає зовсім —
    None (стара «Облік СП 6», книга, складена руками), є не всі — відмова з назвою відсутнього аркуша. `docs` — документи
    книги (ключ документа акта — дата, номер, «підрозділ» вибуття без «кому»), `alias` — tblRename старої книги.
    Відмови — у `problems` з аркушем і рядком."""
    have = [s for s in SHEETS if s in wb.sheetnames]
    if not have:
        return None
    if len(have) < len(SHEETS):
        problems += [f"Немає аркуша «{s}»: аркуші актів сухпаю книга несе всі чотири або жодного"
                     for s in SHEETS if s not in have]
        return None
    acts = _read_acts(wb["Акти"], alias, problems)
    lines = _read_lines(wb["Рядки актів"], alias, problems)
    days = _read_days(wb["Дні актів"], alias, problems)
    signers = _read_signers(wb["Підписанти актів"], alias, problems)
    _check(acts, lines, days, signers, docs, problems)

    def bare(rows):
        return [{k: v for k, v in r.items() if k != "row"} for r in rows]
    return {"acts": bare(acts), "act_lines": bare(lines), "act_days": bare(days), "act_signers": bare(signers)}


def _qty(v):
    return B._comma(float(v))


def off_rules(model, keys=None):
    """Рядки «Зауваження» про акти книги (`keys` — лише ці ключі), кількість яких у «Списання» не за правилами акта сухпаю
    (`sp_acts.by_rules` — ті самі, що в складання акта): пайків не ROUNDUP(людино-дні ÷ дободач), підігрівачів не по
    одному на пайок з підігрівачем, води не за нормою дат. Прийом бере кількість з книги без позначки «не за нормою»
    (рішення власника 23.10.2026); рядок називає акт, рядок акта, кількість у книзі й за правилами. Підігрівачі й вода —
    разом за роллю: програма бере їх з кількох кодів тієї самої позиції; ролі, якої в акті немає (не обрали підігрівача
    чи води), правило не додає."""
    per = {i["code"]: i["per_ration"] for i in model["items"]}
    qty = _writeoffs(model["docs"])
    lines, days = {}, {}
    for x in model["act_lines"]:
        lines.setdefault(_key(x), []).append(x)
    for x in model["act_days"]:
        days.setdefault(_key(x), []).append([x["day"], x["persons"]])
    out = []
    for a in model["acts"]:
        k = _key(a)
        if keys is not None and k not in keys:
            continue
        have, said = qty.get(k, {}), []
        rations = [x for x in lines.get(k, []) if x["role"] == RATION]
        pers = [sp_acts.per_ration(per.get(x["code"])) for x in rations]
        each, heated, water = sp_acts.by_rules([(p, x["pdays"], x["heater"]) for p, x in zip(pers, rations)
                                                if p is not None], sorted(days.get(k, [])))
        each = iter(each)
        for p, x in zip(pers, rations):
            got = have.get(x["code"], 0)
            if p is None:
                said.append(f"Пайок {x['code']}: у книзі {_qty(got)}, у позиції немає дободача")
                continue
            want = next(each)
            if abs(got - want) > 1e-9:
                said.append(f"Пайок {x['code']}: у книзі {_qty(got)}, за правилами {want}")
        # Пайок без дободача — підігрівачів за правилом не порахувати: про сам пайок рядок уже є.
        for role, want in (("підігрівач", None if None in pers else heated), ("вода", water / 1000)):
            codes = [x["code"] for x in lines.get(k, []) if x["role"] == role]
            got = sum(have.get(c, 0) for c in codes)
            if codes and want is not None and abs(got - want) > 1e-9:
                said.append(f"{role.capitalize()} {', '.join(codes)}: у книзі {_qty(got)}, за правилами {_qty(want)}")
        if said:
            out.append(f"{label(k)}: кількість не за правилами, прийом бере її з книги. " + " ".join(s + "." for s in said))
    return out


def remarks(have, model):
    """«Зауваження» вікна прийому про кількість не за правилами (`off_rules`) — лише для актів книги, нових чи змінених
    проти бази: дані акта (`_sigs`) чи кількість його документа в «Списання». Акт, уже прийнятий таким, — не новина:
    своя книга програми зауважень не дає. `have` — модель бази, `model` — книги з назвами місць бази."""
    if model.get("acts") is None:
        return []
    old, new = _sigs(have), _sigs(model)
    was, now = _writeoffs(have["docs"]), _writeoffs(model["docs"])
    return off_rules(model, {k for k in new if new[k] != old.get(k) or now.get(k) != was.get(k)})


# ---------------------------------------------------------------- відмінності й прийом

def _sigs(m):
    """{ключ акта: підпис} — одержувач, примітка, папери за роллю, рядки (код, роль, людино-дні, з підігрівачем, норма) у
    порядку запису, дні за датою, підписанти за роллю й порядком."""
    acts = {_key(a): ((a["party"], a["note"], tuple(sorted((p, r.get("no") or "", r.get("date") or "")
                                                            for p, r in a["papers"].items()))), [], [], [])
            for a in m.get("acts") or []}
    for part, n, row in (("act_lines", 1, lambda x: (x["code"], x["role"], x["pdays"], x["heater"],
                                                       None if x["norm"] is None else float(x["norm"]))),
                         ("act_days", 2, lambda x: (x["day"], x["persons"])),
                         ("act_signers", 3, lambda x: (x["role"], x["sort"], x["pos"], x["rank"], x["name"]))):
        for x in m.get(part) or []:
            if _key(x) in acts:
                acts[_key(x)][n].append(row(x))
    return {k: (*head, tuple(ln), tuple(sorted(d)), tuple(sorted(s))) for k, (head, ln, d, s) in acts.items()}


HOW_TO_CHANGE = ("Книга без аркушів актів рядків акта не змінює: змініть акт у вікні акта («Сухпай» → «Акти списання») "
                 "або на аркушах актів книги з програми.")


def _refuse_changed_lines(have, kept):
    """Книга без аркушів актів рядків акта не змінює (`_restore`): документ акта, що в книзі дістав код, якого акт не має,
    чи втратив пайок, лишив би акт, якого книга програми назад уже не прочитала б (рядок на кожен код документа,
    людино-дні пайків — сума осіб). Відмова (`BookError`) називає акт — ще в перегляді. `kept` — документи вибуття книги
    за ключем акта (`_writeoffs`)."""
    lines = {}
    for x in have["act_lines"]:
        lines.setdefault(_key(x), []).append(x)
    problems = []
    for k in sorted(_sigs(have)):
        codes = kept.get(k)
        if codes is None:                         # документа книга не несе — акт іде разом з ним
            continue
        mine = lines.get(k, [])
        own = {x["code"] for x in mine}
        more = [c for c in codes if c not in own]
        less = [x["code"] for x in mine if x["role"] == RATION and x["code"] not in codes]
        if more:
            what = f"код {more[0]}, якого" if len(more) == 1 else f"коди {', '.join(more)}, яких"
            problems.append(f"{label(k)}: у документі акта є {what} немає в акті. {HOW_TO_CHANGE}")
        if less:
            what = f"пайка {less[0]}" if len(less) == 1 else f"пайків {', '.join(less)}"
            problems.append(f"{label(k)}: у документі акта немає {what}. {HOW_TO_CHANGE}")
    if problems:
        raise B.BookError(problems)


def diff(have, model):
    """Акти сухпаю, що зміняться з прийомом книги: нові, змінені й прибрані — за ключем акта, кожен перелік за ключем.
    `have` — модель бази (`book_xlsx.model_from_db`), `model` — книги з назвами місць, як їх знає база. Книга без аркушів
    актів нових і змінених не дає, а прибрані — акти бази, документа яких (вибуття без «кому» з тим самим ключем) у книзі
    немає: вони підуть разом з документом, і вікно прийому це каже; документ акта, якому така книга додала код чи
    прибрала пайок, — відмова (`_refuse_changed_lines`)."""
    old = _sigs(have)
    if model.get("acts") is None:
        kept = _writeoffs(model["docs"])
        _refuse_changed_lines(have, kept)
        return {"new": [], "changed": [], "removed": [label(k) for k in sorted(old) if k not in kept]}
    new = _sigs(model)
    return {"new": [label(k) for k in sorted(new.keys() - old.keys())],
            "changed": [label(k) for k in sorted(new.keys() & old.keys()) if new[k] != old[k]],
            "removed": [label(k) for k in sorted(old.keys() - new.keys())]}


def canonical(con):
    """Порівнювана проєкція актів сухпаю бази (`book_xlsx.canonical`)."""
    return sorted(_sigs(model(con)).items(), key=repr)


def held(con):
    """Акти сухпаю бази за ключем документа (дата, номер, місце) — до того, як прийом перепише документи: документ, рядки
    якого книга змінила, записується наново, і дані акта йшли б разом зі старим (як `op_book._held_acts`)."""
    out = {}
    for r in con.execute("""SELECT a.*, d.doc_date, d.number, s.name AS sub FROM act a
                              JOIN document d ON d.id = a.document_id JOIN subdivision s ON s.id = d.from_subdivision_id
                             WHERE a.kind = ?""", (sp_acts.KIND,)):
        doc = r["document_id"]

        def rows(sql, doc=doc):
            return [tuple(x) for x in con.execute(sql, (doc,))]
        out[(r["doc_date"], " ".join(str(r["number"]).split()), r["sub"])] = {
            "doc": doc, "act": {k: r[k] for k in ACT_FIELDS},
            "papers": rows("SELECT role, number, doc_date FROM act_paper WHERE document_id = ?"),
            "signers": rows("SELECT role, sort, position, rank, name FROM act_signer WHERE document_id = ? "
                            "ORDER BY role, sort"),
            "items": rows("""SELECT n.code, i.norm, i.role, i.people_days, i.heater FROM act_item i
                               JOIN nomen n ON n.id = i.nomen_id WHERE i.document_id = ? ORDER BY i.rowid"""),
            "days": rows("SELECT on_date, persons FROM act_day WHERE document_id = ? ORDER BY on_date")}
    return out


def _restore(con, held, ids):
    """Книга без аркушів актів: акт бази лишається за своїм документом. Документа з ключем акта книга не несе — акт пішов з
    ним; документ той самий — акт на місці; документ записано наново — акт стає до нього з тими самими даними, а рядки
    акта — лише ті, коди яких є в рядках нового документа (як гілка без аркуша «Акти» в `op_book._apply_acts`)."""
    alive = {r[0] for r in con.execute("SELECT document_id FROM act")}
    for key, old in held.items():
        doc = ids.get(key)
        if doc is None or doc in alive:
            continue
        con.execute(f"INSERT INTO act(document_id, {', '.join(ACT_FIELDS)}) VALUES(?{', ?' * len(ACT_FIELDS)})",
                    (doc, *(old["act"][k] for k in ACT_FIELDS)))
        con.executemany("INSERT INTO act_paper(document_id, role, number, doc_date) VALUES(?, ?, ?, ?)",
                        [(doc, *p) for p in old["papers"]])
        con.executemany("INSERT INTO act_signer(document_id, role, sort, position, rank, name) VALUES(?, ?, ?, ?, ?, ?)",
                        [(doc, *s) for s in old["signers"]])
        codes = {r[0] for r in con.execute("SELECT n.code FROM document_line l JOIN nomen n ON n.id = l.nomen_id "
                                           "WHERE l.document_id = ?", (doc,))}
        con.executemany("""INSERT INTO act_item(document_id, nomen_id, norm, role, people_days, heater)
                           VALUES(?, (SELECT id FROM nomen WHERE code = ?), ?, ?, ?, ?)""",
                        [(doc, *i) for i in old["items"] if i[0] in codes])
        con.executemany("INSERT INTO act_day(document_id, on_date, persons) VALUES(?, ?, ?)",
                        [(doc, *d) for d in old["days"]])


def _settle_q1(con):
    """«1прод» рядка документа акта сухпаю, що дорівнює розрахунку за днями акта (`sp_acts.q1_label`), — знову програмне
    (`q1prod_fixed` = 0): книга несе «1прод» як вписане людиною, і виправлення акта після кругу вже не перераховувало б
    кварталів. Інше вписане лишається вписаним."""
    acts = {r[0]: (" ".join(str(r[1]).split()), r[2]) for r in con.execute(
        "SELECT a.document_id, d.number, d.doc_date FROM act a JOIN document d ON d.id = a.document_id WHERE a.kind = ?",
        (sp_acts.KIND,))}
    days, water, qty = {}, {}, {}
    for doc, on, n in con.execute("SELECT document_id, on_date, persons FROM act_day ORDER BY document_id, on_date"):
        if doc in acts:
            days.setdefault(doc, []).append([on, n])
    for doc, code in con.execute("SELECT i.document_id, n.code FROM act_item i JOIN nomen n ON n.id = i.nomen_id "
                                 "WHERE i.role = 'вода'"):
        water.setdefault(doc, set()).add(code)
    for doc, code, q in con.execute("SELECT l.document_id, n.code, SUM(l.qty_milli) FROM document_line l "
                                    "JOIN nomen n ON n.id = l.nomen_id GROUP BY l.document_id, n.code"):
        if doc in acts:
            qty[(doc, code)] = q
    for doc, code, note, q1 in con.execute("SELECT document_id, code, note, q1prod FROM doc_line_extra "
                                           "WHERE q1prod_fixed = 1").fetchall():
        if doc in days and (doc, code) in qty and q1 == sp_acts.q1_label(
                *acts[doc], days[doc], qty[(doc, code)] / 1000, code, code in water.get(doc, ())):
            con.execute("UPDATE doc_line_extra SET q1prod_fixed = 0 WHERE document_id = ? AND code = ? AND note = ?",
                        (doc, code, note))


def apply(con, model, held):
    """Акти сухпаю з книги: аркуші є — акти виду «sp» стають такими, як у книзі (акт — до документа з тим самим ключем,
    дані — op_acts._write_act, дні — act_day), решти актів сухпаю немає; аркушів немає — акти бази лишаються за своїми
    документами (_restore). Далі — «1прод» рядків актів (_settle_q1). Відмова — BookError, транзакцію відкочує
    book_xlsx.apply."""
    problems = []
    mine = B._book_doc_ids(con, "СП")
    ids = {(d, " ".join(str(no).split()), src): i for j, d, no, src, dst, i in state_db.doc_ids(con)
           if j == "writeoffs" and not dst and i in mine}
    if model.get("acts") is None:
        _restore(con, held, ids)
    else:
        by = {p: {} for p in B.ACT_PARTS[1:]}
        for p in by:
            for r in model[p]:
                by[p].setdefault(_key(r), []).append(r)
        written = []
        for a in model["acts"]:
            k, doc = _key(a), ids.get(_key(a))
            if doc is None:                       # читання вже звірило акт з «Списання» книги — лише на всяк випадок
                problems.append(f"{label(k)}: документа акта в базі немає")
                continue
            try:                                  # SaveError — словами з ключем акта, а не 500
                days = sp_acts.days_of([[d["day"], d["persons"]] for d in by["act_days"].get(k, [])])
                if not days:
                    raise state_db.SaveError("немає днів рапорту.")
                op_acts._write_act(con, doc, {"kind": sp_acts.KIND, "paperRoles": sp_acts.PAPER_ROLES,
                                              "series": sp_acts.SERIES, "point": None, "from": days[0][0],
                                              "to": days[-1][0], "rations": sum(n for _, n in days),
                                              "party": a["party"], "lines": by["act_lines"].get(k, [])},
                                   {"papers": a["papers"], "note": a["note"], "signers": by["act_signers"].get(k, [])})
            except state_db.SaveError as e:
                problems.append(f"{label(k)}: {e}")
                continue
            con.execute("DELETE FROM act_day WHERE document_id = ?", (doc,))
            con.executemany("INSERT INTO act_day(document_id, on_date, persons) VALUES(?, ?, ?)",
                            [(doc, d, n) for d, n in days])
            written.append(doc)
        con.execute("DELETE FROM act WHERE kind = ? AND document_id NOT IN (SELECT value FROM json_each(?))",
                    (sp_acts.KIND, json.dumps(written)))
    _settle_q1(con)
    if problems:
        raise B.BookError(problems)
