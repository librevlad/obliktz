# -*- coding: utf-8 -*-
"""Пакет паперів акта сухпаю в теці справи: «<номер>.xlsx» — накладна (вимога, Додаток 25 Інструкції 440), роздавальна
відомість (Додаток 57) і акт списання (Додаток 10 до наказу №3 від 02.01.2024). Тека — <тека актів сухпаю>\\<рік>\\
<N> квартал\\<номер> від <дд.мм.рррр> - <підрозділ>, як у справах служби; якщо вже є тека, назва якої починається з
номера, — вона.

Бланк — desktop/templates/sp-akt.xlsx («Накладна» й «Акт» — ті самі, що в акта ОП/МЗ; заповнює їх той самий код
`op_act_export`). Змінне — з бази: дані акта, рядки документа, знімок підписантів, реквізити частини. Прізвища одержувачів
у роздавальній вписують від руки з рапорту — програма їх не знає; рядків стільки, скільки найбільше осіб за день рапорту.
Програма переписує лише свій файл і лише таким, яким сама його записала (відбиток у файлах акта, як у пакета акта ОП/МЗ):
правлений руками (скажімо, роздавальна з уписаними прізвищами) чи складений руками під тим самим іменем —
op_act_export.PackageConflict, переписати можна лише явною згодою (`force`). Витяг, скан рапорту й інше в теці програма не
чіпає. Відкритий в Excel файл — запис поруч із часом в імені.

Як у паперах служби (рішення власника 23.10.2026): рядки різних одиниць підсумовуються («всього» акта, «Всього
передано: … одиниці» накладної), «Постачальник» роздавальної — частина з реквізитів, «Відповідальний одержувач»
накладної — посада; акт — на двох сторінках."""
import datetime as dt
import os
import re

import excel_export
import op_act_export as X
import sp_acts
import stable_hash
from db.uk_number import money_words, qty_words

TEMPLATE = "sp-akt.xlsx"
SCAN_FOLDER = "spActs"
RZ_FIRST, RZ_SLOTS, RZ_TOTAL = 24, 136, 160
RZ_COLS = "DEFGHIJKLMNOPQR"                          # 15 граф найменувань роздавальної
NONE = "---"
ACT_PAGE2, ACT_LAST, ACT_COLS = 36, 100, 52          # друга сторінка акта: з рядка 36 (розрив бланка) до 100; графи A:AZ
CONCLUSION_ROW = 76                                  # висновок комісії (без зсуву таблиці)


def act_dir(root, a):
    """Тека акта: <рік>\\<N> квартал\\<номер> від <дата> - <підрозділ>; наявна тека з тим самим номером (без різниці
    регістру, як її бачить Windows) — вона."""
    base = os.path.join(root, a["date"][:4], f"{(int(a['date'][5:7]) + 2) // 3} квартал")
    if os.path.isdir(base):
        pat = re.compile(rf"^{re.escape(excel_export.safe_name(a['number']))}(?:$|[\s\-_.])", re.IGNORECASE)
        for name in sorted(os.listdir(base)):
            if pat.match(name) and os.path.isdir(os.path.join(base, name)):
                return os.path.join(base, name)
    tail = f" - {a['label']}" if a["label"] else ""
    return os.path.join(base, excel_export.safe_name(f"{a['number']} від {X._dmy(a['date'])}{tail}"))


def conflicts(con, root, a):
    """Файл пакета («<номер>.xlsx») у теці акта, якого програма таким не писала: правлений руками після неї чи складений
    руками під тим самим іменем. `a` — акт чи запит вікна (number, date, label). Немає файла — не конфлікт."""
    number = " ".join(str(a.get("number") or "").split())
    name = excel_export.safe_name(number) + ".xlsx"
    path = os.path.join(act_dir(root, dict(a, number=number)), name)
    if os.path.isfile(path) and not set(stable_hash.shas(path)) & X._own_shas(con, name, SCAN_FOLDER):
        return [name]
    return []


def _paper(a, role):
    rec = a["papers"].get(role) or {}
    return " ".join(str(rec.get("no") or "").split()).lstrip("№").strip(), X._dmy(rec.get("date") or "")


def _order(a):
    """Номер і дата витягу з наказу для паперів: пункт наказу («301 п.18») — лише в книзі СП, у паперах — №301."""
    no, day = _paper(a, "витяг")
    return re.split(r"\s+п\.?\s*\d", no)[0], day


def _ground(a):
    no, day = _order(a)
    return f"наказ №{no} від {day}" if no else ""


def _units(total_milli):
    """«одиниця» за числом, як у «Всього передано» паперів служби: 21 одиниця, 63 одиниці, 65 одиниць; дріб — одиниці."""
    whole, milli = divmod(int(total_milli), 1000)
    if milli:
        return "одиниці"
    if whole % 10 == 1 and whole % 100 != 11:
        return "одиниця"
    return "одиниці" if whole % 10 in (2, 3, 4) and whole % 100 not in (12, 13, 14) else "одиниць"


def _invoice(ws, a, req, sig):
    no, day = _paper(a, "накладна")
    units = {l["uom"] for l in a["lines"]}
    total = sum(l["qty_milli"] for l in a["lines"])
    rcv = X._who(sig, "одержувач")
    excel_export.fill_invoice_sheet(ws, {
        "no": no or a["number"], "date": a["papers"].get("накладна", {}).get("date") or a["date"],
        "operation": "Розподіл (видача)", "basis": _ground(a), "sender": a["sub"], "receiver": a["party"] or a["label"],
        "place": req["place"],
        "lines": [{"name": l["name"], "code": l["fes"], "uom": l["uom"], "cat": "", "price": l["price"], "qty": l["qty"],
                   "note": ""} for l in a["lines"]],
        # Різні одиниці (комплекти, штуки, кілограми) — сумою «одиниць», як у накладній служби.
        "qty_words": f"{qty_words(total)} {units.pop() if len(units) == 1 else _units(total)}" if a["lines"] else "",
        "sum_words": money_words(sum(round(l["sum"] * 100) for l in a["lines"])),
        "sender_person": X._who(sig, "МВО"), "receiver_person": rcv,
        "chief": X._who(sig, "начальник служби"), "accountant": {}, "fes": X._who(sig, "бухгалтер"),
        "unit": {k: req[k] for k in ("legalName", "edrpou", "serviceFull", "validDays")}})
    ws["A17"] = f"Відповідальний одержувач: {rcv['pos'] or a['party'] or a['label']}"      # посада, як у службі


def _rozdavalna(ws, a, req, sig):
    no, day = _paper(a, "роздавальна")
    when = dt.datetime.fromisoformat(a["papers"].get("роздавальна", {}).get("date") or a["date"])
    ws["A7"] = f"{ws['A7'].value or ''}{no or a['number']}"
    ws["J8"] = ws["G11"] = ws["T11"] = when
    # Постачальник — частина з реквізитів («А0000», як у роздавальній служби); реквізитів немає — місце акта.
    ws["E11"], ws["I11"], ws["M11"] = no or a["number"], _ground(a) or None, req["code"] or a["sub"]
    ws["P11"] = f"{a['party'] or a['label']} (о/с)"
    lines = a["lines"][:len(RZ_COLS)]
    for i, col in enumerate(RZ_COLS):
        l = lines[i] if i < len(lines) else None
        for row, v in ((14, l and l["name"]), (16, l and (l["fes"] or NONE)), (18, l and l["uom"]), (20, NONE),
                       (22, l and l["price"])):
            ws[f"{col}{row}"] = v if l else NONE
        ws[f"{col}{RZ_TOTAL}"] = l["qty"] if l else None
        if l:
            ws[f"{col}{RZ_TOTAL}"].number_format = X._qfmt([l["qty"]])
    people = max([n for _d, n, _o in a["days"]] or [1])
    shown = min(max(people, 1), RZ_SLOTS)
    X._fit(ws, RZ_FIRST, RZ_SLOTS, shown)
    for r in range(RZ_FIRST, RZ_FIRST + shown):
        for i, col in enumerate(RZ_COLS):
            if i >= len(lines):
                ws[f"{col}{r}"] = NONE
        ws[f"T{r}"] = when
    total = sum(l["qty_milli"] for l in lines)
    words = qty_words(total)
    ws[f"S{RZ_TOTAL}"] = words[:1].lower() + words[1:]
    rcv, chief = X._who(sig, "одержувач"), X._who(sig, "начальник служби")
    ws["A162"], ws["E162"], ws["P162"] = rcv["pos"] or None, rcv["rank"] or None, rcv["name"] or None
    ws["A164"], ws["E164"], ws["P164"] = chief["pos"] or None, chief["rank"] or None, chief["name"] or None
    ws["A168"] = when


def _act(ws, a, req, sig):
    unit = X._unit(req)
    nak_no, nak_day = _paper(a, "накладна")
    rz_no, rz_day = _paper(a, "роздавальна")
    rep_no, rep_day = _paper(a, "рапорт")
    order, order_day = _order(a)
    rep = f", вх. №{rep_no}" if rep_no else ""
    att = [f"НАКЛАДНА (вимога) №{nak_no or a['number']} від {nak_day or X._dmy(a['date'])} року;",
           f"РОЗДАВАЛЬНА (ЗДАВАЛЬНА) ВІДОМІСТЬ №{rz_no or a['number']} від {rz_day or X._dmy(a['date'])} року;",
           f"РАПОРТ командира {a['party'] or a['label']} {unit}{rep}" + (f" від {rep_day} року;" if rep_day else ";"),
           f"ВИТЯГ ІЗ НАКАЗУ командира {unit} (по стройовій частині) №{order or '____'} від {order_day or '____'} року."]
    conclusion = (f"Висновок комісії: на підставі наказу по стройовій частині №{order or '____'} від "
                  f"{order_day or '____'} року — списати з обліку {unit}.")
    # Кількість акта сухпаю вже за правилами (пайки — людино-дні пайка ÷ дободач, підігрівачі — до пайків з підігрівачем,
    # вода — особи × норма дати, «інше» — вписане): «належить за нормою» — списане, «економії» немає. Пул норми на д/д
    # (op_norms.pool_key) — лише в актів ОП/МЗ: рядок без пулу аркуш «Акт» так і пише (op_act_export._shares).
    X._act(ws, dict(a, lines=[dict(l, pool=None) for l in a["lines"]]), req, sig, att=att, conclusion=conclusion,
           ground=_ground(a) or "Списання", sum_mixed=True)
    extra = max(0, len(a["lines"]) - X.ACT_SLOTS)             # на стільки рядків таблиця розсунула бланк (X._grow)
    _fit_rows(ws, extra)
    _two_pages(ws, ACT_LAST + extra)


def _fit_rows(ws, extra):
    """Висновок комісії й рядок МВО — заввишки за своїм текстом: бланк дає їм по два-три рядки, а в акта сухпаю вони
    здебільшого в один (як у паперах служби) — друга сторінка не губить місця на порожніх рядках."""
    for row in (CONCLUSION_ROW + extra, X.MVO_ROW + extra):
        cell = ws.cell(row=row, column=1)
        rng = next((m for m in ws.merged_cells.ranges if cell.coordinate in m), None)
        if not cell.value or rng is None:
            continue
        width = sum(excel_export._col_width(ws, c) for c in range(rng.min_col, rng.max_col + 1)) * X.PT_PER_UNIT
        size = float(cell.font.sz or 12)
        ws.row_dimensions[row].height = round(X._lines(str(cell.value), width, size) * size * X.LINE_PT, 1)


def _two_pages(ws, last):
    """Акт сухпаю — на двох сторінках (рішення власника 23.10.2026): шапка — на першій, таблиця, додатки, висновок і
    підписи — на другій (рядки ACT_PAGE2…last). Друга сторінка вміщується за масштабом бланка «одна сторінка завширшки»
    — він і лишається; ні (рядків майна багато, посади на два рядки) — масштаб менший, щоб уміщувалась: так само, як
    сторінку рахує inventory_export.Pager (одиниця ширини графи на папері — CHAR_PT, запас 3 %)."""
    from inventory_export import CHAR_PT, PAPER                  # noqa: PLC0415
    m = ws.page_margins
    w, h = PAPER[::-1]                                           # альбомна А4
    widths = sum(excel_export._col_width(ws, c) for c in range(1, ACT_COLS + 1))
    wide = min(100, int(100 * (w - (m.left + m.right) * 72) / (widths * CHAR_PT)))
    need = sum(excel_export._row_pt(ws, r) for r in range(ACT_PAGE2, last + 1))
    tall = int(100 * (h - (m.top + m.bottom) * 72) * 0.97 / need) if need else wide
    if tall < wide:
        if ws.sheet_properties.pageSetUpPr is not None:
            ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = max(10, tall)                      # менше 10 % Excel не друкує


def write_package(con, doc_id, root, force=False):
    """Пакет акта сухпаю в теку справи → {"folder", "files", "locked"} (locked — збережені поруч із часом в імені).
    Файл пакета в теці правлено чи складено руками — PackageConflict, нічого не записано; `force` — переписати все одно."""
    from openpyxl import load_workbook                          # noqa: PLC0415
    a = sp_acts.act_of(con, doc_id)
    if not force:
        changed = conflicts(con, root, a)
        if changed:
            raise X.PackageConflict(changed)
    req = X.requisites(con, a["date"])
    sig = {(s["role"], int(s["sort"])): s for s in a["signers"]}
    folder = act_dir(root, a)
    os.makedirs(folder, exist_ok=True)
    wb = load_workbook(excel_export.template_path(TEMPLATE))
    excel_export.clean_workbook(wb)
    _invoice(wb["Накладна"], a, req, sig)
    _rozdavalna(wb["Роздавальна"], a, req, sig)
    _act(wb["Акт"], a, req, sig)
    files = [X._save(wb, folder, a["number"])]
    X._register(con, a, root, files, SCAN_FOLDER)
    own = {excel_export.safe_name(a["number"])}
    return {"folder": folder, "files": files,
            "locked": [f for f in files if os.path.splitext(os.path.basename(f))[0] not in own]}
