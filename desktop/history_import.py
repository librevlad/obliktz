# -*- coding: utf-8 -*-
"""Імпорт історії з Excel: шаблон, який служба заповнює зі своїх журналів, і читання
заповненого файла.

Нова служба починає з чистої бази, а її історія з 2022 року лежить у власних книгах
Excel. Вносити сотні документів руками довго, тому програма дає шаблон на три аркуші —
«Підрозділи», «Номенклатура», «Документи» — і читає його назад.

Тут лише файл: запис шаблона (`save_template`) і розбір заповненого (`read_history`).
Що з прочитаного можна провести, вирішує сторінка — тими самими правилами, що й форма
документа: цей модуль обліку не знає.
"""
import io
import os
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Times New Roman"
F_BODY = Font(name=FONT, size=11)
F_HEAD = Font(name=FONT, size=11, bold=True)
F_TITLE = Font(name=FONT, size=13, bold=True)
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E8ECE0")
MAX_ROWS = 50000

KINDS = ["Прихід", "Переміщення", "Вибуття", "Знищення"]

# Аркуші шаблона: ключ -> (назва, [(ключ графи, заголовок, ширина)]). Заголовок у файлі
# читається за початком слова, тож «Ціна» і «Ціна, грн» — та сама графа.
SHEETS = {
    "subs": ("Підрозділи", [("name", "Назва", 34), ("parent", "Підпорядкований", 30), ("type", "Вид", 18),
                            ("closed", "Закритий", 11), ("note", "Примітка", 40)]),
    "items": ("Номенклатура", [("code", "Код", 10), ("name", "Найменування", 60), ("unit", "Одиниця виміру", 12),
                               ("group", "Розділ 21/Прод", 14), ("price", "Ціна, грн", 14),
                               ("nonrev", "Необоротний актив", 14), ("fes", "Номер ФЕС", 16), ("note", "Примітка", 30)]),
    "docs": ("Документи", [("kind", "Вид", 14), ("type", "Тип документа", 20), ("no", "Номер", 18), ("date", "Дата", 12),
                           ("from", "Від кого", 30), ("to", "Кому", 30), ("code", "Код", 10),
                           ("name", "Найменування", 50), ("qty", "Кількість", 11), ("price", "Ціна, грн", 14),
                           ("basis", "Підстава", 30), ("note", "Примітка", 30), ("act", "Акт списання", 16)]),
}
# З чого починається заголовок графи у файлі. Довші початки стоять раніше коротших.
STARTS = {
    "subs": [("name", "назва"), ("parent", "підпорядк"), ("type", "вид"), ("closed", "закрит"), ("note", "приміт")],
    "items": [("code", "код"), ("name", "наймен"), ("unit", "одиниц"), ("unit", "од."), ("group", "розділ"),
              ("price", "ціна"), ("nonrev", "необорот"), ("fes", "номер фес"), ("fes", "фес"), ("note", "приміт")],
    "docs": [("kind", "вид"), ("type", "тип"), ("no", "номер"), ("no", "№"), ("date", "дата"), ("from", "від кого"),
             ("to", "кому"), ("code", "код"), ("name", "наймен"), ("qty", "кільк"), ("qty", "к-сть"), ("price", "ціна"),
             ("basis", "підстав"), ("note", "приміт"), ("act", "акт")],
}
NEED = {"subs": {"name"}, "items": {"code", "name"}, "docs": {"kind", "no", "date", "code", "qty"}}
TITLE_STARTS = {"subs": "підрозділ", "items": "номенклатур", "docs": "документ"}

# Кожен рядок пояснень — один рядок аркуша, без переносів: книгу пише не Excel, і висоти
# рядка під перенесений текст він сам не підбере.
HOWTO = [
    "Імпорт історії: як заповнити файл",
    "",
    "1. Аркуш «Підрозділи». Рядок — підрозділ.",
    "   Назва, кому підпорядкований (назва з цього аркуша або з програми), вид.",
    "   Вид: бригада, батальйон, рота, взвод/відділення, склад, служба.",
    "   Розформованому поставте «так» у графі «Закритий»: в історії він лишиться, у нових документах його не буде.",
    "   Підрозділи, які вже є в програмі, лишіть як є.",
    "",
    "2. Аркуш «Номенклатура». Рядок — позиція.",
    "   Код, найменування, одиниця виміру, розділ табеля 21/Прод, ціна.",
    "   Позиція має одну ціну: майно за іншою ціною заведіть іншим кодом.",
    "   Позиції, які вже є в програмі, лишіть як є.",
    "",
    "3. Аркуш «Документи». Рядок — одне найменування в документі.",
    "   Рядки з тим самим видом, датою, номером і маршрутом складають один документ.",
    "   Рядок із порожніми видом, номером і датою належить документу вище.",
    "   Прихід — від постачальника, з іншої частини, перенос залишків. «Від кого» — постачальник, «Кому» — підрозділ (порожньо — склад).",
    "   Переміщення — між підрозділами. «Від кого» й «Кому» — підрозділи.",
    "   Вибуття — списання або передача в іншу частину. «Від кого» — підрозділ, «Кому» — частина-одержувач, якщо майно передано.",
    "   Знищення — рапорт про знищення. «Від кого» — підрозділ, «Акт списання» — номер акта, якщо він уже є.",
    "   Залишки на початок обліку — прихід із типом «Перенос залишків» на кожного утримувача.",
    "   Ціна потрібна лише в приході; без неї програма бере ціну з аркуша «Номенклатура».",
    "   Дата — як 15.03.2022. Кількість — число. Код — як на аркуші «Номенклатура».",
    "",
    "4. Програма вносить файл цілком або не вносить нічого: помилки вона показує за аркушем і номером рядка.",
    "5. Заводські номери, людей, МВО, штат, звірки й інвентаризації вносять у самій програмі.",
]


class NotHistory(ValueError):
    """Файл не схожий на шаблон імпорту історії."""


# ----------------------------------------------------------------- шаблон
def build_template(spec: dict) -> Workbook:
    """Книга-шаблон. spec: {subs: [[назва, підпорядкований, вид, закритий, примітка]], items: [[код,
    найменування, одиниця, розділ, ціна, необоротний, номер ФЕС, примітка]], subKinds: [...],
    groups: [[код, назва розділу]]} — довідники програми, щоб їх було видно у файлі."""
    wb = Workbook()
    how = wb.active
    how.title = "Як заповнювати"
    how.column_dimensions["A"].width = 150
    lines = list(HOWTO)
    groups = [g for g in (spec.get("groups") or []) if g and g[0]]
    if groups:
        lines += ["", "Розділи табеля 21/Прод для аркуша «Номенклатура»:"]
        lines += [f"   {g[0]} — {g[1]}" if len(g) > 1 and g[1] else f"   {g[0]}" for g in groups]
    for r, text in enumerate(lines, 1):
        how.cell(r, 1, text).font = F_TITLE if r == 1 else F_BODY

    for key in ("subs", "items", "docs"):
        title, cols = SHEETS[key]
        ws = wb.create_sheet(title)
        for c, (name, label, width) in enumerate(cols, 1):
            head = ws.cell(1, c, label)
            head.font, head.fill, head.border = F_HEAD, HEAD_FILL, BOX
            head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            column = ws.column_dimensions[get_column_letter(c)]
            column.width = width
            # Код і номер — текст («00123» не стає числом 123), дата — дата.
            if name in ("code", "fes", "no"):
                column.number_format = "@"
            elif name == "date":
                column.number_format = "DD.MM.YYYY"
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        for r, row in enumerate(spec.get(key) or [], 2):
            for c, v in enumerate(list(row)[:len(cols)], 1):
                cell = ws.cell(r, c, v if v not in ("", None) else None)
                cell.font = F_BODY
                if cols[c - 1][0] in ("code", "fes", "no"):
                    cell.number_format = "@"
    docs = wb[SHEETS["docs"][0]]
    _list(docs, "A", KINDS)
    kinds = [str(k) for k in (spec.get("subKinds") or []) if k]
    if kinds:
        _list(wb[SHEETS["subs"][0]], "C", kinds)
    _list(wb[SHEETS["subs"][0]], "D", ["так", "ні"])
    _list(wb[SHEETS["items"][0]], "F", ["так", "ні"])
    wb.active = wb.index(docs)
    pr = wb.properties
    pr.creator = pr.lastModifiedBy = "Облік ТЗ ПС"
    pr.created = pr.modified = datetime.now()
    return wb


def _list(ws, col, values):
    """Перелік на вибір у графі: Excel підказує значення, власне слово не забороняє."""
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{MAX_ROWS}")


def save_template(spec: dict, folder: str) -> str:
    """Записує шаблон у теку вивантажень і повертає шлях до файла."""
    from excel_export import safe_name                              # noqa: PLC0415
    os.makedirs(folder, exist_ok=True)
    name = safe_name(spec.get("file") or "Шаблон імпорту історії")
    path = os.path.join(folder, name + ".xlsx")
    wb = build_template(spec)
    try:
        wb.save(path)
    except PermissionError:                  # файл із таким іменем відкрито в Excel
        path = os.path.join(folder, f"{name} {datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path


# ----------------------------------------------------------------- читання
def _norm(v):
    return " ".join(str(v if v is not None else "").split()).lower()


def _text(v):
    """Текст клітинки: код 10101, який Excel тримає числом 10101.0, — «10101»."""
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    return " ".join(str(v).split())


DMY = re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$")
YMD = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})")


def _date(v):
    """Дата клітинки як «2022-03-15»; що датою не є — порожньо."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 36526 <= v < 73051:
        # Дата в клітинці без формату дати — число днів від 1900 року (тут — роки 2000–2099).
        return from_excel(v).date().isoformat()
    text = _text(v)
    m = DMY.match(text)
    parts = (m[3], m[2], m[1]) if m else None
    if not parts:
        m = YMD.match(text)
        parts = (m[1], m[2], m[3]) if m else None
    if not parts:
        return ""
    try:
        return date(int(parts[0]), int(parts[1]), int(parts[2])).isoformat()
    except ValueError:
        return ""


def _number(v):
    """Число клітинки: «1 250,50» і 1250.5 — те саме; порожньо — None, не число — "?"."""
    if v is None or isinstance(v, bool) or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return "?"


def _columns(row, key):
    """Графи аркуша за рядком заголовків: {ключ: номер стовпця}."""
    cols = {}
    for c, v in enumerate(row):
        text = _norm(v)
        if not text:
            continue
        for name, start in STARTS[key]:
            if name not in cols and text.startswith(start):
                cols[name] = c
                break
    return cols


def _find(wb, key):
    """Аркуш і його заголовки: спершу за назвою аркуша, потім — будь-який із потрібними графами."""
    sheets = sorted(wb.worksheets, key=lambda ws: not _norm(ws.title).startswith(TITLE_STARTS[key]))
    for ws in sheets:
        named = _norm(ws.title).startswith(TITLE_STARTS[key])
        if key != "docs" and not named:
            continue                          # довідники — лише зі своїх аркушів: графи «Код» і «Назва» є всюди
        grid = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            grid.append(list(row))
            if i >= MAX_ROWS:
                break
        for hr, row in enumerate(grid[:20]):
            cols = _columns(row, key)
            if NEED[key] <= set(cols):
                return ws.title, cols, grid[hr + 1:], hr + 2
    return None, {}, [], 0


def read_history(data: bytes) -> dict:
    """Рядки трьох аркушів шаблона: {sheets: {subs, items, docs: назва аркуша або None},
    subs: [...], items: [...], docs: [...]}. У кожному рядку `row` — номер рядка в Excel.
    Дата — «2022-03-15» (а що не прочиталось — лишається в `dateText`), кількість і ціна —
    числа (нечислове — "?"). Рядок документів без виду, номера й дати продовжує документ вище."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                        # noqa: BLE001
        raise NotHistory("файл не відкривається як книга Excel (.xlsx)") from e
    try:
        out = {"sheets": {}}
        for key in ("subs", "items", "docs"):
            title, cols, rows, first = _find(wb, key)
            out["sheets"][key] = title
            out[key] = _rows(key, cols, rows, first) if title else []
    finally:
        wb.close()
    if not out["sheets"]["docs"]:
        raise NotHistory("це не шаблон імпорту історії: у файлі немає аркуша «Документи» з графами "
                         "«Вид», «Номер», «Дата», «Код» і «Кількість»")
    return out


def _rows(key, cols, grid, first):
    out = []
    names = [name for name, _, _ in SHEETS[key][1]]
    prev = None

    def at(row, name):
        c = cols.get(name)
        return row[c] if c is not None and c < len(row) else None

    for i, row in enumerate(grid):
        if all(_text(at(row, name)) == "" for name in names):
            continue
        rec = {"row": first + i}
        for name in names:
            v = at(row, name)
            if name == "date":
                rec["date"] = _date(v)
                rec["dateText"] = _text(v)
            elif name in ("qty", "price"):
                rec[name] = _number(v)
            else:
                rec[name] = _text(v)
        if key == "docs":
            # Продовження документа: шапку людина вписала один раз, далі — лише рядки майна.
            head = ("kind", "type", "no", "from", "to")
            if prev and not rec["kind"] and not rec["no"] and not rec["dateText"]:
                rec.update({k: prev[k] for k in head + ("date", "dateText", "basis", "act")
                            if not rec.get(k)})
                rec["same"] = True
            prev = rec
        out.append(rec)
    return out
