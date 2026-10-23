# -*- coding: utf-8 -*-
"""Пакет паперів акта списання ОП/МЗ у теці справи: «<номер>.xlsx» (накладна, рапорт, розрахунок, акт)
і «Звірка <номер>.xlsx»; тека — <тека актів>\\<рік>\\Посуд\\<номер> чи …\\Миючі і серветки\\<номер>, а якщо
вже є тека, назва якої починається з номера, — вона.

Бланки — desktop/templates/op-akt.xlsx і op-zvirka.xlsx (бланки служби без даних частини). Змінне — з бази:
дані акта, рядки документа, знімок підписантів, реквізити частини, дні д/д пункту. Програма переписує лише
ці два свої файли — і лише такими, якими сама їх записала (відбиток у файлах акта): правлений руками чи
складений руками файл під тим самим іменем — PackageConflict, переписати можна лише явною згодою (`force`).
Скани й інше в теці програма не чіпає й не видаляє. Файл лягає цілим (тимчасовий «….part» поруч стає на місце
одним кроком); зайнятий на мить — ще кілька спроб, відкритий в Excel — запис поруч із часом в імені.
"""
import contextlib
import datetime as dt
import json
import os
import re
import tempfile
import time
from copy import copy
from decimal import ROUND_FLOOR, Decimal

import excel_export
import op_acts
import op_norms
import stable_hash
from db.uk_number import money_words, number_words, qty_words

TEMPLATE, ZVIRKA = "op-akt.xlsx", "op-zvirka.xlsx"
ACT_FIRST, ACT_SLOTS, ACT_TOTAL = 39, 29, 68
CALC_FIRST, CALC_SLOTS, CALC_TOTAL = 7, 3, 10
ZV_FIRST, ZV_SLOTS = 16, 30                # рядок 46 — відступ перед підписами
MEMBER_ROWS = (81, 84, 87)                 # рядки членів комісії в акті (без зсуву таблиці)
MONTHS = ("січня", "лютого", "березня", "квітня", "травня", "червня", "липня", "серпня", "вересня",
          "жовтня", "листопада", "грудня")
EMPTY = {"pos": "", "rank": "", "name": ""}
TRIES, PAUSE = 5, 0.1                      # скільки разів і через скільки секунд пробувати стати на місце файла пакета
# Ширина знака Times New Roman у частках розміру шрифту — за класом знака: великі ширші за малі («Ім'я ПРІЗВИЩЕ» — майже
# самі великі), широкі (Ж М Ш Щ Ю Ф) — окремо; цифри, пробіл і розділові — як виміряно. Виміряно 21.10 на шрифті й на
# друку Excel об'єднаних клітинок бланків (312 рядків): оцінка не дала жодного рядка менше, ніж надруковано; запас
# GLYPH_MARGIN — у бік зайвого рядка. Ширина графи: одиниця Excel — 7 px (Calibri 11 книги) = 5,25 pt; рядок тексту —
# 1,3 розміру шрифту (бланк: 15,6 для 12 pt). Excel об'єднаних клітинок сам не підганяє.
WIDE_UP, WIDE_LO, NARROW = set("ЖМШЩЮФЫMW"), set("жмшщюфыmw"), set(",.;:'’!|/()")
GLYPH_MARGIN, PT_PER_UNIT, LINE_PT = 1.03, 5.25, 1.3
MVO_ROW = 92                               # рядок МВО в акті: опис з посадою й званням зліва, ім'я — AZ, лінія підпису V:AZ


def requisites(con, day):
    """Реквізити частини на дату: як на екрані реквізитів (правки частини поверх бази) і дислокація."""
    conf = dict(con.execute("SELECT key, value FROM settings"))
    row = con.execute("SELECT value FROM app_setting WHERE key = 'unit'").fetchone()
    try:
        mine = json.loads(row[0]) if row else {}
    except ValueError:
        mine = {}
    mine = mine if isinstance(mine, dict) else {}

    def pick(key, base, default=""):
        return str(mine.get(key) or "").strip() or str(conf.get(base) or "").strip() or default

    legal = pick("legalName", "unit_legal_name", "Військова частина")
    code = re.search(r"[АA]\d{4}", legal)
    place = con.execute("SELECT place FROM unit_location WHERE valid_from <= ? ORDER BY valid_from DESC LIMIT 1",
                        (day,)).fetchone()
    try:
        days = int(pick("validDays", "invoice_valid_days", "1"))
    except ValueError:
        days = 1
    return {"legalName": legal, "code": code.group(0) if code else "", "edrpou": pick("edrpou", "unit_edrpou"),
            "serviceFull": pick("serviceFull", "service_full"), "validDays": days,
            "place": place[0] if place else "", "actForm": str(mine.get("actForm") or "").strip()}


def act_dir(root, series, year, number):
    """Тека акта: наявна, назва якої починається з номера (так, як номер стоїть в імені теки — `safe_name`: «/» → «-»),
    інакше нова «<номер>»."""
    base = os.path.join(root, str(year), op_acts.SERIES[series]["folder"])
    own = excel_export.safe_name(number)
    if os.path.isdir(base):
        pat = re.compile(rf"^{re.escape(own)}(?:$|[\s\-_.])", re.IGNORECASE)
        for name in sorted(os.listdir(base)):
            if pat.match(name) and os.path.isdir(os.path.join(base, name)):
                return os.path.join(base, name)
    return os.path.join(base, own)


class PackageConflict(Exception):
    """Файл пакета в теці правлений чи складений руками (не той, яким програма його записала): питання до людини, а не
    тихий перезапис. `changed` — імена таких файлів."""

    def __init__(self, changed):
        super().__init__(" і ".join(f"«{n}»" for n in changed) + " змінено після програми — перезаписати?")
        self.changed = list(changed)


def _own_shas(con, name, folder=None):                  # folder — тека пакета виду: «opActs», в акта сухпаю «spActs»
    """Відбитки, якими програма записувала файл пакета з таким іменем (записи про файли актів, тека «opActs»); ім'я — без
    різниці регістру, як його бачить Windows (номер перенумерували лише регістром — файл той самий)."""
    return {sha for file_name, sha in con.execute("SELECT file_name, sha256 FROM app_file WHERE folder = ? AND sha256 IS NOT NULL",
                                                  (folder or SCAN_FOLDER,)) if file_name.casefold() == name.casefold()}


def conflicts(con, root, series, year, number):
    """Файли пакета акта (`<номер>.xlsx`, `Звірка <номер>.xlsx`), що лежать у теці акта, а програма такими їх не писала:
    правлені руками після неї чи складені руками під тим самим іменем. Немає файла — не конфлікт."""
    folder = act_dir(root, series, year, number)
    out = []
    for name in op_acts.package_names(number):
        path = os.path.join(folder, name)
        if os.path.isfile(path) and not set(stable_hash.shas(path)) & _own_shas(con, name):
            out.append(name)
    return out


def _dmy(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}" if iso else ""


def _long(iso):
    d = dt.date.fromisoformat(iso)
    return f"“{d.day:02d}” {MONTHS[d.month - 1]} {d.year} р."


def _n(q):
    return str(int(q)) if float(q).is_integer() else f"{q:.3f}".rstrip("0").replace(".", ",")


def _items_words(n):
    w = number_words(n, "m")
    if n % 10 == 1 and n % 100 != 11:
        w = w[:-len("один")] + "одне"
    tail = "найменування" if n % 10 in (1, 2, 3, 4) and n % 100 not in (11, 12, 13, 14) else "найменувань"
    return f"Всього {n} ({w}) {tail}"


def _who(sig, role, sort=0):
    return sig.get((role, sort)) or EMPTY


def _pr(p):
    return ", ".join(x for x in (p["pos"], p["rank"]) if x)


def _limit(a, line):
    per = op_norms.per_milli(line["norm"])
    if per is None:
        return 0
    step = op_norms.STEP_MILLI.get(line["group"], 1)
    return int((Decimal(a["rations"]) * per).to_integral_value(rounding=ROUND_FLOOR)) // step * step / 1000


def _unit(req):
    """«військової частини А0000»; коду частини немає — просто «військової частини»."""
    return " ".join(x for x in ("військової частини", req["code"]) if x)


def _shares(a):
    """«Належить за нормою» по рядках: норма пулу (рядок норми наказу 390 — op_norms.pool_key, як у розрахунку акта)
    ділиться між його кодами в їхньому порядку — кожному списане, останньому — решта норми пулу. Рядок, чий рядок норми не
    визначено (R60: число є в кількох рядках наказу, рядка не вибрано; так буває лише в акті, складеному раніше), — без
    частки норми: належить те, що списано."""
    out, pools = [0] * len(a["lines"]), {}
    for i, l in enumerate(a["lines"]):
        key = l["pool"] if "pool" in l else (l["group"], l["norm"])
        if key is None:
            out[i] = l["qty"]
            continue
        pools.setdefault(key, []).append(i)
    for idx in pools.values():
        limit, used = _limit(a, a["lines"][idx[0]]), 0
        for k, i in enumerate(idx):
            q = a["lines"][i]["qty"]
            out[i] = q if k < len(idx) - 1 else max(q, round(limit - used, 3))
            used += q
    return out


def _mid(text):
    """Назва посеред речення — з малої літери («що одноразовий посуд — …»), крім назви, що починається скороченням: дві
    великі літери поспіль чи велика й цифра або «-» («ТВН-12 термос», «Т-34», «PRO Миючий»)."""
    t = str(text or "")
    if len(t) < 2 or not t[0].isupper() or t[1].isupper() or t[1].isdigit() or t[1] == "-":
        return t
    return t[0].lower() + t[1:]


def _glyph(ch):
    """Ширина знака в частках розміру шрифту (клас знака, з запасом GLYPH_MARGIN)."""
    if ch == " ":
        w = 0.25
    elif ch.isdigit():
        w = 0.5
    elif ch.isupper():
        w = 1.0 if ch in WIDE_UP else 0.72
    elif ch.islower():
        w = 0.78 if ch in WIDE_LO else 0.5
    else:
        w = 0.3 if ch in NARROW else 0.5
    return w * GLYPH_MARGIN


def _text_pt(text, size):
    return sum(_glyph(ch) for ch in text) * size


def _lines(text, width_pt, size):
    """Скільки рядків займе текст у графі шириною `width_pt`: перенос, як в Excel, — після пробілу й після дефіса; шматок,
    довший за рядок, ділиться за знаками."""
    total = 0
    for para in str(text).split("\n"):
        lines, cur, first = 1, 0.0, True
        for word in para.split(" "):
            parts = word.split("-")
            for i, part in enumerate(parts):
                piece = part + ("-" if i < len(parts) - 1 else "")
                w = _text_pt(piece, size)
                gap = _glyph(" ") * size if (i == 0 and not first) else 0.0
                if first:
                    cur, first = w, False
                elif cur + gap + w <= width_pt:
                    cur += gap + w
                else:
                    lines += 1
                    cur = w
                while cur > width_pt:
                    lines += 1
                    cur -= width_pt
        total += lines
    return total


def _fit_text(ws, ref):
    """Рядок об'єднаної клітинки, у яку програма вписала посаду й звання чи ім'я, — заввишки, щоб текст уміщався: Excel
    об'єднаних клітинок сам не підганяє, і довге «…, молодший сержант» чи «Іван ВОЛОДИМИРЕНКО» обрізалося б. Рядків тексту —
    за шириною знаків (_lines) проти ширини об'єднання; текст на один рядок висоти бланка не змінює, на два й більше — рядок
    росте (не менший за бланковий), перенос слів — увімкнено, шрифт не менший за бланковий. Клітинку, що не об'єднана,
    текст сам переходить у сусідні порожні — її не чіпаємо."""
    cell = ws[ref]
    text = str(cell.value or "")
    rng = next((m for m in ws.merged_cells.ranges if cell.coordinate in m), None)
    if not text or rng is None:
        return
    width = sum(excel_export._col_width(ws, c) for c in range(rng.min_col, rng.max_col + 1)) * PT_PER_UNIT
    size = float(cell.font.sz or 11)
    lines = _lines(text, width, size)
    have = sum(excel_export._row_pt(ws, r) for r in range(rng.min_row, rng.max_row + 1))
    need = lines * size * LINE_PT
    if lines > 1:
        if need > have + 0.5:
            last = rng.max_row
            ws.row_dimensions[last].height = round(excel_export._row_pt(ws, last) + need - have, 1)
        align = copy(cell.alignment)
        align.wrap_text = True
        cell.alignment = align


def _grown_area(area, extra):
    """Область друку бланка, що виросла на `extra` рядків униз."""
    m = re.search(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$", str(area or ""))
    return f"{m.group(1)}{m.group(2)}:{m.group(3)}{int(m.group(4)) + extra}" if m else area


def _grow(ws, first, slots, total_row, n, cols):
    """Більше рядків, ніж заготовлено в бланку: розсунути таблицю, як у накладній."""
    extra = max(0, n - slots)
    if extra:
        excel_export._shift_down(ws, total_row, extra)
        for r in range(first + slots, first + slots + extra):
            ws.row_dimensions[r].height = ws.row_dimensions[first].height
            for c in range(1, cols + 1):
                excel_export._copy_style(ws.cell(row=first, column=c), ws.cell(row=r, column=c))
    return extra


def _qfmt(values):
    """Формат кількості: цілі (комплекти, штуки) — без дробу, як у паперах ОП; кілограми — до грама."""
    return "#,##0" if all(float(v).is_integer() for v in values) else "#,##0.000"


def _tall(ws, row, name, per_line=44):
    """Назва довша за рядок графи — рядок вищий: у паперах служби довгі назви йшли у два рядки."""
    lines = max(1, -(-len(str(name or "")) // per_line))
    if lines > 1:
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 15 * lines)


def _fit(ws, first, slots, n):
    """Заготовлених рядків більше, ніж треба: зайві ховаються, як у паперах служби."""
    for r in range(first, first + slots):
        ws.row_dimensions[r].hidden = r >= first + n


def _issue(a):
    """Накладна пакета — (номер, дата), як її друкує аркуш «Накладна»: лише для пункту, місце якого — склад, і коли
    папір «накладна» є; номер — паперу чи номер акта, дата — паперу чи початок періоду. Інакше None."""
    if "накладна" not in a["papers"] or a["sub_kind"] != "склад":
        return None
    paper = a["papers"].get("накладна") or {}
    return paper.get("no") or a["number"], paper.get("date") or a["from"]


def _words(a):
    """(що, чого, для чого) — як папери пакета називають вміст акта: за групами його рядків (op_acts.wording)."""
    return op_acts.wording({l["group"] for l in a["lines"]} or op_acts.SERIES[a["series"]]["groups"])


def _invoice(ws, a, req, sig, place):
    no, day = _issue(a)
    units = {l["uom"] for l in a["lines"]}
    extra = excel_export.fill_invoice_sheet(ws, {
        "no": no, "date": day, "round_lines": True, "blank_mixed_total": True,
        "operation": "Розподіл (видача)", "basis": "Забезпечення", "sender": a["sub"],
        "receiver": a["hall"] or a["pointName"], "place": place,
        "lines": [{"name": l["name"], "code": l["fes"], "uom": l["uom"], "cat": "", "price": l["price"],
                   "qty": l["qty"], "note": ""} for l in a["lines"]],
        "qty_words": f"{qty_words(sum(l['qty_milli'] for l in a['lines']))} {units.pop()}" if len(units) == 1 else "",
        "sum_words": money_words(sum(round(l["sum"] * 100) for l in a["lines"])),
        "sender_person": _who(sig, "МВО"), "receiver_person": _who(sig, "начальник їдальні"),
        "chief": _who(sig, "начальник служби"), "accountant": {}, "fes": _who(sig, "бухгалтер"),
        "unit": {k: req[k] for k in ("legalName", "edrpou", "serviceFull", "validDays")}})
    for row in (30, 38, 40):                                 # імена підписантів накладної
        _fit_text(ws, f"G{row + extra}")


def _report(ws, a, req, sig):
    chief, hall = _who(sig, "начальник служби"), _who(sig, "начальник їдальні")
    to = chief["pos"]
    ws["O4"] = ("Начальнику " + to[len("Начальник "):]) if to.lower().startswith("начальник ") else \
        (to or f"Начальнику продовольчої служби {_unit(req)}")
    issue = _issue(a)                    # накладна — та сама, що в пакеті: її номер і дата, а немає її — і посилання немає
    by = f" за накладною (вимогою) №{issue[0]} від {_dmy(issue[1])}" if issue else ""
    what = ", ".join(f"{_mid(l['name'])} — {_n(l['qty'])} {l['uom']}" for l in a["lines"])
    ws["A15"] = (f"        Дійсним доповідаю, що {what}, видані мені{by}, в період з {_dmy(a['from'])} по "
                 f"{_dmy(a['to'])} були використані для {_words(a)[2]} "
                 f"військовослужбовцям {_unit(req)}.")
    ws["A24"], ws["A26"], ws["P26"] = hall["pos"] or None, hall["rank"] or None, hall["name"] or None
    ws["A27"] = _dmy((a["papers"].get("рапорт") or {}).get("date") or a["to"])
    for ref in ("O4", "A24", "A26", "P26"):
        _fit_text(ws, ref)


def _calc(ws, a, req, sig):
    paper = a["papers"].get("розрахунок") or {}
    ws["H2"] = f"№ {a['number']} від {_dmy(paper.get('date') or a['date'])}"
    ws["A4"] = "по використанню " + _words(a)[1]
    ws["E6"] = f"{_dmy(a['from'])} - {_dmy(a['to'])}"
    extra = _grow(ws, CALC_FIRST, CALC_SLOTS, CALC_TOTAL, len(a["lines"]), 10)
    _fit(ws, CALC_FIRST, CALC_SLOTS, len(a["lines"]))
    shares = _shares(a)
    for i, l in enumerate(a["lines"]):
        r = CALC_FIRST + i
        norm = (l["norm"] or 0) * (100 if l["uom"] == "кг" else 1)
        for col, v in zip("ABCDEFGIJ", (i + 1, l["name"], l["uom"], round(norm, 6), a["rations"], shares[i],
                                         l["qty"], l["qty"], None)):
            ws[f"{col}{r}"] = v
        for col in "FGI":
            ws[f"{col}{r}"].number_format = _qfmt([ws[f"{col}{r}"].value])
        _tall(ws, r, l["name"], 30)
    ws[f"A{CALC_TOTAL + extra}"] = _items_words(len(a["lines"])).replace("Всього", "ВСЬОГО", 1)
    chief = _who(sig, "начальник служби")
    ws[f"A{13 + extra}"] = chief["pos"] or None
    ws[f"A{14 + extra}"] = f"{chief['rank']} ______________________ {chief['name']}".strip()


def _act(ws, a, req, sig, att=None, conclusion=None, ground="Списання", sum_mixed=False):
    """Аркуш «Акт» (Додаток 10). Додатки (`att`, до чотирьох), висновок комісії й підстава (`ground`) — від виду акта;
    без них — як в акта ОП/МЗ за нормами 390. `sum_mixed` — «всього» кількостей і рядків різних одиниць (акт сухпаю, як у
    паперах служби); в акта ОП/МЗ штуки й кілограми не додаються."""
    issue = _issue(a)
    ws["AL5"] = req["actForm"] or None
    cmd = _who(sig, "командир")
    ws["AG9"], ws["AG10"] = cmd["pos"] or None, f"{cmd['rank']}________________{cmd['name']}"
    _fit_text(ws, "AG9")
    _fit_text(ws, "AG10")
    ws["AG12"] = "___.___.20___"
    ws["AB29"], ws["Q34"], ws["X34"] = a["number"], a["number"], _dmy(a["date"])
    ws["AD34"], ws["AM34"], ws["AW34"] = ground, _dmy(a["date"]), req["code"] or None
    n = len(a["lines"])
    hidden = sorted(r for r, d in ws.row_dimensions.items() if d.hidden and r >= ACT_TOTAL)
    extra = _grow(ws, ACT_FIRST, ACT_SLOTS, ACT_TOTAL, n, 53)
    if extra:                       # рядки майна понад бланк — з об'єднаннями першого, решта бланка — нижче
        slot = [m for m in ws.merged_cells.ranges if m.min_row == m.max_row == ACT_FIRST]
        for r in range(ACT_FIRST + ACT_SLOTS, ACT_FIRST + ACT_SLOTS + extra):
            for m in slot:
                ws.merge_cells(start_row=r, start_column=m.min_col, end_row=r, end_column=m.max_col)
        for r in hidden:
            ws.row_dimensions[r].hidden = False
        for r in hidden:
            ws.row_dimensions[r + extra].hidden = True
        ws.print_area = _grown_area(ws.print_area, extra)
    for r in range(ACT_FIRST + 1, ACT_FIRST + min(n, ACT_SLOTS)):      # у бланку служби рядки нижче мали інші формати
        for c in range(1, 54):
            excel_export._copy_style(ws.cell(row=ACT_FIRST, column=c), ws.cell(row=r, column=c))
    _fit(ws, ACT_FIRST, ACT_SLOTS, n)
    last = ACT_FIRST + max(n, 1) - 1
    # «Належить використати за нормою» — частка норми рядка, те саме число, що «належить за нормою» Розрахунку (_shares);
    # «Фактично використано» — списане; залишку менше за норму — різниця йде в «Економію» (кількість і сума), а «---» там
    # лише тоді, коли вона нуль (рішення власника 21.10: «9 000, економія 1 500»). Частка не менша за списане (_shares) —
    # перевитрат акт програми не має.
    shares = _shares(a)
    saved = [round(share - line["qty"], 3) for share, line in zip(shares, a["lines"])]
    for i, l in enumerate(a["lines"]):
        r = ACT_FIRST + i
        # Сума рядка — до копійки, як рахує програма (реєстр, сума прописом у накладній): «Всього» — сума видимих сум.
        vals = {"A": i + 1, "E": l["name"], "Q": l["fes"] or "--", "S": l["uom"], "U": "--", "W": l["price"],
                "Z": l["qty"], "AC": f"=ROUND(W{r}*Z{r},2)", "AG": shares[i], "AK": "---", "AN": "---",
                "AR": saved[i] if saved[i] > 0 else "---", "AU": f"=ROUND(AR{r}*W{r},2)" if saved[i] > 0 else "---",
                "AY": "---"}
        for col, v in vals.items():
            ws[f"{col}{r}"] = v
        ws[f"Z{r}"].number_format = ws[f"AG{r}"].number_format = _qfmt([l["qty"], shares[i]])
        ws[f"AC{r}"].number_format = "#,##0.00"
        if saved[i] > 0:
            ws[f"AR{r}"].number_format, ws[f"AU{r}"].number_format = _qfmt([saved[i]]), "#,##0.00"
        _tall(ws, r, l["name"])
    t = ACT_TOTAL + extra
    ws[f"A{t}"] = _items_words(len(a["lines"]))
    ws[f"Z{t}"], ws[f"AC{t}"], ws[f"AG{t}"] = \
        f"=SUM(Z{ACT_FIRST}:Z{last})", f"=SUM(AC{ACT_FIRST}:AC{last})", f"=SUM(AG{ACT_FIRST}:AG{last})"
    ws[f"Z{t}"].number_format = ws[f"AG{t}"].number_format = _qfmt([l["qty"] for l in a["lines"]] + shares)
    ws[f"AC{t}"].number_format = "#,##0.00"
    ws[f"AK{t}"] = ws[f"AN{t}"] = "---"
    if any(s > 0 for s in saved):
        ws[f"AR{t}"], ws[f"AU{t}"] = f"=SUM(AR{ACT_FIRST}:AR{last})", f"=SUM(AU{ACT_FIRST}:AU{last})"
        ws[f"AR{t}"].number_format, ws[f"AU{t}"].number_format = _qfmt([s for s in saved if s > 0]), "#,##0.00"
    else:
        ws[f"AR{t}"] = ws[f"AU{t}"] = "---"
    if len({l["uom"] for l in a["lines"]}) > 1 and not sum_mixed:     # штуки й кілограми не додаються
        ws[f"Z{t}"] = ws[f"AG{t}"] = None
        if any(s > 0 for s in saved):
            ws[f"AR{t}"] = None
    zv = a["papers"].get("звірка") or {}
    period = f"за період з {_dmy(a['from'])} по {_dmy(a['to'])} року"
    if att is None:
        att = ([f"НАКЛАДНА (вимога) №{issue[0]} від {_dmy(issue[1])} року;"] if issue else []) + [
            f"РАПОРТ-ДОНЕСЕННЯ {period};", f"АКТ ЗВІРКИ №{zv.get('no') or a['number']} від {_dmy(zv.get('date') or a['date'])} {period};",
            f"РОЗРАХУНОК {period}."]
    for i in range(4):
        ws[f"A{71 + extra + i}"] = f"{i + 1}. {att[i]}" if i < len(att) else None
    if conclusion is None:
        basis = (f"накладної (вимоги) №{issue[0]} від {_dmy(issue[1])}, " if issue else "")
        conclusion = (f"Висновок комісії: на підставі {basis}акта звірки №{zv.get('no') or a['number']} {period} та "
                      f"розрахунку за нормами наказу МОУ №390 — {_words(a)[0]} списати з обліку {_unit(req)}.")
    ws[f"A{76 + extra}"] = conclusion
    # Члени комісії — ті, хто є, підряд з першого рядка («Члени комісії:» бланка — на ньому): немає першого — другий
    # стає на його місце, а не лишає порожній рядок посередині; порожні рядки для руки — лише в кінці.
    members = [sig[k] for k in sorted(k for k in sig if k[0] == "член комісії")][:len(MEMBER_ROWS)]
    for row, p in [(78, _who(sig, "голова комісії"))] + [
            (r, members[i] if i < len(members) else EMPTY) for i, r in enumerate(MEMBER_ROWS)]:
        name = f"{'AJ' if row == 78 else 'AS'}{row + extra}"
        ws[f"F{row + extra}"], ws[name] = _pr(p) or None, p["name"] or None
        _fit_text(ws, f"F{row + extra}")
        _fit_text(ws, name)
    mvo = _who(sig, "МВО")
    _mvo_line(ws, MVO_ROW + extra, "Відповідальна особа за зберігання і утримання військового майна" +
              (f", {_mid(_pr(mvo))}" if _pr(mvo) else ""), mvo["name"])
    for row, role in ((95, "начальник служби"), (98, "бухгалтер")):
        p = _who(sig, role)
        ws[f"A{row + extra}"], ws[f"AZ{row + extra}"] = _pr(p) or None, p["name"] or None


def _mvo_line(ws, row, text, name):
    """Рядок МВО акта (R62): опис із посадою й званням — об'єднана клітинка зліва з переносом, рядок росте під текст
    (_fit_text), тож довге «…5 батальйону ТрО …, молодший сержант» не обрізається; ім'я — у своїй AZ, праворуч на лінії
    підпису. До AY об'єднати не можна: ім'я в AZ (4 од.) займає порожні графи зліва від себе, і під об'єднанням його було б
    обрізано, — тож опис сягає лише до граф, яких імені досить (не лівіше U). Лінія підпису V:AZ лишається, як у бланку."""
    ws[f"AZ{row}"] = name or None
    room, last = _text_pt(name or "", float(ws[f"AZ{row}"].font.sz or 12)) + 6, 51
    have = excel_export._col_width(ws, 52) * PT_PER_UNIT
    while have < room and last > 21:
        have += excel_export._col_width(ws, last) * PT_PER_UNIT
        last -= 1
    # Об'єднання (openpyxl) стирає межі граф і дає першій межу знизу останньої — межі повертаються, як у бланку: лінія
    # підпису лише V:AZ, під описом зліва — нічого.
    lines = {c: copy(ws.cell(row=row, column=c).border) for c in range(1, last + 1)}
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    for c, border in lines.items():
        ws.cell(row=row, column=c).border = border
    cell = ws.cell(row=row, column=1)
    cell.value = text
    align = copy(cell.alignment)
    align.wrap_text = True
    cell.alignment = align
    _fit_text(ws, f"A{row}")


def _zvirka(ws, a, req, sig):
    cmd, zv = _who(sig, "командир"), a["papers"].get("звірка") or {}
    ws["C3"], ws["C4"] = cmd["pos"] or None, f"{cmd['rank']} _________ {cmd['name']}".strip()
    _fit_text(ws, "C3")
    _fit_text(ws, "C4")
    ws["C5"] = _long(zv.get("date") or a["date"])
    ws["C5"].alignment = copy(ws["C3"].alignment)
    ws["A8"] = f"АКТ ЗВІРКИ №{zv.get('no') or a['number']}"
    ws["A8"].alignment = copy(ws["A10"].alignment)
    ws["A10"] = (f"обліку руху осіб, які перебували на харчуванні в {a['title'] or a['pointName']}\n"
                 f" у військовій частині {req['code']}").rstrip()
    ws["B11"] = f"з {_dmy(a['from'])} по {_dmy(a['to'])} року"
    extra = _grow(ws, ZV_FIRST, ZV_SLOTS, ZV_FIRST + ZV_SLOTS, len(a["days"]), 7)
    _fit(ws, ZV_FIRST, ZV_SLOTS, len(a["days"]))
    for i, (day, n, order) in enumerate(a["days"]):
        r = ZV_FIRST + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = \
            dt.datetime.fromisoformat(day), (int(order) if str(order).isdigit() else order or None), n, n, "немає"
    for row, role in ((47, "начальник штабу"), (50, "начальник групи персоналу"), (53, "начальник логістики"),
                      (56, "начальник служби")):
        p = _who(sig, role)
        ws[f"A{row + extra}"] = p["pos"] or None
        ws[f"A{row + 1 + extra}"] = (p["rank"] + ("" if row in (53, 56) else " " + "_" * 30)).strip() or None
        ws[f"E{row + 1 + extra}"] = p["name"] or None
        for ref in (f"A{row + extra}", f"A{row + 1 + extra}", f"E{row + 1 + extra}"):
            _fit_text(ws, ref)


def _place(part, path):
    """Тимчасовий файл — на місце файла пакета одним кроком (окремо: перевірки підміняють саме цей крок). Файл на мить
    тримає антивірус чи індексатор — ще кілька спроб з паузою (TRIES × PAUSE); так і не відпустили (відкритий в Excel) —
    PermissionError, і пакет лягає поруч із часом в імені (_save)."""
    for left in range(TRIES - 1, -1, -1):
        try:
            os.replace(part, path)
            return
        except PermissionError:
            if not left:
                raise
            time.sleep(PAUSE)


def _save(wb, folder, name):
    """Файл пакета лягає цілим: книга пишеться поруч тимчасовим «….part», отримує імена Excel цього ПК і стає на місце
    одним кроком — збій посеред запису лишає попередній файл як був і не лишає ні половини, ні «.part». Файл, відкритий у
    Excel, заміні не піддається: пакет лягає поруч із часом в імені."""
    path = os.path.join(folder, excel_export.safe_name(name) + ".xlsx")
    fd, part = tempfile.mkstemp(dir=folder, suffix=".part")
    os.close(fd)
    try:
        wb.save(part)
        with contextlib.suppress(OSError):
            import excel_names                                   # noqa: PLC0415
            excel_names.localize_file(part)
        try:
            _place(part, path)
        except PermissionError:
            path = os.path.join(folder, f"{excel_export.safe_name(name)} {dt.datetime.now():%H%M%S}.xlsx")
            _place(part, path)
    except BaseException:
        with contextlib.suppress(OSError):
            os.remove(part)
        raise
    return path


SCAN_FOLDER = "opActs"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _register(con, a, root, files, folder=SCAN_FOLDER):
    """Файли пакета — у файли документа акта (як скани): ключ документа на сторінці, тека пакета (`folder`: «opActs»,
    у акта сухпаю — «spActs»), шлях від неї, відбиток. Повторне складання пакета заміняє записи тих самих файлів."""
    key = f"wr|{a['date']}|{a['number']}|{a['sub']}|"
    con.execute("BEGIN IMMEDIATE")
    try:
        # Ті самі файли — і тоді, коли номер змінився лише регістром («ОП…» → «оп…»): Windows бачить той самий файл.
        names = {os.path.basename(f).casefold() for f in files}
        gone = [pos for pos, name in con.execute("SELECT pos, file_name FROM app_file WHERE target = ?", (key,))
                if name.casefold() in names]
        con.executemany("DELETE FROM app_file WHERE pos = ?", [(p,) for p in gone])
        top = con.execute("SELECT COALESCE(MAX(pos), 0) FROM app_file").fetchone()[0]
        for n, f in enumerate(files, top + 1):
            with open(f, "rb") as fh:
                data = fh.read()
            con.execute("""INSERT INTO app_file(pos, id, target, doc_date, doc_no, file_name, rel_path, mime, size_bytes,
                           folder, sha256) VALUES(?, NULL, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (n, key, a["date"], a["number"], os.path.basename(f),
                         os.path.relpath(f, root).replace(os.sep, "/"), XLSX, len(data), folder,
                         stable_hash.stable_sha_bytes(data)))
        con.execute("COMMIT")
    except Exception:
        con.execute("ROLLBACK")
        raise


def _thousands(n):
    return f"{n:,}".replace(",", " ")


def write_package(con, doc_id, root, force=False):
    """Пакет акта в теку справи → {"folder", "files", "locked"} (locked — збережені поруч із часом в імені).

    Не складається (нічого не записано), коли д/д пункту за період уже не ті, з яких акт складено (звірка з нових днів
    не зійшлася б із розрахунком акта: SaveError), і коли файл пакета в теці правлено чи складено руками
    (PackageConflict; `force` — переписати все одно)."""
    from openpyxl import load_workbook                          # noqa: PLC0415
    a = op_acts.act_of(con, doc_id)
    now = sum(n for _, n, _ in a["days"])
    if a["rations"] and now and now != a["rations"]:
        raise op_acts.SaveError(f"Д/д пункту за період змінилися після складання акта: в акті {_thousands(a['rations'])}, "
                                f"у дободачах {_thousands(now)}. Відкрийте акт і «Зберегти» — папери складуться з новими д/д.")
    if not force:
        changed = conflicts(con, root, a["series"], a["date"][:4], a["number"])
        if changed:
            raise PackageConflict(changed)
    req = requisites(con, a["date"])
    sig = {(s["role"], int(s["sort"])): s for s in a["signers"]}
    folder = act_dir(root, a["series"], a["date"][:4], a["number"])
    os.makedirs(folder, exist_ok=True)
    wb = load_workbook(excel_export.template_path(TEMPLATE))
    excel_export.clean_workbook(wb)
    issue = _issue(a)
    if issue:                            # місце складання накладної — дислокація на дату накладної, не акта
        _invoice(wb["Накладна"], a, req, sig, requisites(con, issue[1])["place"])
    else:
        del wb["Накладна"]
    _report(wb["Рапорт"], a, req, sig)
    _calc(wb["Розрахунок"], a, req, sig)
    _act(wb["Акт"], a, req, sig)
    files = [_save(wb, folder, a["number"])]
    wz = load_workbook(excel_export.template_path(ZVIRKA))
    excel_export.clean_workbook(wz)
    _zvirka(wz["Звірка"], a, req, sig)
    files.append(_save(wz, folder, f"Звірка {a['number']}"))
    _register(con, a, root, files)
    own = {excel_export.safe_name(a["number"]), excel_export.safe_name(f"Звірка {a['number']}")}
    return {"folder": folder, "files": files,
            "locked": [f for f in files if os.path.splitext(os.path.basename(f))[0] not in own]}
