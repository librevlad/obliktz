# -*- coding: utf-8 -*-
"""Добовка служби — лише читання: д/д пунктів харчування, номери витягів, підписанти на дату.

Аркуші (назви — як у книзі служби):
- «Пункти харчування»: рядок шапки — «період», «дата», «разом», далі назви пунктів харчування; далі рядок
  на день (дата в колонці «дата»), між тижнями — рядки «за період …» без дати. День заповнено, коли «разом»
  більше нуля; нуль, порожньо чи від'ємне (служба заздалегідь вписує в майбутні дні -10) — день ще порожній, і в
  дані не йде ні він, ні нулі інших пунктів. Нуль пункту в заповненому дні — справжній нуль.
- «Добовка»: рядок «Добова відомість» — дати по колонках, рядок «Витяг з наказу №» — номер витягу на дату,
  рядок «разом» — усього на харчуванні за день. Добовку заповнено до останнього дня, де «разом» більше нуля:
  пізніші дні «Пунктів харчування» того самого файла — порожні, а номери витягів беруться лише заповнених днів
  (у книзі далі до кінця року стоїть формула «попередній + 1»).
- «Підписанти»: рядок 1 — групи («управління», «батальйон» …) над першою колонкою групи, рядок 2 — ролі
  («командир», «НШ», «комісія · голова» …), рядок 3 — «посада», «звання», «ПІБ»; далі рядок на дату
  (дата в другій колонці). На кожну роль — три колонки.

- «Сухпай»: дні й особи на сухпаї по підрозділах, рапорт і акт дня — для актів сухпаю (`read_sukhpai`, `report_days`)
  і «1прод» (`suhpay_days`).

Те саме приходить вставкою: Ctrl+A, Ctrl+C на аркуші Excel дає текст із табуляціями, аркуш упізнається
за шапкою; шматок без шапки — «дата | число» (рівно два стовпці) або стовпець чисел пункту з дня `start`
(не далі `end`). Стовпець іде за будовою аркуша «Пункти харчування»: після неділі — рядок «за період», після 13-го,
26-го, 39-го тижня року ПС і після 31.12 — рядок кварталу; ці рядки не дні й пропускаються за місцем (кожен названо
в `problems`). Підписані підсумки («за період …», «разом», «… квартал») пропускає підпис. Дату без року («дд.мм» — так
Excel копіює клітинку у форматі «дд/мм») вставка датує роком дня `start`; дата з «/» і роком («7/12/2026» —
так копіює Excel з англійськими налаштуваннями, місяць попереду) неоднозначна — відмова словами.
Чого база не прийме (від'ємне число, понад state_db.RATIONS_MAX, рік поза роками обліку), — слова в `problems`,
а не дані. Аркуш, шапку якого впізнано, але з якого нічого не прочитано, дає текст у `problems`.
У Добовку програма не пише.
"""
import datetime as dt
import glob
import io
import os
import re

import state_db

# Ролі акта: (роль, порядок, звідки — спільна з групи «управління» чи з групи пункту, підписи в рядку 2).
ROLES = [
    ("командир", 0, "common", ("командир",)),
    ("голова комісії", 0, "common", ("комісія · голова",)),
    ("член комісії", 1, "common", ("комісія · 1",)),
    ("член комісії", 2, "common", ("комісія · 2",)),
    ("член комісії", 3, "common", ("комісія · 3",)),
    ("МВО", 0, "point", ("мво",)),
    ("начальник служби", 0, "common", ("нпс",)),
    ("бухгалтер", 0, "common", ("начфін",)),
    ("начальник їдальні", 0, "point", ("начїд", "нач.їд")),
    ("начальник штабу", 0, "common", ("нш",)),
    ("начальник групи персоналу", 0, "none", ()),
    ("начальник логістики", 0, "common", ("нл",)),
]
COMMON = "управління"
_DMY = re.compile(r"^(\d{1,2})([./])(\d{1,2})(?:([./])(\d{4}|\d{2}))?$")
_YMD = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")
# Рядок підсумку, а не дня: «за період 11.2» (після кожної неділі), «I квартал» (латиницею чи кирилицею), «разом», «всього».
_SUBTOTAL = re.compile(r"^(за період|разом|всього|усього|[ivxі]{1,4} квартал)\b")
# Аркуш упізнано, а даних немає: що сказати (ключ — розділ результату).
_NOTHING = {"points": "Аркуш «Пункти харчування» впізнано, але жодного дня не прочитано.",
            "orders": "Аркуш «Добовка» впізнано, але жодного номера витягу не прочитано.",
            "signers": "Аркуш «Підписанти» впізнано, але жодного рядка не прочитано.",
            "sukhpai": "Аркуш «Сухпай» впізнано, але жодного дня з особами не прочитано."}
_NO_YEAR = " Якщо в датах немає року, вставте у вікні акта: рік візьметься з періоду."
_UNKNOWN = ("Не впізнано: вставте аркуш «Пункти харчування», «Підписанти» чи «Добовка» або стовпці «дата | число» "
            "чи стовпець чисел.")


def _low(v):
    return re.sub(r"\s+", " ", str(v or "").replace("\xa0", " ")).strip().lower()


def _dmy(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}"


def _parse_day(v, year=None):
    """(ISO, чому ні) — дата клітинки чи тексту вставки. «дд.мм.рррр», «дд.мм.рр» (рік 20рр), «рррр-мм-дд», «дд.мм» і
    «дд/мм» (рік — `year`). Чому ні: "slash" — дата з «/» і роком (місяць/день чи день/місяць — не вгадуємо), "year" — рік
    поза роками обліку (state_db.DATE_YEARS: «12.07.0026» — рік з опискою), None — не дата."""
    if isinstance(v, (dt.datetime, dt.date)):
        parts = (v.year, v.month, v.day)
    else:
        s = str(v or "").strip()
        m = _DMY.match(s)
        if m:
            day, sep1, month, sep2, y = m.groups()
            if y and "/" in (sep1, sep2):
                return None, "slash"
            y = (2000 + int(y) if len(y) == 2 else int(y)) if y else year
            parts = (int(y), int(month), int(day)) if y else None
        else:
            m = _YMD.match(s)
            parts = (int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None
    if not parts:
        return None, None
    try:
        iso = dt.date(*parts).isoformat()
    except ValueError:
        return None, None
    if not state_db.DATE_YEARS[0] <= parts[0] <= state_db.DATE_YEARS[1]:
        return None, "year"
    return iso, None


def _day(v, year=None):
    """Дата клітинки чи тексту вставки → ISO, інакше None (_parse_day без причини)."""
    return _parse_day(v, year)[0]


def _num(v):
    """Число клітинки чи тексту («3 664», «3664,0») → int, інакше None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        try:
            return int(round(v))
        except (ValueError, OverflowError):
            return None
    s = str(v or "").replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return int(round(float(s))) if s else None
    except (ValueError, OverflowError):
        return None


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _empty():
    return {"kind": "unknown", "points": {}, "orders": {}, "signers": {}, "values": {}, "sukhpai": {}, "problems": []}


def _say(out, text):
    if text not in out["problems"]:
        out["problems"].append(text)


def _shown(v):
    """Дата клітинки так, як її бачить людина: «21.10.1900», а не «1900-10-21 00:00:00»."""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.day:02d}.{v.month:02d}.{v.year:04d}"
    return _text(v)


def _odd_date(out, v, why, where=""):
    """Слова про дати, яких не взято (одне речення на аркуш і причину, з першим прикладом): неоднозначні («/» і рік)
    чи з роком-опискою."""
    head = f"{where}: дати" if where else "Дати"
    if why == "slash":
        text = (f"{head} з «/» і роком (як «{_shown(v)}») неоднозначні: місяць/день чи день/місяць. "
                "Скопіюйте дати у форматі дд.мм.рррр.")
    else:
        text = f"{head}, де рік з опискою (як «{_shown(v)}»), пропущено."
    stem = text.split("(як")[0]
    if not any(p.startswith(stem) for p in out["problems"]):
        out["problems"].append(text)


def _refused(n):
    """Чого база не прийме як д/д дня (state_db._save_rations): слово, інакше None."""
    if n < 0:
        return "від'ємне число"
    if n > state_db.RATIONS_MAX:
        return "завелике число"
    return None


def _subtotal(cells):
    """Рядок підсумку («за період …», «разом», «всього») — не день."""
    return any(_SUBTOTAL.match(_low(c)) for c in cells)


def _points(rows, out, year=None, last=None):
    """Рядки аркуша «Пункти харчування» → out["points"]; True, якщо шапку знайдено.
    Шапка — «період», «дата», «разом»: «Підписанти» теж починається з «період», «дата», але далі «управління».
    День заповнено, коли «разом» більше нуля й день не пізніший за `last` (останній заповнений день аркуша «Добовка»
    того самого файла); порожній день у дані не йде зовсім. Число пункту, якого база не прийме, — слова, а не дані."""
    at = next((i for i, r in enumerate(rows) if len(r) > 2 and _low(r[0]) == "період" and _low(r[1]) == "дата"
               and _low(r[2]) == "разом"), None)
    if at is None:
        return False
    names = [(i, _text(h)) for i, h in enumerate(rows[at]) if i > 2 and _text(h)]
    for r in rows[at + 1:]:
        if len(r) < 2 or _subtotal(r[:1]):
            continue
        day, why = _parse_day(r[1], year)
        if not day:
            if why:
                _odd_date(out, r[1], why, "Пункти харчування")
            continue
        total = _num(r[2]) if len(r) > 2 else None
        if total is None or total <= 0 or (last and day > last):
            continue                                     # день ще не заповнено
        for i, name in names:
            cell = r[i] if i < len(r) else None
            n = _num(cell)
            if n is None and _text(cell):
                _say(out, f"Пункти харчування, {_dmy(day)}, «{name}»: «{_text(cell)}» — не число, не взято.")
                continue
            n = 0 if n is None else n
            why = _refused(n)
            if why:
                _say(out, f"Пункти харчування, {_dmy(day)}, «{name}»: {n} — {why}, не взято.")
                continue
            out["points"].setdefault(name, {})[day] = n
    return True


def _dob_totals(rows, year=None):
    """«Разом» аркуша «Добовка» по днях ({дата: число чи None}); None — рядка «разом» (чи рядка дат) немає."""
    dates = next((r for r in rows if r and _low(r[0]).startswith("добова відомість")), None)
    total = next((r for r in rows if r and _low(r[0]) == "разом"), None)
    if dates is None or total is None:
        return None
    out = {}
    for i, v in enumerate(dates):
        day = _day(v, year) if i else None
        if day:
            out[day] = _num(total[i]) if i < len(total) else None
    return out


def _orders(rows, out, year=None, filled=None):
    """Номери витягів аркуша «Добовка» → out["orders"]; True, якщо рядки дат і витягів знайдено. Номер — лише дня, який
    Добовку заповнено: «разом» аркуша більше нуля (є такий рядок) чи день заповнено в «Пунктах харчування» того ж файла
    (`filled`); номери далі — формула «попередній + 1» до кінця року, а не витяги."""
    dates = next((r for r in rows if r and _low(r[0]).startswith("добова відомість")), None)
    nums = next((r for r in rows if r and _low(r[0]) == "витяг з наказу №"), None)
    if dates is None or nums is None:
        return False
    totals = _dob_totals(rows, year)
    for i, v in enumerate(dates):
        if not i:
            continue
        day, why = _parse_day(v, year)
        if not day:
            if why:
                _odd_date(out, v, why, "Добовка")
            continue
        if i < len(nums) and _text(nums[i]):
            if totals is not None and not (totals.get(day) or 0) > 0 and day not in (filled or ()):
                continue                                 # Добовку цього дня ще не заповнено
            out["orders"][day] = _text(nums[i])
    return True


def _signers(rows, out, year=None):
    at = next((i for i, r in enumerate(rows) if i >= 2 and "посада" in [_low(c) for c in r]
               and "командир" in [_low(c) for c in rows[i - 1]]), None)
    if at is None:
        return False
    groups, roles, group = rows[at - 2], rows[at - 1], ""
    cols = []
    for c in range(2, len(roles)):
        if c < len(groups) and _text(groups[c]):
            group = _low(groups[c])
        if _text(roles[c]):
            cols.append((c, f"{group}|{_low(roles[c])}"))
    for r in rows[at + 1:]:
        day, why = _parse_day(r[1], year) if len(r) > 1 else (None, None)
        if not day:
            if why:
                _odd_date(out, r[1], why, "Підписанти")
            continue
        rec = {}
        for c, key in cols:
            pos, rank, name = (_text(r[c + k]) if c + k < len(r) else "" for k in range(3))
            if pos or rank or name:
                rec[key] = {"pos": pos, "rank": rank, "name": name}
        if rec:                                  # порожній рядок дати не заступає заповнений раніший
            out["signers"][day] = rec
    return True


SHEETS = (("Пункти харчування", _points, "points"), ("Добовка", _orders, "orders"), ("Підписанти", _signers, "signers"))


def read_file(data):
    """Книга Добовки (байти .xlsx) → дні пунктів, номери витягів, підписанти. Аркуш «Добовка» читається першим: його
    «разом» каже, до якого дня Добовку заповнено."""
    import openpyxl                                              # noqa: PLC0415
    out = _empty()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                       # noqa: BLE001 — будь-який зіпсований файл
        print(f"Добовка: файл не прочитано — {type(e).__name__}: {e}")
        out["problems"].append("Файл не є книгою Excel (.xlsx).")
        return out
    try:
        if not any(title in wb.sheetnames for title, _, _ in SHEETS):
            out["problems"].append("Немає аркушів «Пункти харчування», «Добовка», «Підписанти».")
            return out
        rows = {}
        for title, _, _ in SHEETS:
            if title not in wb.sheetnames:
                out["problems"].append(f"Немає аркуша «{title}».")
                continue
            try:
                rows[title] = [list(r) for r in wb[title].iter_rows(values_only=True)]
            except Exception as e:                               # noqa: BLE001 — зіпсований аркуш
                print(f"Добовка: аркуш «{title}» не прочитано — {type(e).__name__}: {e}")
                out["problems"].append(f"Аркуш «{title}» не прочитано.")
        totals = _dob_totals(rows["Добовка"]) if "Добовка" in rows else None
        last = max((d for d, n in totals.items() if (n or 0) > 0), default=None) if totals else None
        steps = {"points": lambda r: _points(r, out, last=last),
                 "orders": lambda r: _orders(r, out, filled={d for days in out["points"].values() for d in days}),
                 "signers": lambda r: _signers(r, out)}
        for title, _, what in SHEETS:
            if title not in rows:
                continue
            try:
                found = steps[what](rows[title])
            except Exception as e:                               # noqa: BLE001 — аркуш несподіваної будови
                print(f"Добовка: аркуш «{title}» не прочитано — {type(e).__name__}: {e}")
                out["problems"].append(f"Аркуш «{title}» не прочитано.")
                continue
            if not found:
                out["problems"].append(f"Аркуш «{title}»: шапку не впізнано.")
            elif not out[what]:
                out["problems"].append(_NOTHING[what])
        if out["points"] or out["orders"] or out["signers"]:
            out["kind"] = "book"
    finally:
        wb.close()
    return out


# Рядок кварталу аркуша «Пункти харчування»: квартал — 13 тижнів року ПС (тижні Пн–Нд від тижня, у якому 1 січня, як у
# ПС_2024 і prod/ps_calendar); рядок I, II, III кварталу стоїть після «за період» 13-го, 26-го й 39-го тижня, IV — після
# «за період» останнього дня року (31.12): так у книзі служби 2026 (рядки 104, 209, 314, 424).
_QUARTER_WEEKS = {13: "I", 26: "II", 39: "III"}


def _ps_week(day):
    """Номер тижня року ПС (1…53): тиждень Пн–Нд належить року своєї неділі, 1-й — той, у якому 1 січня."""
    monday = day - dt.timedelta(days=day.weekday())
    jan1 = dt.date((monday + dt.timedelta(days=6)).year, 1, 1)
    return (monday - (jan1 - dt.timedelta(days=jan1.weekday()))).days // 7 + 1


def _layout(first):
    """Рядки аркуша «Пункти харчування» від дня `first` по черзі, як їх копіює Excel: ("day", дата, None); після неділі й
    після 31.12 — ("week", перший день тижня на аркуші, останній); після 13-го, 26-го, 39-го тижня року ПС —
    ("quarter", «I» / «II» / «III», неділя), після 31.12 — ще ("quarter", «IV», 31.12)."""
    day = first
    while True:
        yield "day", day, None
        year_end = (day.month, day.day) == (12, 31)
        if day.weekday() == 6 or year_end:
            yield "week", max(day - dt.timedelta(days=day.weekday()), dt.date(day.year, 1, 1)), day
            quarter = "IV" if year_end else _QUARTER_WEEKS.get(_ps_week(day))
            if quarter:
                yield "quarter", quarter, day
        day += dt.timedelta(days=1)


def _quarter_start(label, last):
    """Перший день кварталу `label` (I…IV), що закінчується днем `last`: квартал — 13 тижнів року ПС, I — від 01.01."""
    if label == "IV":
        jan1 = dt.date(last.year, 1, 1)
        return jan1 - dt.timedelta(days=jan1.weekday()) + dt.timedelta(weeks=39)
    return max(last - dt.timedelta(days=13 * 7 - 1), dt.date(last.year, 1, 1))


def _off_layout(rows):
    """Перший рядок підсумку, що не тримає суми своїх днів: (номер рядка, що там мало бути, сума днів, що в рядку), інакше
    None. «За період» тижня, вставленого цілим, дорівнює сумі його днів, вставленого частково — не менший за неї (у книзі
    служби 2026 так усі 477 тижневих клітинок); рядок кварталу служба вписує наприкінці кварталу — порожній чи нуль
    не перевіряється, інакше не менший за суму вставлених днів кварталу. Так видно стовпець не за будовою аркуша (рядок
    підсумку зсунуто чи його немає), і вставка не зсуває днів мовчки."""
    days = {a: (n or 0) for _, kind, a, _, _, n in rows if kind == "day"}
    for no, kind, a, b, text, n in rows:
        if kind == "week":
            span = [a + dt.timedelta(days=i) for i in range((b - a).days + 1)]
            what = f"підсумок тижня «за період» {_dmy(a.isoformat())}–{_dmy(b.isoformat())}"
        elif kind == "quarter":
            first = _quarter_start(a, b)
            span = [first + dt.timedelta(days=i) for i in range((b - first).days + 1)]
            what = f"підсумок «{a} квартал»"
        else:
            continue
        have = [d for d in span if d in days]
        total, v = sum(days[d] for d in have), n or 0
        whole = kind == "week" and len(have) == len(span)
        if (whole and v != total) or (v < total and (kind == "week" or v)):
            return no, what, total, text
    return None


def _column(lines, out, first, end):
    """Стовпець чисел пункту з дня `first`. Без підписів підсумків це стовпець аркуша «Пункти харчування»: рядки йдуть за
    будовою аркуша (_layout), тож «за період» після неділі й рядок кварталу пропускаються за місцем, а не за числом (і з
    нулем у них), і кожен пропущений названо в `problems`. Рядок підсумку, що не тримає суми своїх днів (_off_layout), —
    стовпець не за будовою аркуша: вставка не береться зовсім, слова називають рядок. Підписані підсумки («за період …»,
    «разом», «… квартал») — інший шматок: їх пропускає підпис, а дні йдуть поспіль. Порожній рядок — день без числа (не
    зсув); дні після `end` не беруться; чого база не прийме, — слова з номером рядка."""
    labelled = any(_subtotal(r) for r in lines)
    slots = _layout(dt.date.fromisoformat(first))
    top = dt.date(state_db.DATE_YEARS[1], 12, 31)
    rows = []
    for no, r in enumerate(lines, 1):
        text = r[0].strip() if r else ""
        if labelled:
            if _subtotal(r):
                _say(out, f"Рядок {no}: {' '.join(c.strip() for c in r if c.strip())} — підсумок, пропущено.")
                continue
            kind, day, last = next(s for s in slots if s[0] == "day")
        else:
            kind, day, last = next(slots)
        if kind == "day" and day > top:
            _say(out, f"Дні після {_dmy(top.isoformat())} — поза роками обліку: не взято.")
            break
        rows.append((no, kind, day, last, text, _num(text) if text else None))
    off = _off_layout(rows)
    if off:
        no, what, total, text = off
        out["problems"][:] = [f"Рядок {no}: за будовою аркуша «Пункти харчування» тут {what} (сума днів у стовпці — {total}), "
                              f"а в рядку {text or 'порожньо'}. Стовпець не за будовою аркуша: вставте стовпець пункту від "
                              "першого дня періоду з рядками «за період» або два стовпці «дата | число»."]
        return
    beyond = 0
    for no, kind, day, last, text, n in rows:
        if kind != "day":
            what = (f"підсумок тижня «за період» ({_dmy(day.isoformat())}–{_dmy(last.isoformat())})" if kind == "week"
                    else f"підсумок «{day} квартал»")
            _say(out, f"Рядок {no}: {text or 'порожньо'} — {what}, пропущено.")
            continue
        if end and day.isoformat() > end:
            beyond += 1
            continue
        if n is None:
            continue
        why = _refused(n)
        if why:
            _say(out, f"Рядок {no}: {n} — {why}, не взято.")
            continue
        out["values"][day.isoformat()] = n
    if beyond:
        _say(out, f"Днів після кінця періоду ({_dmy(end)}): {beyond} — не взято.")


def _pairs(lines, out, year):
    """«дата | число» — рівно два стовпці. Рядок без дати чи підписаний підсумок («за період») — не день, пропускається;
    рядок, де після числа ще стовпці (скажімо, «дата | разом | пункт»), — відмова з номером рядка: інакше «разом» мовчки
    ставав би д/д пункту. Повертає True, якщо вставка — такий шматок."""
    seen = [(no, r) for no, r in enumerate(lines, 1) if any(c.strip() for c in r)]
    # Рядок без дати (підсумок «за період»: дата в ньому порожня) — не день і нічого не зсуває: форму шматка кажуть
    # рядки з датою.
    dated = [(no, r, *_parse_day(r[0], year)) for no, r in seen if not _subtotal(r) and r[0].strip()]
    if not dated or not all((d or why) and len(r) >= 2 for _, r, d, why in dated):
        return False
    wide = [(no, r) for no, r, _, _ in dated if any(c.strip() for c in r[2:])]
    odd = [(no, r, why) for no, r, d, why in dated if why]
    if wide:
        no, r = wide[0]
        shown = " | ".join(c.strip() for c in r if c.strip())
        out["problems"].append(f"Рядок {no} «{shown}»: більше двох стовпців. Вставте два стовпці «дата | число» "
                               "або стовпець чисел пункту.")
        return True
    if odd:
        for no, r, why in odd:
            _odd_date(out, r[0], why)
        return True
    values = {}
    for no, r, d, _ in dated:
        n = _num(r[1])
        if n is None:
            if r[1].strip():
                _say(out, f"Рядок {no}: «{r[1].strip()}» — не число, не взято.")
            continue
        why = _refused(n)
        if why:
            _say(out, f"Рядок {no}: {n} — {why}, не взято.")
            continue
        values[d] = n
    out["kind"], out["values"] = "pairs", values
    return True


def parse_paste(text, start=None, end=None):
    """Вставлене з Excel (рядки — переноси, клітинки — табуляції) → те саме, що read_file, або шматок одного пункту:
    «дата | число» чи стовпець чисел із дня `start` (не далі `end`)."""
    out = _empty()
    first, why = _parse_day(start) if start else (None, None)
    if start and not first:
        _say(out, f"Перший день періоду «{_text(start)}» — не дата або рік з опискою.")
    last = _day(end) if end else None
    year = int(first[:4]) if first else None
    lines = [line.split("\t") for line in str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    filled = [i for i, r in enumerate(lines) if any(c.strip() for c in r)]
    lines = lines[filled[0]:filled[-1] + 1] if filled else []    # порожні рядки — лише по краях
    rows = [r for r in lines if any(c.strip() for c in r)]
    for title, fn, what in SHEETS:
        if fn(rows, out, year):
            out["kind"] = what
            if not out[what]:
                out["problems"].append(_NOTHING[what] + ("" if year else _NO_YEAR))
            return out
    if _sukhpai(rows, out, year):                            # аркуш «Сухпай» — для актів сухпаю
        out["kind"] = "sukhpai"
        if not out["sukhpai"]:                               # дні аркуша — дати з роком, а вікно акта сухпаю періоду не шле:
            out["problems"].append(_NOTHING["sukhpai"])      # поради «рік візьметься з періоду» тут немає
        return out
    if _pairs(lines, out, year):
        return out
    # Стовпець чисел поспіль: порожній рядок усередині — день без числа, рядок підсумку — не день.
    body = [r for r in lines if not _subtotal(r)]
    if body and all(len(r) == 1 or not any(c.strip() for c in r[1:]) for r in body) \
            and all(_num(r[0]) is not None or not r[0].strip() for r in body):
        if not first:
            _say(out, "Стовпець чисел без дат: вкажіть перший день періоду.")
            return out
        out["kind"] = "column"
        _column(lines, out, first, last)
        return out
    _say(out, _UNKNOWN)
    return out


def signers_for(signers, day, group):
    """Підписанти акта на дату `day` для пункту `group`: (дата рядка, [{role, sort, pos, rank, name}])."""
    dates = sorted(d for d in signers if d <= str(day))
    if not dates:
        return None, []
    rec, grp, out = signers[dates[-1]], _low(group), []
    for role, sort, src, labels in ROLES:
        if src == "none":
            continue
        g = COMMON if src == "common" else grp
        hit = next((rec[f"{g}|{lab}"] for lab in labels if f"{g}|{lab}" in rec), None)
        if hit:
            out.append({"role": role, "sort": sort, **hit})
    return dates[-1], out


def newest(folder):
    """Найсвіжіша за часом збереження «Добовка*.xlsx» теки; None — теки чи файла немає."""
    if not folder or not os.path.isdir(folder):
        return None
    files = [p for p in glob.glob(os.path.join(glob.escape(folder), "Добовка*.xlsx"))
             if not os.path.basename(p).startswith("~$")]

    def saved(p):
        try:
            return os.path.getmtime(p)
        except OSError:                                          # файл зник між пошуком і читанням
            return -1.0

    return max(files, key=saved) if files else None


# ----------------------------------------------------------------------------- аркуш «Сухпай»
# Рядок груп (батальйони: «Упр», «1 БАТ» …) над першою колонкою групи, рядок підрозділів («період», «дата», «№ витягу»…,
# далі підрозділ над першою колонкою своєї трійки), рядок полів — «осіб», «рапорт», «акт» (у «усього» — лише «осіб»);
# далі рядок на день (дата в колонці «дата»), між тижнями — «за період …» без дати. Підрозділ — «група підрозділ»
# («2 БАТ 4 рота»): той самий підрозділ («РВП», «ВМТЗ») є в кількох батальйонах. «Усього» — сума інших, не читається.

def _sukhpai(rows, out, year=None):
    """Рядки аркуша «Сухпай» → out["sukhpai"] = {підрозділ: {"block", "sub", "days": {дата: [осіб, рапорт, акт]}}}
    (день — лише з особами, рапортом чи актом); True, якщо шапку знайдено. Шапка — за мітками, а не за місцем: рядок
    полів з «осіб» під рядком підрозділів з «дата» (графи «рапорт» може й не бути: старі копії Добовки її не мали); поля
    підрозділу — «осіб», «рапорт», «акт» у його графах (до наступного підрозділу чи групи), у будь-якому порядку; нова
    група починає підрозділи наново, група без підрозділу — сама собі підрозділ."""
    at = next((i for i, r in enumerate(rows) if i >= 2 and "осіб" in [_low(c) for c in r]
               and "дата" in [_low(c) for c in rows[i - 1]]), None)
    if at is None:
        return False
    groups, subs, fields = rows[at - 2], rows[at - 1], rows[at]
    date_col = next(i for i, v in enumerate(subs) if _low(v) == "дата")
    first = next(i for i, v in enumerate(fields) if _low(v) == "осіб")
    cols, group, sub = {}, "", ""
    for c in range(first, len(fields)):
        if c < len(groups) and _text(groups[c]):
            group, sub = _text(groups[c]), ""                # нова група — підрозділи наново
        if c < len(subs) and _text(subs[c]):
            sub = _text(subs[c])
        field = _low(fields[c])
        if field not in ("осіб", "рапорт", "акт") or not group or _low(sub) == "усього":
            continue
        unit = cols.setdefault(f"{group} {sub}" if sub else group, {"block": group, "sub": sub or group})
        unit.setdefault(field, c)
    found = {}
    for r in rows[at + 1:]:
        day = _day(r[date_col], year) if date_col < len(r) else None
        if not day:
            continue
        for label, c in cols.items():
            if "осіб" not in c:
                continue
            n = _num(r[c["осіб"]]) if c["осіб"] < len(r) else None
            rep = _text(r[c["рапорт"]]) if "рапорт" in c and c["рапорт"] < len(r) else ""
            act = _text(r[c["акт"]]) if "акт" in c and c["акт"] < len(r) else ""
            if n or rep or act:
                unit = found.setdefault(label, {"block": c["block"], "sub": c["sub"], "days": {}})
                unit["days"][day] = [n or 0, rep, act]
    out["sukhpai"].update((label, found[label]) for label in cols if label in found)     # порядок колонок аркуша
    return True


def read_sukhpai(data):
    """Аркуш «Сухпай» книги Добовки (байти .xlsx) → {"sukhpai": {…}, "problems": […]}."""
    import openpyxl                                              # noqa: PLC0415
    out = {"sukhpai": {}, "problems": []}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                       # noqa: BLE001 — будь-який зіпсований файл
        print(f"Добовка: файл не прочитано — {type(e).__name__}: {e}")
        out["problems"].append("Файл не є книгою Excel (.xlsx).")
        return out
    try:
        if "Сухпай" not in wb.sheetnames:
            out["problems"].append("Немає аркуша «Сухпай».")
            return out
        rows = [list(r) for r in wb["Сухпай"].iter_rows(values_only=True)]
    finally:
        wb.close()
    if not _sukhpai(rows, out):
        out["problems"].append("Аркуш «Сухпай»: шапку не впізнано.")
    elif not out["sukhpai"]:
        out["problems"].append(_NOTHING["sukhpai"])
    return out


def _acts_in(cell):
    """Клітинка «акт» → номери актів: «СП2026-116,СП2026-104» → [СП2026-116, СП2026-104]."""
    return [x for x in (" ".join(p.split()) for p in re.split(r"[,;\n]+", str(cell or ""))) if x]


def report_days(sukhpai, report):
    """Дні рапорту на аркуші «Сухпай»: номер рапорту — окремим числом у клітинці «рапорт» («Вх.№41207 від 23.09.2026»;
    41207 — не 4120 і не 41207/1). → {"label", "block", "sub", "days": [[дата, осіб]] (дні з особами), "acts": номери
    актів, уже вписаних у ці дні, "reportDate": дата рапорту з клітинки (чи ""), "text": клітинка, "others": підрозділи,
    де той самий рапорт ще є} або None. Рапорт у кількох підрозділах — перший за порядком колонок аркуша."""
    no = " ".join(str(report or "").split()).lstrip("№").strip()
    if not no:
        return None
    pat = re.compile(r"(?<![\d/])" + re.escape(no) + r"(?![\d/])")
    hits = []
    for label, unit in sukhpai.items():
        days = sorted((d, v) for d, v in unit["days"].items() if v[1] and pat.search(v[1]))
        if days:
            hits.append((label, unit, days))
    if not hits:
        return None
    label, unit, days = hits[0]
    acts = []
    for _, v in days:
        acts += [a for a in _acts_in(v[2]) if a not in acts]
    m = re.search(r"від\s*(\d{1,2}[./]\d{1,2}[./]\d{2,4})", days[0][1][1])
    return {"label": label, "block": unit["block"], "sub": unit["sub"], "days": [[d, v[0]] for d, v in days if v[0]],
            "acts": acts, "reportDate": (_day(m.group(1)) or "") if m else "", "text": days[0][1][1],
            "others": [x[0] for x in hits[1:]]}
