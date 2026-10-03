# -*- coding: utf-8 -*-
"""Книга обліку наявності та руху військового майна — Додаток 47 до
Інструкції з обліку військового майна у Збройних Силах України (п. 12 розд. IV).

Розкладка сторінки — як у бланку книги служби («Додаток 47 v1.6.0», шаблон
альбомної сторінки): над таблицею найменування й код, під ним нормативний
запас; далі графи 1–7 — дата запису, найменування, номер і дата документа,
постачальник (одержувач), надійшло, вибуло; «перебуває згідно з документами»
— усього й за категоріями 1–5; і те саме «у тому числі» — на складі й у
кожному підрозділі. Сторінка книги — одна позиція; уся книга — титул, зміст
і по аркушу на позицію.

У графах складу й підрозділів записується залишок лише там, де його змінив
цей документ, — як у паперовій книзі: чинний залишок — останнє число графи.
Категорії служба не веде, тож їхні графи лишаються порожніми.

spec: {"file", "unit": "А0000", "service": "продовольча служба",
       "book": true — титул і зміст перед сторінками,
       "items": [{"code", "name", "uom", "price", "prices" — «3 × 1 250,00; 2 × 980,00», коли партій кілька,
                  "blocks": ["склад", "5 б ТрО"],
                  "rows": [{"rec", "type", "no", "date", "party",
                            "in", "out", "total", "cells": [залишок | null]}]}]}
"""
import math
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

FONT = "Times New Roman"
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CATS = 5                   # категорії (сорти) 1–5
LEFT = 7                   # графи 1–7: дата запису … вибуло
TOTAL_COL = LEFT + 1       # «перебуває згідно з документами»
BLOCK_COL = TOTAL_COL + 1 + CATS
HEAD_TOP, HEAD_NUM = 4, 8  # шапка таблиці й рядок з номерами граф
FIRST = HEAD_NUM + 1
FMT_INT, FMT_QTY, FMT_DATE = "#,##0", "#,##0.###", "dd.mm.yyyy"
WIDTHS = [10.5, 15, 14, 10.5, 22, 7.5, 7.5]
W_TOTAL, W_CAT = 7.5, 3.3
COLS_PER_PAGE = 26         # стільки граф уміщує альбомний А4 без дрібного шрифту


def F(size=10, bold=False, italic=False):
    return Font(name=FONT, size=size, bold=bold, italic=italic)


UP = Alignment(horizontal="center", vertical="center", wrap_text=True, text_rotation=90)
MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOP_L = Alignment(horizontal="left", vertical="top", wrap_text=True)
TOP_C = Alignment(horizontal="center", vertical="top", wrap_text=True)
TOP_R = Alignment(horizontal="right", vertical="top")


def _date(iso):
    try:
        return datetime.strptime(iso or "", "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


def _qty(cell, v):
    if v is None or v == "":
        return
    cell.value = v
    cell.number_format = FMT_INT if float(v).is_integer() else FMT_QTY
    cell.alignment = TOP_R


def _box(ws, r1, c1, r2, c2, text=None, font=None, align=MID):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = BOX
    cell = ws.cell(row=r1, column=c1)
    if text is not None:
        cell.value = text
    cell.font, cell.alignment = font or F(9), align
    if r2 > r1 or c2 > c1:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _stock_head(ws, col, title):
    """Графи «усього» й «з них за категоріями (сортами)» 1–5; назва — над ними."""
    if title is not None:
        _box(ws, 5, col, 5, col + CATS, title, F(9, bold=True))
    _box(ws, 6, col, 7, col, "усього", align=UP)
    _box(ws, 6, col + 1, 6, col + CATS, "з них за категоріями (сортами)", F(8))
    for k in range(1, CATS + 1):
        _box(ws, 7, col + k, 7, col + k, k, F(8))


def _line(ws, r, c1, c2, text, font, align):
    ws.cell(row=r, column=c1, value=text)
    ws.cell(row=r, column=c1).font = font
    ws.cell(row=r, column=c1).alignment = align
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)


def _fits(widths, first, last, text, size=11):
    """Чи вміститься рядок у графи з `first` по `last`.

    Ширина графи в Excel — приблизно стільки символів шрифту за замовчуванням.
    Для меншого кегля символів входить більше, для більшого — менше.
    """
    room = sum(widths[c - 1] for c in range(first, last + 1)) * (11.0 / size)
    return room >= len(text) + 1


def _span_for(widths, last, text, size, keep_left):
    """Скільки граф ліворуч віддати заголовку, щоб він не обрізався."""
    first = last
    while first > keep_left and not _fits(widths, first, last, text, size):
        first -= 1
    return first


def item_page(ws, item: dict) -> None:
    blocks = item.get("blocks") or []
    ncols = BLOCK_COL - 1 + len(blocks) * (1 + CATS)
    widths = WIDTHS + [W_TOTAL] + [W_CAT] * CATS + [W_TOTAL, *[W_CAT] * CATS] * len(blocks)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Над таблицею: найменування з кодом і нормативний запас — як у бланку.
    # Область під код беремо таку, щоб напис не обрізало: у найширшому варіанті
    # книги праворуч стоять найвужчі графи категорій.
    code_text = f"Код номенклатури  {item.get('code') or ''}"
    code_from = _span_for(widths, ncols, code_text, 11, LEFT + 1)
    name = item.get("name") or item.get("code") or ""
    _line(ws, 1, 1, code_from - 1, name, F(12, bold=True), Alignment(horizontal="center"))
    for c in range(1, code_from):
        ws.cell(row=1, column=c).border = Border(bottom=THIN)
    _line(ws, 1, code_from, ncols, code_text, F(11), Alignment(horizontal="right"))
    _line(ws, 2, 1, code_from - 1, "(найменування військового майна, індекс, номер креслення)",
          F(8, italic=True), Alignment(horizontal="center", vertical="top"))
    extra = " · ".join(x for x in (
        f"одиниця виміру: {item['uom']}" if item.get("uom") else "",
        f"ціни партій: {item['prices']} грн" if item.get("prices")
        else f"облікова ціна: {item['price']:,.2f} грн".replace(",", " ").replace(".", ",")
        if item.get("price") else "") if x)
    _line(ws, 3, 1, LEFT, "Нормативний запас: мінімальний ________, максимальний ________",
          F(10), Alignment(horizontal="left"))
    if extra:
        _line(ws, 3, LEFT + 1, ncols, extra, F(9), Alignment(horizontal="right"))

    for c, text in enumerate(("Дата запису", "Найменування документа", "Номер документа",
                              "Дата документа", "Постачальник (одержувач)", "Надійшло",
                              "Вибуло"), 1):
        _box(ws, HEAD_TOP, c, 7, c, text, align=UP)
    _box(ws, HEAD_TOP, TOTAL_COL, 5, TOTAL_COL + CATS, "Перебуває згідно з документами",
         F(9, bold=True))
    _stock_head(ws, TOTAL_COL, None)
    head_h = 18
    if blocks:
        over = "У тому числі на складі (у підрозділах, військових частинах)"
        _box(ws, HEAD_TOP, BLOCK_COL, HEAD_TOP, ncols, over, F(9, bold=True))
        # Один-два блоки — місця вгорі мало: пускаємо заголовок у кілька рядків
        # і піднімаємо висоту, інакше Excel обріже його на друці.
        room = sum(widths[c - 1] for c in range(BLOCK_COL, ncols + 1)) * (11.0 / 9)
        lines = max(1, -(-len(over) // max(8, int(room))))
        if lines > 1:
            cell = ws.cell(row=HEAD_TOP, column=BLOCK_COL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            head_h = 13 * lines + 5
        for i, b in enumerate(blocks):
            _stock_head(ws, BLOCK_COL + i * (1 + CATS), "На складі" if b == "склад" else b)
    for c in range(1, ncols + 1):
        _box(ws, HEAD_NUM, c, HEAD_NUM, c, c, F(8))
    for r, h in ((4, head_h), (5, 30), (6, 28), (7, 16), (8, 13)):
        ws.row_dimensions[r].height = h

    rows = item.get("rows") or []
    for n, row in enumerate(rows):
        r = FIRST + n
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border, cell.font = BOX, F(10)
        for c, key in ((1, "rec"), (4, "date")):
            when = _date(row.get(key))
            cell = ws.cell(row=r, column=c, value=when or row.get(key) or None)
            cell.number_format, cell.alignment = FMT_DATE, TOP_C
        ws.cell(row=r, column=2, value=row.get("type") or None).alignment = TOP_L
        ws.cell(row=r, column=3, value=row.get("no") or None).alignment = TOP_C
        ws.cell(row=r, column=5, value=row.get("party") or None).alignment = TOP_L
        _qty(ws.cell(row=r, column=6), row.get("in") or None)
        _qty(ws.cell(row=r, column=7), row.get("out") or None)
        _qty(ws.cell(row=r, column=TOTAL_COL), row.get("total", 0))
        for i, v in enumerate(row.get("cells") or []):
            if i < len(blocks):
                _qty(ws.cell(row=r, column=BLOCK_COL + i * (1 + CATS)), v)
    if not rows:
        _box(ws, FIRST, 1, FIRST, ncols, "Руху за документами немає", F(10, italic=True))

    # Друк: альбомний А4, шапка — на кожній сторінці; широка сторінка книги
    # ділиться на аркуші по ширині, і графи 1–5 повторюються зліва на кожному.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = max(1, math.ceil(ncols / COLS_PER_PAGE))
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_title_rows = f"{HEAD_TOP}:{HEAD_NUM}"
    if ncols > COLS_PER_PAGE:
        ws.print_title_cols = "A:E"
    # Координата рядком, не коміркою: у сторінці без граф підрозділів ця
    # комірка лежить в об'єднанні, і openpyxl на об'єкті об'єднаної комірки падав.
    ws.freeze_panes = f"{get_column_letter(LEFT - 1)}{FIRST}"


def title_page(ws, spec: dict) -> None:
    """Титул книги — як аркуш «Звіт» бланка служби."""
    for c in range(1, 13):
        ws.column_dimensions[get_column_letter(c)].width = 11
    right = Alignment(horizontal="left")
    for r, text in ((2, "Додаток 47"), (3, "до Інструкції з обліку військового"),
                    (4, "майна у Збройних Силах України"), (5, "(пункт 12 розділу IV)")):
        _line(ws, r, 9, 12, text, F(11), right)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _line(ws, 9, 1, 12, "КНИГА", F(20, bold=True), center)
    _line(ws, 10, 1, 12, "обліку наявності та руху", F(14, bold=True), center)
    _line(ws, 11, 1, 12, "військового майна (служба забезпечення)", F(14, bold=True), center)
    _line(ws, 15, 3, 10, spec.get("service") or "продовольча служба", F(14, bold=True), center)
    _line(ws, 16, 3, 10, "(служба забезпечення)", F(9, italic=True), center)
    _line(ws, 19, 3, 10, f"військова частина {spec.get('unit') or ''}".strip(), F(14, bold=True),
          center)
    _line(ws, 20, 3, 10, "(військова частина)", F(9, italic=True), center)
    for c in range(3, 11):
        ws.cell(row=15, column=c).border = Border(bottom=THIN)
        ws.cell(row=19, column=c).border = Border(bottom=THIN)
    _line(ws, 26, 8, 12, "Розпочато «____» ______________ 20___ р.", F(12), right)
    _line(ws, 28, 8, 12, "Закінчено «____» ______________ 20___ р.", F(12), right)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def contents_page(ws, items, titles) -> None:
    """Зміст: код майна й аркуш (сторінка) книги, клац — на сторінку."""
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 14
    _line(ws, 1, 1, 4, "ЗМІСТ", F(13, bold=True), Alignment(horizontal="center"))
    for c, text in enumerate(("№ з/п", "Код майна", "Найменування військового майна",
                              "Сторінка книги"), 1):
        _box(ws, 3, c, 3, c, text, F(10, bold=True))
    for n, (it, title) in enumerate(zip(items, titles), 1):
        r = 3 + n
        vals = (n, it.get("code"), it.get("name"), title)
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border, cell.font = BOX, F(10)
            cell.alignment = TOP_L if c == 3 else TOP_C
        link = ws.cell(row=r, column=4)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{title}'!A1", display=title)
        link.font = Font(name=FONT, size=10, color="1F4E9A", underline="single")
    ws.print_title_rows = "3:3"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _title(code, used) -> str:
    base = "".join(ch if ch not in '\\/:*?"<>|[]' else "-" for ch in str(code or "позиція"))[:31]
    title, n = base, 2
    while title in used:
        title = f"{base[:27]} ({n})"
        n += 1
    used.add(title)
    return title


def build_book47(spec: dict) -> Workbook:
    items = spec.get("items") or []
    if not items:
        raise ValueError("немає жодної позиції для книги")
    from excel_export import clean_workbook                     # noqa: PLC0415
    wb = clean_workbook(Workbook())
    wb.remove(wb.active)
    used = {"Титул", "Зміст"}
    titles = [_title(it.get("code"), used) for it in items]
    if spec.get("book"):
        title_page(wb.create_sheet("Титул"), spec)
        contents_page(wb.create_sheet("Зміст"), items, titles)
    for it, title in zip(items, titles):
        item_page(wb.create_sheet(title), it)
    return wb


def save_book47(spec: dict, folder: str) -> str:
    from excel_export import safe_name                              # noqa: PLC0415
    os.makedirs(folder, exist_ok=True)
    name = safe_name(spec.get("file") or "Книга обліку № 47")
    path = os.path.join(folder, name + ".xlsx")
    wb = build_book47(spec)
    try:
        wb.save(path)
    except PermissionError:
        path = os.path.join(folder, f"{name} {datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path
