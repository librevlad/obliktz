# -*- coding: utf-8 -*-
"""Інвентаризація в Excel: описи й протокол за формами наказу Мінфіну від 17.06.2015
№572, акт інвентаризації — так, як його складає служба, а також план проведення,
реєстр паперів, які підкомісія служби подає за наказом про інвентаризацію, і відомості
про її хід для щотижневої доповіді.

Опис необоротних активів має шістнадцять граф, опис запасів — дванадцять; і
текст над таблицею, і підсумки прописом, і підписи — за бланком. Усе, що
залежить від обліку (рядки, суми, прописом), рахує сторінка програми й передає
готовим: цей модуль лише розкладає це на аркуш А4 так, щоб його можна було
підписати без правок. Сторінки теж розкладає він, а не Excel (Pager): блок підписів
лишається на одній сторінці, продовження таблиці починається номерами граф.

    {"kind": "inventory", "file": "...", "unit": "А0000", "legal_name": "Військова частина А0000", "edrpou": "00000000",
     "date": "2026-12-25", "start": "2026-11-01", "end": "2026-12-25",
     "orderDate": "2026-10-01", "orderNo": "100",
     "accFixed": "1116 «…», …", "accStock": "1812 «…», …",
     "head": {"pos", "rank", "name"}, "members": [...], "chief": {...}, "buh": {...},
     "descriptions": [{"sheet", "type": "fixed"|"stock", "title", "where", "mvo": {...}, "cmdr": {...},
                       "rows": [...], "totals": {...}, "words": {...}}],
     "act": {...} | None,
     "orderItem": "3.7", "service": "… служби …", "place": "місце складання",
     "blank": true — робочий опис: фактичну наявність вписує комісія,
     "plan": {"agree": {...}, "approve": {...}, "rows": [{"place", "mvo", "kinds", "date", "who", "mark"}]} | None,
     "protocol": {"approve": {...}, "date", "totals": {"fixedQty", "fixedSum", "stockQty", "stockSum"},
                  "rows": [{"name", "mvo", "qty", "sum", "reason", "loss": true — втрата за рапортом}]} | None,
     "register": {"fes": {...}, "rows": [{"doc", "date", "who", "sheets", "note"}]} | None,
     "progress": {"date", "lines": [підсумок словами], "totals": {"subs", "done", "descs", "lines", "sum", "diffs"},
                  "rows": [{"sub", "mvo", "descs", "lines", "sum", "plan", "done", "diffs", "who"}]} | None}
"""
import math
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

FONT = "Times New Roman"
THIN = Side(style="thin", color="000000")
MED = Side(style="medium", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINE = Border(bottom=THIN)
F = lambda size=10, bold=False, italic=False: Font(name=FONT, size=size, bold=bold, italic=italic)  # noqa: E731
A_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
A_LT = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_R = Alignment(horizontal="right", vertical="center", wrap_text=True)
A_J = Alignment(horizontal="justify", vertical="top", wrap_text=True)
FMT_QTY = "#,##0.###"
FMT_INT = "0"
FMT_MONEY = "#,##0.00"
FMT_PRICE = "#,##0.00###"
PAPER = (595.28, 841.89)          # А4 у пунктах
CHAR_PT = 5.6                     # одиниця ширини графи на папері
ROW_PT = 15.0                     # рядок, якому висоту не задано

MONTHS = ["січня", "лютого", "березня", "квітня", "травня", "червня", "липня",
          "серпня", "вересня", "жовтня", "листопада", "грудня"]

# Графи бланка: (заголовок групи, заголовок графи, ширина в символах).
FIXED = [
    ("", "№\nз/п", 5), ("", "Найменування, стисла характеристика та призначення об’єкта", 30),
    ("", "Рік випуску (будівництва) чи дата придбання (введення в експлуатацію) та виготовлювач", 13),
    ("Номер", "інвентарний/ номенклатурний", 11), ("Номер", "заводський", 10), ("Номер", "паспорта", 8),
    ("", "Один. вимір.", 7),
    ("Фактична наявність", "кількість", 9), ("Фактична наявність", "первісна (переоцінена) вартість", 13),
    ("", "Відмітка про вибуття", 9),
    ("За даними бухгалтерського обліку", "кількість", 9),
    ("За даними бухгалтерського обліку", "первісна (переоцінена) вартість", 13),
    ("За даними бухгалтерського обліку", "сума зносу (накопиченої амортизації)", 11),
    ("За даними бухгалтерського обліку", "балансова вартість", 11),
    ("За даними бухгалтерського обліку", "строк корисного використання", 10),
    ("", "Інші відомості", 16),
]
STOCK = [
    ("", "№\nз/п", 5), ("", "Рахунок, субрахунок", 10),
    ("Матеріальні цінності", "найменування, вид, сорт, група", 34),
    ("Матеріальні цінності", "номенклатурний номер (за наявності)", 12),
    ("", "Одиниця виміру", 8),
    ("Фактична наявність", "кількість", 10), ("Фактична наявність", "вартість", 11),
    ("Фактична наявність", "сума", 13),
    ("За даними бухгалтерського обліку", "кількість", 10),
    ("За даними бухгалтерського обліку", "вартість", 11),
    ("За даними бухгалтерського обліку", "сума", 13),
    ("", "Інші відомості", 16),
]


def date_uk(iso: str, year_suffix: str = "р.") -> str:
    """«25» грудня 2026р. — як у бланку."""
    try:
        d = datetime.strptime(iso or "", "%Y-%m-%d")
    except ValueError:
        return "«___» ____________ 20__р."
    return f"«{d.day}» {MONTHS[d.month - 1]} {d.year}{year_suffix}"


def date_words(iso: str) -> str:
    """25 грудня 2026 року — у тексті акта."""
    try:
        d = datetime.strptime(iso or "", "%Y-%m-%d")
    except ValueError:
        return "___ ____________ 20__ року"
    return f"{d.day} {MONTHS[d.month - 1]} {d.year} року"


def date_dots(iso: str) -> str:
    """25.12.2026 — у таблицях плану й реєстру; порожнє, якщо дати немає."""
    try:
        return datetime.strptime(iso or "", "%Y-%m-%d").strftime("%d.%m.%Y")
    except ValueError:
        return ""


def unit_gen(spec: dict) -> str:
    return f"військової частини {spec['unit']}" if spec.get("unit") else "військової частини"


def order_ref(spec: dict, item: bool = True) -> str:
    """«від 24.09.2026 № 4773, пункт 3.7» — посилання на наказ у плані й реєстрі."""
    out = f"від {date_dots(spec.get('orderDate')) or '__.__.____'} № {spec.get('orderNo') or '_____'}"
    if item and spec.get("orderItem"):
        out += f", пункт {spec['orderItem']}"
    return out


def rank_name(p: dict) -> str:
    """«майор ПЕТРЕНКО Т. Г.» — у складі комісії."""
    p = p or {}
    return " ".join(x for x in ((p.get("rank") or "").strip(), (p.get("name") or "").strip()) if x)


def blank_date(spec: dict, suffix: str = "р.") -> str:
    year = str(spec.get("end") or spec.get("date") or "")[:4] or "20__"
    return f"«____» ______________ {year} {suffix}"


def person_line(p: dict) -> str:
    """«Посада, звання» — підпис у бланку; порожнє, якщо людини не вказано."""
    p = p or {}
    pos = (p.get("pos") or "").strip().rstrip(",")
    rank = (p.get("rank") or "").strip()
    return ", ".join(x for x in (pos, rank) if x)


class Sheet:
    """Аркуш із розміткою по графах: пише текст у злиті діапазони й підбирає
    висоту рядка під довгий текст — сам Excel злитих клітинок не розтягує."""

    def __init__(self, ws, widths):
        self.ws = ws
        self.widths = widths
        self.row = 1
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def span_width(self, c1, c2):
        return sum(self.widths[c1 - 1:c2])

    def put(self, r, c1, c2, text, font=None, align=A_L, border=None, fit=True, fmt=None):
        ws = self.ws
        cell = ws.cell(row=r, column=c1, value=text)
        cell.font = font or F()
        cell.alignment = align
        if fmt:
            cell.number_format = fmt
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        if border:
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = border
        if fit and isinstance(text, str) and text:
            self.fit(r, text, self.span_width(c1, c2), (font or F()).sz or 10)
        return cell

    @staticmethod
    def need(text, width_chars, size, tight=0.8):
        """Висота рядка під текст у клітинці такої ширини."""
        # Ширина колонки в Excel — у символах шрифту 11 пт; дрібніший шрифт вміщує більше.
        # Запас на перенос за словами й на широкі великі літери: краще зайвий рядок висоти,
        # ніж обрізана назва («Шафа (скриня) холодильна…» займала три рядки, а не два).
        per_line = max(1.0, width_chars * tight * 11 / size)
        lines = sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))
        return lines * size * 1.32 + 2

    def fit(self, r, text, width_chars, size, tight=0.8):
        cur = self.ws.row_dimensions[r].height or ROW_PT
        self.ws.row_dimensions[r].height = max(cur, self.need(text, width_chars, size, tight))

    def sign(self, r, left_label, person, c_label, c_pos, c_sign, c_name, date_text=None, name_key="name"):
        """Рядок підпису: [мітка або дата] посада | підпис | ім'я, під ними — дрібні пояснення."""
        if left_label:
            self.put(r, 1, c_label, left_label, F(10), A_L)
        elif date_text is not None:
            self.put(r, 1, c_label, date_text, F(10), A_C, border=LINE)
        self.put(r, c_pos[0], c_pos[1], person_line(person), F(10), Alignment(
            horizontal="center", vertical="bottom", wrap_text=True), border=LINE)
        self.put(r, c_sign[0], c_sign[1], "", F(10), A_C, border=LINE, fit=False)
        self.put(r, c_name[0], c_name[1], (person or {}).get(name_key) or "", F(10),
                 Alignment(horizontal="center", vertical="bottom", wrap_text=True), border=LINE)
        self.put(r + 1, c_pos[0], c_pos[1], "(посада)", F(7), A_C, fit=False)
        self.put(r + 1, c_sign[0], c_sign[1], "(підпис)", F(7), A_C, fit=False)
        self.put(r + 1, c_name[0], c_name[1], "(ініціали, прізвище)", F(7), A_C, fit=False)
        self.ws.row_dimensions[r + 1].height = 10
        return r + 2


class Pager:
    """Сторінки розкладає програма, а не Excel: рядки, які мають лишитися разом (підпис і
    пояснення під ним, шапка таблиці з першим рядком), не рвуться між сторінками, а
    продовження таблиці починається рядком із номерами граф. Кожен рядок аркуша проходить
    через нього по порядку й дістає задану висоту — тоді сторінка в Excel така сама, як тут.
    Орієнтацію й поля аркуша задають до нього."""

    def __init__(self, ws, widths):
        self.ws = ws
        m = ws.page_margins
        w, h = PAPER[::-1] if ws.page_setup.orientation == "landscape" else PAPER
        # «Одна сторінка завширшки»: Excel зменшує аркуш до цілого відсотка, за якого графи вміщуються.
        scale = min(100, int(100 * (w - (m.left + m.right) * 72) / (sum(widths) * CHAR_PT))) / 100
        # Запас 3 %: висоту рядка Excel округлює до точок принтера, масштаб — до відсотка.
        self.room = (h - (m.top + m.bottom) * 72) / scale * 0.97
        self.y = 0.0
        self.done = 0

    def height(self, r):
        return _height(self.ws, r)

    def fits(self, need):
        return not self.y or self.y + need <= self.room

    def page(self, r):
        """Нова сторінка починається з рядка r."""
        self.flow(r - 1)
        self.ws.row_breaks.append(Break(id=r - 1))
        self.y = 0.0

    def flow(self, upto):
        """Рядки до upto включно: між ними сторінку рвати можна."""
        for r in range(self.done + 1, upto + 1):
            self.keep(r, r)

    def keep(self, first, last, extra=0.0):
        """Рядки, які лишаються на одній сторінці; extra — висота рядка, який стане за ними
        й має бути на тій самій сторінці."""
        self.flow(first - 1)
        need = sum(self.height(r) for r in range(first, last + 1))
        if not self.fits(need + extra):
            self.page(first)
        self.y += need
        self.done = last

    def table(self, r, need, head, extra=0.0):
        """Рядок таблиці висотою need. Повертає рядок аркуша, у який його писати: якщо
        сторінка скінчилась, нова починається рядком із номерами граф (його пише head).
        extra — висота того, що стане за рядком і має бути з ним на одній сторінці."""
        self.flow(r - 1)
        if not self.fits(need + extra):
            self.page(r)
            head(r)
            self.y += self.height(r)
            r += 1
        self.ws.row_dimensions[r].height = need
        self.y += need
        self.done = r
        return r


def _height(ws, r):
    """Висота рядка; якої не задано — задає за шрифтом: в Excel вона така сама, як у розкладці."""
    dim = ws.row_dimensions[r]
    if dim.height is None:
        sizes = [c.font.sz or 11 for c in ws[r] if c.value is not None]
        dim.height = max([ROW_PT] + [sz * 1.32 + 2 for sz in sizes])
    return dim.height


def _dry(wb, widths, write):
    """Пише блок на чернетковий аркуш, щоб знати висоту його рядків наперед: чи стане він на
    сторінку разом з останнім рядком таблиці. Повертає висоти рядків і те, що повернув write."""
    probe = wb.create_sheet("_")
    try:
        out = write(Sheet(probe, widths), 1)
        end = out[0] if isinstance(out, tuple) else out
        return [_height(probe, r) for r in range(1, end + 1)], out
    finally:
        wb.remove(probe)


def _qty(cell, v):
    cell.value = v
    if isinstance(v, (int, float)):
        cell.number_format = FMT_INT if float(v).is_integer() else FMT_QTY


def description_sheet(wb, spec: dict, d: dict, used: set):
    """Один інвентаризаційний опис на аркуш."""
    fixed = d.get("type") == "fixed"
    cols = FIXED if fixed else STOCK
    last = len(cols)
    title = _title(d.get("sheet") or ("Необоротні" if fixed else "Запаси"), used)
    ws = wb.create_sheet(title)
    s = Sheet(ws, [w for _, _, w in cols])
    _paper(ws, landscape=True, left=0.4, right=0.4, top=0.5, bottom=0.75)

    # Шапка: ЗАТВЕРДЖЕНО праворуч, установа й код ЄДРПОУ ліворуч.
    right = last - 3
    s.put(1, right, last, "ЗАТВЕРДЖЕНО", F(10), A_L, fit=False)
    s.put(2, right, last, "Наказ Міністерства фінансів України", F(10), A_L, fit=False)
    s.put(3, right, last, "17.06.2015  № 572", F(10), A_L, fit=False)
    s.put(3, 1, 4, spec.get("legal_name") or (f"Військова частина {spec['unit']}" if spec.get("unit") else "Військова частина"),
          F(11), A_C, border=LINE)
    s.put(4, 1, 4, "(установа)", F(7), A_C, fit=False)
    edrpou = str(spec.get("edrpou") or "")
    s.put(5, 1, 2, "Ідентифікаційний код за ЄДРПОУ", F(8), A_L, border=BOX)
    s.put(5, 3, 4, "  ".join(edrpou), F(9), A_C, border=BOX, fit=False)

    s.put(7, 1, last, "ІНВЕНТАРИЗАЦІЙНИЙ ОПИС", F(12, bold=True), A_C)
    s.put(8, 1, last, d.get("title") or "", F(11, bold=True), A_C)
    s.put(9, 1, last, "(основні засоби, нематеріальні активи¹, інші необоротні матеріальні активи, "
          "капітальні інвестиції)" if fixed else "(запаси)", F(7), A_C, fit=False)
    s.put(10, 1, last, date_uk(spec.get("date")), F(10), A_C, fit=False)
    s.put(11, 1, last, "(дата складання)", F(7), A_C, fit=False)

    order = f"від {date_uk(spec.get('orderDate'))}  №{spec.get('orderNo') or '_____'}"
    where = d.get("where") or ""
    if fixed:
        text = (f"На підставі розпорядчого документа {order}, виконано зняття фактичних залишків "
                "основних засобів, нематеріальних активів, інших необоротних матеріальних активів, "
                "капітальних інвестицій (необхідне підкреслити), які обліковуються на субрахунку(ах) "
                f"{spec.get('accFixed') or ''} та зберігаються: {where} станом на {date_uk(spec.get('date'), ' р.')}")
        receipt = ("До початку проведення інвентаризації всі видаткові та прибуткові документи на "
                   "необоротні активи здано в бухгалтерську службу і всі необоротні активи, що надійшли "
                   "на мою відповідальність, оприбутковано, а ті, що вибули, списано.")
    else:
        text = (f"На підставі розпорядчого документа {order} виконано зняття фактичних залишків "
                f"запасів, які обліковуються на субрахунку(ах) {spec.get('accStock') or ''} "
                f"та зберігаються: {where} станом на {date_uk(spec.get('date'), ' р.')}")
        receipt = ("До початку проведення інвентаризації всі видаткові та прибуткові документи на "
                   "матеріальні цінності здано в бухгалтерську службу і всі матеріальні цінності, що "
                   "надійшли на мою відповідальність, оприбутковано, а ті, що вибули, списано.")
    s.put(13, 1, last, text, F(10), A_J)
    s.put(15, 1, last, "Розписка", F(10, bold=True), A_C, fit=False)
    s.put(16, 1, last, "      " + receipt, F(10), A_J)
    s.put(18, 1, last, "Матеріально відповідальна особа:", F(10), A_L, fit=False)
    lay = _sign_cols(last)
    lab = lay[0]
    r = s.sign(19, None, d.get("mvo"), *lay)
    s.put(r + 1, 1, last, f"Інвентаризація: розпочата {date_uk(spec.get('start'))}", F(10), A_L, fit=False)
    s.put(r + 2, 1, last, f"                          закінчена {date_uk(spec.get('end'))}", F(10), A_L,
          fit=False)

    # Таблиця.
    r += 4
    label = r
    s.put(r, 1, last, "При інвентаризації встановлено таке:", F(10), A_L, fit=False)
    h1, h2, h3 = r + 1, r + 2, r + 3
    i = 0
    while i < last:
        group = cols[i][0]
        j = i
        while j + 1 < last and group and cols[j + 1][0] == group:
            j += 1
        if group:
            s.put(h1, i + 1, j + 1, group, F(9), A_C, border=BOX, fit=False)
            for k in range(i, j + 1):
                s.put(h2, k + 1, k + 1, cols[k][1], F(8), A_C, border=BOX, fit=False)
        else:
            s.put(h1, i + 1, i + 1, cols[i][1], F(8), A_C, border=BOX, fit=False)
            ws.merge_cells(start_row=h1, start_column=i + 1, end_row=h2, end_column=i + 1)
            ws.cell(row=h2, column=i + 1).border = BOX
        i = j + 1
    ws.row_dimensions[h1].height = 26
    ws.row_dimensions[h2].height = 64 if fixed else 40

    def numbers(r):
        """Номери граф: під шапкою й першим рядком кожної наступної сторінки таблиці."""
        for k in range(last):
            s.put(r, k + 1, k + 1, k + 1, F(8, bold=True), A_C, border=BOX, fit=False)
        ws.row_dimensions[r].height = ROW_PT

    numbers(h3)
    rows = d.get("rows") or []
    name_w, note_w = s.widths[1 if fixed else 2], s.widths[last - 1]
    needs = [max(ROW_PT, s.need(str(row.get("name") or ""), name_w, 9), s.need(str(row.get("note") or ""), note_w, 9))
             for row in rows]
    pager = Pager(ws, s.widths)
    # Шапка таблиці не лишається внизу сторінки сама: з нею — перший рядок.
    pager.keep(label, h3, extra=needs[0] if needs else ROW_PT)
    # Робочий опис: облік надруковано, фактичну наявність комісія вписує після підрахунку.
    blank = bool(spec.get("blank"))
    rr = h3
    for n, row in enumerate(rows):
        rr = pager.table(rr + 1, needs[n], numbers)
        if blank:
            row = dict(row, fact=None, factSum=None)
        vals = _fixed_values(n + 1, row) if fixed else _stock_values(n + 1, row)
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=c)
            cell.border = BOX
            cell.font = F(9)
            cell.alignment = A_L if (fixed and c in (2, 16)) or (not fixed and c in (3, 12)) else A_C
            if c in _qty_cols(fixed):
                _qty(cell, v)
            elif c in _money_cols(fixed):
                cell.value = v
                if isinstance(v, (int, float)):
                    cell.number_format = FMT_PRICE if (not fixed and c in (7, 10)) else FMT_MONEY
            else:
                cell.value = v if v not in ("", None) else None
    w = d.get("words") or {}
    labels = [("Разом за описом: а) кількість порядкових номерів", w.get("n")),
              ("б) загальна кількість одиниць (фактично)", None if blank else w.get("factQty")),
              ("в) вартість фактична", None if blank else w.get("factSum")),
              ("г) загальна кількість одиниць за даними бухгалтерського обліку", w.get("accQty")),
              ("ґ) вартість за даними бухгалтерського обліку", w.get("accSum"))]
    lab_end = 4 if fixed else 3

    def closing(s, r):
        """Підсумки прописом і підписи — від рядка r. Повертає останній рядок і блоки, які не
        рвуться між сторінками: підсумки прописом, підписи комісії, підписи за майно й облік."""
        first = r
        for text, value in labels:
            # В описі запасів графи вужчі: довгий підпис стає у два рядки, а не обрізається.
            s.put(r, 1, lab_end, text, F(10), A_L, fit=False)
            s.fit(r, text, s.span_width(1, lab_end), 10, tight=1.0)
            s.put(r, lab_end + 1, last - 3, value or "", F(10, italic=True), A_C, border=LINE)
            s.put(r + 1, lab_end + 1, last - 3, "(прописом)", F(7), A_C, fit=False)
            s.ws.row_dimensions[r + 1].height = 10
            r += 2
        words = r - 1
        r += 1
        commission = r
        r = s.sign(r, "Голова комісії", spec.get("head"), *lay)
        for k, m in enumerate(spec.get("members") or []):
            r = s.sign(r, "Члени комісії" if k == 0 else None, m, *lay)
        signed = r - 1
        r += 1
        rest = r
        s.put(r, 1, last, f"Усі цінності, пронумеровані в цьому інвентаризаційному описі з №1 до №{len(rows)}, "
              "перевірено комісією в натурі в моїй присутності та внесено в опис. У зв’язку з цим претензій "
              "до інвентаризаційної комісії не маю. Цінності, перелічені в описі, знаходяться на моєму "
              "відповідальному зберіганні.", F(10), A_J)
        r += 2
        s.put(r, 1, last, "Матеріально відповідальна особа:", F(10), A_L, fit=False)
        r = s.sign(r + 1, None, d.get("mvo"), *lay, date_text=date_uk(spec.get("end")))
        r += 1
        # Наказ про інвентаризацію: опис підписує й командир підрозділу, де її проводили.
        s.put(r, 1, last, "Командир (начальник) підрозділу:", F(10), A_L, fit=False)
        r = s.sign(r + 1, None, d.get("cmdr"), *lay)
        r += 1
        s.put(r, 1, lab, "Інформацію за даними бухгалтерського обліку вніс:", F(10), A_L)
        r = s.sign(r, None, spec.get("buh"), *lay)
        r += 1
        s.put(r, 1, last, "Вказані в цьому описі дані перевірив:", F(10), A_L, fit=False)
        r = s.sign(r + 1, None, spec.get("chief"), *lay, date_text=date_uk(spec.get("end")))
        r += 1
        s.put(r, 1, last, "¹ Графи 11–15 заповнюються бухгалтерською службою." if fixed
              else "Графи 9–11 заповнюються бухгалтерською службою.", F(7), A_L, fit=False)
        return r, [(first, words), (commission, signed), (rest, r)]

    # «Разом» не лишається внизу сторінки без підсумків прописом.
    heights, (_, blocks) = _dry(wb, s.widths, closing)
    tot = pager.table(rr + 1, max(ROW_PT, s.need("Разом", s.widths[0], 9)), numbers,
                      extra=ROW_PT + sum(heights[blocks[0][0] - 1:blocks[0][1]]))
    t = d.get("totals") or {}
    for c in range(1, last + 1):
        cell = ws.cell(row=tot, column=c, value="X")
        cell.font, cell.alignment, cell.border = F(9, bold=True), A_C, BOX
    ws.cell(row=tot, column=1, value="Разом")
    if fixed:
        blanks = (10, 13, 14, 15)
        places = {8: t.get("factQty"), 9: t.get("factSum"), 11: t.get("accQty"), 12: t.get("accSum")}
    else:
        blanks = ()
        places = {6: t.get("factQty"), 8: t.get("factSum"), 9: t.get("accQty"), 11: t.get("accSum")}
    if blank:
        places = {c: (None if c in (8, 9) else v) for c, v in places.items()} if fixed \
            else {c: (None if c in (6, 8) else v) for c, v in places.items()}
    for c in blanks:
        ws.cell(row=tot, column=c).value = None
    for c, v in places.items():
        cell = ws.cell(row=tot, column=c)
        if c in _qty_cols(fixed):
            _qty(cell, v)
        else:
            cell.value = v
            cell.number_format = FMT_MONEY

    r, blocks = closing(s, tot + 2)
    for first, end in blocks:
        pager.keep(first, end)
    ws.print_area = f"A1:{get_column_letter(last)}{r}"
    return ws


def _sign_cols(last):
    """Колонки рядка підпису під ширину таблиці: остання колонка мітки
    («Голова комісії», дата), посада | підпис | ім'я."""
    if last == 16:
        return 2, (3, 7), (9, 12), (13, 16)
    return 3, (4, 7), (8, 9), (10, 12)


def _qty_cols(fixed):
    return {8, 11} if fixed else {6, 9}


def _money_cols(fixed):
    return {9, 12, 14} if fixed else {7, 8, 10, 11}


def _fixed_values(n, r):
    return [n, r.get("name"), r.get("year") or None, r.get("inv") or None, r.get("serial") or None,
            r.get("passport") or None, r.get("uom"), r.get("fact"), r.get("factSum"),
            r.get("mark") or None, r.get("acc"), r.get("accSum"), None, None, None, r.get("note") or None]


def _stock_values(n, r):
    return [n, r.get("account") or None, r.get("name"), r.get("nomen") or None, r.get("uom"),
            r.get("fact"), r.get("price"), r.get("factSum"), r.get("acc"), r.get("price"),
            r.get("accSum"), r.get("note") or None]


def _title(name, used):
    base = "".join("-" if ch in '\\/:*?"<>|[]' else ch for ch in str(name))[:31] or "Опис"
    title, n = base, 2
    while title in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def act_sheet(wb, spec: dict, a: dict, used: set):
    """Акт щорічної інвентаризації — книжковий аркуш, як паперовий акт служби."""
    ws = wb.create_sheet(_title("Акт", used))
    s = Sheet(ws, [14, 14, 14, 14, 14, 14])
    last = 6
    _paper(ws, landscape=False, left=1.0, right=0.6, top=None, bottom=None, pages=False)
    pager = Pager(ws, s.widths)
    ap = a.get("approve") or {}
    s.put(1, 4, 6, "ЗАТВЕРДЖУЮ", F(12), A_L, fit=False)
    s.put(2, 4, 6, ap.get("pos") or "", F(12), A_L)
    s.put(3, 4, 6, f"{ap.get('rank') or ''} ______ {ap.get('name') or ''}".strip(), F(12), A_L)
    s.put(4, 4, 6, date_uk(ap.get("date")), F(12), A_L, fit=False)
    s.put(6, 1, last, "АКТ", F(14, bold=True), A_C, fit=False)
    s.put(7, 1, last, a.get("title") or "", F(14, bold=True), A_C)
    s.put(9, 1, last, "      " + (a.get("intro") or ""), F(12), A_J)
    r = 11
    s.put(r, 1, 2, "Голова комісії:", F(12, bold=True), A_LT, fit=False)
    s.put(r, 3, last, a.get("head") or "", F(12), A_LT)
    r += 1
    s.put(r, 1, 2, "Члени комісії:", F(12, bold=True), A_LT, fit=False)
    s.put(r, 3, last, a.get("members") or "", F(12), A_LT)
    r += 2
    for line in (a.get("basis"), a.get("prev"), a.get("asof")):
        if line:
            s.put(r, 1, last, line, F(12), A_J)
            r += 1
    r += 1
    s.put(r, 1, last, "За результатами проведеної інвентаризації встановлено:", F(12), A_L, fit=False)
    r += 1
    for line in a.get("findings") or []:
        s.put(r, 1, last, line, F(12), A_J)
        r += 1
    r += 1
    for line in (a.get("result") or "").split("\n"):
        if line.strip():
            s.put(r, 1, last, line, F(12), A_J)
            r += 1
    r += 1
    if a.get("attachments"):
        g = r
        s.put(r, 1, last, "До акту додається:", F(12), A_L, fit=False)
        r += 1
        for line in a["attachments"]:
            s.put(r, 1, last, line, F(12), A_J)
            r += 1
        pager.keep(g, r - 1)
        r += 1

    def signer(r, label, p):
        g = r
        if label:
            s.put(r, 1, last, label, F(12, bold=True), A_L, fit=False)
            r += 1
        s.put(r, 1, 4, (p or {}).get("pos") or "", F(12), A_LT)
        r += 1
        s.put(r, 1, 3, (p or {}).get("rank") or "", F(12), A_L, fit=False)
        s.put(r, 4, last, (p or {}).get("name") or "", F(12), A_R, fit=False)
        pager.keep(g, r)
        return r + 2

    r = signer(r, "Голова комісії:", spec.get("head"))
    for k, m in enumerate(spec.get("members") or []):
        r = signer(r, "Члени комісії:" if k == 0 else None, m)
    ack = a.get("ack") or []
    for k, p in enumerate(ack):
        r = signer(r, "З актом інвентаризації ознайомлений:" if k == 0 else None, p)

    pager.flow(r)
    ws.print_area = f"A1:F{r}"
    return ws


def _approve_block(s, r, c1, c2, head, label, who, spec, small=None):
    """«ЗАТВЕРДЖУЮ» чи «ПОГОДЖЕНО»: посада, звання з місцем підпису й іменем, дата від руки."""
    who = who or {}
    s.put(r, c1, c2, head, F(12), A_L, fit=False)
    s.put(r + 1, c1, c2, label or who.get("pos") or "", F(12), A_L)
    if small:
        s.put(r + 2, c1, c2, small, F(7), A_C, fit=False)
        r += 1
    s.put(r + 2, c1, c2, f"{who.get('rank') or ''} ____________ {who.get('name') or ''}".strip(), F(12), A_L)
    s.put(r + 3, c1, c2, blank_date(spec, "року"), F(12), A_L, fit=False)
    return r + 4


def _paper(ws, landscape, left=None, right=None, top=0.6, bottom=0.7, pages=True):
    """Аркуш А4, одна сторінка завширшки; поля — до розкладки сторінок, вона рахує від них."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = left if left is not None else 0.5 if landscape else 1.0
    ws.page_margins.right = right if right is not None else 0.4 if landscape else 0.6
    if top is not None:
        ws.page_margins.top = top
    if bottom is not None:
        ws.page_margins.bottom = bottom
    if pages:
        ws.page_margins.footer = 0.3
        ws.oddFooter.center.text = "&P із &N"


def _table_need(s, values, size=10, floor=ROW_PT):
    """Висота рядка таблиці під найдовший текст у його графах."""
    return max([floor] + [s.need(v, s.widths[c], size) for c, v in enumerate(values) if isinstance(v, str) and v])


def _table_row(s, r, values, font=None, aligns=None, floor=ROW_PT):
    """Рядок таблиці з рамками; висота — під найдовший текст."""
    for c, v in enumerate(values, 1):
        cell = s.ws.cell(row=r, column=c, value=v if v not in ("",) else None)
        cell.font = font or F(10)
        cell.border = BOX
        cell.alignment = (aligns or {}).get(c, A_C)
        if isinstance(v, float):
            cell.number_format = FMT_MONEY
        elif isinstance(v, int) and not isinstance(v, bool):
            cell.number_format = FMT_INT
    s.ws.row_dimensions[r].height = _table_need(s, values, (font or F(10)).sz or 10, floor)


PLAN = [("№\nз/п", 5), ("Підрозділ, місце зберігання майна", 40), ("Матеріально відповідальна особа", 28),
        ("Вид майна (описи)", 20), ("Дата проведення", 13), ("Відповідальний член підкомісії", 28),
        ("Відмітка про виконання", 14)]


def plan_sheet(wb, spec: dict, p: dict, used: set):
    """План проведення інвентаризації майна служби: де, у кого, що, коли й хто з підкомісії."""
    ws = wb.create_sheet(_title("План", used))
    s = Sheet(ws, [w for _, w in PLAN])
    last = len(PLAN)
    _paper(ws, landscape=True)
    pager = Pager(ws, s.widths)
    unit = unit_gen(spec)
    _approve_block(s, 1, 1, 3, "ПОГОДЖЕНО", "Голова інвентаризаційної комісії", p.get("agree"), spec)
    approve = p.get("approve") or {}
    _approve_block(s, 1, 5, last, "ЗАТВЕРДЖУЮ", approve.get("pos") or f"Командир {unit}", approve, spec)
    s.put(6, 1, last, "ПЛАН", F(14, bold=True), A_C, fit=False)
    s.put(7, 1, last, "проведення інвентаризації майна", F(12, bold=True), A_C, fit=False)
    s.put(8, 1, last, f"{spec.get('service') or ''} {unit} станом на {date_dots(spec.get('date'))}".strip(),
          F(12, bold=True), A_C)
    s.put(9, 1, last, f"(наказ командира {unit} {order_ref(spec)})", F(11), A_C)
    r = 11
    _table_row(s, r, [h for h, _ in PLAN], F(10, bold=True), floor=44)

    def numbers(r):
        _table_row(s, r, list(range(1, last + 1)), F(8, bold=True))

    numbers(r + 1)
    left = {2: A_L, 3: A_L, 4: A_L, 6: A_L}
    rows = [[n, row.get("place") or None, row.get("mvo") or None, row.get("kinds") or None,
             date_dots(row.get("date")) or None, row.get("who") or None, row.get("mark") or None]
            for n, row in enumerate(p.get("rows") or [], 1)]
    needs = [_table_need(s, vals, 10, 20) for vals in rows]

    def closing(s, r):
        return s.sign(r, "Голова інвентаризаційної підкомісії", spec.get("head"), 2, (3, 4), (5, 5), (6, 7)) - 1

    tail = ROW_PT + sum(_dry(wb, s.widths, closing)[0])
    pager.keep(r, r + 1, extra=needs[0] if needs else 0)
    r += 1
    for n, (vals, need) in enumerate(zip(rows, needs), 1):
        # Підпис не лишається на сторінці сам: з ним — останній рядок плану.
        r = pager.table(r + 1, need, numbers, extra=tail if n == len(rows) else 0)
        _table_row(s, r, vals, F(10), left, floor=20)
    end = closing(s, r + 2)
    pager.keep(r + 2, end)
    ws.print_area = f"A1:{get_column_letter(last)}{end}"
    return ws


# Графи пропозицій протоколу: (перший рядок шапки, другий, третій, ширина).
PROTOCOL = [
    ("№\nз/п", "", "", 5), ("Найменування цінностей", "", "", 30), ("Матеріально відповідальна особа", "", "", 20),
    ("Результати згідно зі звіряльними відомостями: лишки (+), нестачі (–)", "", "кількість", 9),
    ("Результати згідно зі звіряльними відомостями: лишки (+), нестачі (–)", "", "сума", 12),
    ("Зарахування пересортиці", "лишки, зараховані у покриття нестач", "кількість", 9),
    ("Зарахування пересортиці", "лишки, зараховані у покриття нестач", "сума", 11),
    ("Зарахування пересортиці", "нестачі, покриті лишками", "кількість", 9),
    ("Зарахування пересортиці", "нестачі, покриті лишками", "сума", 11),
    ("Списання в межах норм природного убутку", "", "кількість", 9),
    ("Списання в межах норм природного убутку", "", "сума", 11),
    ("Списання понаднормових нестач і втрат", "", "кількість", 9),
    ("Списання понаднормових нестач і втрат", "", "сума", 11),
    ("Оприбуткування остаточних лишків", "", "кількість", 9), ("Оприбуткування остаточних лишків", "", "сума", 12),
    ("Остаточні нестачі", "", "кількість", 9), ("Остаточні нестачі", "", "сума", 12),
    ("Причина лишків, нестач", "", "", 26),
]
# Рядки «встановлено в наявності», яких у майні служби немає: у бланку вони є, тож прочерк.
PROTOCOL_REST = ["готівка у валюті: національній", "іноземній", "грошові кошти у валюті: національній", "іноземній",
                 "грошових документів", "бланків документів суворої звітності",
                 "балансова вартість фінансових інвестицій", "дебіторської заборгованості",
                 "кредиторської заборгованості", "активи в дорозі"]


def protocol_sheet(wb, spec: dict, p: dict, used: set):
    """Протокол інвентаризаційної комісії — типова форма наказу Мінфіну №572."""
    ws = wb.create_sheet(_title("Протокол", used))
    s = Sheet(ws, [w for *_, w in PROTOCOL])
    last = len(PROTOCOL)
    _paper(ws, landscape=True)
    pager = Pager(ws, s.widths)
    unit = unit_gen(spec)
    right = last - 5
    s.put(1, right, last, "ЗАТВЕРДЖЕНО", F(10), A_L, fit=False)
    s.put(2, right, last, "Наказ Міністерства фінансів України", F(10), A_L, fit=False)
    s.put(3, right, last, "17.06.2015  № 572", F(10), A_L, fit=False)
    s.put(3, 1, 5, spec.get("legal_name") or unit.replace("військової частини", "Військова частина"),
          F(11), A_C, border=LINE)
    s.put(4, 1, 5, "(установа)", F(7), A_C, fit=False)
    s.put(5, 1, 3, "Ідентифікаційний код за ЄДРПОУ", F(8), A_L, border=BOX)
    s.put(5, 4, 5, "  ".join(str(spec.get("edrpou") or "")), F(9), A_C, border=BOX, fit=False)
    approve = p.get("approve") or {}
    r = _approve_block(s, 5, right, last, "ЗАТВЕРДЖУЮ", approve.get("pos") or f"Командир {unit}", approve, spec,
                       small="(керівник установи або уповноважена особа)")
    r += 1
    s.put(r, 1, last, "ПРОТОКОЛ", F(14, bold=True), A_C, fit=False)
    s.put(r + 1, 1, last, "інвентаризаційної комісії", F(12, bold=True), A_C, fit=False)
    s.put(r + 2, 1, last, f"{spec.get('service') or ''} {unit}".strip(), F(12, bold=True), A_C)
    r += 4
    s.put(r, 1, 5, date_uk(p.get("date")) if p.get("date") else blank_date(spec), F(11), A_C, border=LINE, fit=False)
    s.put(r, right, last, spec.get("place") or "", F(11), A_C, border=LINE, fit=False)
    s.put(r + 1, 1, 5, "(дата складання)", F(7), A_C, fit=False)
    s.put(r + 1, right, last, "(місце складання)", F(7), A_C, fit=False)
    r += 3
    order = f"від {date_uk(spec.get('orderDate'))}  №{spec.get('orderNo') or '_____'}"
    s.put(r, 1, last, f"      На підставі розпорядчого документа {order} інвентаризація проводилася станом на "
          f"{date_uk(spec.get('date'), ' р.')} комісією у складі:", F(11), A_J)
    r += 1
    s.put(r, 1, 3, "Голова:", F(11), A_L, fit=False)
    s.put(r, 4, last, rank_name(spec.get("head")), F(11), A_L, border=LINE)
    r += 1
    for k, m in enumerate(spec.get("members") or []):
        if k == 0:
            s.put(r, 1, 3, "Члени комісії:", F(11), A_L, fit=False)
        s.put(r, 4, last, rank_name(m), F(11), A_L, border=LINE)
        r += 1
    s.put(r, 4, last, "(звання, прізвище, ініціали)", F(7), A_C, fit=False)
    r += 2
    s.put(r, 1, last, "За даними інвентаризації встановлено в наявності:", F(11), A_L, fit=False)
    r += 1
    t = p.get("totals") or {}

    def found(r, label, qty, total):
        s.put(r, 1, 5, label, F(11), A_L)
        s.put(r, 6, 9, "загальною кількістю одиниць", F(11), A_R, fit=False)
        cell = s.put(r, 10, 11, qty, F(11), A_C, border=LINE, fit=False)
        _qty(cell, qty)
        s.put(r, 12, 14, "на суму (грн)", F(11), A_R, fit=False)
        s.put(r, 15, 17, total, F(11), A_C, border=LINE, fit=False, fmt=FMT_MONEY)
        return r + 1

    r = found(r, "балансова вартість необоротних активів", t.get("fixedQty") or 0, float(t.get("fixedSum") or 0))
    r = found(r, "запасів", t.get("stockQty") or 0, float(t.get("stockSum") or 0))
    for label in PROTOCOL_REST:
        s.put(r, 1, 5, label, F(11), A_L)
        s.put(r, 6, 17, "—", F(11), A_C, border=LINE, fit=False)
        r += 1
    r += 1
    rows = p.get("rows") or []
    if rows:
        # Таблиця розбіжностей — з нової сторінки: шапка й перші рядки лишаються разом.
        pager.page(r)
    intro = r
    s.put(r, 1, last, "      Пропозиції щодо врегулювання виявлених розбіжностей між фактичною наявністю матеріальних "
          "активів та зобов’язань і даними бухгалтерського обліку, які наводяться у звіряльних відомостях та актах "
          "інвентаризації:", F(11), A_J)
    r += 1
    if not rows:
        s.put(r, 1, last, "Розбіжностей між фактичною наявністю і даними бухгалтерського обліку не виявлено.",
              F(11), A_L, fit=False)
        r += 1
    h1, h2, h3 = r, r + 1, r + 2
    i = 0
    while i < last:
        top = PROTOCOL[i][0]
        j = i
        while j + 1 < last and PROTOCOL[j + 1][0] == top:
            j += 1
        if i == j:                                           # графа на всю висоту шапки
            s.put(h1, i + 1, i + 1, top, F(8), A_C, border=BOX, fit=False)
            ws.merge_cells(start_row=h1, start_column=i + 1, end_row=h3, end_column=i + 1)
            for rr in (h2, h3):
                ws.cell(row=rr, column=i + 1).border = BOX
        else:
            s.put(h1, i + 1, j + 1, top, F(8), A_C, border=BOX, fit=False)
            k = i
            while k <= j:
                mid = PROTOCOL[k][1]
                m = k
                while m + 1 <= j and PROTOCOL[m + 1][1] == mid:
                    m += 1
                if mid:
                    s.put(h2, k + 1, m + 1, mid, F(8), A_C, border=BOX, fit=False)
                    for c in range(k, m + 1):
                        s.put(h3, c + 1, c + 1, PROTOCOL[c][2], F(8), A_C, border=BOX, fit=False)
                else:                                        # без середнього рядка: «кількість» і «сума» на два рядки
                    for c in range(k, m + 1):
                        s.put(h2, c + 1, c + 1, PROTOCOL[c][2], F(8), A_C, border=BOX, fit=False)
                        ws.merge_cells(start_row=h2, start_column=c + 1, end_row=h3, end_column=c + 1)
                        ws.cell(row=h3, column=c + 1).border = BOX
                k = m + 1
        i = j + 1
    ws.row_dimensions[h1].height = 40
    ws.row_dimensions[h2].height = 30
    ws.row_dimensions[h3].height = 14

    def numbers(r):
        _table_row(s, r, list(range(1, last + 1)), F(8, bold=True))

    r = h3 + 1
    numbers(r)
    total = {5: 0.0, 13: 0.0, 15: 0.0, 17: 0.0}
    left = {2: A_L, 3: A_L, 18: A_L}
    table = []
    for n, row in enumerate(rows, 1):
        qty, money = row.get("qty") or 0, float(row.get("sum") or 0)
        vals = [n, row.get("name"), row.get("mvo") or None, qty, money] + [None] * 12 + [row.get("reason") or None]
        if qty > 0:
            vals[13], vals[14] = qty, money
            total[15] += money
        elif qty < 0 and row.get("loss"):
            # Знищене за рапортом, ще не списане: втрата, яку спише акт, а не остаточна нестача.
            vals[11], vals[12] = -qty, -money
            total[13] += -money
        elif qty < 0:
            vals[15], vals[16] = -qty, -money
            total[17] += -money
        total[5] += money
        table.append(vals)
    if not rows:
        table.append(["—"] * last)
    needs = [_table_need(s, vals, 9) for vals in table]
    # Шапка таблиці не лишається внизу сторінки сама: з нею — перший рядок.
    pager.keep(intro, r, extra=needs[0])
    for vals, need in zip(table, needs):
        r = pager.table(r + 1, need, numbers)
        _table_row(s, r, vals, F(9), left if rows else None)
        for c in (4, 12, 14, 16):
            _qty(ws.cell(row=r, column=c), vals[c - 1])
    vals = ["Разом", "Х"] + [None] * 15 + ["Х"]
    if rows:
        for c, v in total.items():
            vals[c - 1] = round(v, 2) if c == 5 or v else None

    def closing(s, r):
        for label in ("Рішення щодо заборгованості, за якою строк позовної давності минув:",
                      "Рішення щодо готівки, грошових коштів, цінних паперів, грошових документів і бланків "
                      "документів суворої звітності:"):
            s.put(r, 1, 9, label, F(11), A_L)
            s.put(r, 10, last, "—", F(11), A_C, border=LINE, fit=False)
            r += 1
        lay = (3, (4, 9), (10, 13), (14, 18))
        r = s.sign(r + 1, "Голова комісії", spec.get("head"), *lay)
        for k, m in enumerate(spec.get("members") or []):
            r = s.sign(r, "Члени комісії" if k == 0 else None, m, *lay)
        return r - 1

    # Рішення й підписи комісії не лишаються на сторінці самі: з ними — рядок «Разом».
    r = pager.table(r + 1, _table_need(s, vals, 9), numbers, extra=ROW_PT + sum(_dry(wb, s.widths, closing)[0]))
    _table_row(s, r, vals, F(9, bold=True))
    end = closing(s, r + 2)
    pager.keep(r + 2, end)
    ws.print_area = f"A1:{get_column_letter(last)}{end}"
    return ws


REGISTER = [("№\nз/п", 5), ("Найменування документа (акт, інвентаризаційний опис, протокол)", 38), ("Номер, дата", 14),
            ("Матеріально відповідальна особа, місцезнаходження", 36), ("Кількість аркушів", 11), ("Примітка", 14)]


def register_sheet(wb, spec: dict, g: dict, used: set):
    """Реєстр актів та описів, з яким підкомісія здає папери начальникові ФЕС на перевірку."""
    ws = wb.create_sheet(_title("Реєстр", used))
    s = Sheet(ws, [w for _, w in REGISTER])
    last = len(REGISTER)
    _paper(ws, landscape=False)
    pager = Pager(ws, s.widths)
    unit = unit_gen(spec)
    s.put(1, 1, last, "РЕЄСТР", F(14, bold=True), A_C, fit=False)
    s.put(2, 1, last, "актів інвентаризації з інвентаризаційними описами, що подаються на перевірку начальнику "
          f"фінансово-економічної служби – головному бухгалтеру {unit}", F(12, bold=True), A_C)
    s.put(3, 1, last, f"(пункт 7 наказу командира {unit} {order_ref(spec, item=False)})", F(11), A_C)
    item = f" (пункт {spec['orderItem']} наказу)" if spec.get("orderItem") else ""
    s.put(5, 1, last, f"Інвентаризаційна підкомісія{item}: {spec.get('service') or ''} {unit}".strip(), F(11), A_L)
    r = 7
    _table_row(s, r, [h for h, _ in REGISTER], F(10, bold=True), floor=44)

    def numbers(r):
        _table_row(s, r, list(range(1, last + 1)), F(8, bold=True))

    numbers(r + 1)
    rows = g.get("rows") or []
    left = {2: A_L, 4: A_L, 6: A_L}
    table = [[n, row.get("doc"), date_dots(row.get("date")) or None, row.get("who") or None,
              row.get("sheets") or None, row.get("note") or None] for n, row in enumerate(rows, 1)]
    needs = [_table_need(s, vals, 10, 20) for vals in table]
    sheets = sum(int(x.get("sheets") or 0) for x in rows)

    def closing(s, r):
        s.put(r, 1, last, f"Усього документів {len(rows)}, аркушів {sheets or '______'}.", F(11), A_L, fit=False)
        r += 2
        s.put(r, 1, last, "Здав:", F(11, bold=True), A_L, fit=False)
        r = s.sign(r + 1, "Голова підкомісії", spec.get("head"), 2, (3, 3), (4, 4), (5, 6))
        s.put(r, 1, 3, blank_date(spec), F(11), A_L, fit=False)
        r += 2
        s.put(r, 1, last, "Прийняв:", F(11, bold=True), A_L, fit=False)
        r = s.sign(r + 1, None, g.get("fes"), 2, (1, 3), (4, 4), (5, 6))
        s.put(r, 1, 3, blank_date(spec), F(11), A_L, fit=False)
        return r

    tail = ROW_PT + sum(_dry(wb, s.widths, closing)[0])
    pager.keep(r, r + 1, extra=needs[0] if needs else 0)
    r += 1
    for n, (vals, need) in enumerate(zip(table, needs), 1):
        # «Здав» і «Прийняв» не лишаються на сторінці самі: з ними — останній рядок реєстру.
        r = pager.table(r + 1, need, numbers, extra=tail if n == len(table) else 0)
        _table_row(s, r, vals, F(10), left, floor=20)
    end = closing(s, r + 2)
    pager.keep(r + 2, end)
    ws.print_area = f"A1:{get_column_letter(last)}{end}"
    return ws


PROGRESS = [("№\nз/п", 5), ("Підрозділ", 26), ("Матеріально відповідальна особа", 28), ("Описів", 8), ("Рядків", 8),
            ("Сума за обліком, грн", 15), ("Дата за планом", 12), ("Проведено", 12), ("Розбіжності, рядків", 13),
            ("Відповідальний член підкомісії", 26)]


def progress_sheet(wb, spec: dict, p: dict, used: set):
    """Відомості про хід інвентаризації: у яких підрозділах її проведено на день доповіді."""
    ws = wb.create_sheet(_title("Хід", used))
    s = Sheet(ws, [w for _, w in PROGRESS])
    last = len(PROGRESS)
    _paper(ws, landscape=True)
    pager = Pager(ws, s.widths)
    unit = unit_gen(spec)
    s.put(1, 1, last, "ВІДОМОСТІ", F(14, bold=True), A_C, fit=False)
    s.put(2, 1, last, "про хід проведення інвентаризації майна", F(12, bold=True), A_C, fit=False)
    s.put(3, 1, last, f"{spec.get('service') or ''} {unit} станом на {date_dots(p.get('date')) or '__.__.____'}".strip(),
          F(12, bold=True), A_C)
    s.put(4, 1, last, f"(наказ командира {unit} {order_ref(spec)}; інвентаризація станом на "
          f"{date_dots(spec.get('date'))}, {date_dots(spec.get('start'))}–{date_dots(spec.get('end'))})", F(11), A_C)
    r = 6
    for line in p.get("lines") or []:
        s.put(r, 1, last, line, F(11), A_L)
        r += 1
    r += 1
    _table_row(s, r, [h for h, _ in PROGRESS], F(10, bold=True), floor=44)

    def numbers(r):
        _table_row(s, r, list(range(1, last + 1)), F(8, bold=True))

    numbers(r + 1)
    left = {2: A_L, 3: A_L, 10: A_L}
    # Інше місце плану (продукти, вода) обліку в програмі не має: сума порожня, не нуль.
    rows = [[n, row.get("sub"), row.get("mvo") or None, row.get("descs") or None, row.get("lines") or None,
             float(row["sum"]) if row.get("sum") is not None else None,
             date_dots(row.get("plan")) or None, date_dots(row.get("done")) or None,
             row.get("diffs") or None, row.get("who") or None] for n, row in enumerate(p.get("rows") or [], 1)]
    needs = [_table_need(s, vals, 10, 20) for vals in rows]
    t = p.get("totals") or {}
    total = [None, None, None, t.get("descs") or None, t.get("lines") or None, float(t.get("sum") or 0), None,
             f"{t.get('done') or 0} з {t.get('subs') or len(rows)}", t.get("diffs") or None, None]

    def closing(s, r):
        s.put(r, 1, 3, blank_date(spec), F(11), A_L, fit=False)
        return s.sign(r + 2, "Голова інвентаризаційної підкомісії", spec.get("head"), 3, (4, 6), (7, 8), (9, 10)) - 1

    pager.keep(r, r + 1, extra=needs[0] if needs else 0)
    r += 1
    for vals, need in zip(rows, needs):
        r = pager.table(r + 1, need, numbers)
        _table_row(s, r, vals, F(10), left, floor=20)
    # Дата й підпис не лишаються на сторінці самі: з ними — рядок «Разом».
    r = pager.table(r + 1, _table_need(s, total, 10, 20), numbers, extra=ROW_PT + sum(_dry(wb, s.widths, closing)[0]))
    _table_row(s, r, total, F(10, bold=True), floor=20)
    s.put(r, 1, 3, "Разом", F(10, bold=True), A_L, border=BOX, fit=False)
    end = closing(s, r + 2)
    pager.keep(r + 2, end)
    ws.print_area = f"A1:{get_column_letter(last)}{end}"
    return ws


def build_inventory(spec: dict):
    descs = spec.get("descriptions") or []
    papers = [k for k in ("plan", "act", "protocol", "register", "progress") if spec.get(k)]
    if not descs and not papers:
        raise ValueError("немає жодного опису для вивантаження")
    from excel_export import clean_workbook                     # noqa: PLC0415
    wb = clean_workbook(Workbook())
    wb.remove(wb.active)
    used = set()
    # Папери підкомісії — попереду, у тому порядку, в якому їх складають; описи — за ними.
    if spec.get("plan"):
        plan_sheet(wb, spec, spec["plan"], used)
    if spec.get("act"):
        act_sheet(wb, spec, spec["act"], used)
    if spec.get("protocol"):
        protocol_sheet(wb, spec, spec["protocol"], used)
    if spec.get("register"):
        register_sheet(wb, spec, spec["register"], used)
    if spec.get("progress"):
        progress_sheet(wb, spec, spec["progress"], used)
    for d in descs:
        description_sheet(wb, spec, d, used)
    wb.active = 0
    return wb
