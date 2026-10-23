# -*- coding: utf-8 -*-
"""Книги «Облік СП» і «Облік РП» — запасний шлях книг сухпаю й реторт-пакетів, які програма веде сама.

Програма вивантажує в книгу все, що знає про книгу обліку, і приймає її назад: коли програма
недоступна, книгу ведуть руками, а потім програма — навіть чиста — приймає її й має ту саму базу.
Будова кожної книги (`Layout`) — як у книзі, якою служба вела облік до програми: ті самі аркуші,
таблиці з A1, ті самі колонки в тому самому порядку (за ними працюють скрипти служби, частина — за
позицією колонки); колонки, яких у старій книзі не було, — лише після наявних. Формули книги
лишаються формулами (обчислювані колонки таблиць: рядок, вписаний руками, отримує їх сам), але кожна
несе й значення на момент вивантаження: скрипти читають значення. Вивантаження детерміноване: та
сама база дає ті самі значення клітинок (скрипт служби щогодини порівнює відбитки рядків).

Книга «Облік СП» несе ще й акти сухпаю — чотири аркуші після «Списання» (`sp_book_acts`, будова — `Layout.acts`).
Книга «Облік ОП» має ту саму основу, але свій модуль (`op_book.py`) з даними актів списання; об'єднати її
з цим рушієм — окремою задачею, коли код актів буде в основній гілці.

Книгу, яку програма сама записала в теку книги чи прийняла, вона пам'ятає своєю (відбитки байтів і вмісту, як у
книги ОП): вивантаження перезаписує таку книгу без звірки з базою.
"""
import contextlib
import datetime
import hashlib
import json
import os
import re
import sqlite3
import traceback
import zipfile
from dataclasses import dataclass
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

import stable_hash
from state_db import FES_STATUSES as STATUSES   # стани ФЕС — сам перелік програми, не його копія: новий стан приймає й книга

# Своїх книг пам'ятаємо не більше стількох відбитків байтів і стількох відбитків вмісту (новіші спереду).
KEEP_SHAS = 10
ARCHIVE = "архів"
JOURNALS = ("incoming", "movement", "writeoffs")
# Частини актів сухпаю в моделі книги СП (`sp_book_acts`): акти, рядки актів, дні рапорту, підписанти.
ACT_PARTS = ("acts", "act_lines", "act_days", "act_signers")
TITLES = {"incoming": ("Прибуток", "tblIncoming"), "movement": ("Переміщення", "tblMovement"),
          "writeoffs": ("Списання", "tblWriteoffs")}

ITEM_BASE = ("код", "найменування", "категорія", "од.вим.", "ціна", "поз", "тип", "дободач", "примітка")
IN_BASE = ("док_тип", "док_дата", "док_номер", "від_кого", "кому", "найменування", "код", "найм", "од_вим",
           "ціна", "кількість", "категорія", "сума", "примітка")
MV_BASE = ("док_тип", "док_дата", "док_номер", "від_кого", "кому", "найменування", "код", "од_вим", "ціна",
           "кількість", "категорія", "зал_відпр", "зал_одерж", "нм", "одвим")
WR_BASE = ("док_тип", "док_дата", "док_номер", "підрозділ", "найменування", "код", "транз_найменування",
           "кількість", "категорія", "статус", "примітка", "рапорт", "наказ")
SUB_COLS = ("назва", "підпорядкований", "вид", "закритий", "псевдоніми")


@dataclass(frozen=True)
class Layout:
    """Будова книги обліку: що в ній за колонки й як позиція книги лягає в групу бази."""
    book: str                    # книга бази: «ОП», «СП», «РП»
    title: str                   # назва книги, вона ж ім'я файла: «Облік ОП»
    sha_setting: str             # app_setting (серед state_db.SERVER_SETTINGS): відбитки своїх книг — `remember_sha`
    group_of_cat: tuple          # ((категорія, група), …)
    item_cols: tuple
    in_cols: tuple
    mv_cols: tuple
    wr_cols: tuple
    types: tuple = ()            # перелік «типи» аркуша «Довідка» (без «архів»): «тип» позиції — вид рядка 1/прод
    supply: bool = False         # аркуш «Забезпеченість» і рядок добо-дач на «Залишки»
    acts: bool = False           # аркуші актів сухпаю: «Акти», «Рядки актів», «Дні актів», «Підписанти актів»

    @property
    def cats(self):
        return dict(self.group_of_cat)

    @property
    def cat_of_group(self):
        return {g: c for c, g in self.group_of_cat}

    @property
    def alias_source(self):
        return f"книга {self.book}"

    @property
    def file_name(self):
        return f"{self.title}.xlsx"

    def cols(self, journal):
        return {"incoming": self.in_cols, "movement": self.mv_cols, "writeoffs": self.wr_cols}[journal]


SP = Layout(
    book="СП", title="Облік СП", sha_setting="sp_book_sha",
    group_of_cat=(("СП", "СП.1"), ("ПК", "СП.2"), ("Т", "СП.3"), ("ТВ", "СП.4")),
    types=("ПНСП", "ПНСП-Р", "ДПНП", "ДПНП-Р", "іноземні добові", "іноземні раціони", "тара"),
    item_cols=ITEM_BASE,
    in_cols=IN_BASE + ("придатний до", "скан", "статус"),
    mv_cols=MV_BASE + ("примітка", "статус", "скан"),
    wr_cols=WR_BASE + ("1прод", "скан", "кому"),
    supply=True, acts=True)

RP = Layout(
    book="РП", title="Облік РП", sha_setting="rp_book_sha",
    group_of_cat=(("РП", "РП.1"),),
    types=("реторт-пакети",),
    item_cols=ITEM_BASE,
    in_cols=IN_BASE + ("скан", "статус"),
    mv_cols=MV_BASE + ("примітка", "статус", "скан"),
    wr_cols=WR_BASE + ("скан", "кому"))

LAYOUTS = {"СП": SP, "РП": RP}


def _me(table, col):
    return f"{table}[[#This Row],[{col}]]"


def _rest_formula(side):
    """Залишок місця з колонки side («від_кого» чи «кому») на дату рядка — «зал_відпр»/«зал_одерж»."""
    code, place, day = _me("tblMovement", "код"), _me("tblMovement", side), _me("tblMovement", "док_дата")
    return (f'SUMIFS(tblIncoming[кількість],tblIncoming[код],{code},tblIncoming[кому],{place},'
            f'tblIncoming[док_дата],"<="&{day})'
            f'+SUMIFS(tblMovement[кількість],tblMovement[код],{code},tblMovement[кому],{place},'
            f'tblMovement[док_дата],"<="&{day})'
            f'-SUMIFS(tblMovement[кількість],tblMovement[код],{code},tblMovement[від_кого],{place},'
            f'tblMovement[док_дата],"<="&{day})'
            f'-SUMIFS(tblWriteoffs[кількість],tblWriteoffs[код],{code},tblWriteoffs[підрозділ],{place},'
            f'tblWriteoffs[док_дата],"<="&{day})')


def calc(layout):
    """Обчислювані колонки книги: та сама формула в описі колонки таблиці й у клітинках."""
    def lookup(table, n):
        return f"VLOOKUP({_me(table, 'код')},tblItems[[код]:[поз]],{n},FALSE)"
    return {
        "Номенклатура": {"поз": f'{_me("tblItems", "код")}&" "&{_me("tblItems", "найменування")}&" "&'
                                f'{_me("tblItems", "ціна")}'},
        "Прибуток": {"код": f"INT(TRIM(LEFT({_me('tblIncoming', 'найменування')},5)))",
                     "од_вим": lookup("tblIncoming", 4), "ціна": lookup("tblIncoming", 5),
                     "категорія": lookup("tblIncoming", 3),
                     "сума": f"{_me('tblIncoming', 'ціна')}*{_me('tblIncoming', 'кількість')}"},
        "Переміщення": {"код": f"INT(TRIM(LEFT({_me('tblMovement', 'найменування')},5)))",
                        "од_вим": lookup("tblMovement", 4), "ціна": lookup("tblMovement", 5),
                        "категорія": lookup("tblMovement", 3),
                        "зал_відпр": _rest_formula("від_кого"), "зал_одерж": _rest_formula("кому")},
        "Списання": {"код": f"INT(TRIM(LEFT({_me('tblWriteoffs', 'найменування')},5)))",
                     "категорія": lookup("tblWriteoffs", 3)},
    }


def _num(v):
    """Ціле — цілим: у книзі 9000, а не 9000.0."""
    v = float(v or 0)
    return int(v) if v.is_integer() else round(v, 6)


def _code_cell(code):
    return int(code) if str(code).isdigit() else str(code)


def _pos(it):
    """«поз» книги: «код найменування ціна» — ціна з комою, як її показує Excel українською."""
    price = ("%g" % float(it["price"] or 0)).replace(".", ",")
    return f"{it['code']} {it['name']} {price}"


def _per_ration_cell(v):
    """«дободач»: 1/3 (раціон на одне приймання їжі) — формулою «=1/3», як пише людина; решта — числом."""
    if v is None:
        return None, None
    if abs(float(v) - 1 / 3) < 1e-9:
        return "=1/3", 1 / 3
    return _num(v), None


# ---------------------------------------------------------------- база → модель

def model_from_db(con, layout):
    """Усе, що база знає про книгу обліку, — у вигляді книги."""
    import state_db                                              # noqa: PLC0415 — модуль сервера
    items = []
    for r in con.execute("""
            SELECT n.id, n.code, n.name, g.code AS grp, u.code AS uom, n.per_ration, n.note, n.archived_at,
                   n.prod1_kind, COALESCE(p.price_kop, n.app_price_kop) AS price_kop
              FROM nomen n JOIN nomen_group g ON g.id = n.group_id JOIN uom u ON u.id = n.uom_id
              LEFT JOIN nomen_last_price p ON p.nomen_id = n.id
             WHERE g.book = ? ORDER BY n.code""", (layout.book,)):
        it = {"code": r["code"], "name": r["name"], "cat": layout.cat_of_group[r["grp"]], "uom": r["uom"],
              "price": (r["price_kop"] or 0) / 100,
              "type": r["prod1_kind"] or "",
              "archived": bool(r["archived_at"]), "per_ration": r["per_ration"], "note": r["note"] or ""}
        items.append(it)
    codes = {i["code"] for i in items}
    subs = [{"name": r["name"], "parent": r["parent"] or "", "kind": r["kind"], "closed": not r["is_active"],
             "aliases": [a[0] for a in con.execute("SELECT name FROM subdivision_alias WHERE subdivision_id = ? "
                                                   "AND source = ? ORDER BY id", (r["id"], layout.alias_source))]}
            for r in con.execute("""SELECT s.id, s.name, p.name AS parent, k.code AS kind, s.is_active
                                      FROM subdivision s JOIN subdivision_kind k ON k.id = s.kind_id
                                      LEFT JOIN subdivision p ON p.id = s.parent_id ORDER BY s.sort, s.id""")]
    fes, meta = state_db._load_doc_fes(con), state_db._load_doc_meta(con)
    extra = {(str(doc), code, note): rec for doc, lines in state_db._load_line_extra(con).items()
             for key, rec in lines.items() for code, _, note in [key.partition("|")]}
    docs, at = [], {}
    for journal, rows in state_db._load_docs(con).items():
        for r in rows:
            date, paper, no, src, dst, code, qty, note = r[:8]
            doc_id = r[10] if len(r) > 10 else None
            if str(code) not in codes:
                continue
            key = (journal, doc_id) if doc_id else (journal, date, no, src, dst)
            d = at.get(key)
            if d is None:
                f, m = fes.get(str(doc_id), {}), meta.get(str(doc_id), {})
                d = at[key] = {"journal": journal, "type": paper or "", "date": date, "no": str(no), "src": src or "",
                               "dst": dst or "", "status": f.get("status", ""), "report": m.get("report", ""),
                               "order": m.get("order", ""), "scan": m.get("scan", ""), "lines": [], "_id": doc_id}
                docs.append(d)
            same = next((x for x in d["lines"] if x["code"] == str(code) and x["note"] == (note or "")), None)
            if same:                                              # рядки однієї позиції по партіях — один рядок книги
                same["qty"] = _num(float(same["qty"]) + float(qty))
            else:
                rec = extra.get((str(doc_id), str(code), note or ""), {})
                d["lines"].append({"code": str(code), "qty": _num(qty), "note": note or "", "best": rec.get("best", ""),
                                   "q1": rec.get("q1", ""), "fixed": bool(rec.get("fixed"))})
    for d in docs:
        d.pop("_id", None)
    renames = [[a[0], a[1]] for a in con.execute(
        "SELECT a.name, s.name FROM subdivision_alias a JOIN subdivision s ON s.id = a.subdivision_id "
        "WHERE a.source = ? ORDER BY a.id", (layout.alias_source,))]
    model = {"items": items, "subs": subs, "renames": renames, "docs": docs}
    if layout.acts:
        import sp_book_acts                                      # noqa: PLC0415 — модуль імпортує цей
        model.update(sp_book_acts.model(con))
    return model


def _balances(model, as_of):
    """{(код, місце): к-сть} на кінець дня as_of — як SUMIFS аркуша «Залишки»."""
    out = {}
    for d in model["docs"]:
        if as_of and d["date"] > as_of:
            continue
        for ln in d["lines"]:
            q = float(ln["qty"] or 0)
            if d["journal"] == "incoming":
                out[(ln["code"], d["dst"])] = out.get((ln["code"], d["dst"]), 0) + q
            elif d["journal"] == "movement":
                out[(ln["code"], d["src"])] = out.get((ln["code"], d["src"]), 0) - q
                out[(ln["code"], d["dst"])] = out.get((ln["code"], d["dst"]), 0) + q
            else:
                out[(ln["code"], d["src"])] = out.get((ln["code"], d["src"]), 0) - q
    return out


# ---------------------------------------------------------------- модель → книга

def _table(ws, name, cols, rows_n, calc_cols=None):
    """Таблиця з A1; колонки з `calc_cols` — обчислювані: Excel сам продовжує їх на новий рядок."""
    ref = f"A1:{get_column_letter(len(cols))}{max(rows_n + 1, 2)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    t._initialise_columns()
    for col, title in zip(t.tableColumns, cols):
        col.name = title
        if calc_cols and title in calc_cols:
            col.calculatedColumnFormula = TableFormula(attr_text=calc_cols[title])
    ws.add_table(t)


def _doc_order(d):
    return (d["date"], d["no"], d["src"], d["dst"])


def _letter(cols, title):
    return get_column_letter(cols.index(title) + 1)


def _day(v):
    return datetime.date.fromisoformat(v) if v else None


def write_book(model, path, layout, as_of=None, headcount=None):
    """Книга з моделі. `headcount` — особи на харчуванні на дату (для аркуша «Забезпеченість» книги СП)."""
    as_of = as_of or datetime.date.today().isoformat()
    items = sorted(model["items"], key=lambda i: (i["archived"], int(i["code"]) if str(i["code"]).isdigit() else 0,
                                                  str(i["code"])))
    by_code = {i["code"]: i for i in items}
    places = [s["name"] for s in model["subs"] if not s["closed"]]
    formulas = calc(layout)
    wb = Workbook()
    cached = {}                                                   # аркуш -> {клітинка: значення}

    def put(ws, ref, formula, value):
        ws[ref] = formula
        cached.setdefault(ws.title, {})[ref] = value

    # Довідка: місця, типи, статуси; перейменування; контрагенти.
    ws = wb.active
    ws.title = "Довідка"
    ws.append(["Підрозділи", "типи", "статус"])
    types = list(layout.types) + [ARCHIVE]
    n_ref = max(len(places), len(types), len(STATUSES))
    for n in range(n_ref):
        ws.append([places[n] if n < len(places) else None, types[n] if n < len(types) else None,
                   STATUSES[n] if n < len(STATUSES) else None])
    _table(ws, "Таблиця4", ["Підрозділи", "типи", "статус"], n_ref)
    ws["H1"], ws["I1"] = "стара назва", "нова назва"
    for n, (old, new) in enumerate(model["renames"], 2):
        ws.cell(n, 8, old)
        ws.cell(n, 9, new)
    ws.add_table(Table(displayName="tblRename", ref=f"H1:I{max(len(model['renames']) + 1, 2)}"))

    # Номенклатура.
    ws = wb.create_sheet("Номенклатура")
    ws.append(list(layout.item_cols))
    pos_col = _letter(layout.item_cols, "поз")
    per_col = layout.item_cols.index("дободач") + 1
    for r, it in enumerate(items, 2):
        row = {"код": _code_cell(it["code"]), "найменування": it["name"], "категорія": it["cat"], "од.вим.": it["uom"],
               "ціна": _num(it["price"]), "тип": ARCHIVE if it["archived"] else (it["type"] or None),
               "примітка": it["note"] or None}
        ws.append([row.get(c) for c in layout.item_cols])
        put(ws, f"{pos_col}{r}", "=" + formulas["Номенклатура"]["поз"], _pos(it))
        per, per_value = _per_ration_cell(it["per_ration"])
        if per_value is None:
            ws.cell(r, per_col, per)
        else:
            put(ws, f"{get_column_letter(per_col)}{r}", per, per_value)
    _table(ws, "tblItems", layout.item_cols, len(items), formulas["Номенклатура"])

    # Журнали.
    rest_of = {}
    for journal in JOURNALS:
        title, table = TITLES[journal]
        cols = layout.cols(journal)
        ws = wb.create_sheet(title)
        ws.append(list(cols))
        f = {k: "=" + v for k, v in formulas[title].items()}
        r = 1
        for d in sorted((d for d in model["docs"] if d["journal"] == journal), key=_doc_order):
            for ln in d["lines"]:
                r += 1
                it = by_code[ln["code"]]
                head = {"док_тип": d["type"], "док_дата": _day(d["date"]), "док_номер": d["no"],
                        "статус": d["status"] or None, "скан": d["scan"] or None, "примітка": ln["note"] or None}
                values = {"код": _code_cell(it["code"]), "од_вим": it["uom"], "ціна": _num(it["price"]),
                          "категорія": it["cat"]}
                if journal == "incoming":
                    row = dict(head, від_кого=d["src"], кому=d["dst"], найменування=_pos(it), найм=_pos(it),
                               кількість=_num(ln["qty"]), рапорт=d["report"] or None, наказ=d["order"] or None)
                    row["придатний до"] = _day(ln.get("best"))
                    values["сума"] = round(float(it["price"]) * float(ln["qty"]), 2)
                elif journal == "movement":
                    row = dict(head, від_кого=d["src"], кому=d["dst"], найменування=_pos(it), кількість=_num(ln["qty"]),
                               нм=_pos(it), одвим=it["uom"])
                    if d["date"] not in rest_of:             # залишки на дату рахуються раз на дату, а не на кожний рядок
                        rest_of[d["date"]] = _balances(model, d["date"])
                    bal = rest_of[d["date"]]
                    values["зал_відпр"] = _num(bal.get((ln["code"], d["src"]), 0))
                    values["зал_одерж"] = _num(bal.get((ln["code"], d["dst"]), 0))
                else:
                    row = dict(head, підрозділ=d["src"], найменування=_pos(it), кількість=_num(ln["qty"]),
                               рапорт=d["report"] or None, наказ=d["order"] or None, кому=d["dst"] or None)
                    row["1прод"] = ln.get("q1") or None
                ws.append([row.get(c) for c in cols])
                for name, formula in f.items():
                    put(ws, f"{_letter(cols, name)}{r}", formula, values[name])
                ws.cell(r, 2).number_format = "dd.mm.yyyy"
                if "придатний до" in cols and row.get("придатний до"):
                    ws.cell(r, cols.index("придатний до") + 1).number_format = "dd.mm.yyyy"
        _table(ws, table, cols, r - 1, formulas[title])
        _validations(ws, journal, cols, r)
    if layout.acts:
        import sp_book_acts                                      # noqa: PLC0415 — модуль імпортує цей
        sp_book_acts.write(wb, model, put)

    _rests_sheet(wb, model, items, places, as_of, put, layout)
    _moves_sheet(wb, model, items, as_of)
    if layout.supply:
        _supply_sheet(wb, model, items, places, as_of, headcount)

    ws = wb.create_sheet("Підрозділи")
    ws.append(list(SUB_COLS))
    for s in model["subs"]:
        ws.append([s["name"], s["parent"] or None, s["kind"], "так" if s["closed"] else None,
                   ", ".join(s["aliases"]) or None])
    _table(ws, "tblSubs", SUB_COLS, len(model["subs"]))

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    _put_cached(path, cached)


def _validations(ws, journal, cols, last):
    """Списки, як у старій книзі: позиції, місця, статуси."""
    def lst(formula):
        return DataValidation(type="list", formula1=formula, allow_blank=True)
    pos = lst("Номенклатура!$F$2:$F$2000")
    place = lst("Довідка!$A$2:$A$200")
    status = lst(f"Довідка!$C$2:$C${len(STATUSES) + 1}")
    for dv in (pos, place, status):
        ws.add_data_validation(dv)
    end = max(last, 1) + 500

    def rng(title, title2=None):
        a, b = _letter(cols, title), _letter(cols, title2 or title)
        return f"{a}2:{b}{end}"
    pos.add(rng("найменування"))
    if journal == "incoming":
        place.add(rng("кому"))
    elif journal == "movement":
        place.add(rng("від_кого", "кому"))
    else:
        place.add(rng("підрозділ"))
    if "статус" in cols:
        status.add(rng("статус"))


def _rest_cell(col, r):
    """Залишок позиції рядка r у місці з шапки колонки col на дату D2 — як аркуш «Залишки» старої книги."""
    return (f"=SUMIFS(Прибуток!$K:$K,Прибуток!$G:$G,$A{r},Прибуток!$E:$E,{col}$4,Прибуток!$B:$B,\"<=\"&$D$2)"
            f"+SUMIFS(Переміщення!$J:$J,Переміщення!$G:$G,$A{r},Переміщення!$E:$E,{col}$4,Переміщення!$B:$B,\"<=\"&$D$2)"
            f"-SUMIFS(Переміщення!$J:$J,Переміщення!$G:$G,$A{r},Переміщення!$D:$D,{col}$4,Переміщення!$B:$B,\"<=\"&$D$2)"
            f"-SUMIFS(Списання!$H:$H,Списання!$F:$F,$A{r},Списання!$D:$D,{col}$4,Списання!$B:$B,\"<=\"&$D$2)")


def _rests_sheet(wb, model, items, places, as_of, put, layout):
    """«Залишки»: D2 — дата, шапка в рядку 4 (за нею працюють скрипти служби й ручна картка). У книзі СП
    рядок 3 — добо-дачі по місцях (залишок × «дободач»)."""
    ws = wb.create_sheet("Залишки")
    ws["D1"] = "Залишки на дату:"
    ws["D2"] = datetime.date.fromisoformat(as_of)
    ws["D2"].number_format = "dd.mm.yyyy"
    for c, title in enumerate(["код", "кат", "найм", "разом"] + places, 1):
        ws.cell(4, c, title)
    bal = _balances(model, as_of)
    last = get_column_letter(4 + len(places))
    live = [i for i in items if not i["archived"]]
    for r, it in enumerate(live, 5):
        ws.cell(r, 1, _code_cell(it["code"]))
        ws.cell(r, 2, it["cat"])
        ws.cell(r, 3, _pos(it))
        tot = 0.0
        for c, place in enumerate(places, 5):
            col = get_column_letter(c)
            v = bal.get((it["code"], place), 0)
            tot += v
            put(ws, f"{col}{r}", _rest_cell(col, r), _num(v))
        put(ws, f"D{r}", f"=SUM(E{r}:{last}{r})" if places else "=0", _num(tot))
    if layout.supply:
        ws["C3"] = "добо-дачі"
        total = 0.0
        for c, place in enumerate(places, 5):
            dd = sum(bal.get((it["code"], place), 0) * float(it["per_ration"] or 0) for it in live)
            total += dd
            ws.cell(3, c, _num(round(dd, 3)))
        ws["D3"] = _num(round(total, 3))


def _moves_sheet(wb, model, items, as_of, place="склад"):
    """«Рух ТМЦ»: місце й період — початковий, надійшло, вибуло, кінцевий (як у старій книзі).
    Значення без формул: ручна картка бере звідси лише підсумки."""
    ws = wb.create_sheet("Рух ТМЦ")
    start = as_of[:4] + "-01-01"
    ws["C3"], ws["D3"], ws["E3"], ws["F3"] = "період", datetime.date.fromisoformat(start), \
        datetime.date.fromisoformat(as_of), place
    ws["D3"].number_format = ws["E3"].number_format = "dd.mm.yyyy"
    for c, title in enumerate(["код", "кат", "найм", "початковий", "надійшло", "вибуло", "кінцевий"], 1):
        ws.cell(4, c, title)
    before = (datetime.date.fromisoformat(start) - datetime.timedelta(days=1)).isoformat()
    a, b = _balances(model, before), _balances(model, as_of)
    got, gone = {}, {}
    for d in model["docs"]:
        if not start <= d["date"] <= as_of:
            continue
        for ln in d["lines"]:
            q = float(ln["qty"] or 0)
            if d["journal"] in ("incoming", "movement") and d["dst"] == place:
                got[ln["code"]] = got.get(ln["code"], 0) + q
            if d["journal"] in ("movement", "writeoffs") and d["src"] == place:
                gone[ln["code"]] = gone.get(ln["code"], 0) + q
    for r, it in enumerate([i for i in items if not i["archived"]], 5):
        c = it["code"]
        for col, v in enumerate([_code_cell(c), it["cat"], it["name"], _num(a.get((c, place), 0)),
                                 _num(got.get(c, 0)), _num(gone.get(c, 0)), _num(b.get((c, place), 0))], 1):
            ws.cell(r, col, v)


def _supply_sheet(wb, model, items, places, as_of, headcount, days=3):
    """«Забезпеченість» книги СП: добо-дачі по місцях на дату й потреба «днів × осіб» — значеннями
    (ручна картка донесення бере звідси числа; осіб на харчуванні програма бере з Добовки)."""
    ws = wb.create_sheet("Забезпеченість")
    ws["A1"], ws["B1"] = "станом на", datetime.date.fromisoformat(as_of)
    ws["B1"].number_format = "dd.mm.yyyy"
    ws["A2"], ws["B2"] = "потреба, діб", days
    ws["A3"], ws["B3"] = "осіб на харчуванні", headcount
    for c, title in enumerate(["місце", "добо-дачі"], 1):
        ws.cell(5, c, title)
    bal = _balances(model, as_of)
    live = [i for i in items if not i["archived"] and float(i["per_ration"] or 0) > 0]
    total = 0.0
    r = 5
    for place in places:
        dd = sum(bal.get((it["code"], place), 0) * float(it["per_ration"]) for it in live)
        if abs(dd) < 1e-9:
            continue
        r += 1
        total += dd
        ws.cell(r, 1, place)
        ws.cell(r, 2, _num(round(dd, 3)))
    ws.cell(r + 1, 1, "разом")
    ws.cell(r + 1, 2, _num(round(total, 3)))
    if headcount:
        need = days * float(headcount)
        ws.cell(r + 2, 1, "потреба")
        ws.cell(r + 2, 2, _num(round(need, 3)))
        ws.cell(r + 3, 1, "забезпеченість, %")
        ws.cell(r + 3, 2, round(total * 100 / need, 1) if need else None)


_CELL = re.compile(r'<c r="(?P<ref>[A-Z]+\d+)"(?P<attrs>[^>]*)><f>(?P<f>.*?)</f><v></v></c>', re.S)


def _put_cached(path, cached):
    """openpyxl пише формули без значень (`<f>…</f><v></v>`): дописуємо значення, які порахувала
    програма, — книга, ще не відкрита в Excel, дає скриптам числа й текст, а не порожнечу."""
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheets = _sheet_paths(files)
    for title, values in cached.items():
        part = sheets[title]
        xml = files[part].decode("utf-8")

        def fill(m, values=values):
            v = values.get(m.group("ref"))
            if v is None:
                return m.group(0)
            if isinstance(v, str):
                return f'<c r="{m.group("ref")}"{m.group("attrs")} t="str"><f>{m.group("f")}</f><v>{escape(v)}</v></c>'
            return f'<c r="{m.group("ref")}"{m.group("attrs")}><f>{m.group("f")}</f><v>{v}</v></c>'

        files[part] = _CELL.sub(fill, xml).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)


def _sheet_paths(files):
    """Назва аркуша → шлях його XML у книзі (атрибути тегу <sheet> — у будь-якому порядку)."""
    wbx = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    target = {}
    for tag in re.findall(r"<Relationship\b[^>]*>", rels):
        rid, tgt = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
        if rid and tgt:
            target[rid.group(1)] = tgt.group(1)
    out = {}
    for tag in re.findall(r"<sheet\b[^>]*>", wbx):
        name, rid = re.search(r'\bname="([^"]+)"', tag), re.search(r'\br:id="([^"]+)"', tag)
        if name and rid:
            t = target[rid.group(1)].lstrip("/")
            out[_unescape(name.group(1))] = t if t.startswith("xl/") else "xl/" + t
    return out


def _unescape(s):
    return s.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'").replace("&amp;", "&")


def export_book(con, path, layout, as_of=None, headcount=None, own=False):
    """Книга з бази — цілою або ніяк: пишеться поруч («Облік СП.xlsx.part» — той, що лишився після обриву, наступний
    запис перезаписує) і стає на місце одним кроком. Збій посеред запису лишає попередню книгу як була; книга, відкрита
    в Excel, заміні не піддається — PermissionError.

    `own` — книга в теці книги: програма запам'ятовує її своєю (`remember_sha`) — відбиток байтів і вмісту (вона ж,
    прочитана назад). Не читається назад — лише байтів, а у відповіді `readBack` — що не так (до 50 рядків), як у книги
    ОП. Без `own` (книга для рушія в тимчасовій теці) пам'ять про свої книги не змінюється."""
    m = model_from_db(con, layout)
    path = os.fspath(path)
    part = path + ".part"
    sha = mark = problems = None
    try:
        write_book(m, part, layout, as_of, headcount)
        if own:
            sha = stable_hash.stable_sha(part)
            with open(part, "rb") as f:
                data = f.read()
            try:
                mark = fingerprint(read_book(data, layout))
            except BookError as e:                # через дані самої бази назад не читається: впізнаємо лише за байтами
                problems = e.problems[:50]
            except Exception:                     # noqa: BLE001 — несподіване при читанні назад: книга однаково стає
                traceback.print_exc()             # людині — речення, назва винятку й текст помилки — лише в консоль
                problems = ["Книгу назад не прочитано."]
        os.replace(part, path)
    except BaseException:
        with contextlib.suppress(OSError):
            os.remove(part)
        raise
    if own:
        with contextlib.suppress(sqlite3.Error):  # не запам'ятали — наступного разу книгу звірять із базою
            remember_sha(con, layout, sha, mark)
    out = {"path": path, "items": len(m["items"]), "docs": len(m["docs"]),
           "lines": sum(len(d["lines"]) for d in m["docs"])}
    if problems:
        out["readBack"] = problems
    return out


# ---------------------------------------------------------------- своя книга

def fingerprint(model):
    """Відбиток вмісту книги: sha256 канонічного JSON її моделі (без звіту перенесення, ключі відсортовані). Той самий
    вміст — той самий відбиток, хоч би байти файла були інші: книгу відкрили в Excel і зберегли."""
    text = json.dumps({k: v for k, v in model.items() if k != "report"}, sort_keys=True, ensure_ascii=False,
                      separators=(",", ":"), default=str)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _own(con, layout):
    """Що програма пам'ятає про свої книги цієї будови: `sha` — відбитки байтів (швидкий шлях), `print` — вмісту."""
    row = con.execute("SELECT value FROM app_setting WHERE key = ?", (layout.sha_setting,)).fetchone()
    try:
        got = json.loads(row[0]) if row and row[0] else {}
    except ValueError:
        got = {}
    got = got if isinstance(got, dict) else {}
    return {key: [x for x in got[key] if isinstance(x, str) and x] if isinstance(got.get(key), list) else []
            for key in ("sha", "print")}


def own_shas(con, layout):
    """Відбитки байтів книг, які програма сама записала чи прийняла (новіші спереду; не більше KEEP_SHAS)."""
    return _own(con, layout)["sha"]


def own_prints(con, layout):
    """Відбитки вмісту таких книг (новіші спереду; не більше KEEP_SHAS)."""
    return _own(con, layout)["print"]


def remember_sha(con, layout, sha=None, fingerprint=None):
    """Запам'ятати книгу, яку програма записала чи прийняла: відбиток байтів і/або вмісту, не більше KEEP_SHAS кожного."""
    own = _own(con, layout)
    for key, value in (("sha", sha), ("print", fingerprint)):
        if value:
            own[key] = [value] + [x for x in own[key] if x != value]
        own[key] = own[key][:KEEP_SHAS]
    con.execute("INSERT INTO app_setting(key, value) VALUES(?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (layout.sha_setting, json.dumps(own)))


def own_book(con, path, layout):
    """(своя, модель): чи файл на диску — книга, яку програма сама записала чи прийняла, і книга, прочитана для відповіді
    (None, коли читати не довелось). Спершу відбиток байтів (збіг за будь-яким із двох — stable_hash.shas); далі вміст:
    книгу відкрили в Excel і зберегли — байти інші, а дані ті самі. Незчитна книга — BookError (вона не своя)."""
    if set(stable_hash.shas(path)) & set(own_shas(con, layout)):
        return True, None
    with open(path, "rb") as f:
        model = read_book(f.read(), layout)
    return fingerprint(model) in set(own_prints(con, layout)), model


# ---------------------------------------------------------------- читання й прийом

class BookError(ValueError):
    """Книга не приймається: перелік причин з аркушем і рядком («Прибуток, рядок 12: …»)."""

    def __init__(self, problems):
        super().__init__("; ".join(problems[:5]))
        self.problems = list(problems)


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v or "").strip()
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    return ""


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _code_of(name, fallback=None):
    """Код — з початку тексту «найменування» (як INT(LEFT(…,5)) у книзі), інакше з колонки «код»."""
    m = re.match(r"\s*(\d+)", str(name or ""))
    if m:
        return m.group(1)
    return _text(fallback) if fallback not in (None, "") else ""


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [_text(h).lower() for h in next(it, [])]
    for n, row in enumerate(it, 2):
        yield n, {h: row[i] if i < len(row) else None for i, h in enumerate(head) if h}


def _number(v):
    """Число з клітинки: 0,07 і 0.07; «1/3» (так пишуть «дободач» раціону) — дріб."""
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(",", ".")
    m = re.match(r"^(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)$", s)
    if m and float(m.group(2)):
        return float(m.group(1)) / float(m.group(2))
    try:
        return float(s)
    except ValueError:
        return None


HEAD_NAMES = {"type": "док_тип", "status": "статус", "report": "рапорт", "order": "наказ", "scan": "скан"}


def _comma(v, places=None):
    """Число з комою, як його пише людина: 0,07474; з places — рівно стільки знаків."""
    text = f"{v:.{places}f}" if places is not None else ("%.10g" % v)
    return text.replace(".", ",")


def _status(text):
    """Стан книги за переліком без огляду на регістр: «їде на фес» → «їде на ФЕС»; інакше None."""
    t = _text(text).lower()
    return next((s for s in STATUSES if s.lower() == t), None)


def _item_of(r, where, layout, legacy, problems, report):
    code = _code_of(r.get("код"))
    if not code:
        return None
    kind = _text(r.get("тип"))
    cat = _text(r.get("категорія")).upper()
    archived = kind.lower() == ARCHIVE
    grp = layout.cats.get(cat)
    if grp is None:
        problems.append(f"{where}: невідомий тип «{kind}» і категорія «{_text(r.get('категорія'))}»")
        return None
    price = _number(r.get("ціна"))
    price = 0.0 if price is None else price
    if abs(price * 100 - round(price * 100)) > 1e-6:
        if legacy:
            report.append(f"{where}: ціна {_comma(price)} — до копійок {_comma(round(price, 2), 2)} "
                          "(облік веде ціни в копійках)")
            price = round(price, 2)
        else:
            problems.append(f"{where}: ціна {_comma(price)} з частками копійки — облік веде ціни в копійках")
    it = {"code": code, "name": _text(r.get("найменування")), "cat": layout.cat_of_group[grp],
          "uom": _text(r.get("од.вим.")) or "шт", "price": price,
          "type": "" if archived else kind,
          "archived": archived, "per_ration": _number(r.get("дободач")), "note": _text(r.get("примітка"))}
    return it


def read_book(data, layout, legacy=False):
    """Модель з книги (значення клітинок, не формули). Помилки — з аркушем і рядком.

    `legacy` — перше перенесення старої книги: розбіжні поля документа зводяться до першого
    непорожнього, стан з «примітки» переходить у «статус» (без огляду на регістр), «транз_найменування»
    — у «примітку», ціна з частками копійки округлюється; кожна така правка — у model["report"]."""
    import io                                                    # noqa: PLC0415
    from openpyxl import load_workbook                           # noqa: PLC0415
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:                                       # noqa: BLE001 — будь-який зіпсований файл
        raise BookError([f"файл не читається як книга Excel: {e}"]) from e
    problems, report = [], []
    missing = [s for s in ("Номенклатура", "Прибуток", "Переміщення", "Списання") if s not in wb.sheetnames]
    if missing:
        raise BookError([f"немає аркуша «{s}»" for s in missing])
    renames, subs = [], []
    if "Довідка" in wb.sheetnames:
        renames = _reference(wb["Довідка"])
    alias = {old: new for old, new in renames if old and new} if legacy else {}
    items = []
    for n, r in _rows(wb["Номенклатура"]):
        it = _item_of(r, f"Номенклатура, рядок {n}", layout, legacy, problems, report)
        if it:
            items.append(it)
    codes = {i["code"] for i in items}
    docs, at = [], {}
    for journal in JOURNALS:
        title = TITLES[journal][0]
        for n, r in _rows(wb[title]):
            name = r.get("найменування")
            if not _text(name) and not r.get("док_дата") and not _text(r.get("док_номер")):
                continue
            where = f"{title}, рядок {n}"
            code = _code_of(name, r.get("код"))
            date = _date(r.get("док_дата"))
            no = _text(r.get("док_номер"))
            qty = _number(r.get("кількість")) or 0
            if not date:
                problems.append(f"{where}: немає дати")
            if not no:
                problems.append(f"{where}: немає номера документа")
            if qty <= 0:
                problems.append(f"{where}: кількість має бути більшою за нуль")
            elif not legacy and abs(qty * 1000 - round(qty * 1000)) > 1e-6:
                problems.append(f"{where}: кількість {_comma(qty)} — облік веде кількість до тисячних")
            if code not in codes:
                problems.append(f"{where}: позиції «{_text(name)}» немає в «Номенклатура»")
            status_text = _text(r.get("статус"))
            note = _text(r.get("примітка"))
            status = _status(status_text) if legacy else (status_text if status_text in STATUSES else None)
            if legacy and status and status != status_text:
                report.append(f"{where}: стан «{status_text}» → «{status}»")
            if legacy and not status_text and _status(note):
                status = _status(note)
                report.append(f"{where}: стан «{note}» стояв у «примітка» — перенесено в «статус»")
                note = ""
            elif legacy and not status_text and note.lower() == "переробити":
                status = "переробка"
                report.append(f"{where}: «переробити» в «примітка» — стан «переробка»")
                note = ""
            if status_text and status is None:
                problems.append(f"{where}: невідомий статус «{status_text}»")
            alt = _text(r.get("транз_найменування"))
            if alt:
                if legacy:
                    report.append(f"{where}: «транз_найменування» «{alt}» перенесено в «примітка»")
                    note = (note + " " + alt).strip()
                else:
                    problems.append(f"{where}: колонку «транз_найменування» програма не веде — перенесіть текст у «примітка»")
            if legacy and journal == "movement":
                nm = _text(r.get("нм"))
                if nm and _text(name) and nm != _text(name):
                    report.append(f"{where}: «нм» «{nm}» не збігається з «найменування» — узято «найменування»")
            if journal == "incoming":
                src, dst = _text(r.get("від_кого")), _text(r.get("кому")) or "склад"
            elif journal == "movement":
                src, dst = _text(r.get("від_кого")), _text(r.get("кому"))
            else:
                src, dst = _text(r.get("підрозділ")), _text(r.get("кому"))
            if journal != "incoming" and src in alias:
                src = alias[src]
            if journal != "writeoffs" and dst in alias:
                dst = alias[dst]
            head = {"type": _text(r.get("док_тип")), "status": status or "", "report": _text(r.get("рапорт")),
                    "order": _text(r.get("наказ")), "scan": _text(r.get("скан"))}
            key = (journal, date, no, src, dst)
            d = at.get(key)
            if d is None:
                d = at[key] = {"journal": journal, "date": date, "no": no, "src": src, "dst": dst, **head, "lines": []}
                docs.append(d)
            elif any(d[k] != v for k, v in head.items()):
                differ = [HEAD_NAMES[k] for k, v in head.items() if d[k] != v]
                if legacy:
                    # Акт на кілька рот несе кілька рапортів і пунктів наказу — їх не губимо: через «; ».
                    # Стан і вид документа одні на документ — узято перше непорожнє.
                    joined = []
                    for k, v in head.items():
                        if not d[k] and v:
                            d[k] = v
                        elif v and d[k] != v and k in ("report", "order"):
                            parts = d[k].split("; ")
                            if v not in parts:
                                d[k] = "; ".join(parts + [v])
                                joined.append(HEAD_NAMES[k])
                    first = [x for x in differ if x not in joined]
                    if joined:
                        report.append(f"{where}: у рядках документа №{no} різні {', '.join(joined)} — записано всі через «; »")
                    if first:
                        report.append(f"{where}: у рядках документа №{no} різні {', '.join(first)} — узято перше непорожнє")
                else:
                    problems.append(f"{where}: у рядках одного документа різні {', '.join(differ)}")
            best = r.get("придатний до")
            ln = {"code": code, "qty": _num(qty), "note": note, "best": _date(best) if best not in (None, "") else ""}
            if best not in (None, "") and not ln["best"]:
                problems.append(f"{where}: «придатний до» — не дата")
            ln["q1"] = _text(r.get("1прод"))
            ln["fixed"] = bool(ln["q1"])
            d["lines"].append(ln)
    if "Підрозділи" in wb.sheetnames:
        for n, r in _rows(wb["Підрозділи"]):
            name = _text(r.get("назва"))
            if name:
                subs.append({"name": name, "parent": _text(r.get("підпорядкований")),
                             "kind": _text(r.get("вид")) or "склад", "closed": bool(_text(r.get("закритий"))),
                             "aliases": [a.strip() for a in _text(r.get("псевдоніми")).split(",") if a.strip()]})
    parts = None
    if layout.acts:                                              # акти сухпаю книги СП; аркушів немає — так і сказати
        import sp_book_acts                                      # noqa: PLC0415 — модуль імпортує цей
        parts = sp_book_acts.read(wb, docs, alias, problems)
        if parts is None:
            report.append(sp_book_acts.NO_SHEETS)
    if problems:
        raise BookError(problems)
    model = {"items": items, "subs": subs, "renames": renames, "docs": docs}
    if parts is not None:
        model.update(parts)
    if legacy:
        for old, new in sorted(alias.items()):
            report.append(f"місце «{old}» → «{new}» за tblRename")
    if legacy or report:                                         # «Зауваження» звичайної книги, «Поправки» старої
        model["report"] = report
    return model


def _reference(ws):
    """tblRename (стара назва → нова назва) з аркуша «Довідка»."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    head = grid[0] if grid else []
    pos = {_text(h).lower(): i for i, h in enumerate(head) if h}
    renames = []

    def cell(row, key):
        i = pos.get(key)
        return _text(row[i]) if i is not None and i < len(row) else ""

    for row in grid[1:]:
        if cell(row, "стара назва"):
            renames.append([cell(row, "стара назва"), cell(row, "нова назва")])
    return renames


def _doc_label(d):
    return f"{TITLES[d['journal']][0]} · {d['date']} · №{d['no']} · {d['src'] or '—'} → {d['dst'] or '—'}"


def _doc_sig(d):
    return (d["type"], d["status"], d["report"], d["order"], d["scan"],
            tuple(sorted((ln["code"], float(ln["qty"]), ln["note"], ln.get("best", ""), ln.get("q1", ""))
                         for ln in d["lines"])))


def _doc_key(d):
    return (d["journal"], d["date"], d["no"], d["src"], d["dst"])


def _names(con):
    """Назва підрозділу за назвою чи псевдонімом."""
    out = {r[0]: r[0] for r in con.execute("SELECT name FROM subdivision")}
    for alias, name in con.execute("SELECT a.name, s.name FROM subdivision_alias a "
                                   "JOIN subdivision s ON s.id = a.subdivision_id"):
        out.setdefault(alias, name)
    return out


def _resolved(con, model):
    """Модель з назвами місць, як їх знає база (псевдонім → назва підрозділу): у документах і актах сухпаю."""
    names = _names(con)

    def fix(d):
        d = dict(d)
        if d["journal"] != "incoming":
            d["src"] = names.get(d["src"], d["src"])
        if d["journal"] != "writeoffs":
            d["dst"] = names.get(d["dst"], d["dst"])
        return d
    out = dict(model, docs=[fix(d) for d in model["docs"]])
    for part in ACT_PARTS:
        if model.get(part) is not None:
            out[part] = [dict(x, sub=names.get(x["sub"], x["sub"])) for x in model[part]]
    return out


def _foreign_owners(con, layout):
    """{код: книга} позицій бази, які належать не книзі `layout.book`, а іншій (ТЗ, ОП, СП, РП)."""
    return dict(con.execute("SELECT n.code, g.book FROM nomen n JOIN nomen_group g ON g.id = n.group_id "
                            "WHERE g.book <> ?", (layout.book,)))


def _dormant_foreign(con, layout):
    """{код: книга} сплячих позицій ОП, СП і РП, крім книги `layout.book`: архівних і без жодного рядка документів у базі.
    Такий рядок коду не тримає — його віддають діючій позиції старої книги (`claim_dormant`). Позиції ТЗ не віддаються
    ніколи: каталог ТЗ несе номери ФЕС і прив'язки до форм."""
    return dict(con.execute("""SELECT n.code, g.book FROM nomen n JOIN nomen_group g ON g.id = n.group_id
                                WHERE g.book IN ('ОП', 'СП', 'РП') AND g.book <> ? AND n.archived_at IS NOT NULL
                                  AND NOT EXISTS (SELECT 1 FROM document_line l WHERE l.nomen_id = n.id)""",
                            (layout.book,)))


def _refuse_foreign_codes(con, model, layout):
    """Код один на всю номенклатуру: позицію іншої книги (ТЗ, ОП, СП, РП) ця книга не забирає, а документи тієї позиції не
    переписує. Відмова — переліком «Номенклатура: код N належить книзі «X»», до будь-якого запису. Виняток один —
    перенесення старої книги: спляча позиція, що віддала свій код (`model["claimed"]`, див. `claim_dormant`); віддача
    перевіряється тут щоразу, тож позиція, яка встигла отримати рядки, вже не віддається."""
    owner = _foreign_owners(con, layout)
    claimed = model.get("claimed")
    if claimed:
        asleep = _dormant_foreign(con, layout)
        owner = {c: b for c, b in owner.items() if not (c in claimed and c in asleep)}
    codes = dict.fromkeys(str(i["code"]) for i in model["items"])
    problems = [f"Номенклатура: код {c} належить книзі «{owner[c]}»" for c in codes if c in owner]
    if problems:
        raise BookError(problems)


def diff(con, model, layout):
    """Що зміниться в базі, коли прийняти книгу: документи, позиції, нові місця. Код позиції іншої книги — відмова
    (`BookError`): перегляд її не показує як нову, а прийом не забирає."""
    _refuse_foreign_codes(con, model, layout)
    have = model_from_db(con, layout)
    model = _resolved(con, model)
    mine, theirs = {_doc_key(d): d for d in model["docs"]}, {_doc_key(d): d for d in have["docs"]}
    out = {"docs": {"new": [_doc_label(d) for k, d in mine.items() if k not in theirs],
                    "changed": [_doc_label(d) for k, d in mine.items() if k in theirs and _doc_sig(d) != _doc_sig(theirs[k])],
                    "removed": [_doc_label(d) for k, d in theirs.items() if k not in mine]}}
    hi = {i["code"]: i for i in have["items"]}
    out["items"] = {"new": [i["code"] for i in model["items"] if i["code"] not in hi],
                    "changed": [i["code"] for i in model["items"] if i["code"] in hi and i != hi[i["code"]]],
                    "removed": [c for c in hi if c not in {i["code"] for i in model["items"]}]}
    names = _names(con)
    used = {d["dst"] for d in model["docs"] if d["journal"] != "writeoffs"} | {
        d["src"] for d in model["docs"] if d["journal"] != "incoming"}
    out["subs"] = sorted(n for n in used if n and n not in names)
    if layout.acts:                                              # акти сухпаю книги СП — за ключем документа акта
        import sp_book_acts                                      # noqa: PLC0415 — модуль імпортує цей
        out["acts"] = sp_book_acts.diff(have, model)
        # Кількість нового чи зміненого акта не за правилами — не відмова: рядки «Зауваження» вікна прийому.
        out["remarks"] = sp_book_acts.remarks(have, model)
    out["same"] = not any(out["docs"][k] for k in out["docs"]) and not any(
        out["items"][k] for k in out["items"]) and not out["subs"] and not any(out.get("acts", {}).values())
    return out


def apply(con, model, layout):
    """Книга обліку в базі стає точно такою, як у файлі. Інші книги не змінюються; місця й постачальники
    лише додаються. Усе або нічого: відмова лишає базу як була."""
    import state_db                                              # noqa: PLC0415
    import sp_book_acts                                          # noqa: PLC0415 — модуль імпортує цей
    plan = diff(con, model, layout)
    con.execute("BEGIN IMMEDIATE")
    try:
        _ensure_subs(con, model, layout)
        model = _resolved(con, model)
        _ensure_items(con, model, layout)
        # Акти сухпаю — до запису документів: документ акта, рядки якого книга змінила, записується наново, і дані акта
        # пішли б за старим каскадом (sp_book_acts.apply ставить їх до нового).
        held = sp_book_acts.held(con) if layout.acts else None
        before = state_db.negative_days(con)
        ours = {i["code"] for i in model["items"]} | state_db.book_codes(con, layout.book)
        docs = {j: [r for r in rows if str(r[5]) not in ours] for j, rows in state_db._load_docs(con).items()}
        price = {i["code"]: i["price"] for i in model["items"]}
        for d in model["docs"]:
            for ln in d["lines"]:
                docs[d["journal"]].append([d["date"], d["type"], d["no"], d["src"], d["dst"], ln["code"], ln["qty"],
                                           ln["note"], price[ln["code"]] if d["journal"] == "incoming" else 0])
        try:
            state_db._save_docs(con, docs, seed_drop_limit=False)
        except state_db.SaveError as e:
            raise BookError([str(e)]) from e
        new = {k: v for k, v in state_db.negative_days(con).items() if k not in before}
        if new:
            raise BookError([state_db._negative_words(con, new)])
        _apply_doc_fields(con, model, layout)
        if layout.acts:
            sp_book_acts.apply(con, model, held)
        _drop_unused_items(con, {i["code"] for i in model["items"]}, layout)
        con.execute("COMMIT")
    except BaseException:
        con.execute("ROLLBACK")
        raise
    return plan


def _ensure_subs(con, model, layout):
    """Місця з аркуша «Підрозділи» (у чистій програмі — усе дерево) і назви з документів. Колишня назва перейменованого
    підрозділу (псевдонім) — той самий підрозділ, а не новий: книга, вивантажена до перейменування, несе стару назву."""
    kinds = dict(con.execute("SELECT code, id FROM subdivision_kind"))
    have = dict(con.execute("SELECT name, id FROM subdivision"))
    alias_of = dict(con.execute("SELECT name, subdivision_id FROM subdivision_alias"))
    aliases = set(alias_of)
    sort = con.execute("SELECT COALESCE(MAX(sort), 0) FROM subdivision").fetchone()[0] or 0
    for s in model["subs"]:
        if s["name"] in have or s["name"] in alias_of:
            continue
        if s["kind"] not in kinds:
            kinds[s["kind"]] = con.execute("INSERT INTO subdivision_kind(code, name) VALUES(?, ?)",
                                           (s["kind"], s["kind"])).lastrowid
        sort += 1
        have[s["name"]] = con.execute(
            "INSERT INTO subdivision(name, kind_id, parent_id, sort, is_active) VALUES(?, ?, ?, ?, ?)",
            (s["name"], kinds[s["kind"]], have.get(s["parent"]) or alias_of.get(s["parent"]), sort,
             0 if s["closed"] else 1)).lastrowid
        for a in s["aliases"]:
            if a not in aliases and a not in have:
                con.execute("INSERT INTO subdivision_alias(subdivision_id, name, source) VALUES(?, ?, ?)",
                            (have[s["name"]], a, layout.alias_source))
                aliases.add(a)
    for old, new in model["renames"]:
        if new in have and old not in have and old not in aliases:
            con.execute("INSERT INTO subdivision_alias(subdivision_id, name, source) VALUES(?, ?, ?)",
                        (have[new], old, layout.alias_source))
            aliases.add(old)
    # Назва з документа, якої немає ні в довіднику, ні серед псевдонімів, — помилка, а не новий підрозділ.
    unknown = set()
    for d in model["docs"]:
        names = (d["dst"],) if d["journal"] == "incoming" else (d["src"], d["dst"]) if d["journal"] == "movement" \
            else (d["src"],)
        unknown |= {n for n in names if n and n not in have and n not in aliases}
    if unknown:
        # У старій книзі аркуша «Підрозділи» немає: місце заводять у програмі чи зіставляють у вікні перенесення.
        how = ("додайте його в аркуш «Підрозділи» книги чи заведіть у програмі («Довідники» → «Підрозділи» → «+ Підрозділ»)"
               if model["subs"] else "заведіть такий підрозділ («Довідники» → «Підрозділи» → «+ Підрозділ») або зіставте його з наявним")
        raise BookError([f"невідоме місце «{n}»: {how}" for n in sorted(unknown)])


def _ensure_items(con, model, layout):
    _refuse_foreign_codes(con, model, layout)       # ще раз під замком запису: після перегляду код міг зайняти хтось інший
    groups = dict(con.execute("SELECT code, id FROM nomen_group"))
    uoms = dict(con.execute("SELECT code, id FROM uom"))
    cats = layout.cats
    for it in model["items"]:
        if it["uom"] not in uoms:
            uoms[it["uom"]] = con.execute("INSERT INTO uom(code, name) VALUES(?, ?)", (it["uom"], it["uom"])).lastrowid
        grp = groups[cats[it["cat"]]]
        row = con.execute("SELECT id, archived_at FROM nomen WHERE code = ?", (it["code"],)).fetchone()
        archived = ((row["archived_at"] if row and row["archived_at"] else datetime.date.today().isoformat())
                    if it["archived"] else None)
        kind1 = it["type"] or None
        vals = (it["name"], grp, uoms[it["uom"]], it["per_ration"], it["note"] or None, archived,
                int(round(float(it["price"]) * 100)) or None, kind1)
        if row is None:
            con.execute("INSERT INTO nomen(code, name, group_id, uom_id, per_ration, note, archived_at, app_price_kop, "
                        "prod1_kind, source) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, 'program')", (it["code"], *vals))
        else:
            con.execute("UPDATE nomen SET name = ?, group_id = ?, uom_id = ?, per_ration = ?, note = ?, archived_at = ?, "
                        "app_price_kop = ?, prod1_kind = ? WHERE id = ?", (*vals, row["id"]))
            if str(it["code"]) in model.get("claimed", ()):      # рядок прийшов з іншої книги: її рядок норми 390 не лишається
                con.execute("UPDATE nomen SET norm_line = NULL WHERE id = ?", (row["id"],))


def _book_doc_ids(con, book):
    return {r[0] for r in con.execute("""SELECT DISTINCT l.document_id FROM document_line l
                                           JOIN nomen n ON n.id = l.nomen_id JOIN nomen_group g ON g.id = n.group_id
                                          WHERE g.book = ?""", (book,))}


def _apply_doc_fields(con, model, layout):
    """Стан ФЕС, поля документів і поля рядків книги — як у книзі; документів інших книг не чіпає."""
    import state_db                                              # noqa: PLC0415
    mine = _book_doc_ids(con, layout.book)
    fes = {k: v for k, v in state_db._load_doc_fes(con).items() if int(k) not in mine}
    meta = {k: v for k, v in state_db._load_doc_meta(con).items() if int(k) not in mine}
    extra = {k: v for k, v in state_db._load_line_extra(con).items() if int(k) not in mine}
    ids = {(j, dt, str(no), s, t): doc for j, dt, no, s, t, doc in state_db.doc_ids(con)}
    for d in model["docs"]:
        doc = ids.get(_doc_key(d))
        if doc is None:
            continue
        if d["status"]:
            fes[str(doc)] = {"status": d["status"]}
        if d["report"] or d["order"] or d["scan"]:
            meta[str(doc)] = {"report": d["report"], "order": d["order"], "scan": d["scan"], "col": None}
        for ln in d["lines"]:
            if ln.get("best") or ln.get("q1"):
                extra.setdefault(str(doc), {})[f"{ln['code']}|{ln['note']}"] = {
                    "best": ln.get("best", ""), "q1": ln.get("q1", ""), "fixed": bool(ln.get("fixed"))}
    state_db._save_doc_fes(con, fes)
    state_db._save_doc_meta(con, meta)
    state_db._save_line_extra(con, extra)


def _drop_unused_items(con, keep, layout):
    """Позиції книги, яких у книзі немає й на які не посилається жоден документ, — геть."""
    for nid, code in con.execute("""SELECT n.id, n.code FROM nomen n JOIN nomen_group g ON g.id = n.group_id
                                     WHERE g.book = ?""", (layout.book,)).fetchall():
        if code in keep or con.execute("SELECT 1 FROM document_line WHERE nomen_id = ? LIMIT 1", (nid,)).fetchone():
            continue
        con.execute("DELETE FROM nomen_report_line WHERE nomen_id = ?", (nid,))
        con.execute("DELETE FROM nomen_form_skip WHERE nomen_id = ?", (nid,))
        con.execute("DELETE FROM fes_item_map WHERE nomen_id = ?", (nid,))
        con.execute("DELETE FROM nomen WHERE id = ?", (nid,))


def canonical(con, layout):
    """Порівнювана проєкція книги в базі: модель і залишки по кодах і місцях на кожну дату руху.
    Партії книги програма розкладає сама за порядком надходження — у порівняння вони не входять."""
    m = model_from_db(con, layout)
    dates = sorted({d["date"] for d in m["docs"]})
    rests = {dt: sorted((k, round(v, 3)) for k, v in _balances(m, dt).items() if abs(v) > 1e-9) for dt in dates}
    docs = sorted((_doc_key(d), _doc_sig(d)) for d in m["docs"])
    out = {"items": sorted(m["items"], key=lambda i: i["code"]), "docs": docs, "rests": rests,
           "subs": [(s["name"], s["parent"], s["kind"], s["closed"]) for s in m["subs"]]}
    if layout.acts:
        import sp_book_acts                                      # noqa: PLC0415 — модуль імпортує цей
        out["acts"] = sp_book_acts.canonical(con)
    return out


# ---------------------------------------------------------------- перше перенесення

def legacy_places(model):
    """Назви місць, які називають документи старої книги (після tblRename), за абеткою."""
    used = set()
    for d in model["docs"]:
        if d["journal"] == "incoming":
            used.add(d["dst"])
        elif d["journal"] == "movement":
            used.update((d["src"], d["dst"]))
        else:
            used.add(d["src"])
    return sorted(n for n in used if n)


_BAT = re.compile(r"^(\d+)\s*бат(?:\s+трО)?$", re.I)
_BBS = re.compile(r"^(\d+)\s*ббс$", re.I)


def suggest_places(places, subs, layout):
    """Пропозиція таблиці місць для вікна перенесення: місце, яке вже є в програмі, — саме воно; батальйон
    старої книги («4 бат», «2 бат ТрО», «1 ббс») — його «· ВМТЗ» (рішення власника 10.10 для сухпаю), лише
    якщо такий підрозділ у програмі є. Решта — порожньо: зіставляє людина."""
    names = set(subs)
    out = {}
    for p in places:
        if p in names:
            out[p] = p
            continue
        target = None
        m = _BAT.match(p)
        if m:
            target = f"{m.group(1)} б ТрО · ВМТЗ"
        m = _BBS.match(p)
        if m:
            target = f"{m.group(1)} б БпС · ВМТЗ"
        out[p] = target if target in names else ""
    return out


def _round_quantities(model, report):
    """Облік веде кількість до тисячних, а стара книга мала й п'ять знаків (19,00313 кг). Рядок
    округлюється до тисячних; вибуття, яке в книзі забирало все до нуля, забирає рівно стільки,
    скільки після округлення лишилося (інакше залишок ішов би в мінус на 0,001)."""
    order = {"incoming": 0, "movement": 1, "writeoffs": 2}
    exact, rounded = {}, {}
    for d in sorted(model["docs"], key=lambda d: (d["date"], order[d["journal"]])):
        for ln in d["lines"]:
            q = float(ln["qty"])
            r = round(q + 1e-12, 3)
            src = None if d["journal"] == "incoming" else (ln["code"], d["src"])
            dst = (ln["code"], d["dst"]) if d["journal"] != "writeoffs" else None
            if src is not None:
                have_r, have_e = rounded.get(src, 0.0), exact.get(src, 0.0)
                if r > have_r + 1e-9 and q <= have_e + 1e-9 and r - have_r <= 0.0015:
                    r = round(have_r, 3)
            if abs(r - q) > 1e-9:
                report.append(f"{TITLES[d['journal']][0]} №{d['no']} від {d['date']}, код {ln['code']}: "
                              f"кількість {_comma(q)} → {_comma(r)} (облік веде тисячні)")
                ln["qty"] = _num(r)
            if src is not None:
                rounded[src] = rounded.get(src, 0.0) - r
                exact[src] = exact.get(src, 0.0) - q
            if dst is not None:
                rounded[dst] = rounded.get(dst, 0.0) + r
                exact[dst] = exact.get(dst, 0.0) + q


def drop_foreign_archived(con, model, layout):
    """Перенесення старої книги: архівна позиція без жодного рядка документів, код якої в програмі вже належить іншій
    книзі, не переноситься — код лишається тій книзі (в одній книзі позиція діюча, в іншій — лише сплячий архівний
    рядок). Кожна така позиція — рядок звіту, і ці рядки стоять першими: звіт старої книги довгий, а пропуск позиції
    людина має побачити. Діюча позиція й позиція з рядками лишаються в моделі, і `diff` відмовляє, як завжди. Книга,
    яку записала програма, цього кроку не проходить: там чужий код — завжди відмова. Модель повертається без таких
    позицій (немає таких — та сама), вхідна не змінюється; викликати до `diff` і `apply`, під тим самим замком
    запису."""
    owner = _foreign_owners(con, layout)
    used = {str(ln["code"]) for d in model["docs"] for ln in d["lines"]}
    kept, notes = [], []
    for it in model["items"]:
        code = str(it["code"])
        if not (it["archived"] and code in owner and code not in used):
            kept.append(it)
            continue
        note = f"Номенклатура: архівну позицію з кодом {code} книги «{owner[code]}» не перенесено (рядків у неї немає)"
        if note not in notes:
            notes.append(note)
    if not notes:
        return model
    return dict(model, items=kept, report=notes + list(model.get("report", [])))


def claim_dormant(con, model, layout):
    """Перенесення старої книги: спляча позиція іншої книги (ОП, СП чи РП: архівна, без рядків документів у базі), код якої
    має діюча позиція старої книги, віддає код цій книзі. Прийом перегрупує її рядок у групу книги й перепише найменування,
    ціну й решту полів, як для власних позицій, а рядок норми 390 книги ОП скине. Кожна така позиція — рядок звіту, і ці
    рядки стоять першими. Позиція старої книги, що сама спляча, нічого не забирає: вона поступається
    (`drop_foreign_archived`), і код лишається за позицією бази. Діюча позиція іншої книги й позиція з рядками лишаються за
    своєю книгою — відмова, як завжди (`_refuse_foreign_codes`); книга, яку записала програма, цього кроку не проходить.
    Модель повертається з `claimed` ({код: книга, що віддала}); немає чого віддавати — та сама, вхідна не змінюється.
    Викликати до `diff` і `apply`, під тим самим замком запису."""
    asleep = _dormant_foreign(con, layout)
    used = {str(ln["code"]) for d in model["docs"] for ln in d["lines"]}
    claimed = {}
    for it in model["items"]:
        code = str(it["code"])
        if code in asleep and (not it["archived"] or code in used):
            claimed[code] = asleep[code]
    if not claimed:
        return model
    notes = [f"Номенклатура: архівну позицію {book} з кодом {code} (без рядків) передано книзі «{layout.book}»"
             for code, book in claimed.items()]
    return dict(model, claimed=claimed, report=notes + list(model.get("report", [])))


def normalize_legacy(model, mapping, layout=None):
    """Стара книга → модель програми: місця — за таблицею, яку погодив власник (стара назва →
    підрозділ програми; порожнє — лишити як є). Кожна правка — у model["report"]."""
    report = list(model.get("report", []))
    mapping = {k: v for k, v in (mapping or {}).items() if v and k != v}
    for d in model["docs"]:
        for side in ("src", "dst"):
            if d[side] in mapping and not (d["journal"] == "incoming" and side == "src") \
                    and not (d["journal"] == "writeoffs" and side == "dst"):
                d[side] = mapping[d[side]]
    for part in ACT_PARTS:                                       # акти сухпаю — за тим самим місцем, що й їхні документи
        for x in model.get(part) or []:
            if x["sub"] in mapping:
                x["sub"] = mapping[x["sub"]]
    for old, new in sorted(mapping.items()):
        report.append(f"місце «{old}» → «{new}»")
    _round_quantities(model, report)
    model["report"] = report
    return model
