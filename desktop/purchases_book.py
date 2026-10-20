# -*- coding: utf-8 -*-
"""Книга «Закупівлі» — запасний шлях: усе, що програма знає про кошти й закупівлі, у книзі Excel
без макросів; книгу ведуть руками й приймають назад — у ту саму або в чисту програму.

Записи книги пов'язані ключами, які бачить людина: рядок коштів — «рік · назва», закупівля —
«рік · тема», документ — «вид · номер · дата · контрагент». Ідентифікатори програми в книгу
не йдуть: при прийомі запис, знайдений за ключем, зберігає свій id (і з ним — файли), новий
отримує новий. Позначки «не було» називають записи їхніми ключами (`ver:<закупівля · версія>:<папір>`),
а не номерами, і в чистій програмі стають номерами її власних записів. Обчислювані колонки — формули
зі значенням на момент вивантаження. Аркуш «Частки за програмою» (частки, які документ без своїх бере на сторінці від
документа-основи, — `doc_parts`) програма складає сама при кожному вивантаженні й при прийомі не читає.

Однакові документи (той самий ключ) розрізняє число в кінці — «… #2», «… #3»; документи пишуться за датою, а за рівної дати
за id. Відбитки книги, яку програма сама записала чи прийняла — байтів (швидкий шлях) і вмісту (книга лишається своєю й
після «Зберегти» в Excel: байти інші, дані ті самі), — лежать у app_setting («pur_book_sha», серед state_db.SERVER_SETTINGS):
така книга на диску — не руками правлена, її перезаписують без звірки з базою.

Людина (відповідальний, кому) — «ім'я | id». Номерам людей в іншій базі вірити не можна: id приймають лише разом з іменем
(однаково очищеним з обох боків), інакше людину шукають за іменем; якої немає чи яких кілька — книгу не приймають (BookError з
аркушем і рядком), а не відкидають людину мовчки й не віддають чужій. Людину без імені за іменем не шукають; її id вірять лише
у книзі, яку записала ця сама база й яку не правили. `pid` підписанта лишається, лише коли в базі є людина з таким id і таким
самим іменем, як у підписі (текст підпису лишається завжди).

Усе, що книга читає, звіряється з тим, що вона пише: текст у модель приходить з обох боків в одному вигляді
(`_canon_model`), тож щойно вивантажена книга не виглядає внесеною руками, а круг «база → книга → база»
нічого не губить. Гроші — цілі копійки (у клітинках гривні, читаються точно: `Decimal`, а не `float * 100`)."""
import base64
import collections
import contextlib
import datetime
import decimal
import hashlib
import json
import os
import re
import secrets
import sqlite3
import zlib

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment

import purchase_papers
import purchases_db
import stable_hash
import state_db
from op_book import BookError, _date, _me, _put_cached, _rows, _table, _text

# Підписи акта введення: затверджує, голова, члени комісії й бухгалтерія (extra.approver, head, members, accountant).
ROLES = ("затверджує", "голова", "член", "бухгалтерія")
SINGLE_ROLES = ("затверджує", "голова", "бухгалтерія")
HELP = ["Книга «Закупівлі»: кошти року, закупівлі, документи й файли.",
        "Рядок коштів — «рік · назва»; закупівля — «рік · тема»; документ — «вид · номер · дата · контрагент».",
        "Суми — гривні з копійками; дати — дд.мм.рррр; колонки з формулами не правлять.",
        "Підписант — «посада | звання | ім'я | особа» (особу можна лишити порожньою; номер особи приймають, лише коли в "
        "програмі є людина з таким номером і таким самим іменем, інакше його очищають, а текст підпису лишається).",
        "Документ без частки йде цілим до рядка коштів у колонці «рядок коштів».",
        "Аркуш «Частки за програмою» програма складає сама при кожному вивантаженні: документ без своїх «Частки» й без рядка "
        "коштів бере частки документа-основи («до документа»), розкладені на свою суму; формули «Коштів» лічать і їх. "
        "Його не правлять: при прийомі програма його не читає.",
        "У закупівлі з кількома рядками коштів документу, внесеному руками, потрібні свої рядки в «Частки»: «Частки за "
        "програмою» з'являються лише при вивантаженні, тож до прийому книги формули його не лічать.",
        "Нові рядки дописують у кінець таблиці; книгу потім приймають у програмі: «Закупівлі» → «Прийняти книгу…».",
        "Свої колонки й аркуші програма не зберігає.",
        "Документ без номера в ключі має ще й суму: «вид · б/н · дата · контрагент · сума».",
        "Однакові документи розрізняє число в кінці ключа: другий — «… #2», третій — «… #3» (за порядком рядків у «Документи»).",
        "Колонку «попередні редакції» заповнює програма (список у форматі JSON); руками її не правлять.",
        "Таблицю «Документи» не сортують, коли в ній є «… #2»: близнюки поміняються рядками.",
        "Людина (відповідальний, кому) — «ім'я | id»; id можна не писати: людину знайдуть за іменем, а якщо такої в програмі "
        "нема, книгу не приймуть.",
        "Людину без імені («—») за іменем не знайти: її клітинку приймають, лише поки книга така, як її записала програма.",
        "Значення переліків — рівно такі:",
        "вид руху коштів: " + ", ".join(purchases_db.MOVE_KINDS) + ";",
        "позначка закупівлі: " + ", ".join(x for x in purchases_db.CLOSED if x) + " (або порожньо);",
        "відповідь штабу: " + ", ".join(x for x in purchases_db.ANSWERS if x) + " (або порожньо);",
        "правило ціни: " + ", ".join(purchases_db.PRICE_RULES) + "; вид пропозиції: " + ", ".join(purchases_db.OFFER_KINDS) + ";",
        "вид документа: " + ", ".join(purchases_db.DOC_KINDS) + ";",
        "стан акта введення: " + ", ".join(x for x in purchases_db.DOC_STATES if x) + " (або порожньо);",
        "статус ФЕС: " + ", ".join(state_db.FES_STATUSES) + " (або порожньо);",
        "роль у комісії: " + ", ".join(ROLES) + ".",
        "«До чого» у «Файли»: рядок коштів, версія, документ, реєстр; у «Не було»: рядок коштів, закупівля, версія.",
        "Пункт у «Не було» називає запис ключем: ver:<закупівля · версія>:<папір>, bud:<рядок коштів>:<папір>, "
        "doc:<документ>:<папір>, rcpt:in|дата|номер|від кого|кому:<папір>."]
COLS = {
    "Кошти": ["рік", "назва", "фонд", "громада", "КПКВ", "КЕКВ", "код видатків", "призначення", "угода №",
              "угода від", "рішення №", "рішення від", "повернути до", "звіт до", "примітка", "ключ", "затверджено",
              "профінансовано", "зобов'язання", "використано", "залишок коштів", "без договору"],
    "Рух коштів": ["рядок коштів", "дата", "вид", "сума", "підстава", "примітка"],
    "Закупівлі": ["рік", "тема", "що купуємо", "відповідальний", "з економії", "перероблено з", "позначка", "тека",
                  "примітка", "ключ"],
    "Гроші закупівель": ["закупівля", "рядок коштів", "сума"],
    "Версії": ["закупівля", "версія", "причина", "підпис командира", "супровід №", "супровід від", "відповідь",
               "відповідь №", "відповідь від", "вимоги", "передано закупівельникам", "командир", "ФЕС",
               "начальник служби", "виконавець", "телефон"],
    "Позиції": ["закупівля", "версія", "№", "найменування", "технічні вимоги", "од.", "кількість", "правило ціни",
                "ціна", "сума", "КЕКВ", "код", "код ДК", "рядок коштів", "код номенклатури", "примітка"],
    "Пропозиції": ["закупівля", "версія", "позиція", "№", "вид", "продавець", "посилання", "ціна", "дата перевірки",
                   "КП №", "КП від", "примітка"],
    "Документи": ["вид", "номер", "дата", "закупівля", "контрагент", "кому", "сума", "ПДВ", "рядок коштів",
                  "до документа", "строк поставки", "дійсна до", "використана", "повернута", "до сплати", "процедура",
                  "стан", "статус ФЕС", "дата статусу", "реєстр", "№ проводки", "місце", "постачальник в акті",
                  "перелік документації", "оглянуто", "де знаходиться", "попередні редакції", "скану немає", "примітка"],
    "Рядки документів": ["документ", "№", "найменування", "код", "од.", "кількість", "ціна", "сума", "номер ФЕС",
                         "рік випуску", "№ паспорта", "строк корисного використання"],
    "Частки": ["документ", "вид документа", "рядок коштів", "сума"],
    "Надходження": ["документ", "АП №", "АП від", "від кого", "кому"],
    "Комісії": ["документ", "роль", "посада", "звання", "ім'я", "особа"],
    "Файли": ["до чого", "ключ", "що", "тека", "шлях", "файл", "sha256", "розмір", "склала програма", "дата",
              "№ документа", "тип", "додано"],
    "Не було": ["до чого", "ключ", "пункт", "причина"],
}
TABLES = {"Кошти": "tblBudget", "Рух коштів": "tblMoves", "Закупівлі": "tblPurchases", "Гроші закупівель": "tblPurFunds",
          "Версії": "tblVersions", "Позиції": "tblPositions", "Пропозиції": "tblOffers", "Документи": "tblDocs",
          "Рядки документів": "tblDocLines", "Частки": "tblShares", "Надходження": "tblReceipts", "Комісії": "tblBoard",
          "Файли": "tblFiles", "Не було": "tblMissing"}
# Аркуш, який програма складає сама при кожному вивантаженні й при прийомі не читає (у COLS його немає): частки документів, що не
# мають ні своїх «Частки», ні «рядка коштів», — такі, як їх бере сторінка (`doc_parts`: частки документа-основи, розкладені на
# суму документа). Колонки — ті самі, що в «Частки»; формули «Коштів» лічать обидва аркуші.
SHARES_CALC, SHARES_CALC_TABLE = "Частки за програмою", "tblSharesCalc"
MONEY_COLS = {"сума", "ПДВ", "ціна", "затверджено", "профінансовано", "зобов'язання", "використано", "залишок коштів",
              "без договору"}
DATE_COLS = {"угода від", "рішення від", "повернути до", "звіт до", "дата", "підпис командира", "супровід від",
             "відповідь від", "передано закупівельникам", "дата перевірки", "КП від", "строк поставки", "дійсна до",
             "використана", "повернута", "до сплати", "дата статусу", "АП від", "додано"}
INT_COLS = {"рік", "№", "позиція", "розмір"}
QTY_COLS = {"кількість"}
# Текст, який база тримає як є (переноси рядків, подвійні пробіли): книга повертає його до символа. Подана версія
# не правиться, тож технічні вимоги, що змінилися від зайвого пробілу, не дали б прийняти книгу.
LONG_COLS = {"технічні вимоги", "перелік документації", "постачальник в акті", "місце", "оглянуто", "де знаходиться",
             "попередні редакції", "телефон", "причина", "пункт"}
VERSIONS_COL = "попередні редакції"   # extra.versions акта введення: список JSON, який пише й читає програма
CELL_LIMIT = 32000                    # клітинка Excel тримає до 32767 знаків, довше він не відкриє
SHA_SETTING = "pur_book_sha"          # app_setting: відбитки книги, яку програма сама записала чи прийняла (state_db.SERVER_SETTINGS)
KEEP_SHAS = 10
# Імена файлів і відбитки: без згортання пробілів (подвійний пробіл в імені файла — це інший файл).
RAW_COLS = {"тека", "шлях", "файл", "sha256"}
MONEY_FORMAT, DATE_FORMAT, TEXT_FORMAT = "#,##0.00", "DD.MM.YYYY", "@"
FILE_OWNERS = ("рядок коштів", "версія", "документ", "реєстр")
MISSING_OWNERS = ("рядок коштів", "закупівля", "версія")
MAX_UAH_DIGITS = 16                     # понад це (за порядком) сума — не сума; гроші не більші за purchase_papers.MAX_KOP


def _ref(table, col):
    """Клітинка рядка таблиці в формулі. Лапка в назві колонки («зобов'язання») у структурованому посиланні
    подвоюється, інакше Excel не прочитає формулу."""
    return _me(table, col.replace("'", "''"))


def _sumifs(table, col, *pairs):
    cond = ",".join(f"{table}[{c}],{v}" for c, v in pairs)
    return f"SUMIFS({table}[{col}],{cond})"


B = "tblBudget"
COMMITTED = "зобов'язання"            # колонка з лапкою: у виразі f-рядка її не напишеш без зворотної риски


def _sum_by(table, key_col, kind_col, kind):
    return _sumifs(table, "сума", (key_col, _ref(B, "ключ")), (kind_col, '"' + kind + '"'))


def _docs_on_line(kind):
    """Документи виду `kind` на рядку коштів: цілі (колонка «рядок коштів»), свої частки й частки за програмою."""
    return (f'{_sum_by("tblDocs", "рядок коштів", "вид", kind)}'
            f'+{_sum_by("tblShares", "рядок коштів", "вид документа", kind)}'
            f'+{_sum_by(SHARES_CALC_TABLE, "рядок коштів", "вид документа", kind)}')


CALC = {
    "Кошти": {
        "ключ": f'{_ref(B, "рік")}&" · "&{_ref(B, "назва")}',
        "затверджено": (f'{_sum_by("tblMoves", "рядок коштів", "вид", "затверджено")}'
                        f'+{_sum_by("tblMoves", "рядок коштів", "вид", "зміна")}'),
        "профінансовано": (f'{_sum_by("tblMoves", "рядок коштів", "вид", "профінансовано")}'
                           f'-{_sum_by("tblMoves", "рядок коштів", "вид", "повернуто")}'),
        COMMITTED: _docs_on_line("договір"),
        "використано": _docs_on_line("оплата"),
        "залишок коштів": f'{_ref(B, "профінансовано")}-{_ref(B, "використано")}',
        "без договору": f'{_ref(B, "профінансовано")}-{_ref(B, COMMITTED)}',
    },
    "Закупівлі": {"ключ": f'{_ref("tblPurchases", "рік")}&" · "&{_ref("tblPurchases", "тема")}'},
    "Позиції": {"сума": f'ROUND({_ref("tblPositions", "кількість")}*{_ref("tblPositions", "ціна")},2)'},
    "Рядки документів": {"сума": f'ROUND({_ref("tblDocLines", "кількість")}*{_ref("tblDocLines", "ціна")},2)'},
}


# ---------------------------------------------------------------------------------- гроші й кількість

def _grn(kop):
    return round(int(kop or 0) / 100, 2)


def milli(qty):
    """Кількість у тисячних, точно: 1.115 → 1115 (через текст числа, а не `1.115 * 1000`). Кількість має не більше трьох
    знаків після коми."""
    return round(decimal.Decimal(repr(qty)) * 1000)


def line_kop(qty_milli, price_kop):
    """Сума рядка в копійках: кількість (тисячні) × ціна (копійки), округлена до копійки від нуля — як `ROUND(кількість*ціна;2)`
    в Excel. Усе цілими числами; сторінка (`lineKop`), книга й папери рахують однаково."""
    p = int(qty_milli) * int(price_kop)
    return (abs(p) + 500) // 1000 * (1 if p >= 0 else -1)


# ------------------------------------------------------------------- частки документа — як на сторінці

def _number(v):
    """Сума так, як її бере сторінка (`+x || 0`): число лишається числом, решта (порожньо, NaN, нескінченність) — 0. База
    тримає суми цілими копійками, тож тексту тут не буває."""
    if isinstance(v, bool) or not isinstance(v, (int, float)) or v != v or v in (float("inf"), float("-inf")):
        return 0
    return v


def _div_round(n, d):
    """n / d до цілого, найближче; половина — від нуля (`divRound` сторінки). d ненульовий."""
    r = (2 * abs(n) + abs(d)) // (2 * abs(d))
    return -r if (n < 0) != (d < 0) else r


def scale_parts(parts, kop):
    """Ті самі пропорції на іншу суму — `scaleParts` сторінки, цілими числами: залишок округлення бере остання частка з вагою того
    самого знаку, що й сума ваг; без ваг (сума нуль) — уся сума на перший рядок."""
    if not parts:
        return []
    weight = [int(_number(p.get("kop"))) for p in parts]          # int() відкидає дріб до нуля, як Math.trunc
    total = sum(weight)
    if total == 0:
        return [{"line": parts[0].get("line"), "kop": kop}]
    whole, sign = int(_number(kop)), 1 if total > 0 else -1
    last = len(parts) - 1
    while weight[last] * sign <= 0:                               # сума ваг не нуль, тож така частка є
        last -= 1
    left, out = whole, []
    for i, p in enumerate(parts):
        if i == last:
            out.append(None)
            continue
        k = _div_round(whole * weight[i], total)
        left -= k
        out.append({"line": p.get("line"), "kop": k})
    out[last] = {"line": parts[last].get("line"), "kop": left}
    return out


def doc_parts(doc, docs, purchases, depth=0):
    """Частки документа між рядками коштів — `docParts` сторінки: свої; без них — частки документа-основи (ВН — як договір),
    розкладені на суму документа (`scale_parts`), не глибше чотирьох основ; без основи — цілим до єдиного рядка закупівлі;
    інакше — не віднесено (порожньо). `docs`, `purchases` — записи сторінки за id."""
    own = [p for p in doc.get("parts") or [] if p and p.get("line")]
    if own:
        return [{"line": p["line"], "kop": _number(p.get("kop"))} for p in own]
    up = docs.get(doc.get("parent")) if doc.get("parent") and depth < 4 else None
    if up is not None and up.get("id") != doc.get("id"):
        inherited = doc_parts(up, docs, purchases, depth + 1)
        if inherited:
            return scale_parts(inherited, _number(doc.get("kop")))
    pu = purchases.get(doc.get("purchase"))
    lines = [f.get("line") for f in pu.get("funds") or []] if pu else []
    return [{"line": lines[0], "kop": _number(doc.get("kop"))}] if len(lines) == 1 else []


def _first_by_id(rows):
    """Записи за id; з однаковим id — перший (як `find` сторінки)."""
    out = {}
    for r in rows:
        out.setdefault(r.get("id"), r)
    return out


_MONEY = decimal.Context(prec=100)           # точність із запасом: у числі не більше 64 знаків, а помноження на 100 — точне


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _shown(v):
    """Значення клітинки для повідомлення про помилку: коротко й однією стрічкою."""
    if isinstance(v, datetime.datetime):
        v = v.date()
    s = v.isoformat() if isinstance(v, datetime.date) else (repr(v) if isinstance(v, float) else str(v))
    s = " ".join(s.split())
    return s if len(s) <= 40 else s[:37] + "…"


def _decimal(v):
    """Число клітинки точно: float — через найкоротший запис (`repr`), текст — без пробілів, зокрема нерозривних, і з комою як
    крапкою. Не число, bool, NaN і нескінченність — None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        d = decimal.Decimal(v)
    elif isinstance(v, float):
        d = decimal.Decimal(repr(v))
    elif isinstance(v, str):
        s = "".join(v.split()).replace(",", ".")
        if len(s) > 64:
            return None
        try:
            d = decimal.Decimal(s)
        except decimal.InvalidOperation:
            return None
    else:
        return None
    return d if d.is_finite() else None


def _kop(v, where, problems):
    """Сума клітинки (гривні) → цілі копійки, точно. Не число — «не число»; дробова копійка — «точність до копійки»: ні
    мовчазного відкидання, ні округлення (11882.67 у `float` — це 11882.67, а не 11882.669999…)."""
    if _blank(v):
        return 0
    d = _decimal(v)
    if d is None:
        problems.append(f"{where}: сума «{_shown(v)}» — не число")
        return 0
    if d != 0 and d.adjusted() > MAX_UAH_DIGITS:
        problems.append(f"{where}: сума «{_shown(v)}» — завелика")
        return 0
    if d != 0 and d.adjusted() < -2:                              # менше копійки: помножувати нема чого
        problems.append(f"{where}: сума «{_shown(v)}» — точність до копійки")
        return 0
    k = _MONEY.multiply(d, 100)
    kop = int(k)
    if decimal.Decimal(kop) != k:
        problems.append(f"{where}: сума «{_shown(v)}» — точність до копійки")
        return 0
    if abs(kop) > purchase_papers.MAX_KOP:
        problems.append(f"{where}: сума «{_shown(v)}» — завелика")
        return 0
    return kop


def _int(v, col, where, problems):
    if _blank(v):
        return 0
    d = _decimal(v)
    if d is None:
        problems.append(f"{where}: {col} «{_shown(v)}» — не число")
        return 0
    if d != d.to_integral_value() or abs(d) > 10 ** 12:
        problems.append(f"{where}: {col} «{_shown(v)}» — потрібне ціле число")
        return 0
    return int(d)


def _qty(v, col, where, problems):
    """Кількість: ціле лишається цілим, дробове — `float` (як у сторінки); не більше трьох знаків після коми."""
    if _blank(v):
        return 0
    d = _decimal(v)
    if d is None:
        problems.append(f"{where}: {col} «{_shown(v)}» — не число")
        return 0
    if d != 0 and d.adjusted() > 12:
        problems.append(f"{where}: {col} «{_shown(v)}» — завелика")
        return 0
    if d != 0 and (d.adjusted() < -3 or _MONEY.multiply(d, 1000) != _MONEY.multiply(d, 1000).to_integral_value()):
        problems.append(f"{where}: {col} «{_shown(v)}» — точність до тисячних")
        return 0
    return int(d) if d == d.to_integral_value() else float(d)


def _day(v, col, where, problems):
    """Дата клітинки → «рррр-мм-дд»; чого не прочитано як дату (і року поза межами обліку) — помилка, а не порожнє поле."""
    if _blank(v):
        return ""
    iso = _date(v)
    try:
        day = datetime.date.fromisoformat(iso) if iso else None
    except ValueError:
        day = None
    low, high = state_db.DATE_YEARS
    if day is None:
        problems.append(f"{where}: {col} «{_shown(v)}» — не дата, потрібно дд.мм.рррр")
        return ""
    if not low <= day.year <= high:
        problems.append(f"{where}: {col} «{_shown(v)}» — рік поза межами {low}–{high}")
        return ""
    return iso


# ------------------------------------------------------------------------------------------ тексти

def _clean(s):
    """Знаки керування (\\x0b із Word тощо) в XML не бувають: їх у книгу не пишуть і з неї не читають."""
    return ILLEGAL_CHARACTERS_RE.sub("", s)


def _long(v):
    """Текст як є: переноси рядків, пробіли. Лише кінці рядків приводяться до \\n."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _clean(str(v)).replace("\r\n", "\n").replace("\r", "\n")


def _raw(v):
    return _long(v).strip()


def _line(v):
    """Звичайний текст: пробіли згорнуто (як у полях сторінки)."""
    return _text(_clean(v) if isinstance(v, str) else v)


def _canon(col, v):
    """Вигляд текстового значення в моделі — той самий, що дає читання клітинки, звідки б значення не прийшло."""
    if not isinstance(v, str) or col in DATE_COLS:
        return v
    return _long(v) if col in LONG_COLS else _raw(v) if col in RAW_COLS else _line(v)


def _canon_model(m):
    return {name: [{c: _canon(c, v) for c, v in row.items()} for row in rows] for name, rows in m.items()}


def _versions_text(versions):
    """extra.versions акта введення → клітинка: стислий JSON; порожній список — порожня клітинка. Довге (клітинка Excel
    тримає до 32767 знаків, довшу він не відкриє) стискається: «gz:» і base64 від zlib."""
    if not isinstance(versions, list) or not versions:
        return ""
    text = json.dumps(versions, ensure_ascii=False, separators=(",", ":"))
    if len(text) <= CELL_LIMIT:
        return text
    packed = "gz:" + base64.b64encode(zlib.compress(text.encode("utf-8"), 9)).decode("ascii")
    if len(packed) > CELL_LIMIT:
        raise ValueError("попередні редакції акта не вміщаються в клітинку Excel")
    return packed


def _versions_of(text):
    """Клітинка → список; порожня — порожній список; ValueError, коли це не список у форматі JSON. Стиснене розпаковується
    з межею (клітинка — вхід із мережі): понад 8 МБ — відмова."""
    text = (text or "").strip()
    if not text:
        return []
    try:
        if text.startswith("gz:"):
            unpack = zlib.decompressobj()
            raw = unpack.decompress(base64.b64decode(text[3:], validate=True), 8 << 20)
            if not unpack.eof:                                   # дійшли до межі, не до кінця потоку, або потік обірваний
                raise ValueError("завелике або обірване")
            text = raw.decode("utf-8")
        got = json.loads(text)
    except (ValueError, zlib.error) as e:                        # JSONDecodeError, UnicodeDecodeError, binascii.Error
        raise ValueError("не JSON") from e
    if not isinstance(got, list):
        raise ValueError("не список")
    return got


def _value(col, v, where, problems):
    if col == VERSIONS_COL:
        text = _long(v)
        try:
            _versions_of(text)
        except ValueError:
            problems.append(f"{where}: {col} — не список у форматі JSON (цю колонку заповнює програма)")
            return ""
        return text
    if col in MONEY_COLS:
        return _kop(v, where, problems)
    if col in DATE_COLS:
        return _day(v, col, where, problems)
    if col in INT_COLS:
        return _int(v, col, where, problems)
    if col in QTY_COLS:
        return _qty(v, col, where, problems)
    if isinstance(v, (datetime.datetime, datetime.date)):
        problems.append(f"{where}: {col} «{_shown(v)}» — Excel зробив із тексту дату: задайте клітинці текстовий формат "
                        "і введіть значення знову")
        return ""
    return _long(v) if col in LONG_COLS else _raw(v) if col in RAW_COLS else _line(v)


# ------------------------------------------------------------------------------------------- ключі

def _ddmm(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}" if iso else ""


def line_key(year, label):
    return _text(f"{year} · {label}")


def pur_key(year, title):
    return _text(f"{year} · {title}")


def ver_key(purchase_key, label):
    return _text(f"{purchase_key} · {label}")


def doc_key(kind, no, date, party, kop):
    key = " · ".join([kind, no or "б/н", _ddmm(date), party or ""])
    return _text(key + (f" · {_grn(kop):.2f}" if not no else ""))


def suffixed(keys):
    """Ключі документів у порядку рядків: однакові розрізняє число в кінці — другий «… #2», третій «… #3». Порядок і правило
    однакові в книзі, яку пише програма, і в читанні: посилання з інших аркушів влучають у свій документ. Ключ, що вже
    є (скажімо, контрагент, названий «… #2»), пропускається."""
    taken, count, out = set(keys), collections.Counter(), []
    for key in keys:
        count[key] += 1
        n = count[key]
        if n == 1:
            out.append(key)
            continue
        while f"{key} #{n}" in taken:
            n += 1
        taken.add(f"{key} #{n}")
        out.append(f"{key} #{n}")
    return out


def doc_keys(rows):
    """Повні ключі рядків «Документи» (рядки книги), у порядку рядків."""
    return suffixed([doc_key(r["вид"], r["номер"], r["дата"], r["контрагент"], r["сума"]) for r in rows])


def page_docs(docs):
    """Документи сторінки в порядку книги — за датою, а за рівної дати за id, — і їхні повні ключі."""
    ordered = sorted(docs, key=lambda d: (d["date"] or "", d["id"]))
    return ordered, suffixed([doc_key(d["kind"], d["no"], d["date"], d["party"], d["kop"]) for d in ordered])


def _sg_text(sg):
    parts = [str((sg or {}).get(k) or "") for k in ("pos", "rank", "name", "pid")]
    return " | ".join(parts) if any(p.strip() for p in parts) else ""


def _sg_of(text):
    parts = ([p.strip() for p in str(text or "").split("|")] + ["", "", "", ""])[:4]
    return {"pid": parts[3], "pos": parts[0], "rank": parts[1], "name": parts[2]}


def _receipt_key(date, no, party, sub):
    return "|".join(["in", date or "", no or "", party or "", sub or ""])


def _receipts_index(con):
    out = {}
    for r in con.execute("""SELECT d.id, d.doc_date, d.number, COALESCE(c.name, '') AS party, COALESCE(s.name, '') AS sub
                              FROM document d LEFT JOIN counterparty c ON c.id = d.counterparty_id
                              LEFT JOIN subdivision s ON s.id = d.to_subdivision_id"""):
        out[r["id"]] = {"no": purchases_db._t(r["number"]), "date": r["doc_date"] or "",
                        "from": purchases_db._t(r["party"]), "to": purchases_db._t(r["sub"])}
    return out


PLACEHOLDER = "—"      # так state_db._full_name записує людину без імені


def _fold(s):
    """Ім'я для порівняння: те саме очищення, що й у клітинці книги (`_line`: без знаків керування, пробіли згорнуто), без регістру.
    Якби база й клітинка чистились по-різному, власна книга програми не впізнала б власну людину."""
    return _line(s).casefold()


def _nameless(name):
    """Людина без імені: порожньо чи заглушка «—»."""
    return _fold(name) in ("", PLACEHOLDER)


def people_index(con):
    """Люди бази для клітинок «ім'я | id»: `name` — id застосунку → ім'я (очищене, як у клітинці); `id` — будь-який дійсний запис
    id (b<номер> чи ext_id) → id застосунку; `by_name` — ім'я без регістру → id застосунку (одна людина чи кілька). Людей без
    імені в `by_name` немає: за іменем їх не шукають ніколи."""
    name, ident, by_name = {}, {}, {}
    for r in con.execute("SELECT id, ext_id, full_name FROM person ORDER BY id"):
        app = r["ext_id"] or f"b{r['id']}"
        name[app] = _line(r["full_name"])
        ident[f"b{r['id']}"] = app
        if r["ext_id"]:
            ident[r["ext_id"]] = app
        if not _nameless(name[app]):
            by_name.setdefault(_fold(name[app]), []).append(app)
    return {"name": name, "id": ident, "by_name": by_name}


def _person_cell(app_id, people):
    """Людина в клітинці: «ім'я | id». Одного id мало: у чужій базі «b5» — інша людина."""
    return f"{((people or {}).get('name') or {}).get(app_id, '')} | {app_id}".strip() if app_id else ""


def _person_of(cell, people, where, problems):
    """Клітинка «ім'я | id» → id людини в цій базі (порожня клітинка — нікого). Номерам людей з іншої бази вірити не можна:
    1. id є тут і людина з цим id має таке саме ім'я (те саме очищення, без регістру) → вона;
    2. інакше людина з таким іменем — рівно одна → вона; кілька тезок — відмова («допишіть id»); жодної — відмова, а коли id є
       тут, але в іншої людини, — відмова з її іменем;
    3. людина без імені (порожньо чи «—») за іменем не шукається ніколи. Її id вірять лише тоді, коли id є тут, людина тут теж
       без імені І книга — та, яку записала чи прийняла ця сама база й яку відтоді не правили (`people["home"]`: відбиток
       вмісту книги серед запам'ятованих у цій базі). Інакше один чужий id міг би вказати на чужу безіменну людину.
    Відмова — BookError з аркушем і рядком: не мовчазна втрата й не чужа людина. Ім'я без id (набране руками) — ті самі правила."""
    text = _line(cell)
    if not text:
        return ""
    name, ident = (x.strip() for x in text.rsplit("|", 1)) if "|" in text else (text, "")
    people = people or {}
    names = people.get("name") or {}
    app = (people.get("id") or {}).get(ident)
    if _nameless(name):
        if app and people.get("home") and _nameless(names.get(app, "")):
            return app
        problems.append(f"{where}: у клітинці «{text}» людина без імені — її номеру вірять лише в книзі, яку записала ця програма "
                        "й яку не правили; дайте людині ім'я в «Люди й МВО» або очистіть клітинку")
        return ""
    if app and _fold(names.get(app, "")) == _fold(name):
        return app
    found = (people.get("by_name") or {}).get(_fold(name), [])
    if len(found) == 1:
        return found[0]
    if found:
        problems.append(f"{where}: людина «{name}» в програмі не одна — допишіть її id: «{name} | id»")
    elif app:                                      # id є, але це інша людина: назвати, хто саме
        holder = "людина без імені" if _nameless(names.get(app, "")) else "«" + names.get(app, "") + "»"
        problems.append(f"{where}: «{text}» — {ident} тут {holder}; людини «{name}» немає в програмі — "
                        "заведіть її в «Люди й МВО» або виправте клітинку")
    else:
        problems.append(f"{where}: людини «{name}» немає в програмі — заведіть її в «Люди й МВО» або очистіть клітинку")
    return ""


def _signer(sg, people):
    """Підписант із книги. Його `pid` — id людини в базі, де книгу записано; тут він лишається, лише коли така людина є й її ім'я
    таке саме, як у підписанта (те саме очищення; безіменних не зіставляють). Інакше `pid` порожній, а текст підписанта (посада,
    звання, ім'я) лишається як є: чужий номер не вказує на чужу людину."""
    people = people or {}
    app = (people.get("id") or {}).get(sg["pid"])
    same = bool(app) and not _nameless(sg["name"]) and _fold((people.get("name") or {}).get(app, "")) == _fold(sg["name"])
    return dict(sg, pid=sg["pid"] if same else "")


PUR_FILES = ("bud|", "purver|", "purdoc|", "fesreg|")


def _files(con):
    out = []
    for r in con.execute("SELECT * FROM app_file ORDER BY pos"):
        if not str(r["target"]).startswith(PUR_FILES):
            continue
        extra = json.loads(r["extra"]) if r["extra"] else {}
        out.append({"key": r["target"], "date": r["doc_date"] or "", "no": r["doc_no"] or "", "file": r["file_name"],
                    "path": r["rel_path"], "mime": r["mime"] or "", "size": r["size_bytes"] or 0,
                    "folder": r["folder"] or "", "sha": r["sha256"] or "", "added": extra.get("added", ""),
                    "made": bool(extra.get("made")), "extra": extra})
    return out


# Позначка «не було» названа номером запису програми: ver:<id версії>:<папір>, bud:<id рядка>:…, doc:<id документа>:…,
# rcpt:<id акта приймання>:…. У книзі номер замінює ключ запису — номерів чужої бази там не знають.
MARK = re.compile(r"^(ver|bud|doc|rcpt):(.+):([^:]+)$")


def _mark(item, names):
    m = MARK.match(item)
    if m and m.group(2) in names[m.group(1)]:
        return f"{m.group(1)}:{names[m.group(1)][m.group(2)]}:{m.group(3)}"
    return item


def to_model(st, files, rcpt, people=None):
    """Розділи сторінки + файли + акти приймання (+ люди бази, `people_index`) → рядки книги (чиста функція: її читає й diff),
    разом із рядками «Частки за програмою» (SHARES_CALC: їх пише `write_book`, а читання книги не бере)."""
    lines = {ln["id"]: ln for ln in st.get("budget", [])}
    purs = {p["id"]: p for p in st.get("purchases", [])}
    docs = {d["id"]: d for d in st.get("purDocs", [])}
    ordered, doc_names = page_docs(st.get("purDocs", []))          # за датою, а за рівної дати — за id: порядок не залежить від бази
    dkey = {d["id"]: k for d, k in zip(ordered, doc_names)}
    lk = lambda lid: line_key(lines[lid]["year"], lines[lid]["label"]) if lid in lines else ""       # noqa: E731
    pk = lambda pid: pur_key(purs[pid]["year"], purs[pid]["title"]) if pid in purs else ""          # noqa: E731
    dk = lambda d: dkey[d["id"]]                                                                   # noqa: E731
    ver_name = {v["id"]: ver_key(pk(pu["id"]), v["label"]) for pu in st.get("purchases", []) for v in pu["versions"]}
    names = {"ver": ver_name, "bud": {i: lk(i) for i in lines}, "doc": dict(dkey),
             "rcpt": {str(i): _receipt_key(a["date"], a["no"], a["from"], a["to"]) for i, a in rcpt.items()}}
    marks = lambda d: sorted((_mark(item, names), why) for item, why in (d or {}).items())          # noqa: E731
    m = {name: [] for name in COLS}
    m[SHARES_CALC] = []
    doc_by_id, pur_by_id = _first_by_id(st.get("purDocs", [])), _first_by_id(st.get("purchases", []))
    for ln in st.get("budget", []):
        m["Кошти"].append({"рік": ln["year"], "назва": ln["label"], "фонд": ln["fund"], "громада": ln["community"],
                           "КПКВ": ln["kpkv"], "КЕКВ": ln["kekv"], "код видатків": ln["exp"],
                           "призначення": ln["purpose"], "угода №": ln["agreeNo"], "угода від": ln["agreeDate"],
                           "рішення №": ln["decisionNo"], "рішення від": ln["decisionDate"],
                           "повернути до": ln["returnBy"], "звіт до": ln["reportBy"], "примітка": ln["note"]})
        for mv in ln["moves"]:
            m["Рух коштів"].append({"рядок коштів": lk(ln["id"]), "дата": mv["date"], "вид": mv["kind"],
                                    "сума": mv["kop"], "підстава": mv["basis"], "примітка": mv["note"]})
        for item, why in marks(ln.get("missing")):
            m["Не було"].append({"до чого": "рядок коштів", "ключ": lk(ln["id"]), "пункт": item, "причина": why})
    for pu in st.get("purchases", []):
        m["Закупівлі"].append({"рік": pu["year"], "тема": pu["title"], "що купуємо": pu["what"],
                               "відповідальний": _person_cell(pu["person"], people), "з економії": pk(pu["parent"]),
                               "перероблено з": pk(pu["source"]), "позначка": pu["closed"], "тека": pu["folder"],
                               "примітка": pu["note"]})
        for f in pu["funds"]:
            m["Гроші закупівель"].append({"закупівля": pk(pu["id"]), "рядок коштів": lk(f["line"]), "сума": f["kop"]})
        for item, why in marks(pu.get("missing")):
            m["Не було"].append({"до чого": "закупівля", "ключ": pk(pu["id"]), "пункт": item, "причина": why})
        for v in pu["versions"]:
            s = v.get("signers") or {}
            m["Версії"].append({"закупівля": pk(pu["id"]), "версія": v["label"], "причина": v["reason"],
                                "підпис командира": v["signedOn"], "супровід №": v["sentNo"], "супровід від": v["sentOn"],
                                "відповідь": v["answer"], "відповідь №": v["answerNo"], "відповідь від": v["answerOn"],
                                "вимоги": v["answerNote"], "передано закупівельникам": v["handedOn"],
                                "командир": _sg_text(s.get("cmdr")), "ФЕС": _sg_text(s.get("fes")),
                                "начальник служби": _sg_text(s.get("chief")), "виконавець": _sg_text(s.get("exec")),
                                "телефон": str((s.get("exec") or {}).get("phone") or "")})
            for item, why in marks(v.get("missing")):
                m["Не було"].append({"до чого": "версія", "ключ": ver_name[v["id"]], "пункт": item, "причина": why})
            for n, p in enumerate(v["positions"], 1):
                m["Позиції"].append({"закупівля": pk(pu["id"]), "версія": v["label"], "№": n, "найменування": p["name"],
                                     "технічні вимоги": p["specs"], "од.": p["uom"], "кількість": p["qty"],
                                     "правило ціни": p["rule"], "ціна": p["priceKop"], "КЕКВ": p["kekv"],
                                     "код": p["exp"], "код ДК": p["dk"], "рядок коштів": lk(p["line"]),
                                     "код номенклатури": p["code"], "примітка": p["note"]})
                for k, o in enumerate(p["offers"], 1):
                    m["Пропозиції"].append({"закупівля": pk(pu["id"]), "версія": v["label"], "позиція": n, "№": k,
                                            "вид": o["kind"], "продавець": o["seller"], "посилання": o["url"],
                                            "ціна": o["priceKop"], "дата перевірки": o["checked"], "КП №": o["no"],
                                            "КП від": o["date"], "примітка": o["note"]})
    for d in ordered:
        pu = purs.get(d["purchase"])
        single = pu["funds"][0]["line"] if pu and len(pu["funds"]) == 1 else ""
        whole = lk(d["parts"][0]["line"]) if len(d["parts"]) == 1 else ("" if d["parts"] else lk(single))
        # Документ без своїх часток сторінка відносить за основою чи за єдиним рядком закупівлі (`doc_parts`): що «рядок коштів»
        # цього не каже, несуть «Частки за програмою», а «рядок коштів» тоді порожній — інакше формули лічили б документ двічі.
        calc = [] if d["parts"] else doc_parts(d, doc_by_id, pur_by_id)
        if whole and calc == [{"line": single, "kop": _number(d["kop"])}]:
            calc = []
        elif calc:
            whole = ""
        ex = d.get("extra") or {}
        m["Документи"].append({"вид": d["kind"], "номер": d["no"], "дата": d["date"], "закупівля": pk(d["purchase"]),
                               "контрагент": d["party"], "кому": _person_cell(d["person"], people), "сума": d["kop"], "ПДВ": d["vatKop"],
                               "рядок коштів": whole, "до документа": dk(docs[d["parent"]]) if d["parent"] in docs else "",
                               "строк поставки": d["deliverBy"], "дійсна до": d["validTo"], "використана": d["usedOn"],
                               "повернута": d["returnedOn"], "до сплати": d["payOn"], "процедура": d["procedure"],
                               "стан": d["state"], "статус ФЕС": d["fes"]["status"], "дата статусу": d["fes"]["date"],
                               "реєстр": d["fes"]["register"], "№ проводки": d["fes"]["ref"],
                               "місце": str(ex.get("place") or ""), "постачальник в акті": str(ex.get("supplier") or ""),
                               "перелік документації": str(ex.get("docs") or ""), "оглянуто": str(ex.get("seen") or ""),
                               "де знаходиться": str(ex.get("location") or ""),
                               VERSIONS_COL: _versions_text(ex.get("versions")),
                               "скану немає": d["missing"], "примітка": d["note"]})
        for n, ln in enumerate(d["lines"], 1):
            m["Рядки документів"].append({"документ": dk(d), "№": n, "найменування": ln["name"], "код": ln["code"],
                                          "од.": ln["uom"], "кількість": ln["qty"], "ціна": ln["priceKop"],
                                          "номер ФЕС": ln["fesNo"], "рік випуску": ln["year"], "№ паспорта": ln["passport"],
                                          "строк корисного використання": ln["life"]})
        for p in (d["parts"] if len(d["parts"]) > 1 else []):
            m["Частки"].append({"документ": dk(d), "вид документа": d["kind"], "рядок коштів": lk(p["line"]), "сума": p["kop"]})
        for p in calc:
            m[SHARES_CALC].append({"документ": dk(d), "вид документа": d["kind"], "рядок коштів": lk(p["line"]),
                                   "сума": p["kop"]})
        for r in d["receipts"]:
            a = rcpt.get(r["doc"]) or {}
            m["Надходження"].append({"документ": dk(d), "АП №": a.get("no", ""), "АП від": a.get("date", ""),
                                     "від кого": a.get("from", ""), "кому": a.get("to", "")})
        for role, sg in [("затверджує", ex.get("approver")), ("голова", ex.get("head"))] + \
                [("член", x) for x in ex.get("members") or []] + [("бухгалтерія", ex.get("accountant"))]:
            if sg:
                m["Комісії"].append({"документ": dk(d), "роль": role, "посада": str(sg.get("pos") or ""),
                                     "звання": str(sg.get("rank") or ""), "ім'я": str(sg.get("name") or ""),
                                     "особа": str(sg.get("pid") or "")})
    for f in files:
        kind, _, rest = f["key"].partition("|")
        ident, _, what = rest.partition("|")
        owner = {"bud": ("рядок коштів", lk(ident)), "purver": ("версія", ver_name.get(ident, "")),
                 "purdoc": ("документ", dk(docs[ident]) if ident in docs else ""),
                 "fesreg": ("реєстр", ident)}[kind]
        if not owner[1]:                                  # файл запису, якого вже немає: книга його назвати не може
            continue
        m["Файли"].append({"до чого": owner[0], "ключ": owner[1], "що": what, "тека": f["folder"], "шлях": f["path"],
                           "файл": f["file"], "sha256": f["sha"], "розмір": f["size"],
                           "склала програма": "так" if f.get("made") else "", "дата": f["date"],
                           "№ документа": f["no"], "тип": f["mime"], "додано": f["added"]})
    m = _canon_model(m)
    m["help"] = HELP
    return m


def model_from_db(con):
    return to_model(purchases_db.load(con), _files(con), _receipts_index(con), people_index(con))


# -------------------------------------------------------------------------------------- вивантаження

def _date_cell(iso):
    if not iso:
        return None
    try:
        return datetime.date.fromisoformat(iso)
    except ValueError:
        return iso                       # те, що база тримає не датою, видно в клітинці — а прочитати його назад не вдасться


def _put(cell, col, v):
    if col in MONEY_COLS:
        cell.value = None if v in ("", None) else _grn(v)
        cell.number_format = MONEY_FORMAT
    elif col in DATE_COLS:
        cell.value = _date_cell(v)
        if isinstance(cell.value, datetime.date):
            cell.number_format = DATE_FORMAT
    elif col in INT_COLS or col in QTY_COLS:
        cell.value = None if v in ("", None) else v
    else:
        cell.value = None if v in ("", None) else (_clean(v) if isinstance(v, str) else v)
        if isinstance(cell.value, str):
            cell.data_type = "s"            # «=…» у примітці — текст, а не формула (і не чужа формула в книзі служби)
        cell.number_format = TEXT_FORMAT    # «040», «1/26», «1.0» Excel без цього зробив би числом чи датою
        if col in LONG_COLS and col != VERSIONS_COL:         # довгий JSON не переносять: рядок виріс би до межі висоти
            cell.alignment = Alignment(wrap_text=True, vertical="top")       # переноси рядків видно


def write_book(model, path):
    """Книга на `path`: формули з кешованими значеннями. Пише й дописує значення в той самий файл — тож `path` має бути
    тимчасовим (`export_book` кладе його на місце лише цілим)."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    cached = {}
    calc_values = _calc_values(model)
    for name, cols in COLS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        calc = CALC.get(name, {})
        for r, row in enumerate(model[name], 2):
            for c, col in enumerate(cols, 1):
                cell = ws.cell(r, c)
                if col in calc:
                    cell.value = "=" + calc[col]
                    cached.setdefault(name, {})[cell.coordinate] = calc_values[name][r - 2][col]
                    if col in MONEY_COLS:
                        cell.number_format = MONEY_FORMAT
                    continue
                _put(cell, col, row.get(col, ""))
        _table(ws, TABLES[name], cols, len(model[name]), calc)
        for c, col in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 40 if col in LONG_COLS else max(10, min(40, len(col) + 4))
    # Частки за програмою: їх рахує програма (`to_model`), а модель, складена не з бази (історія закупівель), їх не має —
    # тоді аркуш порожній, але є: на його таблицю посилаються формули «Коштів».
    ws, cols, rows = wb.create_sheet(SHARES_CALC), COLS["Частки"], model.get(SHARES_CALC) or []
    ws.append(cols)
    for r, row in enumerate(rows, 2):
        for c, col in enumerate(cols, 1):
            _put(ws.cell(r, c), col, row.get(col, ""))
    _table(ws, SHARES_CALC_TABLE, cols, len(rows))
    for c, col in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = max(10, min(40, len(col) + 4))
    ws = wb.create_sheet("Як заповнювати")
    for text in model.get("help", HELP):
        ws.append([text])
    ws.column_dimensions["A"].width = 120
    wb.save(path)
    _put_cached(path, cached)


def _calc_values(model):
    """Значення обчислюваних колонок — так, як їх порахує Excel."""
    out = {name: [{} for _ in model[name]] for name in COLS}
    for i, ln in enumerate(model["Кошти"]):
        key = line_key(ln["рік"], ln["назва"])
        moves = [mv for mv in model["Рух коштів"] if mv["рядок коштів"] == key]
        s = lambda kinds: sum(mv["сума"] for mv in moves if mv["вид"] in kinds)                     # noqa: E731
        whole = lambda kind: sum(d["сума"] for d in model["Документи"] if d["вид"] == kind and d["рядок коштів"] == key)  # noqa: E731
        share = lambda kind: sum(p["сума"] for p in model["Частки"] + (model.get(SHARES_CALC) or [])            # noqa: E731
                                 if p["вид документа"] == kind and p["рядок коштів"] == key)
        appr, fin = s({"затверджено", "зміна"}), s({"профінансовано"}) - s({"повернуто"})
        com, used = whole("договір") + share("договір"), whole("оплата") + share("оплата")
        out["Кошти"][i] = {"ключ": key, "затверджено": _grn(appr), "профінансовано": _grn(fin),
                           "зобов'язання": _grn(com), "використано": _grn(used), "залишок коштів": _grn(fin - used),
                           "без договору": _grn(fin - com)}
    for i, pu in enumerate(model["Закупівлі"]):
        out["Закупівлі"][i] = {"ключ": pur_key(pu["рік"], pu["тема"])}
    for name in ("Позиції", "Рядки документів"):
        for i, r in enumerate(model[name]):
            out[name][i] = {"сума": _grn(line_kop(milli(r["кількість"] or 0), r["ціна"] or 0))}
    return out


def export_book(con, path):
    """Книга з бази — на `path`, цілою або ніяк: пишеться в тимчасовий файл поруч (разом з кешованими значеннями формул) і
    стає на місце одним кроком. Збій лишає попередню книгу як була; книга, відкрита в Excel, — PaperConflict(busy).
    Відбитки книги (байтів і вмісту) беруться з тимчасового файла, коли він готовий, а на місце ще не став, і
    запам'ятовуються (`remember_sha`): `own_book` пізнає її й не питає, чи правили її руками. Тимчасовий файл читається назад
    (`read_book`): якщо через дані самої бази він не читається, книга однаково стає на місце, у відповіді `readBack` — що не
    так (до 50 рядків), а відбиток вмісту не запам'ятовується (її впізнають лише за байтами, доки її не зберегли в Excel)."""
    m = model_from_db(con)
    path = os.fspath(path)
    sealed = {}

    def write(part):
        write_book(m, part)
        sealed["part"] = part

    def seal():                                   # книга цілком готова (імена Excel вже в ній), на місце ще не стала
        with contextlib.suppress(Exception):      # не вдалося — книга однаково стане, лише її не впізнають як свою
            sealed["sha"] = stable_hash.stable_sha(sealed["part"])
            with open(sealed["part"], "rb") as f:
                data = f.read()
            try:                                  # читаємо її назад так, як читатиме прийом
                sealed["print"] = fingerprint(read_book(data))
            except BookError as e:                # через дані самої бази (двійники приходів…): книга однаково стане, але назад
                sealed["problems"] = e.problems[:50]          # не читається — відбитка вмісту нема, а що не так, скажемо

    purchase_papers._write_atomic(path, write, before=seal)
    if sealed.get("sha"):
        with contextlib.suppress(sqlite3.Error):  # не запам'ятали — наступного разу книгу звірять із базою
            remember_sha(con, sealed["sha"], sealed.get("print"))
    out = {"path": path, "purchases": len(m["Закупівлі"]), "docs": len(m["Документи"])}
    if sealed.get("problems"):
        out["readBack"] = sealed["problems"]
    return out


def fingerprint(model):
    """Відбиток вмісту книги: sha256 канонічного JSON її даних (без номерів рядків і довідки, ключі відсортовані). Той самий
    вміст дає той самий відбиток, хоч би байти файла були інші: книгу відкрили в Excel і зберегли."""
    text = json.dumps(_strip(model), sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _own(con):
    """Що програма пам'ятає про свої книги: `sha` — відбитки байтів (швидкий шлях), `print` — відбитки вмісту. Лежить в
    app_setting під ключем із `state_db.SERVER_SETTINGS`: запис вікна цього не стирає й у стан вікна воно не йде."""
    row = con.execute("SELECT value FROM app_setting WHERE key = ?", (SHA_SETTING,)).fetchone()
    try:
        got = json.loads(row[0]) if row and row[0] else {}
    except ValueError:
        got = {}
    if isinstance(got, list):                     # перший вигляд значення: самі відбитки байтів
        got = {"sha": got}
    got = got if isinstance(got, dict) else {}
    return {key: [x for x in got[key] if isinstance(x, str) and x] if isinstance(got.get(key), list) else []
            for key in ("sha", "print")}


def own_shas(con):
    """Відбитки байтів книг, які програма сама записала чи прийняла (новіші спереду; не більше KEEP_SHAS)."""
    return _own(con)["sha"]


def own_prints(con):
    """Відбитки вмісту таких книг (новіші спереду; не більше KEEP_SHAS)."""
    return _own(con)["print"]


def remember_sha(con, sha=None, fingerprint=None):
    """Запам'ятати книгу, яку програма записала чи прийняла: відбиток байтів і/або вмісту, не більше KEEP_SHAS кожного."""
    own = _own(con)
    for key, value in (("sha", sha), ("print", fingerprint)):
        if value:
            own[key] = [value] + [x for x in own[key] if x != value]
        own[key] = own[key][:KEEP_SHAS]
    con.execute("INSERT INTO app_setting(key, value) VALUES(?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (SHA_SETTING, json.dumps(own)))


def own_book(con, path):
    """(своя, модель): чи файл на диску — той, який програма сама записала чи прийняла, і книга, прочитана для відповіді (або
    None, коли читати не довелось). Спершу швидкий шлях — відбиток байтів (R25: збіг за будь-яким із двох); далі вміст:
    книгу відкрили в Excel і зберегли — байти інші, а дані ті самі. Незчитна книга — BookError (вона не своя)."""
    if set(stable_hash.shas(path)) & set(own_shas(con)):
        return True, None
    with open(path, "rb") as f:
        model = read_book(f.read())
    return fingerprint(model) in set(own_prints(con)), model


def is_own(con, path):
    try:
        return own_book(con, path)[0]
    except BookError:
        return False


# --------------------------------------------------------------------------------------------- читання

def read_book(data):
    """Модель з книги (лише значення; формули пропускає). Помилки — з аркушем і рядком."""
    import io                                                    # noqa: PLC0415
    from openpyxl import load_workbook                           # noqa: PLC0415
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:                                       # noqa: BLE001 — будь-який зіпсований файл
        raise BookError([f"файл не читається як книга Excel: {e}"]) from e
    try:
        missing = [s for s in COLS if s not in wb.sheetnames]
        if missing:
            raise BookError([f"немає аркуша «{s}»" for s in missing])
        problems, m = [], {}
        for name, cols in COLS.items():
            m[name] = _read_sheet(wb[name], name, cols, problems)
    finally:
        wb.close()
    m["help"] = HELP
    _check(m, problems)
    if problems:
        raise BookError(problems)
    return m


def _read_sheet(ws, name, cols, problems):
    rows, calc, headed = [], CALC.get(name, {}), False
    try:
        for n, raw in _rows(ws):
            if not headed:                                       # шапка: бракує колонки — її дані мовчки б зникли
                headed = True
                lost = [c for c in cols if c not in calc and c.lower() not in raw]
                problems.extend(f"{name}: немає колонки «{c}»" for c in lost)
                if lost:
                    return rows
            row = {c: raw.get(c.lower()) for c in cols if c not in calc}
            if all(_blank(v) for v in row.values()):
                continue
            where = f"{name}, рядок {n}"
            row = {c: _value(c, v, where, problems) for c, v in row.items()}
            row["_row"] = where
            rows.append(row)
    except BookError:
        raise
    except Exception as e:                                       # noqa: BLE001 — аркуш, який openpyxl не розібрав
        problems.append(f"{name}: аркуш не читається: {e}")
    return rows


# Тека закупівлі з книги: рівнів не більше трьох (наступний займає версія, «…/2.0», а глибше чотирьох ні папір, ні файл не
# лягають), частина шляху — до 80 знаків (стільки лишає завантаження файла, папір — до 120), без знаків, які папір
# (`safe_name`) чи файл (`launcher.BAD_FILE_CHARS`) замінили б на «-», і без дати рррр-мм-дд (папір пише її як дд.мм.рррр).
FOLDER_DEPTH, FOLDER_PART = 3, 80
FOLDER_BAD = re.compile(r'[\\:*?"<>|\[\]\x00-\x1f\x7f]')
FOLDER_DAY = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")                # те саме правило дати, що в `safe_name`


def folder_problem(text):
    """Чому тека закупівлі з книги не годиться (None — годиться). Сторінка бере теку як є, а частину шляху папір чистить одним
    правилом (`excel_export.safe_name`: дата рррр-мм-дд → дд.мм.рррр, знаки → «-», краї, до 120 знаків), а завантаження файла —
    іншим (`launcher.BAD_FILE_CHARS`, краї, до 80 знаків): частина, яку змінило б хоч одне з них, розвела б файли однієї справи
    по різних теках. Тож тека з книги має бути такою, якою її лишають обидва правила."""
    if not text:
        return None
    parts = text.split("/")
    if len(parts) > FOLDER_DEPTH:
        return f"не глибше {FOLDER_DEPTH} рівнів: наступний — для версії"
    for part in parts:
        if not part:
            return "зайва «/»"
        bad = FOLDER_BAD.search(part)
        if bad and bad.group(0) == "\\":
            return "теки розділяє «/», а не «\\»"
        if bad:
            return ("знак керування (табуляція, перенос рядка) в імені теки" if bad.group(0) < " " or bad.group(0) == "\x7f"
                    else f"знак «{bad.group(0)}» в імені теки програма замінила б на «-»")
        if part != part.strip(" ."):
            return f"«{_shown(part)}» починається чи закінчується пробілом або крапкою"
        day = FOLDER_DAY.search(part)
        if day:
            y, mo, d = day.group(0).split("-")
            return f"дату «{day.group(0)}» в імені теки пишуть як {d}.{mo}.{y}"
        if len(part) > FOLDER_PART:
            return f"«{_shown(part)}» довша за {FOLDER_PART} знаків"
    return None


def _check(m, problems):
    """Зв'язки й значення, які бачить лише книга цілком: кожен рядок знаходить свого батька, ключі не повторюються, переліки
    — з дозволених значень. Усе — з аркушем і рядком: рядок, що нікуди не веде, не зникає мовчки."""
    low, high = state_db.DATE_YEARS
    lines, purs, docs, vers, poss = {}, {}, {}, {}, {}
    kind_of = {}

    def once(table, key, row, what):
        if key in table:
            problems.append(f"{row['_row']}: {what} уже є ({table[key]})")
        else:
            table[key] = row["_row"]

    def one_of(row, col, allowed, blank=True):
        v = row[col]
        if v == "" and blank:
            return
        if v not in allowed:
            problems.append(f"{row['_row']}: {col} «{v}» — невідоме значення, потрібне одне з: {', '.join(allowed)}")

    def year_of(row):
        if not low <= row["рік"] <= high:
            problems.append(f"{row['_row']}: рік «{row['рік']}» — потрібен рік від {low} до {high}")

    for r in m["Кошти"]:
        year_of(r)
        for col in ("назва", "фонд"):
            if not r[col]:
                problems.append(f"{r['_row']}: не вказано {col}")
        once(lines, line_key(r["рік"], r["назва"]), r, f"рядок коштів «{line_key(r['рік'], r['назва'])}»")
    for r in m["Закупівлі"]:
        year_of(r)
        if not r["тема"]:
            problems.append(f"{r['_row']}: не вказано тему")
        once(purs, pur_key(r["рік"], r["тема"]), r, f"закупівля «{pur_key(r['рік'], r['тема'])}»")
        one_of(r, "позначка", [x for x in purchases_db.CLOSED if x])
        folder = folder_problem(r["тека"])
        if folder:
            problems.append(f"{r['_row']}: тека «{_shown(r['тека'])}» — {folder}")
    doc_names = doc_keys(m["Документи"])               # однакові документи розрізнені числом у кінці ключа — повтору нема
    for r, key in zip(m["Документи"], doc_names):
        docs[key] = r["_row"]
        kind_of[key] = r["вид"]
        if r["вид"] not in purchases_db.DOC_KINDS:
            problems.append(f"{r['_row']}: невідомий вид документа «{r['вид']}»")
        one_of(r, "стан", [x for x in purchases_db.DOC_STATES if x])
        one_of(r, "статус ФЕС", state_db.FES_STATUSES)

    def ref(sheet, col, known, required=True):
        for r in m[sheet]:
            v = r[col]
            if not v:
                if required:
                    problems.append(f"{r['_row']}: не вказано {col}")
            elif v not in known:
                problems.append(f"{r['_row']}: {col} «{v}» — немає в книзі")

    ref("Закупівлі", "з економії", purs, required=False)
    ref("Закупівлі", "перероблено з", purs, required=False)
    ref("Рух коштів", "рядок коштів", lines)
    ref("Гроші закупівель", "закупівля", purs)
    ref("Гроші закупівель", "рядок коштів", lines)
    ref("Версії", "закупівля", purs)
    ref("Позиції", "закупівля", purs)
    ref("Позиції", "рядок коштів", lines, required=False)
    ref("Пропозиції", "закупівля", purs)
    ref("Документи", "закупівля", purs, required=False)
    ref("Документи", "рядок коштів", lines, required=False)
    ref("Документи", "до документа", docs, required=False)
    ref("Рядки документів", "документ", docs)
    ref("Частки", "документ", docs)
    ref("Частки", "рядок коштів", lines)
    ref("Надходження", "документ", docs)
    ref("Комісії", "документ", docs)
    for r in m["Рух коштів"]:
        one_of(r, "вид", purchases_db.MOVE_KINDS, blank=False)
    for r in m["Версії"]:
        if not r["версія"]:
            problems.append(f"{r['_row']}: не вказано версію (скажімо, 1.0)")
        once(vers, (r["закупівля"], r["версія"]), r, f"версія «{r['версія']}» закупівлі «{r['закупівля']}»")
        one_of(r, "відповідь", [x for x in purchases_db.ANSWERS if x])
    ver_names = {ver_key(p, label) for p, label in vers}
    for r in m["Позиції"]:
        if r["закупівля"] in purs and (r["закупівля"], r["версія"]) not in vers:
            problems.append(f"{r['_row']}: версія «{r['версія']}» закупівлі «{r['закупівля']}» — немає в книзі")
        if r["№"] < 1:
            problems.append(f"{r['_row']}: № позиції має бути від 1")
        once(poss, (r["закупівля"], r["версія"], r["№"]), r, f"позиція № {r['№']} версії «{r['версія']}»")
        one_of(r, "правило ціни", purchases_db.PRICE_RULES, blank=False)
    for r in m["Пропозиції"]:
        if r["закупівля"] in purs:
            if (r["закупівля"], r["версія"]) not in vers:
                problems.append(f"{r['_row']}: версія «{r['версія']}» закупівлі «{r['закупівля']}» — немає в книзі")
            elif (r["закупівля"], r["версія"], r["позиція"]) not in poss:
                problems.append(f"{r['_row']}: позиція № {r['позиція']} версії «{r['версія']}» закупівлі "
                                f"«{r['закупівля']}» — немає в книзі")
        one_of(r, "вид", purchases_db.OFFER_KINDS, blank=False)
    shares = collections.defaultdict(int)
    for r in m["Частки"]:
        shares[r["документ"]] += r["сума"]
        if r["документ"] in kind_of and r["вид документа"] and r["вид документа"] != kind_of[r["документ"]]:
            problems.append(f"{r['_row']}: вид документа «{r['вид документа']}» не збігається з видом документа "
                            f"«{kind_of[r['документ']]}»")
    for r, key in zip(m["Документи"], doc_names):
        if key in shares:
            if shares[key] != r["сума"]:
                problems.append(f"{r['_row']}: частки ({_grn(shares[key]):.2f}) не дають суми документа ({_grn(r['сума']):.2f})")
            if r["рядок коштів"]:
                problems.append(f"{r['_row']}: і рядок коштів, і частки — залиште одне")
    single = collections.Counter()
    for r in m["Комісії"]:
        one_of(r, "роль", ROLES, blank=False)
        if r["роль"] in SINGLE_ROLES:
            single[(r["документ"], r["роль"])] += 1
            if single[(r["документ"], r["роль"])] > 1:
                problems.append(f"{r['_row']}: роль «{r['роль']}» у документа лише одна")
    owners = {"рядок коштів": lines, "версія": ver_names, "документ": docs}
    for r in m["Файли"]:
        one_of(r, "до чого", FILE_OWNERS, blank=False)
        if not r["ключ"]:
            problems.append(f"{r['_row']}: не вказано ключ")
        elif r["до чого"] in owners and r["ключ"] not in owners[r["до чого"]]:
            problems.append(f"{r['_row']}: ключ «{r['ключ']}» — немає в книзі")
    owners = {"рядок коштів": lines, "закупівля": purs, "версія": ver_names}
    marked = {}
    for r in m["Не було"]:
        one_of(r, "до чого", MISSING_OWNERS, blank=False)
        if not r["ключ"]:
            problems.append(f"{r['_row']}: не вказано ключ")
        elif r["до чого"] in owners and r["ключ"] not in owners[r["до чого"]]:
            problems.append(f"{r['_row']}: ключ «{r['ключ']}» — немає в книзі")
        if not r["пункт"]:
            problems.append(f"{r['_row']}: не вказано пункт")
        elif (r["до чого"], r["ключ"], r["пункт"]) in marked:       # база тримає пункт одним записом: другий зник би мовчки
            problems.append(f"{r['_row']}: пункт «{r['пункт']}» цього запису уже є ({marked[(r['до чого'], r['ключ'], r['пункт'])]})")
        else:
            marked[(r["до чого"], r["ключ"], r["пункт"])] = r["_row"]
    came = {}
    for r in m["Надходження"]:
        mark = (r["документ"], r["АП №"], r["АП від"], r["від кого"], r["кому"])
        if mark in came:                                            # той самий прихід двічі в одного документа — один зв'язок у базі
            problems.append(f"{r['_row']}: прихід «{r['АП №']}» цього документа уже є ({came[mark]})")
        else:
            came[mark] = r["_row"]
    for r, key in zip(m["Документи"], doc_names):
        if r["до документа"] and r["до документа"] == key:
            problems.append(f"{r['_row']}: документ не може бути «до документа» сам для себе")


# ------------------------------------------------------------------------------------------- прийом

def _collect_ids(node, out):
    """Усі id, що вже є в записах (у кожному розділі й кожному вкладеному переліку)."""
    if isinstance(node, dict):
        if isinstance(node.get("id"), str):
            out.add(node["id"])
        for v in node.values():
            _collect_ids(v, out)
    elif isinstance(node, list):
        for v in node:
            _collect_ids(v, out)


def _group(rows, *cols):
    out = collections.defaultdict(list)
    for r in rows:
        out[tuple(r[c] for c in cols) if len(cols) > 1 else r[cols[0]]].append(r)
    return out


def from_model(m, have, rcpt_keys, old_files=None, people=None):
    """Рядки книги → розділи сторінки й файли; записи, знайдені за ключем у `have`, зберігають id (а з ними — файли), решта
    отримує новий: «u» і 16 шістнадцяткових цифр, яких немає ніде серед наявних записів і серед щойно виданих. Прихід, якого
    нема в обліку (`rcpt_keys`), — BookError з аркушем і рядком: так він не губиться мовчки."""
    used = set()
    _collect_ids(have, used)

    def new():
        while True:
            i = "u" + secrets.token_hex(8)
            if i not in used:
                used.add(i)
                return i

    problems = []
    old_lines = {line_key(x["year"], x["label"]): x for x in have.get("budget", [])}
    old_purs = {pur_key(x["year"], x["title"]): x for x in have.get("purchases", [])}
    old_ordered, old_names = page_docs(have.get("purDocs", []))      # ті самі порядок і ключі, що в книзі, яку писала програма
    old_docs = dict(zip(old_names, old_ordered))
    moves_of, funds_of = _group(m["Рух коштів"], "рядок коштів"), _group(m["Гроші закупівель"], "закупівля")
    versions_of, positions_of = _group(m["Версії"], "закупівля"), _group(m["Позиції"], "закупівля", "версія")
    offers_of = _group(m["Пропозиції"], "закупівля", "версія", "позиція")
    lines_of, shares_of = _group(m["Рядки документів"], "документ"), _group(m["Частки"], "документ")
    board_of, receipts_of = _group(m["Комісії"], "документ"), _group(m["Надходження"], "документ")

    lid = {line_key(r["рік"], r["назва"]): (old_lines.get(line_key(r["рік"], r["назва"])) or {}).get("id") or new()
           for r in m["Кошти"]}
    pid = {pur_key(r["рік"], r["тема"]): (old_purs.get(pur_key(r["рік"], r["тема"])) or {}).get("id") or new()
           for r in m["Закупівлі"]}
    vid = {}
    for r in m["Версії"]:
        was = {v["label"]: v for v in (old_purs.get(r["закупівля"]) or {}).get("versions", [])}
        vid[(r["закупівля"], r["версія"])] = (was.get(r["версія"]) or {}).get("id") or new()
    row_keys = doc_keys(m["Документи"])
    # Нові документи отримують id за зростанням у порядку рядків: наступне вивантаження (за датою, а за рівної дати — за id)
    # піде в тому самому порядку, тож однакові документи не міняються номерами «#2».
    fresh_ids = iter(sorted(new() for key in row_keys if key not in old_docs))
    did = {key: (old_docs.get(key) or {}).get("id") or next(fresh_ids) for key in row_keys}
    names = {"ver": {ver_key(p, label): i for (p, label), i in vid.items()}, "bud": lid, "doc": did,
             "rcpt": {k: str(i) for k, i in rcpt_keys.items()}}
    missing = {}
    for r in m["Не було"]:
        missing.setdefault((r["до чого"], r["ключ"]), {})[_mark(r["пункт"], names)] = r["причина"]

    budget = []
    for r in m["Кошти"]:
        key = line_key(r["рік"], r["назва"])
        old = list((old_lines.get(key) or {}).get("moves") or [])
        moves = [{"id": old.pop(0)["id"] if old else new(), "kind": mv["вид"], "date": mv["дата"], "kop": mv["сума"],
                  "basis": mv["підстава"], "note": mv["примітка"]} for mv in moves_of.get(key, [])]
        budget.append({"id": lid[key], "year": r["рік"], "label": r["назва"], "fund": r["фонд"], "kpkv": r["КПКВ"],
                       "kekv": r["КЕКВ"], "exp": r["код видатків"], "community": r["громада"],
                       "purpose": r["призначення"], "agreeNo": r["угода №"], "agreeDate": r["угода від"],
                       "decisionNo": r["рішення №"], "decisionDate": r["рішення від"], "returnBy": r["повернути до"],
                       "reportBy": r["звіт до"], "missing": missing.get(("рядок коштів", key), {}),
                       "note": r["примітка"], "moves": moves})
    purchases = []
    for r in m["Закупівлі"]:
        key = pur_key(r["рік"], r["тема"])
        old_ver = {v["label"]: v for v in (old_purs.get(key) or {}).get("versions", [])}
        versions = []
        for v in versions_of.get(key, []):
            ov = old_ver.get(v["версія"]) or {}
            old_pos = {p["name"] + "|" + str(n): p for n, p in enumerate(ov.get("positions", []), 1)}
            positions = []
            for p in positions_of.get((key, v["версія"]), []):
                op = old_pos.get(p["найменування"] + "|" + str(p["№"])) or {}
                old_off = list(op.get("offers") or [])
                offers = [{"id": old_off.pop(0)["id"] if old_off else new(), "kind": o["вид"], "seller": o["продавець"],
                           "url": o["посилання"], "priceKop": o["ціна"], "checked": o["дата перевірки"], "no": o["КП №"],
                           "date": o["КП від"], "note": o["примітка"]}
                          for o in offers_of.get((key, v["версія"], p["№"]), [])]
                positions.append({"id": op.get("id") or new(), "name": p["найменування"], "specs": p["технічні вимоги"],
                                  "uom": p["од."], "qty": p["кількість"], "rule": p["правило ціни"], "priceKop": p["ціна"],
                                  "kekv": p["КЕКВ"], "exp": p["код"], "dk": p["код ДК"],
                                  "line": lid.get(p["рядок коштів"], ""), "code": p["код номенклатури"],
                                  "note": p["примітка"], "offers": offers})
            ex = dict(_signer(_sg_of(v["виконавець"]), people), phone=v["телефон"])
            versions.append({"id": vid[(key, v["версія"])], "label": v["версія"], "reason": v["причина"],
                             "signedOn": v["підпис командира"], "sentNo": v["супровід №"], "sentOn": v["супровід від"],
                             "answer": v["відповідь"], "answerNo": v["відповідь №"], "answerOn": v["відповідь від"],
                             "answerNote": v["вимоги"], "handedOn": v["передано закупівельникам"],
                             "signers": {"cmdr": _signer(_sg_of(v["командир"]), people), "fes": _signer(_sg_of(v["ФЕС"]), people),
                                         "chief": _signer(_sg_of(v["начальник служби"]), people), "exec": ex},
                             "missing": missing.get(("версія", ver_key(key, v["версія"])), {}), "positions": positions})
        purchases.append({"id": pid[key], "year": r["рік"], "title": r["тема"], "what": r["що купуємо"],
                          "person": _person_of(r["відповідальний"], people, r.get("_row", "Закупівлі"), problems),
                          "parent": pid.get(r["з економії"], ""),
                          "source": pid.get(r["перероблено з"], ""), "closed": r["позначка"], "folder": r["тека"],
                          "missing": missing.get(("закупівля", key), {}), "note": r["примітка"],
                          "funds": [{"line": lid.get(f["рядок коштів"], ""), "kop": f["сума"]} for f in funds_of.get(key, [])],
                          "versions": versions})
    by_id = {p["id"]: p for p in purchases}
    docs = []
    sg = lambda x: _signer({"pid": x["особа"], "pos": x["посада"], "rank": x["звання"], "name": x["ім'я"]}, people)  # noqa: E731
    for r, key in zip(m["Документи"], row_keys):
        was = old_docs.get(key) or {}
        old_l = list(was.get("lines") or [])
        lines_ = [{"id": old_l.pop(0)["id"] if old_l else new(), "name": x["найменування"], "code": x["код"],
                   "uom": x["од."], "qty": x["кількість"], "priceKop": x["ціна"], "fesNo": x["номер ФЕС"],
                   "year": x["рік випуску"], "passport": x["№ паспорта"], "life": x["строк корисного використання"]}
                  for x in lines_of.get(key, [])]
        parts = [{"line": lid.get(x["рядок коштів"], ""), "kop": x["сума"]} for x in shares_of.get(key, [])]
        board = board_of.get(key, [])
        extra = dict(was.get("extra") or {})
        for k_, col in (("place", "місце"), ("supplier", "постачальник в акті"), ("docs", "перелік документації"),
                        ("seen", "оглянуто"), ("location", "де знаходиться")):
            if r[col]:
                extra[k_] = r[col]
            else:
                extra.pop(k_, None)
        try:
            versions = _versions_of(r[VERSIONS_COL])
        except ValueError:                                   # читання такого не пропустило; модель, зібрана не з книги, — як порожня
            versions = []
        if versions or r["вид"] == "акт введення":
            extra["versions"] = versions
        else:
            extra.pop("versions", None)
        if board or r["вид"] == "акт введення":              # комісія — рівно та, що в книзі: прибраний рядок прибирає людину
            extra["approver"] = next((sg(x) for x in board if x["роль"] == "затверджує"), {})
            extra["head"] = next((sg(x) for x in board if x["роль"] == "голова"), {})
            extra["members"] = [sg(x) for x in board if x["роль"] == "член"]
            extra["accountant"] = next((sg(x) for x in board if x["роль"] == "бухгалтерія"), {})
        else:
            for k_ in ("approver", "head", "members", "accountant"):
                extra.pop(k_, None)
        receipts = []
        for x in receipts_of.get(key, []):
            rk = _receipt_key(x["АП від"], x["АП №"], x["від кого"], x["кому"])
            if rk in rcpt_keys:
                receipts.append({"doc": rcpt_keys[rk], "key": rk})
            else:
                problems.append(f"{x.get('_row', 'Надходження')}: прихід «{x['АП №']}» від {_ddmm(x['АП від'])} "
                                f"({x['від кого']} → {x['кому']}) не знайдено в обліку")
        docs.append({"id": did[key], "kind": r["вид"], "no": r["номер"], "date": r["дата"],
                     "purchase": pid.get(r["закупівля"], ""), "parent": did.get(r["до документа"], ""),
                     "party": r["контрагент"], "person": _person_of(r["кому"], people, r.get("_row", "Документи"), problems),
                     "kop": r["сума"], "vatKop": r["ПДВ"],
                     "deliverBy": r["строк поставки"], "validTo": r["дійсна до"], "usedOn": r["використана"],
                     "returnedOn": r["повернута"], "payOn": r["до сплати"], "procedure": r["процедура"],
                     "state": r["стан"], "fes": {"status": r["статус ФЕС"], "date": r["дата статусу"],
                                                 "register": r["реєстр"], "ref": r["№ проводки"]},
                     "extra": extra, "missing": r["скану немає"], "note": r["примітка"], "lines": lines_,
                     "parts": _doc_parts(parts, lid.get(r["рядок коштів"], ""), r["сума"], by_id.get(pid.get(r["закупівля"]))),
                     "receipts": receipts})
    files = _files_of(m, lid, vid, did, old_files, problems)
    if problems:
        raise BookError(problems)
    return {"budget": budget, "purchases": purchases, "purDocs": docs}, files


def _files_of(m, lid, vid, did, old_files, problems):
    """Рядки «Файли» → записи app_file. Усе, що знає книга (№ документа, тип, дата додавання), береться з неї; чужі поля
    запису (extra, яких книга не несе) лишаються з наявного запису того самого файла."""
    ver_ids = {ver_key(p, label): i for (p, label), i in vid.items()}
    old = {(f["key"], f["path"]): f for f in old_files or []}
    files = []
    for f in m["Файли"]:
        owner = {"рядок коштів": ("bud", lid.get(f["ключ"])), "версія": ("purver", ver_ids.get(f["ключ"])),
                 "документ": ("purdoc", did.get(f["ключ"])), "реєстр": ("fesreg", f["ключ"])}.get(f["до чого"])
        if not owner or not owner[1]:
            problems.append(f"{f.get('_row', 'Файли')}: ключ «{f['ключ']}» — немає в книзі")
            continue
        key = f"{owner[0]}|{owner[1]}" + (f"|{f['що']}" if f["що"] else "")
        was = old.get((key, f["шлях"])) or {}
        extra = dict(was.get("extra") or {})
        for name, on in (("added", f["додано"]), ("made", f["склала програма"] == "так")):
            if on:
                extra[name] = on
            else:
                extra.pop(name, None)
        files.append({"key": key, "date": f["дата"], "no": f["№ документа"], "file": f["файл"], "path": f["шлях"],
                      "mime": f["тип"], "size": f["розмір"], "folder": f["тека"], "sha": f["sha256"], "extra": extra})
    return files


def _doc_parts(parts, whole_line, kop, purchase):
    """Частки документа з книги: кілька рядків «Частки» — як є; «рядок коштів» — цілим до нього, а якщо
    це і є єдиний рядок закупівлі — без часток (так документ веде сторінка)."""
    if len(parts) > 1:
        return parts
    line = parts[0]["line"] if parts else whole_line
    if not line:
        return []
    only = [f["line"] for f in (purchase or {}).get("funds", [])]
    return [] if only == [line] else [{"line": line, "kop": kop}]


def _strip(model):
    return {k: [{c: v for c, v in r.items() if c != "_row"} for r in rows] for k, rows in model.items() if k != "help"}


def _compare(have, want):
    """Відмінності двох моделей за аркушами: нові, змінені, прибрані записи (за ключами книги)."""
    keyf = {"Кошти": lambda rows: [line_key(r["рік"], r["назва"]) for r in rows],
            "Закупівлі": lambda rows: [pur_key(r["рік"], r["тема"]) for r in rows], "Документи": doc_keys}
    out, same = {}, True
    for name in COLS:
        if name in keyf:
            ha, wa = have.get(name, []), want.get(name, [])
            a, b = dict(zip(keyf[name](ha), ha)), dict(zip(keyf[name](wa), wa))
            res = {"new": sorted(set(b) - set(a)), "removed": sorted(set(a) - set(b)),
                   "changed": sorted(x for x in set(a) & set(b) if a[x] != b[x])}
        else:                                                    # рядки без власного ключа — лише скільки (однакові теж лічать)
            sig = lambda r: json.dumps(r, sort_keys=True, ensure_ascii=False)       # noqa: E731
            a = collections.Counter(sig(r) for r in have.get(name, []))
            b = collections.Counter(sig(r) for r in want.get(name, []))
            more, less = sum((b - a).values()), sum((a - b).values())
            res = {"new": [f"+{more}"] if more else [], "removed": [f"-{less}"] if less else [], "changed": []}
        same = same and not any(res.values())
        out[name] = res
    out["same"] = same
    return out


def _write(con, state, files):
    purchases_db.save(con, state)
    con.execute("DELETE FROM app_file WHERE " + " OR ".join("target LIKE ?" for _ in PUR_FILES),
                tuple(p + "%" for p in PUR_FILES))
    top = con.execute("SELECT COALESCE(MAX(pos), 0) FROM app_file").fetchone()[0]
    for n, f in enumerate(files, top + 1):
        con.execute("""INSERT INTO app_file(pos, id, target, doc_date, doc_no, file_name, rel_path, mime, size_bytes,
                       extra, folder, sha256) VALUES(?, NULL, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (n, f["key"], f["date"] or None, f["no"] or None, f["file"], f["path"], f["mime"] or None,
                     f["size"] or None, json.dumps(f["extra"], ensure_ascii=False) if f["extra"] else None,
                     f["folder"] or None, f["sha"] or None))


def _try(con, model, commit):
    """Книга в базі: записується в одній транзакції, і з неї береться те, чим база СТАЛА, — а не те, що написано в книзі:
    база нормалізує записане (порожню назву, версію без позначки, рядок коштів єдиної лінії закупівлі…), і книга, набрана
    руками, інакше завжди здавалась би зміненою. `commit=False` — «що буде»: усе відкочується. Повертає відмінності між
    базою до й після. Відмова бази — BookError зі словами; нічого не записано."""
    people = people_index(con)
    people["home"] = fingerprint(model) in set(own_prints(con))     # книга — та, що ця база записала чи прийняла (див. _person_of)
    state, files = from_model(model, purchases_db.load(con), purchases_db.receipt_keys(con), _files(con), people)
    before = _strip(model_from_db(con))
    con.execute("BEGIN IMMEDIATE")
    try:
        _write(con, state, files)
        after = _strip(model_from_db(con))
        con.execute("COMMIT" if commit else "ROLLBACK")
    except BaseException as e:
        with contextlib.suppress(sqlite3.Error):
            con.execute("ROLLBACK")
        if isinstance(e, state_db.SaveError):
            raise BookError([str(e)]) from e
        raise
    return _compare(before, after)


def diff(con, model):
    """Що змінить прийом книги в базі, за аркушами: нові, змінені, прибрані записи (за ключами книги). База не змінюється."""
    return _try(con, model, commit=False)


def apply(con, model):
    """Книга → база (усе або нічого): закупівлі в базі стають такими, як у книзі. Повертає те, що змінилось."""
    return _try(con, model, commit=True)
