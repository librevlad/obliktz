# -*- coding: utf-8 -*-
"""Папери закупівлі, які служба складає сама, — з коду, без файлів-шаблонів.

Розрахунки до рапорту (Додаток 1), рапорт, форми 1/фс і 9/фс, супровідний лист, акт введення
в експлуатацію (типова форма Мінфіну № 818) і реєстр документів на ФЕС. Розкладки — за
зразками служби; реквізити частини, підписанти й суми приходять зі сторінки. Файл лягає в
теку з налаштувань під тим самим іменем, що в людини. Файл, змінений після програми (його
відбиток не той, яким програма його записала), без згоди не перезаписується. Порожні бланки
всіх цих паперів лягають у «Закупівлі/Бланки» (`save_blanks`): без програми є з чого почати."""
import contextlib
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, Side

import excel_names
import stable_hash
from excel_export import safe_name

FONT = "Times New Roman"
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
MONEY = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
MONEY9 = "#,##0.00\\ _₽"
MONTHS = ["січня", "лютого", "березня", "квітня", "травня", "червня", "липня", "серпня", "вересня", "жовтня",
          "листопада", "грудня"]
NAMES = {"rozrakhunky": "Розрахунки до рапорту.xlsx", "raport": "Рапорт.docx", "1fs": "1фс.xlsx",
         "9fs": "9фс.xlsx", "suprovid": "Супровід.docx"}


class PaperConflict(Exception):
    """Файл змінено руками або відкрито в Excel чи Word — питання до людини, а не тихий перезапис."""

    def __init__(self, text, changed=False, busy=False):
        super().__init__(text)
        self.changed, self.busy = changed, busy


def sha_of(path):
    """Відбиток файла паперу — той самий, що скрізь у програмі (`stable_hash`): для .xls він не залежить від мітки часу,
    яку Excel дописує в контейнер, коли файл просто відкрили для правки."""
    return stable_hash.stable_sha(path)


MAX_KOP = 10 ** 15                              # 10 трлн грн: до цього цілі копійки в float ще точні (2**53 ≈ 9·10**15)
WHOLE = re.compile(r"-?[0-9]+(?:[.,]0+)?")      # ціле число текстом: «250», «250,00»; «250,5» — ні


def _kop(value):
    """Сума з запиту — цілі копійки: ціле число, число з крапкою, що точно дорівнює цілому (так JSON віддає 1250000.0), або
    ціле число текстом («250», «250,00»); порожнє — нуль. Решта — півкопійки (0.5), дріб від множення (431009.99999999994),
    нескінченність, NaN, слово, true, завелика сума — ValueError, тобто 422: гроші в програмі лише цілі копійки, а мовчазне
    відкидання дробу дало б інший рахунок, ніж на сторінці. Раніше `int()` відкидав дріб, `true` ставало копійкою, а
    `Infinity` давало OverflowError і відповідь 500."""
    if value is None or value == "":
        return 0
    number = None
    if isinstance(value, bool):
        pass
    elif isinstance(value, int):
        number = value
    elif isinstance(value, float):
        number = int(value) if value.is_integer() else None       # нескінченність і NaN цілими не бувають
    elif isinstance(value, str):
        text = value.strip().replace(" ", "")
        number = int(re.split("[.,]", text)[0]) if WHOLE.fullmatch(text) else None
    if number is None:
        shown = repr(value)[:40] if isinstance(value, (float, str)) else type(value).__name__
        raise ValueError(f"сума не в цілих копійках: {shown}")
    if abs(number) > MAX_KOP:
        raise ValueError("сума завелика: не більше 10^15 копійок")
    return number


def grn(kop):
    """Копійки → гривні для клітинки (рівно дві цифри після коми); копійки — цілі (`_kop`). Через неї й `grn_text` проходять
    усі суми всіх паперів, тож перевірка тут охоплює їх усі."""
    return round(_kop(kop) / 100, 2)


def grn_text(kop):
    """«3 500,00» — сума словами-цифрами для тексту; копійки — цілі (`_kop`)."""
    kop = _kop(kop)
    whole, part = divmod(abs(kop), 100)
    return ("-" if kop < 0 else "") + f"{whole:,}".replace(",", " ") + f",{part:02d}"


def ddmmyyyy(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}" if iso and len(iso) >= 10 else ""


def date_words(iso):
    """«20» січня 2026 р. — дата затвердження в акті."""
    if not iso or len(iso) < 10:
        return "«___» ____________ ____ р."
    return f"« {int(iso[8:10])} » {MONTHS[int(iso[5:7]) - 1]} {iso[:4]}р."


def _kekv(v):
    return int(v) if str(v).isdigit() else v


def _text(value):
    """Текст, вставлений зі сторінок, несе знаки керування (\\x0b із Word тощо): в XML їх не буває, і ні openpyxl, ні
    python-docx їх не приймають. Вони невидимі, тож відкидаються, а папір складається."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _cell(ws, ref, value, size=12, bold=False, align=None, box=False, fmt=None):
    c = ws[ref]
    c.value = _text(value)
    c.font = Font(name=FONT, size=size, bold=bold)
    if align is not None:
        c.alignment = align
    if box:
        c.border = BOX
        # Межу об'єднаної клітинки Excel малює з меж усіх її клітинок, а не з лівої верхньої: без цього нижня й права
        # сторони шапки та рядка «Усього» не надрукувалися б.
        for area in ws.merged_cells.ranges:
            if (area.min_row, area.min_col) == (c.row, c.column):
                for line in ws.iter_rows(min_row=area.min_row, max_row=area.max_row,
                                         min_col=area.min_col, max_col=area.max_col):
                    for edge in line:
                        edge.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def _widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _print(ws, landscape, area, scale=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if scale:
        ws.page_setup.scale = scale
    ws.print_area = area


def _signers(ws, row, signers, name_col="G", size=14, under=("D", "(підпис)", "(прізвище, ім’я, по батькові)")):
    """Блоки підписантів: посада / звання + ім'я / підписи під рискою — як у бланках служби."""
    for s in signers:
        _cell(ws, f"A{row}", s.get("pos", ""), size, align=WRAP)
        ws.merge_cells(f"A{row}:E{row}")
        ws.row_dimensions[row].height = 36
        _cell(ws, f"A{row + 1}", s.get("rank", ""), size)
        _cell(ws, f"{name_col}{row + 1}", s.get("name", ""), size)
        _cell(ws, f"{under[0]}{row + 2}", under[1], 9, align=MID)
        _cell(ws, f"{name_col}{row + 2}", under[2], 9, align=MID)
        row += 4
    return row


def write_1fs(spec, path):
    unit, year = spec.get("unit") or {}, spec.get("year") or date.today().year
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист 1"
    _widths(ws, {"A": 16.4, "B": 10.7, "C": 44.6, "D": 13.9, "E": 11.0, "F": 13.0, "G": 15.4, "H": 13.0, "I": 16.4})
    _cell(ws, "I1", "Додаток 2", 13, align=RIGHT)
    ws.merge_cells("F3:I3")
    _cell(ws, "F3", f"Відомчий код розпорядника коштів {unit.get('fundCode', '')}".strip(), 13)
    ws.merge_cells("A5:I5")
    _cell(ws, "A5", "РОЗРАХУНКИ, ЯКІ ОБГРУНТОВУЮТЬ ЗМІНИ ПОКАЗНИКІВ ВИДАТКІВ БЮДЖЕТУ ЗА БЮДЖЕТНОЮ ПРОГРАМОЮ "
                    f"(ПІДПРОГРАМОЮ) КПКВ {spec.get('kpkv', '')}", 13, True, MID)
    ws.row_dimensions[5].height = 36
    ws.merge_cells("A6:I6")
    _cell(ws, "A6", f"НА {year} РІК", 13, True, MID)
    _cell(ws, "I7", "(грн.)", 12, align=RIGHT)
    for col, text in zip("ABCDEF", ["КЕКВ", "Код видатків/вид", "Найменування витрат", "Одиниці виміру", "Кількість",
                                    "Ціна за одиницю"]):
        ws.merge_cells(f"{col}8:{col}9")
        _cell(ws, f"{col}8", text, 12, True, MID, True)
    _cell(ws, "G8", "Вартість", 12, True, MID, True)
    _cell(ws, "G9", "(гр.5 х гр.6)", 12, False, MID, True)
    ws.merge_cells("H8:I8")
    _cell(ws, "H8", "У тому числі за рахунок", 12, True, MID, True)
    _cell(ws, "H9", "загального фонду", 12, False, MID, True)
    _cell(ws, "I9", "спеціального фонду", 12, False, MID, True)
    for n, col in enumerate("ABCDEFGHI", 1):
        _cell(ws, f"{col}10", n, 12, align=MID, box=True)
    general = bool(spec.get("general"))
    r = 11
    # У бланку без позицій лишається один порожній рядок: інакше «Усього» стало б у рядок 11 і складало б само себе.
    # Вартість рядка — до копійки, як позиція на сторінці, у 9/фс і в супроводі: без округлення 2,5 × 33,33 давало б 83,325,
    # і два такі рядки складались би в 166,65 замість 166,66.
    for row in spec.get("rows") or [{}]:
        _cell(ws, f"A{r}", _kekv(row.get("kekv", "")), 12, align=MID, box=True)
        _cell(ws, f"B{r}", str(row.get("exp", "")), 12, align=MID, box=True, fmt="@")
        _cell(ws, f"C{r}", row.get("name", ""), 12, align=WRAP, box=True)
        _cell(ws, f"D{r}", row.get("uom", ""), 12, align=MID, box=True)
        _cell(ws, f"E{r}", row.get("qty"), 12, align=MID, box=True)
        _cell(ws, f"F{r}", grn(row.get("priceKop")), 12, box=True, fmt=MONEY)
        _cell(ws, f"G{r}", f"=ROUND(E{r}*F{r},2)", 12, box=True, fmt=MONEY)
        _cell(ws, f"H{r}", f"=G{r}" if general else None, 12, box=True, fmt=MONEY)
        _cell(ws, f"I{r}", None if general else f"=G{r}", 12, box=True, fmt=MONEY)
        r += 1
    last = max(r - 1, 11)
    ws.merge_cells(f"A{r}:C{r}")
    _cell(ws, f"A{r}", "Усього", 12, True, MID, True)
    for col in "DEF":
        _cell(ws, f"{col}{r}", "х", 12, align=MID, box=True)
    _cell(ws, f"G{r}", f"=SUM(G11:G{last})", 12, True, box=True, fmt=MONEY)
    _cell(ws, f"H{r}", f"=G{r}" if general else "х", 12, True, MID, True, MONEY if general else None)
    _cell(ws, f"I{r}", "х" if general else f"=G{r}", 12, True, MID, True, None if general else MONEY)
    end = _signers(ws, r + 3, spec.get("signers") or [])
    _cell(ws, f"A{end}", f"контактний телефон {unit.get('phone', '')}".strip(), 11)
    _print(ws, True, f"A1:I{end}")
    wb.save(path)


def write_9fs(spec, path):
    unit, year = spec.get("unit") or {}, spec.get("year") or date.today().year
    wb = Workbook()
    ws = wb.active
    ws.title = "Дод11 9фс ЗФ"
    _widths(ws, {"A": 17.4, "B": 13.4, "C": 18.6, "D": 18.9, "E": 15.0, "F": 20.1, "G": 19.1, "H": 26.1, "I": 19.6})
    for ref, text in (("G1", "Додаток 11"), ("G2", "до Правил фінансового забезпечення військових частин, установ та "
                                                  "організацій Збройних Сил України\n(пункт 2.6)"),
                      ("G4", f"Відомчий код розпорядника коштів {unit.get('fundCode', '')} "),
                      ("G6", "Керівник служби забезпечення (відповідальний виконавець)"),
                      ("G8", "(військове звання, підпис, прізвище, ім’я, по батькові)")):
        ws.merge_cells(f"{ref}:I{ref[1:]}")
        _cell(ws, ref, text, 14 if ref != "G8" else 10, align=WRAP)
    ws.row_dimensions[2].height = 54
    _cell(ws, "G3", "Форма 9/фс", 14)
    _cell(ws, "G5", "ПОГОДЖЕНО", 14, True)
    ws.merge_cells("G7:I7")
    for ref in ("G7", "H7", "I7"):                              # лінія підпису: Excel бере межу з усіх клітинок об'єднаної
        ws[ref].border = Border(bottom=THIN)
    _cell(ws, "G9", f"“___” ____________ {year} р.", 14)
    _cell(ws, "I10", "МП", 14)
    for ref, text, bold in (("A11", "ЗАЯВКА", True), ("A12", f"на бюджетні асигнування до кінця {year} року", True),
                            ("A13", f"за   КПКВ {spec.get('kpkv', '')}  ({spec.get('phrase', '')})", True)):
        ws.merge_cells(f"{ref}:I{ref[1:]}")
        _cell(ws, ref, text, 14, bold, MID)
    heads = ["Коди економічної класифікації", "Коди витрат, види", "Призначено з урахуванням змін",
             "Зменшено з донесення                         (форма 10 ФС)", "Отримано",
             "Витрачено на день уточнення  з копійками", "Зареєстровані фінансові зобов’язання",
             "Залишок призначень \n(гр.3-гр.4-гр.6-гр.7) ", "Потрібно допризначити з урахуванням залишку (перевитрат)"]
    for n, (col, text) in enumerate(zip("ABCDEFGHI", heads), 1):
        _cell(ws, f"{col}16", text, 12, True, MID, True)
        _cell(ws, f"{col}17", n, 12, align=MID, box=True)
    ws.row_dimensions[16].height = 90
    ws.print_title_rows = "17:17"

    def zeros(r):
        for col in "CDEFGH":
            _cell(ws, f"{col}{r}", 0, 12, box=True, fmt=MONEY9)

    _cell(ws, "A18", spec.get("kpkv", ""), 12, align=MID, box=True)
    _cell(ws, "B18", " ", 12, box=True)
    zeros(18)
    _cell(ws, "I18", None, 12, box=True)
    r, code_rows = 19, []
    for g in spec.get("groups") or []:
        _cell(ws, f"A{r}", _kekv(g.get("kekv", "")), 12, align=MID, box=True)
        _cell(ws, f"B{r}", None, 12, box=True)
        zeros(r)
        _cell(ws, f"I{r}", None, 12, box=True)
        r += 1
        for c in g.get("codes") or []:
            _cell(ws, f"A{r}", " ", 12, box=True)
            _cell(ws, f"B{r}", str(c.get("exp", "")), 12, align=MID, box=True, fmt="@")
            zeros(r)
            _cell(ws, f"I{r}", grn(c.get("kop")), 12, box=True, fmt=MONEY9)
            code_rows.append(r)
            r += 1
    _cell(ws, f"A{r}", "Разом за \nКПКВ", 12, True, WRAP, True)
    _cell(ws, f"B{r}", " ", 12, box=True)
    zeros(r)
    first, last = (code_rows[0], code_rows[-1]) if code_rows else (r, r)
    _cell(ws, f"I{r}", f"=SUM(I{first}:I{last})" if code_rows else 0, 12, True, box=True, fmt=MONEY9)
    note = max(r + 2, 23)
    ws.merge_cells(f"A{note}:I{note}")
    _cell(ws, f"A{note}", "Пояснювальна записка", 14, True, MID)
    ws.merge_cells(f"A{note + 2}:I{note + 2}")
    _cell(ws, f"A{note + 2}", spec.get("note", ""), 14, align=WRAP)
    ws.row_dimensions[note + 2].height = 60
    row = note + 4
    for s in spec.get("signers") or []:
        _cell(ws, f"A{row}", s.get("pos", ""), 14, align=WRAP)
        ws.merge_cells(f"A{row}:E{row}")
        ws.row_dimensions[row].height = 38
        _cell(ws, f"A{row + 1}", s.get("rank", ""), 14)
        _cell(ws, f"H{row + 1}", s.get("name", ""), 14)
        _cell(ws, f"F{row + 2}", "(підпис)", 10, align=MID)
        _cell(ws, f"H{row + 2}", " (ім’я та прізвище) ", 10, align=MID)
        row += 5
    row -= 1                                                    # виконавець — через один порожній рядок після підписів
    _cell(ws, f"A{row}", "Виконавець ", 12)
    _cell(ws, f"B{row}", spec.get("phone") or unit.get("phone", ""), 12, fmt="@")
    _print(ws, False, f"A1:I{row}", 60)
    wb.save(path)


def write_rozrakhunky(spec, path):
    unit = spec.get("unit") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Додаток до рапорта"
    _widths(ws, {"A": 25.7, "B": 9.7, "C": 52.6, "D": 12.6, "E": 13.4, "F": 34.1})
    _cell(ws, "F1", "Додаток 1", 12, align=RIGHT)
    _cell(ws, "A2", "1.    Перелік предметів закупівлі та технічних вимог до них", 12, True)
    for col, text in zip("ABCDEF", ["Назва товару (роботи, послуги)", "Кількість одиниць",
                                    "Технічні вимоги до предмету закупівлі *", "Орієнтовна ціна, грн.", "Сума, грн.",
                                    "Посилання на аналогічні товари в мережі **"]):
        _cell(ws, f"{col}4", text, 8, True, MID, True)
    r = 5
    for p in spec.get("positions") or []:
        links = [_text(x) for x in (list(p.get("links") or []) + ["", "", ""])[:3]]
        for col in "ABCDE":
            ws.merge_cells(f"{col}{r}:{col}{r + 2}")
        _cell(ws, f"A{r}", p.get("name", ""), 10, align=WRAP, box=True)
        _cell(ws, f"B{r}", p.get("qty"), 10, align=MID, box=True)
        _cell(ws, f"C{r}", p.get("specs", ""), 9, align=WRAP, box=True)
        _cell(ws, f"D{r}", grn(p.get("priceKop")), 10, box=True, fmt="#,##0.00")
        _cell(ws, f"E{r}", grn(p.get("kop")), 10, box=True, fmt="#,##0.00")
        for i, link in enumerate(links):
            c = _cell(ws, f"F{r + i}", link or None, 9, align=WRAP, box=True)
            if link.startswith("http"):
                c.hyperlink = link
        r += 3
    _cell(ws, f"A{r}", "ЗАГАЛЬНА СУМА", 10, True, box=True)
    _cell(ws, f"E{r}", f"=SUM(E5:E{r - 1})" if r > 5 else 0, 10, True, box=True, fmt="#,##0.00")
    chief = spec.get("chief") or {}
    _cell(ws, f"A{r + 2}", "Начальник продовольчої служби тилу логістики", 12)
    _cell(ws, f"A{r + 3}", f"військової частини {unit.get('code', '')}".strip(), 12)
    _cell(ws, f"A{r + 4}", chief.get("rank", ""), 12)
    _cell(ws, f"F{r + 4}", chief.get("name", ""), 12)
    _print(ws, False, f"A1:F{r + 4}", 59)
    # Робочий аркуш: по три ціни на позицію — людина правитиме ціни саме тут.
    wt = wb.create_sheet("Товари")
    wt.append(["Найменування", "ціна", "посилання", "характеристики", "середня", "кількість", "сума"])
    for p in spec.get("positions") or []:
        offers = (list(p.get("offers") or []) + [{}, {}, {}])[:3]
        for i, o in enumerate(offers):
            wt.append([_text(p.get("name", "")) if i == 0 else None, grn(o.get("priceKop")) if o.get("priceKop") else None,
                       _text(o.get("url") or o.get("seller") or None), _text(p.get("specs", "")) if i == 0 else None,
                       grn(p.get("priceKop")) if i == 0 else None, p.get("qty") if i == 0 else None,
                       grn(p.get("kop")) if i == 0 else None])
    _widths(wt, {"A": 40, "B": 12, "C": 50, "D": 40, "E": 12, "F": 10, "G": 14})
    wb.save(path)


def _doc(margins=(3.0, 1.5, 2.0, 2.0)):
    from docx import Document                                    # noqa: PLC0415
    from docx.oxml.ns import qn                                  # noqa: PLC0415
    from docx.shared import Cm, Pt                               # noqa: PLC0415
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(14)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = (Cm(x) for x in margins)
    core = doc.core_properties
    core.author = core.last_modified_by = "Облік ТЗ ПС"
    return doc


def _p(doc, text, align="left", bold=False, size=14, indent=False, after=0):
    from docx.enum.text import WD_ALIGN_PARAGRAPH               # noqa: PLC0415
    from docx.shared import Cm, Pt                               # noqa: PLC0415
    par = doc.add_paragraph()
    par.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                     "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    fmt = par.paragraph_format
    fmt.space_after = Pt(after)
    if indent:
        fmt.first_line_indent = Cm(1.25)
    run = par.add_run(_text(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT
    return par


def _sign_line(doc, s, size=14):
    _p(doc, s.get("pos", ""), "left", size=size)
    _p(doc, f"{s.get('rank', '')}\t\t\t\t\t\t\t\t{s.get('name', '')}", "left", size=size)


def write_raport(spec, path):
    unit = spec.get("unit") or {}
    code = unit.get("code", "")
    doc = _doc()
    _p(doc, f"Командиру військової частини {code}".strip(), "right", after=12)
    _p(doc, "РАПОРТ", "center", True, after=12)
    _p(doc, f"З метою виконання поставлених завдань продовольчою службою тилу логістики військової частини {code}, "
            "існує нагальна необхідність у закупівлі майна відповідно до затвердженого обсягу та переліку закупівель.",
       "justify", indent=True)
    _p(doc, f"Прошу надати вказівку уповноваженій особі з організації та проведення закупівель у військовій частині "
            f"{code} на здійснення закупівлі, відповідно до переліку товарів та технічних вимог до них, викладених у "
            "додатку 1 до цього рапорту.", "justify", indent=True, after=24)
    chief = spec.get("chief") or {}
    _p(doc, "Начальник продовольчої служби тилу логістики")
    _p(doc, f"військової частини {code}".strip())
    _p(doc, f"{chief.get('rank', '')}\t\t\t\t\t\t\t\t{chief.get('name', '')}")
    doc.save(path)


def write_suprovid(spec, path):
    unit, year = spec.get("unit") or {}, spec.get("year") or date.today().year
    code = unit.get("code", "")
    doc = _doc((3.0, 1.5, 2.0, 1.25))
    for text in ("ЗБРОЙНІ СИЛИ УКРАЇНИ", "ВІЙСЬКОВА ЧАСТИНА", code, f"Код ЄДРПОУ {unit.get('edrpou', '')}".strip()):
        _p(doc, text, "left", True, 12)
    _p(doc, f"___ {year} року № ___", "left", size=12)
    if unit.get("place"):
        _p(doc, unit["place"], "left", size=12)
    _p(doc, "", after=6)
    for text in ("Передати засобами СЕДО", f"Командиру військової частини {unit.get('upperHq', '')}".strip(),
                 "(через фінансово-економічне управління,", "через продовольчу службу тилу логістики)"):
        _p(doc, text, "right", size=12)
    _p(doc, "", after=6)
    _p(doc, f"Військова частина {code} направляє на погодження розрахунки, які обґрунтовують зміни показників видатків "
            "бюджету що включаються до проекту кошторису за бюджетною програмою (підпрограмою) "
            f"КПКВ {spec.get('kpkv', '')}  {spec.get('codes', '')} на {year} рік на загальну суму "
            f"{grn_text(spec.get('sumKop'))} грн.", "justify", indent=True, after=12)
    _p(doc, "Додатки:", "justify")
    items = [f"Розрахунки, які обґрунтовують зміни показників видатків бюджету до проекту кошторису за бюджетною "
             f"програмою (підпрограмою) КПКВ {spec.get('kpkv', '')} на {year} рік (форма 1/фс)– на 1 арк.;",
             f"Заявка на бюджетні асигнування за КПКВ {spec.get('kpkv', '')} {spec.get('phrase', '')}  "
             "(форма 9/фс) – на 1 арк."] + list(spec.get("attachments") or [])
    for n, text in enumerate(items, 1):
        _p(doc, f"{n}. {text}", "justify")
    _p(doc, "", after=18)
    _sign_line(doc, spec.get("cmdr") or {})
    doc.save(path)


AKT_MARGINS = (3.0, 1.0, 2.0, 2.0)          # поля акта, см: ліве, праве, верхнє, нижнє
# Ширини таблиць акта в двадцятих пункта; кожна — рівно ширина тексту між полями (21 − 3 − 1 = 17 см): ширша таблиця виходила
# б за аркуш, і права графа не надрукувалась би. Графи головної таблиці — за найдовшим словом шапки при 8 пт і вужчих відступах
# клітинки (CELL_PAD): так жодне слово шапки не рветься посередині.
W_ACT_HEAD = [3800, 2000, 3837]
W_ACT_ITEMS = [1480, 880, 1010, 880, 770, 590, 1110, 1190, 920, 807]
W_ACT_BOOKS = [2600, 2600, 2600, 1837]
CELL_PAD = 57                                # відступ клітинки ліворуч і праворуч, двадцяті пункта (звичайний — 108)
SIGN_BLANK = "_" * 26 + " " + "_" * 19 + " " + "_" * 24      # рядок підпису без даних: посада — підпис — ім'я, для запису рукою
SIGN_FIELDS = (24, 17, 21)                   # довжини ліній під «Прийняв» / «Здав*»: посада, підпис, ініціали та прізвище («_»)
UNDERSCORE_PT = 6.0                          # ширина «_» у Times New Roman 12 пт — пів-кегля, пункти
SIGN_FIRST, SIGN_GAP = 58, 20                # пункти: де починається перша лінія (після «Прийняв», 45 пт) і проміжок між лініями
BLANK_REGISTER_ROWS = 5                      # рядків у рамці в порожньому бланку реєстру
QTY = re.compile(r"[0-9]{1,12}(?:[.,][0-9]+)?")


def _s(obj, key):
    """Поле запису із запиту як текст: немає чи null — порожньо, а не «None»."""
    value = obj.get(key)
    return "" if value is None else str(value)


def _qty(value):
    """Кількість із запиту — число чи текст на кшталт «2» або «0,5» — як Decimal до тисячних. Хибне (слово, нескінченність,
    1e30) — ValueError, тобто 422, а не обрив."""
    text = "0" if value is None or value == "" else str(value).strip().replace(" ", "")
    if not QTY.fullmatch(text):
        raise ValueError(f"кількість «{text[:40]}» не схожа на число")
    return Decimal(text.replace(",", ".")).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def _qty_text(number):
    """«2», «2,5», «0,333»: ціле без коми, дробове — до тисячних, із комою."""
    text = f"{number:f}"
    return (text.rstrip("0").rstrip(".") if "." in text else text).replace(".", ",")


def _tc(table, r, c):
    """Клітинка рядка `r`, графи `c` — не через `table.cell`: python-docx щоразу обходить ним усю таблицю, і сто рядків акта
    складалися б понад пів хвилини. Номер графи тут — місце клітинки в рядку, а не в сітці: надійний, поки в цьому рядку
    ще нічого не зливали по горизонталі."""
    from docx.table import _Cell                                 # noqa: PLC0415
    return _Cell(table._tbl.tr_lst[r].tc_lst[c], table)          # noqa: SLF001


def _grid(doc, rows, widths, borders=True, size=10, bold=(), align=None, body=0, pad=None):
    """Таблиця Word, заповнена текстами `rows`, — клітинками акта ЯТС (`techact_export`): ті самі шрифт, відступи й межі.
    Вони текст не чистять, тож кожен текст із запиту проходить `_text` тут. `bold` — номери рядків жирним; `align` —
    вирівнювання за номером графи у рядках від `body` (решта — по центру); `pad` — відступ клітинки ліворуч і праворуч.

    `_table` робить лише перший рядок: на кожному рядку він обходить усю таблицю (`row.cells`), тож на сотні рядків працював
    би квадратично. Решта дописується `add_row` (ширини — із сітки), і весь текст ставиться за один обхід."""
    from docx.oxml import OxmlElement                            # noqa: PLC0415
    from docx.oxml.ns import qn                                  # noqa: PLC0415
    from docx.shared import Twips                                # noqa: PLC0415
    from docx.table import _Cell                                 # noqa: PLC0415
    from techact_export import _cell, _table                     # noqa: PLC0415
    table = _table(doc, 1, widths, borders=borders)
    for c, w in enumerate(widths):
        table.columns[c].width = Twips(w)
    for _ in rows[1:]:
        table.add_row()
    for r, (tr, row) in enumerate(zip(table._tbl.tr_lst, rows)):  # noqa: SLF001
        for c, (tc, text) in enumerate(zip(tr.tc_lst, row)):
            _cell(_Cell(tc, table), _text("" if text is None else str(text)), size, bold=r in bold,
                  align=(align or {}).get(c, "center") if r >= body else "center")
    props = table._tbl.tblPr                                     # noqa: SLF001
    for extra in props.findall(qn("w:tblLayout"))[1:]:           # `_table` дописує другий tblLayout до вже наявного…
        props.remove(extra)
    if pad is not None:
        margins = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:w"), str(pad))
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        props.append(margins)
    look = props.find(qn("w:tblLook"))                           # …а tblLook за схемою — останній
    if look is not None:
        props.remove(look)
        props.append(look)
    return table


def _sign_text(person):
    """Рядок підпису в акті: «посада звання ___ ім'я». Без жодної з цих даних — порожній рядок із трьома полями для запису
    рукою (бланк, чи комісії ще немає): служба обходиться без програми."""
    pos, rank, name = _s(person, "pos"), _s(person, "rank"), _s(person, "name")
    if not (pos or rank or name):
        return SIGN_BLANK
    return f"{pos} {rank} ___________________ {name}".strip()


def _signature_block(doc, who):
    """«Прийняв» / «Здав*»: три лінії на табуляціях зліва, а під ними підписи на табуляціях по центру — центр кожного підпису
    над серединою своєї лінії. Раніше підписи вирівнювались рядами пробілів, і в Word «(ініціали та прізвище)» ставало не під
    своєю лінією: пробіл і «_» різної ширини, а розкладка пробілів — на одному шрифті. Тексти ті самі, пробіли — табуляції."""
    from docx.enum.text import WD_TAB_ALIGNMENT                  # noqa: PLC0415
    from docx.shared import Pt                                   # noqa: PLC0415
    starts, at = [], SIGN_FIRST
    for n in SIGN_FIELDS:
        starts.append(at)
        at += n * UNDERSCORE_PT + SIGN_GAP
    line = _p(doc, f"{who}\t" + "\t".join("_" * n for n in SIGN_FIELDS), size=12)
    line.paragraph_format.keep_with_next = True
    for start in starts:
        line.paragraph_format.tab_stops.add_tab_stop(Pt(start), WD_TAB_ALIGNMENT.LEFT)
    caption = _p(doc, "\t(посада)\t(підпис)\t(ініціали та прізвище)", size=9)
    for start, n in zip(starts, SIGN_FIELDS):
        caption.paragraph_format.tab_stops.add_tab_stop(Pt(start + n * UNDERSCORE_PT / 2), WD_TAB_ALIGNMENT.CENTER)


def write_akt(spec, path):
    """Акт введення в експлуатацію основних засобів — типова форма Мінфіну № 818."""
    unit = spec.get("unit") or {}
    code = _s(unit, "code")
    day = _s(spec, "date")
    doc = _doc(AKT_MARGINS)
    for text in ("ЗАТВЕРДЖЕНО", "Наказ Міністерства фінансів України", "13.09.2016 № 818"):
        _p(doc, text, "right", size=10)
    _p(doc, _s(unit, "name") or "Військова частина", "left", size=12)
    _p(doc, "(найменування юридичної особи)", "left", size=8)
    _grid(doc, [["Ідентифікаційний код за ЄДРПОУ", _s(unit, "edrpou"), "ЗАТВЕРДЖУЮ:"]], W_ACT_HEAD, borders=False, size=11,
          align={0: "left", 1: "left", 2: "right"})
    appr = spec.get("approver") or {}
    _p(doc, _s(appr, "pos") or f"Командир військової частини {code}".strip(), "right", size=11)
    _p(doc, f"{_s(appr, 'rank')} ______________ {_s(appr, 'name')}".strip(), "right", size=11)
    _p(doc, date_words(day), "right", True, 11, after=8)
    _p(doc, f"Акт введення в експлуатацію основних засобів {_s(spec, 'no')}".strip(), "center", True, 13)
    _p(doc, _s(spec, "place") or _s(unit, "place"), "right", size=13)
    _p(doc, "(місце складання)", "right", size=8)
    number = _s(spec, "no").split("-")[-1]
    rows = [["", "Номер документа", "Дата складання", "", "", "", "", "", "", ""],
            ["", number, f"{ddmmyyyy(day)} р." if day else "", "", "", "", "", "", "", ""],
            ["Інвентарний (номенклатурний) номер", "Кількість", "Первісна вартість", "Сума всього", "Знос", "",
             "Ліквідаційна вартість", "Строк корисного використання", "Рік випуску (побудови)", "Номер паспорта"],
            ["", "", "", "", "за одиницю", "всього", "", "", "", ""],
            [str(n) for n in range(1, 11)]]
    lines = spec.get("lines") or []
    quantities, total_kop = [], 0
    for ln in lines:
        name = _s(ln, "name")
        if ln.get("fesNo"):                                      # лише номер ФЕС: своїх інвентарних номерів у паперах нема
            name = f"{name}, інв. № {ln['fesNo']}"
        qty, kop = _qty(ln.get("qty")), _kop(ln.get("kop"))
        quantities.append(qty)
        total_kop += kop
        rows.append([name, _qty_text(qty), grn_text(ln.get("priceKop")), grn_text(kop), "", "", "", _s(ln, "life"),
                     _s(ln, "year"), _s(ln, "passport")])
    # Порожній бланк: три порожні рядки для запису й «Всього» без цифр (нулі в бланку — шум).
    rows += [[""] * 10 for _ in range(0 if lines else 3)]
    rows.append(["Всього:", _qty_text(sum(quantities, Decimal(0))) if lines else "", "", grn_text(total_kop) if lines else "",
                 "", "", "", "", "", ""])
    items = _grid(doc, rows, W_ACT_ITEMS, size=8, bold=(0, 2, 3, len(rows) - 1), align={0: "left", 2: "right", 3: "right"},
                  body=5, pad=CELL_PAD)
    # Шапка: порожній кут і порожнє поле праворуч від номера й дати; назви граф — на два рядки, «Знос» — над двома графами.
    # Злиття йдуть у такому порядку, що номер графи в рядку лишається місцем клітинки (див. `_tc`).
    _tc(items, 0, 0).merge(_tc(items, 1, 0))
    _tc(items, 0, 3).merge(_tc(items, 1, 9))
    for c in (0, 1, 2, 3, 6, 7, 8, 9):
        _tc(items, 2, c).merge(_tc(items, 3, c))
    _tc(items, 2, 4).merge(_tc(items, 2, 5))
    for text in (f"Проведено огляд: {_s(spec, 'seen')}",
                 f"у момент введення в експлуатацію об’єкт(и) знаходиться(яться): {_s(spec, 'location')}",
                 "Коротка характеристика об’єкта(ів): назві відповідає",
                 "Об’єкт технічним умовам відповідає/не відповідає: технічним умовам відповідає",
                 "Доробка не потрібна/потрібна: не потрібна",
                 "Результати випробування об’єкта(ів): ________________________________",
                 f"Інші відомості: Постачальник {_s(spec, 'supplier')}",
                 "Висновок комісії: порушень не виявлено",
                 f"Перелік документації, що додається: {_s(spec, 'docs')}"):
        _p(doc, text, "justify", size=12)
    # Підпис не відривається від підпису під ним (посада, підпис, ініціали), а заголовок — від свого переліку чи таблиці.
    _p(doc, "Голова комісії:", "justify", size=12).paragraph_format.keep_with_next = True
    _p(doc, _sign_text(spec.get("head") or {}), size=12)
    _p(doc, "Члени комісії:", "justify", size=12).paragraph_format.keep_with_next = True
    for member in spec.get("members") or [{}, {}, {}]:           # комісії нема (бланк): три рядки для запису рукою
        _p(doc, _sign_text(member), size=12)
    for who in ("Прийняв", "Здав*"):
        _signature_block(doc, who)
    _p(doc, "Відмітка бухгалтерської служби про відображення у регістрах бухгалтерського обліку:", "justify",
       size=12).paragraph_format.keep_with_next = True
    _grid(doc, [["Назва облікового регістру", "За дебетом рахунку (субрахунку, коду аналітичного обліку)",
                 "За кредитом рахунку (субрахунку, коду аналітичного обліку)", "Сума"], ["", "", "", ""], ["", "", "", ""]],
          W_ACT_BOOKS, size=9)
    acc = spec.get("accountant") or {}
    _p(doc, "Особа, яка відобразила господарську операцію в бухгалтерському обліку   _____________      "
            "_________________________________", size=12).paragraph_format.keep_with_next = True
    _p(doc, f"«______» ________________ {day[:4] or '____'} р.", size=12)
    _p(doc, _s(acc, "pos"), size=9)
    _p(doc, f"{_s(acc, 'rank')}\t\t\t\t\t\t\t\t{_s(acc, 'name')}", size=9)
    _p(doc, "* Заповнюється при зарахуванні об’єктів незавершених капітальних інвестицій до складу основних засобів "
            "при зміні матеріально відповідальної особи.", size=8)
    doc.save(path)


def write_reestr(spec, path):
    """Реєстр документів, переданих до фінансово-економічної служби, — 7 граф за бланком; внизу «Передав» і «Отримав»."""
    unit = spec.get("unit") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Реєстр"
    _widths(ws, {"A": 7, "B": 34, "C": 18, "D": 13, "E": 12, "F": 10, "G": 16})
    ws.merge_cells("A1:G1")
    stamp = ddmmyyyy(_s(spec, "date"))
    _cell(ws, "A1", "РЕЄСТР\nДОКУМЕНТІВ, ПЕРЕДАНИХ ДО ФІНАНСОВО-ЕКОНОМІЧНОЇ СЛУЖБИ\n"
                    + (f"від {stamp} р." if stamp else "від «___» ________ 20__ р."), 12, True, MID)   # без дати — для запису рукою
    ws.row_dimensions[1].height = 54
    _cell(ws, "B2", f"Фінансово-економічна служба військової частини {_s(unit, 'code')}".strip(), 12)
    _cell(ws, "B3", "Передача документів від продовольчої служби тилу логістики", 12)
    for col, text in zip("ABCDEFG", ["№ п/п", "Назва документа", "№ докум.", "Дата докум.", "кіл-ть примірників",
                                     "кіл-ть аркушів", "Сума документу"]):
        _cell(ws, f"{col}5", text, 11, True, MID, True)
    entries = [[n, _s(row, "name"), _s(row, "no"), ddmmyyyy(_s(row, "date")), row.get("copies") or 1,
                row.get("sheets") or 1, grn(row.get("kop")) if row.get("kop") else None]
               for n, row in enumerate(spec.get("rows") or [], 1)]
    blank = not entries                                          # порожній бланк: рядки в рамці, у які пишуть рукою
    r = 6
    for values in entries or [[None] * 7 for _ in range(BLANK_REGISTER_ROWS)]:
        for col, v in zip("ABCDEFG", values):
            _cell(ws, f"{col}{r}", v, 11, align=WRAP if col == "B" else MID, box=True,
                  fmt="#,##0.00" if col == "G" else None)
        if blank:
            ws.row_dimensions[r].height = 24
        r += 1
    _cell(ws, f"A{r + 1}", "Передав", 12)
    _cell(ws, f"C{r + 1}", _s(spec.get("giver") or {}, "name"), 12)
    _cell(ws, f"A{r + 2}", "Отримав", 12)
    _print(ws, False, f"A1:G{r + 2}")
    wb.save(path)


def to_xls(path):
    """9/фс у форматі бланка штабу (.xls) через Excel цього ПК. Без Excel, без pywin32 і в
    перевірках (OBLIK_NO_OPEN) — None: файл лишається .xlsx. COM цього потоку відкривається й закривається
    парою (CoInitialize — CoUninitialize), хоч би чим скінчилось перетворення."""
    if os.environ.get("OBLIK_NO_OPEN"):
        return None
    try:
        import pythoncom                                         # noqa: PLC0415
        import win32com.client                                   # noqa: PLC0415
        import win32process                                      # noqa: PLC0415
    except ImportError:
        return None
    target = os.path.splitext(path)[0] + ".xls"
    pythoncom.CoInitialize()
    try:
        return _excel_save_as(win32com.client, win32process, path, target)
    finally:
        pythoncom.CoUninitialize()                               # об'єкти Excel уже звільнено: їх тримав лише `_excel_save_as`


def _excel_save_as(client, process, path, target):
    """Книга `path` → .xls `target` окремим Excel, якого не видно. Увесь шлях від запуску Excel — під охороною: збій
    будь-де (і тоді, коли не вдалося навіть дізнатися номер його процесу) дає None, а Excel закривається (Quit), а не лишається
    висіти без вікна. Сторож за дві хвилини знімає Excel, що завис; без номера процесу сторожа нема, тож і перетворення нема."""
    try:
        xl = client.DispatchEx("Excel.Application")
    except Exception:                                            # noqa: BLE001 — Excel не встановлено
        return None
    watch = None
    try:
        pid = process.GetWindowThreadProcessId(xl.Hwnd)[1]
        watch = threading.Timer(120, lambda: subprocess.call(["taskkill", "/F", "/PID", str(pid)]))
        watch.start()
        xl.Visible = False
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Open(os.path.abspath(path))
        try:
            wb.SaveAs(os.path.abspath(target), FileFormat=56)     # xlExcel8
        finally:
            wb.Close(False)
    except Exception:                                            # noqa: BLE001 — лишаємо .xlsx
        return None
    finally:
        with contextlib.suppress(Exception):
            xl.Quit()
        if watch is not None:
            watch.cancel()
    return target if os.path.exists(target) else None


WRITERS = {"rozrakhunky": write_rozrakhunky, "raport": write_raport, "1fs": write_1fs, "9fs": write_9fs,
           "suprovid": write_suprovid, "akt": write_akt, "reestr": write_reestr}


MAX_DIR_PARTS = 4                   # рівнів підтек у папері пакета: стільки ж, скільки й у файла із завантаження
CONTROL = re.compile(r"[\x00-\x1f\x7f]+")       # табуляція, переніс рядка тощо: Windows в імені теки їх не приймає
YEAR = re.compile("[0-9]{4}")                    # лише цифри ASCII: \d прийняв би й цифри інших писемностей
ISO_DAY = re.compile("([0-9]{4})-([0-9]{2})-([0-9]{2})")
ACT_DIR = "Введення в експлуатацію"              # підтека акта в теці «Прибутково-видаткові документи/<рік>»


def _part(text):
    """Частина шляху з тексту сторінки: без знаків керування (вони лишаються від вставлених назв) і без знаків, яких
    Windows в імені не приймає."""
    return safe_name(CONTROL.sub(" ", str(text)))


def _day(value, what):
    """Дата «рррр-мм-дд» із запиту, рівно така (порожня — порожній рядок), інакше ValueError: з хибною датою ім'я файла було
    б сміттям, а місяць 13 ще й зламав би текст акта."""
    if value is None or value == "":
        return ""
    try:
        match = ISO_DAY.fullmatch(value) if isinstance(value, str) else None
        if match:
            date(*map(int, match.groups()))
    except ValueError:
        match = None
    if not match:
        raise ValueError(f"{what} не схожа на дату рррр-мм-дд")
    return value


def _year(spec, day):
    """Рік теки паперу: `year` із запиту, а коли його нема, — перші чотири знаки дати. У шлях потрапляє лише рівно чотири
    цифри, а не те, що надіслано (`..`, диск, цифри інших писемностей, кінцевий переніс рядка, який пропустив би `$`)."""
    year = str(spec.get("year") or day[:4])
    if not YEAR.fullmatch(year):
        raise ValueError("рік паперу не вказано або не схожий на рік: потрібні чотири цифри")
    return year


def _target(spec):
    """(ключ теки, підтека, ім'я файла) паперу. Усе, що видно за самим запитом, відхиляється тут, до першого звернення
    до диска: ValueError — це 422, і такий запит не лишає ні файла, ні порожньої теки."""
    paper = spec.get("paper")
    if paper in NAMES:
        parts = [_part(p) for p in str(spec.get("dir") or "").replace("\\", "/").split("/") if p.strip()]
        if len(parts) > MAX_DIR_PARTS:
            raise ValueError(f"тека закупівлі глибша за {MAX_DIR_PARTS} рівні")
        key, rel, name = "purchases", "/".join(parts), NAMES[paper]
    elif paper == "akt":                                    # акт — у «Прибутково-видаткові документи/<рік>/Введення…»
        day = _day(spec.get("date"), "дата акта")
        key, rel = "receipts", f"{_year(spec, day)}/{ACT_DIR}"
        name = _part(f"Акт введення в експлуатацію {_s(spec, 'no')}".strip()) + ".docx"
    elif paper == "reestr":                                 # реєстр — у «Реєстри на ФЕС/<рік>»
        day = _day(spec.get("date"), "дата реєстру")
        if not day:
            raise ValueError("дату реєстру не вказано")
        key, rel = "fesRegisters", _year(spec, day)
        name = _part(f"Реєстр документів на ФЕС {ddmmyyyy(day)}") + ".xlsx"
    else:
        raise ValueError(f"невідомий папір: {paper!r}")
    rel.encode("utf-8")             # окремий сурогат (JSON \ud83d) не кодується: UnicodeEncodeError — це ValueError
    name.encode("utf-8")
    return key, rel, name


def _ensure_inside(folder, root):
    """Кидає ValueError, коли тека паперу виходить за межі теки служби (нічого не повертає). Після розкриття посилань вона має
    лишатися всередині: з'єднання (junction) чи посилання, що вже лежить у теці служби, не виводить запис геть. Перевірка — до
    першого запису. Інший диск чи шлях, що не складається, — теж «поза текою». Аргументи — (тека паперу, тека служби); у
    `launcher._inside(root, full)` порядок обернений, і вона повертає bool, тож імена не сплутати."""
    real, top = os.path.realpath(folder), os.path.realpath(root)
    try:
        inside = os.path.normcase(os.path.commonpath([real, top])) == os.path.normcase(top)
    except ValueError:
        inside = False
    if not inside:
        raise ValueError("тека паперу виходить за межі теки служби")


def _variants(spec, path):
    """Файли, під якими папір може лежати в теці: 9/фс — .xls (бланк штабу) або .xlsx (де Excel не було), решта — один."""
    if spec.get("paper") == "9fs":
        stem = os.path.splitext(path)[0]
        return [stem + ".xls", stem + ".xlsx"]
    return [path]


def _guard(paths, spec):
    """Кожен наявний файл паперу має бути таким, яким його записала програма (відбиток із запису), інакше — питання до
    людини: «змінено після програми — перезаписати?». Для 9/фс, що може лежати в обох форматах, дивляться на обидва.

    Запис може тримати й стабільний відбиток (нові записи), і звичайний SHA-256 (старі записи й історія): годиться будь-який
    із двох відбитків файла, що лежить у теці.

    Згода («так» — force) має покривати все, що вона зітре, тож питання називає КОЖЕН наявний файл, що не збігається. А коли
    наявні обидва формати 9/фс, «так» лишає один новий файл і видаляє другий (який — стане ясно лише після перетворення):
    про це сказано в питанні."""
    if spec.get("force"):
        return
    recorded = spec.get("prevSha") or ""
    there = [path for path in paths if os.path.exists(path)]
    changed = [path for path in there if recorded not in stable_hash.shas(path)]
    if changed:
        names = " і ".join(f"«{os.path.basename(path)}»" for path in changed)
        note = " (один із двох файлів буде видалено)" if len(there) > 1 else ""
        raise PaperConflict(f"{names} змінено після програми — перезаписати?{note}", changed=True)


TRIES, PAUSE = 5, 0.1               # скільки разів і через скільки секунд пробувати стати на місце паперу


def _busy_retry(action, path):
    """`action()` — заміна файла `path` чи його прибирання. Файл, відкритий у Excel чи Word, їй не піддається — це
    PaperConflict(busy) саме про нього. Але й щойно записаний файл на мить може тримати антивірус чи індексатор Windows,
    тож перш ніж просити закрити, пробуємо кілька разів."""
    for attempt in range(TRIES):
        try:
            return action()
        except PermissionError as e:
            if attempt == TRIES - 1:
                raise PaperConflict(f"«{os.path.basename(path)}» відкрито в Excel чи Word: закрийте й складіть знову",
                                    busy=True) from e
            time.sleep(PAUSE)


def _remove_dirs(made):
    """Прибирає теки зі списку `_make_dirs` (від найглибшої), якщо вони порожні; непорожні й відсутні лишає."""
    for here in made:
        with contextlib.suppress(OSError):
            os.rmdir(here)


def _make_dirs(folder):
    """`os.makedirs`, що повертає створені теки (від найглибшої): коли папір не склався, порожні теки, яких раніше не було,
    прибираються, щоб відмова не лишала по собі порожнього дерева. Так само, коли не вдалося створити саму теку: верхні
    рівні, які `os.makedirs` устиг створити, перш ніж глибший не дався, теж не лишаються."""
    made, here = [], os.path.abspath(folder)
    # lexists, а не exists: з'єднання чи посилання, що висить, — запис у теці, а не «нема». Інакше воно потрапило б до списку
    # створеного тут, і прибирання після відмови стерло б його, хоч воно було ще до запису.
    while not os.path.lexists(here):
        made.append(here)
        up = os.path.dirname(here)
        if up == here:
            break
        here = up
    try:
        os.makedirs(folder, exist_ok=True)
    except BaseException:
        _remove_dirs(made)
        raise
    return made


def _write_atomic(path, write, before=None):
    """`write(шлях)` складає папір у тимчасовий файл поруч, і під своїм іменем він з'являється лише цілим: збій чи
    обрив лишає попередній папір як був і не лишає недописаного. Тимчасовий файл прибирається за будь-якого виходу, а
    теки, яких не було, — якщо вони лишились порожніми.

    Книга .xlsx отримує імена Excel цього ПК (`excel_names`), ще поки вона тимчасовий файл: заголовки й область друку
    переживуть відкриття в Excel з українським інтерфейсом, а відбиток, який поверне `save`, уже остаточний. Папір, відкритий
    у Excel чи Word, заміні не піддається — це PaperConflict(busy).

    `before()` викликається, коли папір цілком готовий, а на місце ще не став: тут прибирають те, що має зникнути разом із
    заміною. Якщо воно (чи заміна) відмовило, усе лишається, як було."""
    folder = os.path.dirname(path)
    made, part = _make_dirs(folder), None
    try:
        fd, part = tempfile.mkstemp(dir=folder, suffix=".part")
        os.close(fd)                                # далі файл відкриває сам запис: двічі відкритим він не пишеться
        write(part)
        if path.lower().endswith(".xlsx"):
            with contextlib.suppress(OSError):
                excel_names.localize_file(part)
        if before:
            before()
        _busy_retry(lambda: os.replace(part, path), path)
    except BaseException:
        if part:
            with contextlib.suppress(OSError):
                os.remove(part)
        _remove_dirs(made)
        raise


def _remove(path):
    """Прибирає застарілий файл паперу (формат, якого цього разу не вийшло) тим самим повтором, що й заміну: зайнятий —
    PaperConflict про нього самого, коли нічого ще не змінено."""
    if os.path.exists(path):
        with contextlib.suppress(FileNotFoundError):
            _busy_retry(lambda: os.remove(path), path)


def _place_9fs(spec, path):
    """9/фс складається в приватній теці поза текою служби: книга, її імена під Excel і перетворення на .xls — усе там, тож
    Excel теки служби не торкається, а проміжного 9фс.xlsx у ній не буває. У теку папір (.xls, а без Excel чи коли
    перетворення не вдалося — .xlsx) потрапляє лише через `_write_atomic`: .xls цілий або ніякий, а відкритий в Excel —
    «зайнято». Формат, якого цього разу не вийшло, прибирається ПЕРЕД тим, як стане новий файл; не прибрався — «зайнято»
    про нього, і нічого не змінено. Повертає шлях до паперу в теці."""
    stem = os.path.splitext(path)[0]
    tmp = tempfile.mkdtemp(prefix="oblik-9fs-")
    try:
        book = os.path.join(tmp, os.path.basename(path))
        WRITERS["9fs"](spec, book)
        # .xls робить Excel цього ПК, а Excel з українським інтерфейсом не впізнає в прочитаній із файла книзі стандартних
        # імен заголовків і області друку: книга отримує його імена до перетворення.
        with contextlib.suppress(OSError):
            excel_names.localize_file(book)
        made = to_xls(book)
        final, stale = (stem + ".xls", stem + ".xlsx") if made else (stem + ".xlsx", stem + ".xls")
        _write_atomic(final, lambda part: shutil.copyfile(made or book, part), before=lambda: _remove(stale))
        return final
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def save(spec, roots):
    """Папір — у теку з налаштувань; відповідь для сторінки: тека, шлях, ім'я, розмір, відбиток.

    Спершу все, що видно за самим запитом (`_target`), далі тека, яка не виходить за межі теки служби (`_ensure_inside`), далі
    охоронець (`_guard`), і лише потім папір складається. Файл у відповіді — той, що справді лежить у теці (для 9/фс це
    .xls або, без Excel, .xlsx), з остаточним відбитком."""
    key, rel, name = _target(spec)
    folder = os.path.join(roots[key], *rel.split("/")) if rel else roots[key]
    _ensure_inside(folder, roots[key])
    path = os.path.join(folder, name)
    _guard(_variants(spec, path), spec)
    if spec.get("paper") == "9fs":
        path = _place_9fs(spec, path)
    else:
        write = WRITERS[spec["paper"]]
        _write_atomic(path, lambda part: write(spec, part))
    name = os.path.basename(path)
    return {"folder": key, "path": "/".join(x for x in (rel, name) if x), "file": name, "size": os.path.getsize(path),
            "sha": sha_of(path), "xls": path.lower().endswith(".xls")}


BLANKS = {"1фс (бланк).xlsx": write_1fs, "9фс (бланк).xlsx": write_9fs,
          "Розрахунки до рапорту (бланк).xlsx": write_rozrakhunky, "Рапорт (бланк).docx": write_raport,
          "Супровід (бланк).docx": write_suprovid, "Акт введення в експлуатацію (бланк).docx": write_akt,
          "Реєстр документів на ФЕС (бланк).xlsx": write_reestr}
BLANKS_DIR = "Бланки"


class _Present(Exception):
    """Бланк з'явився, поки складався другий: він не наш, його не чіпають."""


def _absent(path):
    """`before` для `_write_atomic`: перед самою постановкою на місце бланка вже не має бути."""
    def check():
        if os.path.lexists(path):
            raise _Present(path)
    return check


def save_blanks(roots, unit):
    """Порожні бланки — у «Закупівлі/Бланки», якщо їх там ще немає: без програми є з чого почати. Повертає імена створених.

    Наявний бланк (навіть порожній файл, тека під таким іменем чи висяче посилання — `lexists`, а не `exists`) не чіпається
    ніколи, ні при перевірці, ні в мить постановки на місце. Кожен бланк лягає через `_write_atomic`: цілий або ніякий, без
    .part, а .xlsx — в іменах Excel цього ПК. `unit`, що не об'єкт (None, рядок, число), — як порожній: бланки без реквізитів.
    Помилки — ValueError (тека «Бланки» веде за межі теки закупівель), PaperConflict (зайнято), OSError (не створити):
    рішення за викликачем."""
    folder = os.path.join(roots["purchases"], BLANKS_DIR)
    _ensure_inside(folder, roots["purchases"])
    blank = {"unit": unit if isinstance(unit, dict) else {}, "year": date.today().year, "kpkv": "", "signers": [{}, {}, {}],
             "rows": [], "groups": [], "positions": [], "lines": []}
    made = []
    for name, write in BLANKS.items():
        path = os.path.join(folder, name)
        if os.path.lexists(path):
            continue
        try:
            _write_atomic(path, lambda part, write=write: write(blank, part), before=_absent(path))
        except _Present:
            continue
        made.append(name)
    return made
