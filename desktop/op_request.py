# -*- coding: utf-8 -*-
"""Заявка на посуд одноразового використання, миючі засоби й серветки на 30 днів.

Розрахунок за нормами наказу МОУ №390 від 29.07.2016: одноразовий посуд — набір на одну
добову видачу (п.13: за фактом, не більше осіб на харчуванні), серветки — 3 шт на особу на
добу (Норма №1), рідкий миючий засіб для ручного миття — 240 г на 100 осіб на добу,
з одноразовим посудом зменшується в 1,5 раза (Норма №13, прим. 3): 0,16 кг на 100 осіб.
Кількість округлюється вгору до упаковки. Лист — як заявки служби: «Передати засобами
СЕДО», адресат, текст із посиланням на наказ, підпис командира.
"""
import datetime
import math
import os

NORMS = [
    {"key": "set", "name": "одноразовим посудом", "uom": "комплектів", "per": 1.0, "step": 1000,
     "basis": "наказ МОУ №390, п.13: набір на одну добову видачу"},
    {"key": "liquid", "name": "рідким миючим засобом", "uom": "кг", "per": 0.0016, "step": 5,
     "basis": "наказ МОУ №390, Норма №13, прим. 3: 240 г на 100 осіб на добу ÷ 1,5"},
    {"key": "napkin", "name": "серветками паперовими", "uom": "шт", "per": 3.0, "step": 1000,
     "basis": "наказ МОУ №390, Норма №1: 3 шт на особу на добу"},
]
ORDER = ("наказу МОУ №390 від 29.07.2016 р. «Про затвердження Порядку застосування Норм забезпечення "
         "столово-кухонним посудом, обладнанням, інвентарем та мийними засобами в системі МО України»")


def compute(per_day, days=30):
    """Потреба на `days` днів при `per_day` добових видачах: сирий розрахунок і кількість з упаковкою."""
    out = []
    for n in NORMS:
        raw = round(float(per_day) * n["per"] * int(days), 6)
        qty = math.ceil(raw / n["step"] - 1e-9) * n["step"]
        out.append({**n, "raw": int(raw) if float(raw).is_integer() else round(raw, 2), "qty": qty})
    return out


def _thousands(v):
    return f"{v:,}".replace(",", " ")


def save_request(con, spec, folder):
    """Лист-заявка у Word у теку заявок: {"perDay", "days", "to", "date"} → шлях до файла."""
    from docx import Document                                     # noqa: PLC0415
    from docx.enum.text import WD_ALIGN_PARAGRAPH                 # noqa: PLC0415
    from docx.shared import Pt                                    # noqa: PLC0415
    from reports.form21 import _unit_code                        # noqa: PLC0415
    per_day = float(spec["perDay"])
    days = int(spec.get("days") or 30)
    if per_day <= 0 or days <= 0:
        raise ValueError("вкажіть середні добові видачі й кількість днів")
    day = str(spec.get("date") or datetime.date.today().isoformat())[:10]
    unit = _unit_code(con) or "А____"
    rows = compute(per_day, days)
    cmd = con.execute("""SELECT p.position, p.rank, p.full_name FROM official o JOIN person p ON p.id = o.person_id
                          WHERE o.role = 'командир' AND o.valid_from <= ? AND (o.valid_to IS NULL OR o.valid_to > ?)
                          ORDER BY o.valid_from DESC LIMIT 1""", (day, day)).fetchone()
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(14)
    doc.add_paragraph("Передати засобами СЕДО").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph(str(spec.get("to") or "")).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph()
    p.add_run("ЗАЯВКА").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    what = ", ".join(f"{r['name']} {_thousands(r['qty'])} {r['uom']}" for r in rows)
    pd = int(per_day) if per_day.is_integer() else per_day
    for text in (f"Відповідно до {ORDER}.",
                 f"Прошу забезпечити військову частину {unit} {what}. Із розрахунку в середньому {pd} д/д на {days} днів."):
        doc.add_paragraph(text).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    doc.add_paragraph((cmd["position"] if cmd else "") or f"Командир військової частини {unit}")
    doc.add_paragraph(f"{(cmd['rank'] if cmd else '') or '_' * 12}\t\t\t\t{(cmd['full_name'] if cmd else '') or '_' * 20}")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{unit} Заявка ОП, МЗ, серветки {day}.docx")
    try:
        doc.save(path)
    except PermissionError:
        path = path[:-5] + f" {datetime.datetime.now():%H%M%S}.docx"
        doc.save(path)
    return path
