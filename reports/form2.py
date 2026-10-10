# -*- coding: utf-8 -*-
"""Звіт-заявка 2/прод — з бази програми, за бланком А2788.

Графи 8 і 5 — зі зданого звіту минулого року (так велить методичка: гр.8 = гр.18 минулого
звіту), рух року — з документів обох книг обліку по рядках, до яких прив'язано коди (з
множником одиниць), помилки минулих звітів — записками в гр.14 (+) і гр.17 (−). Гр.18 мусить
дорівнювати обліку на кінець року; різницю програма пропонує закрити запискою.

Пакет — книга Excel за бланком (шаблон desktop/templates/form2_a2788.xlsx без зразкових імен):
аркуш «звіт», розшифровки «графа 12/13/14/16/17» переліком документів (як ручний файл 2026)
і «Перевірки»; окремо — проєкт пояснювальної записки (Word). Ручний файл 2/прод програма не
перезаписує: її файл має свою назву «… (програма)».
"""
import datetime
import io
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook                               # noqa: E402
from openpyxl.styles import Font                                 # noqa: E402

FORM = "2/Прод"
OTHERS = "Майно інших служб"
MILLI = 1000
INPUT_COLS = (5, 10, 20, 21, 22, 23, 24, 25, 26, 29)
DOC_COLS = (12, 13, 14, 16, 17)
# У зібраній програмі бланки лежать у теці templates поруч із кодом (PyInstaller), з джерел — у desktop/templates.
TEMPLATE = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1] / "desktop")) / "templates" / "form2_a2788.xlsx"
SHEET = "звіт"
FIRST_ROW, LAST_BLANK = 23, 1848
WRITE_COLS = (5, 8, 9, 10, 12, 13, 14, 16, 17, 19, 20, 21, 22, 23, 24, 25, 26, 29)
FORMULAS = {"F": "=E{r}*0.05", "G": "=E{r}+F{r}", "K": "=L{r}+M{r}+N{r}", "O": "=P{r}+Q{r}",
            "R": "=H{r}+K{r}-O{r}", "AA": "=R{r}-G{r}", "AB": "=G{r}-R{r}"}
DOC_TITLES = {12: "від заводу (підприємства) - постачальника", 13: "отримано за нарядами органу забезпечення",
              14: "інші надходження", 16: "списано після закінчення термінів експлуатації", 17: "інші витрати"}
LIST_HEAD = ["№ з/п", "Найменування матеріальних засобів", "Одиниця обліку", "Категорія", "Кількість",
             "Документ (назва, номер)", "Дата документа", "Вантажовідправник / вантажоодержувач"]
VERIFY = "рядок — перевірити"


# ---------------------------------------------------------------- рушій

def column_of(journal, paper, party_kind, note, override):
    """Графа документа і звідки правило: вручну, рішення власника, методичка, припущення."""
    if override in DOC_COLS:
        return override, "вручну"
    paper = (paper or "").strip().lower()
    if journal == "incoming":
        if paper == "наряд" or "наряд" in (note or "").lower():
            return 13, "підтверджено"
        if party_kind == "постачальник":
            return 12, "потребує підтвердження"
        if party_kind == "військова частина":
            return 14, "підтверджено"
        return 14, "за методичкою"
    if paper in ("акт технічного стану", "акт зміни якісного стану"):
        return 16, "за методичкою"
    return 17, "потребує підтвердження"


def _own_subs(con):
    """Підрозділи частини без вузла «Майно інших служб» і його гілки (КЕС, ЕТС у звіт не йдуть)."""
    rows = con.execute("SELECT id, name, parent_id FROM subdivision").fetchall()
    kids = {}
    for r in rows:
        kids.setdefault(r["parent_id"], []).append(r["id"])
    out = {r["id"] for r in rows}
    for r in rows:
        if r["name"] == OTHERS:
            stack = [r["id"]]
            while stack:
                x = stack.pop()
                out.discard(x)
                stack += kids.get(x, [])
    return out


def _form_id(con):
    row = con.execute("SELECT id FROM report_form WHERE code = ?", (FORM,)).fetchone()
    return row[0] if row else None


def collect(con, year, as_of=None):
    """Звіт за рік: рядки бланка з графами 5–29, документи по графах, залишок обліку,
    пропозиції записок і перевірки."""
    year = int(year)
    as_of = as_of or f"{year}-12-31"
    d1, d2 = f"{year}-01-01", as_of
    form = _form_id(con)
    rows = {}
    for ln in con.execute("SELECT * FROM report_line WHERE form_id = ? AND row_no IS NOT NULL ORDER BY row_no",
                          (form,)):
        rows[ln["row_no"]] = {"row": ln["row_no"], "name": ln["name"], "section": ln["section"] or "",
                              "uom": ln["uom"] or "", "header": bool(ln["is_header"]),
                              "c": {c: 0.0 for c in range(5, 30)}, "balance": 0.0,
                              "docs": {c: [] for c in DOC_COLS}, "codes": []}
    bind, unsure, unsure_moved = {}, set(), set()
    for r in con.execute("""SELECT m.nomen_id, rl.row_no, m.factor, n.code, m.checked_by FROM nomen_report_line m
                              JOIN report_line rl ON rl.id = m.report_line_id JOIN nomen n ON n.id = m.nomen_id
                             WHERE rl.form_id = ? AND rl.row_no IS NOT NULL""", (form,)):
        bind[r["nomen_id"]] = (r["row_no"], r["factor"], r["code"])
        if not r["checked_by"]:
            unsure.add(r["nomen_id"])
        if r["row_no"] in rows:
            rows[r["row_no"]]["codes"].append(r["code"])
    skip = {r[0] for r in con.execute("SELECT nomen_id FROM nomen_form_skip WHERE form_id = ?", (form,))}
    own = _own_subs(con)
    prev = {(r["row_no"], r["col"]): r["value_milli"] / MILLI
            for r in con.execute("SELECT * FROM form2_submitted WHERE year = ?", (year - 1,))}
    cells = {(r["row_no"], r["col"]): r["value_milli"] / MILLI
             for r in con.execute("SELECT * FROM form2_cell WHERE year = ?", (year,))}
    for (row_no, col), v in prev.items():
        if row_no in rows and col == 18:
            rows[row_no]["c"][8] = v
        if row_no in rows and col == 5 and (row_no, 5) not in cells:
            rows[row_no]["c"][5] = v
    for (row_no, col), v in cells.items():
        if row_no in rows and col in INPUT_COLS:
            rows[row_no]["c"][col] = v
    checks, unbound = [], {}
    for r in con.execute("""
            SELECT d.id, d.number, d.doc_date, COALESCE(d.paper, k.name) AS paper,
                   d.from_subdivision_id AS fsub, d.to_subdivision_id AS tsub, c.name AS party, ck.code AS pkind,
                   d.note, e.form2_col, l.nomen_id, n.code, SUM(l.qty_milli) AS q
              FROM document d JOIN doc_kind k ON k.id = d.kind_id
              JOIN document_line l ON l.document_id = d.id JOIN nomen n ON n.id = l.nomen_id
              LEFT JOIN counterparty c ON c.id = d.counterparty_id
              LEFT JOIN counterparty_kind ck ON ck.id = c.kind_id
              LEFT JOIN document_extra e ON e.document_id = d.id
             WHERE k.affects_stock = 1 AND d.doc_date BETWEEN ? AND ?
             GROUP BY d.id, l.nomen_id
             ORDER BY d.doc_date, d.id""", (d1, d2)):
        fin, tin = r["fsub"] in own, r["tsub"] in own
        if fin and tin:
            continue                                            # усередині частини
        if tin and not fin:
            journal = "incoming"
        elif fin and not tin:
            journal = "writeoffs"
        else:
            continue
        if r["nomen_id"] not in bind:
            if r["nomen_id"] not in skip:
                unbound[r["code"]] = unbound.get(r["code"], 0) + 1
            continue
        row_no, factor, _code = bind[r["nomen_id"]]
        if row_no not in rows:
            continue
        if r["nomen_id"] in unsure:
            unsure_moved.add(r["code"])
        col, state = column_of(journal, r["paper"], r["pkind"], r["note"], r["form2_col"])
        qty = r["q"] / MILLI * factor
        rows[row_no]["c"][col] += qty
        rows[row_no]["docs"][col].append({"id": r["id"], "type": r["paper"] or "", "no": r["number"],
                                          "date": r["doc_date"], "party": r["party"] or "", "qty": qty,
                                          "rule": state, "code": r["code"]})
    for n in con.execute("SELECT * FROM form2_note WHERE year = ? ORDER BY row_no, id", (year,)):
        if n["row_no"] not in rows:
            continue
        rows[n["row_no"]]["c"][n["col"]] += n["qty_milli"] / MILLI
        rows[n["row_no"]]["docs"][n["col"]].append(
            {"id": f"n{n['id']}", "type": "Записка", "no": n["number"] or "", "date": n["doc_date"] or "",
             "party": "", "qty": n["qty_milli"] / MILLI, "rule": n["reason"], "code": ""})
        if not n["number"] or not n["doc_date"]:
            checks.append({"level": "✗", "text": f"рядок {n['row_no']}: записка без номера чи дати"})
    ids = ",".join(str(x) for x in sorted(own)) or "NULL"         # id — цілі з бази, не текст людини
    for nid, (row_no, factor, _code) in bind.items():
        if row_no not in rows:
            continue
        q = con.execute(f"""SELECT COALESCE(SUM(sign * qty_milli), 0) FROM posting
                             WHERE nomen_id = ? AND doc_date <= ? AND subdivision_id IN ({ids})""",
                        (nid, as_of)).fetchone()[0]
        rows[row_no]["balance"] += q / MILLI * factor
    proposals = []
    for r in rows.values():
        c = r["c"]
        c[9] = c[8] - c[10]
        c[11] = c[12] + c[13] + c[14]
        c[15] = c[16] + c[17]
        c[18] = c[8] + c[11] - c[15]
        c[19] = c[18] - c[20]
        c[6] = c[5] * 0.05
        c[7] = c[5] + c[6]
        c[27] = (c[18] - c[26]) - c[7]
        c[28] = c[7] - (c[18] - c[26])
        for k in c:
            c[k] = round(c[k], 6)
        r["balance"] = round(r["balance"], 6)
        if r["header"]:
            continue
        gap = round(c[18] - r["balance"], 6)
        if abs(gap) > 1e-6:
            proposals.append({"row": r["row"], "col": 17 if gap > 0 else 14, "qty": abs(gap),
                              "reason": f"розбіжність звіту {year - 1} з обліком"})
            checks.append({"level": "✗", "text": f"рядок {r['row']} «{r['name']}»: гр.18 {gap:+g} проти обліку"})
        if not r["uom"] and (c[11] or c[15] or c[8]):
            checks.append({"level": "⚠", "text": f"рядок {r['row']} «{r['name']}»: у бланку немає одиниці"})
    for code in sorted(unbound):
        checks.append({"level": "✗", "text": f"код {code}: рух у {year} році, але не прив'язано до рядка 2/Прод"})
    for code in sorted(unsure_moved):
        checks.append({"level": "⚠", "text": f"код {code}: рух у {year} році, прив'язку ще не перевірено («уточнити»)"})
    if not prev:
        checks.append({"level": "⚠", "text": f"немає зданого звіту {year - 1}: гр.8 і гр.5 порожні"})
    return {"year": year, "as_of": as_of, "rows": [rows[k] for k in sorted(rows)], "proposals": proposals,
            "checks": checks, "submittedPrev": bool(prev)}


# ---------------------------------------------------------------- пакет

def _num(v):
    v = round(float(v or 0), 3)
    return int(v) if float(v).is_integer() else v


def folder_of(root, year):
    return os.path.join(root, str(year), "2прод")


def _save_wb(wb, path):
    """Книга з іменами шапки й області друку для Excel цього комп'ютера; файл, відкритий в Excel,
    не затирається — поруч лягає новий, з часом у назві."""
    from desktop.excel_names import save_book                    # noqa: PLC0415
    try:
        save_book(wb, path)
    except PermissionError:
        stem, ext = os.path.splitext(path)
        path = f"{stem} {datetime.datetime.now():%H%M%S}{ext}"
        save_book(wb, path)
    return path


def _save_doc(doc, path):
    try:
        doc.save(path)
    except PermissionError:
        stem, ext = os.path.splitext(path)
        path = f"{stem} {datetime.datetime.now():%H%M%S}{ext}"
        doc.save(path)
    return path


def _commander(con, cut):
    r = con.execute("""SELECT p.position, p.rank, p.full_name FROM official o JOIN person p ON p.id = o.person_id
                        WHERE o.role = 'командир' AND o.valid_from <= ? AND (o.valid_to IS NULL OR o.valid_to > ?)
                        ORDER BY o.valid_from DESC LIMIT 1""", (cut, cut)).fetchone()
    return tuple(str(x or "") for x in r) if r else ("", "", "")


def save_package(con, spec, root):
    """Бланк А2788 з графами з бази, розшифровки 12/13/14/16/17 переліком і «Перевірки»."""
    from reports.form21 import _signers, _unit_code               # noqa: PLC0415
    year = int(spec["year"])
    rep = collect(con, year, spec.get("asOf") or None)
    unit = _unit_code(con)
    cut = (datetime.date.fromisoformat(rep["as_of"]) + datetime.timedelta(days=1)).isoformat()
    wb = load_workbook(TEMPLATE)
    ws = wb[SHEET]
    pos, rank, name = _commander(con, cut)
    subst = {"{Р+1}": str(year + 1), "{Р}": str(year), "{ЧАСТИНА}": unit or "А____",
             "{КОМАНДИР_ПОСАДА}": pos or f"Командир військової частини {unit or 'А____'}",
             "{КОМАНДИР_ЗВАННЯ}": rank or "_" * 15, "{КОМАНДИР_ІМ'Я}": name or "_" * 20}
    for row in ws.iter_rows(max_row=FIRST_ROW - 1):
        for c in row:
            if isinstance(c.value, str) and "{" in c.value:
                v = c.value
                for k, s in subst.items():                         # {Р+1} — раніше за {Р}
                    v = v.replace(k, s)
                c.value = v
    last = LAST_BLANK
    for r in rep["rows"]:
        n = r["row"]
        if n > LAST_BLANK:                                         # власний рядок — після бланка
            ws[f"A{n}"] = f"=SUBTOTAL(103,$B$24:B{n})"
            ws[f"B{n}"], ws[f"C{n}"] = r["name"], r["uom"] or None
            for col, f in FORMULAS.items():
                ws[f"{col}{n}"] = f.format(r=n)
            last = max(last, n)
        if r["header"]:
            continue
        for col in WRITE_COLS:
            v = r["c"].get(col, 0)
            if v:
                ws.cell(n, col, _num(v))
        ws.row_dimensions[n].hidden = False
    end = _signatures(ws, last + 2, _signers(con, cut))
    ws.print_area = f"A1:AC{end}"
    for col in DOC_COLS:
        _list_sheet(wb, col, year, rep)
    _checks_sheet(wb, con, rep)
    folder = folder_of(root, year)
    os.makedirs(folder, exist_ok=True)
    return _save_wb(wb, os.path.join(folder, f"{unit or 'А0000'} {year} 2прод - Звіт (програма).xlsx"))


def _signatures(ws, row, signers):
    """Під таблицею — начальник служби, тоді начальник логістики («МВО й посадовці» на дату звіту)."""
    small = Font(size=8, italic=True)
    for position, rank, name in reversed(signers):
        ws.cell(row, 2, position or "_" * 60)
        ws.cell(row + 1, 2, rank or "_" * 15)
        ws.cell(row + 1, 8, name or "_" * 20)
        ws.cell(row + 2, 2, "(посада, військове звання, підпис, ім'я та ПРІЗВИЩЕ)").font = small
        row += 4
    return row


def _list_sheet(wb, col, year, rep):
    """Розшифровка графи переліком документів — як у ручному файлі 2026."""
    ws = wb.create_sheet(f"графа {col}")
    ws["A1"] = "РОЗШИФРОВКА"
    ws["A2"] = f"до графи {col} «{DOC_TITLES[col]}» звіту 2/прод за {year} рік"
    for c, title in enumerate(LIST_HEAD, 1):
        ws.cell(4, c, title)
    n, at = 0, 5
    for r in rep["rows"]:
        for d in sorted(r["docs"][col], key=lambda d: (d["date"] or "", str(d["no"]))):
            n += 1
            day = datetime.date.fromisoformat(d["date"]) if d["date"] else None
            for c, v in enumerate([n, r["name"], r["uom"] or None, None, _num(d["qty"]),
                                   f"{d['type']} №{d['no']}", day, d["party"] or None], 1):
                ws.cell(at, c, v)
            if day:
                ws.cell(at, 7).number_format = "dd.mm.yyyy"
            at += 1
    ws.print_title_rows = "4:4"
    for letter, width in zip("ABCDEFGH", (6, 60, 10, 10, 12, 28, 12, 28)):
        ws.column_dimensions[letter].width = width


def _checks_sheet(wb, con, rep):
    from reports.form21 import _ua                                # noqa: PLC0415
    ws = wb.create_sheet("Перевірки")
    ws.append(["", "Перевірка"])
    out = list(rep["checks"])
    scanned = {r[0] for r in con.execute("SELECT DISTINCT document_id FROM attachment")} | {
        r[0] for r in con.execute("SELECT document_id FROM document_extra WHERE scan IS NOT NULL AND scan <> ''")}
    seen, guessed = set(), {}
    for r in rep["rows"]:
        for col in DOC_COLS:
            for d in r["docs"][col]:
                if not isinstance(d["id"], int):
                    continue
                if d["rule"] == "потребує підтвердження":
                    guessed.setdefault(col, set()).add(d["id"])
                if d["id"] in seen or d["id"] in scanned:
                    continue
                seen.add(d["id"])
                out.append({"level": "⚠", "text": f"{d['type']} №{d['no']} від {_ua(d['date'])}: немає скану"})
    for col, ids in sorted(guessed.items()):
        out.append({"level": "⚠", "text": f"графа {col}: {len(ids)} док. за правилом, яке ще треба підтвердити"})
    out.append({"level": "⚠", "text": "гр.6 «5 %» стоїть у кожному рядку бланка: для обладнання перевірте, "
                                      "чи потрібні перехідні запаси"})
    for x in out:
        ws.append([x["level"], x["text"]])
    ws.column_dimensions["B"].width = 110


def save_note(con, spec, root):
    """Проєкт пояснювальної записки: кожна записка року — рядок таблиці з причиною."""
    from docx import Document                                     # noqa: PLC0415
    from docx.enum.text import WD_ALIGN_PARAGRAPH                 # noqa: PLC0415
    from docx.shared import Pt                                    # noqa: PLC0415
    from reports.form21 import _signers, _ua, _unit_code         # noqa: PLC0415
    year = int(spec["year"])
    unit = _unit_code(con) or "А____"
    rows = con.execute("""SELECT n.*, rl.name, rl.uom FROM form2_note n
                            LEFT JOIN report_line rl ON rl.row_no = n.row_no
                             AND rl.form_id = (SELECT id FROM report_form WHERE code = ?)
                           WHERE n.year = ? ORDER BY n.row_no, n.id""", (FORM, year)).fetchall()
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(14)
    p = doc.add_paragraph()
    p.add_run("ПОЯСНЮВАЛЬНА ЗАПИСКА").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"до звіту-заявки 2/прод за {year} рік військової частини {unit}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Графа 8 звіту за {year} рік дорівнює графі 18 зданого звіту за {year - 1} рік. Розбіжності "
                      "звітів попередніх років з даними обліку виправлено графами 14 (+) і 17 (−):")
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("№", "Найменування", "Од.", "Графа", "Кількість", "Причина", "Записка")):
        c.text = h
    for i, n in enumerate(rows, 1):
        qty = f"{_num(n['qty_milli'] / MILLI):,}".replace(",", " ")
        vals = (str(i), n["name"] or f"рядок {n['row_no']}", n["uom"] or "", f"{n['col']} ({'+' if n['col'] == 14 else '−'})",
                qty, n["reason"], f"№{n['number'] or '___'} від {_ua(n['doc_date']) if n['doc_date'] else '___'}")
        for c, v in zip(t.add_row().cells, vals):
            c.text = v
    doc.add_paragraph()
    for position, rank, name in reversed(_signers(con, f"{year + 1}-01-01")):
        doc.add_paragraph(position or "_" * 40)
        doc.add_paragraph(f"{rank or '_' * 12}\t\t\t\t{name or '_' * 20}")
    folder = folder_of(root, year)
    os.makedirs(folder, exist_ok=True)
    return _save_doc(doc, os.path.join(folder, f"{unit} {year} 2прод - Пояснювальна записка (проєкт, програма).docx"))


# ---------------------------------------------------------------- файли служби

def _norm_name(s):
    return " ".join(str(s or "").split()).lower()


def read_submitted(con, data):
    """Зданий звіт (бланк А2788): рік — з рядка «… за <рік> рік», графи 5–18 — по рядках; рядок
    береться, лише коли назва збігається з переліком (бланк міг бути іншої редакції)."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    m = re.search(r"за\s+(\d{4})\s+рік", str(ws["A14"].value or ""))
    if not m:
        raise ValueError("не видно року звіту: у клітинці A14 бланка має бути «… за <рік> рік»")
    form = _form_id(con)
    names = dict(con.execute("SELECT row_no, name FROM report_line WHERE form_id = ? AND row_no IS NOT NULL", (form,)))
    values, mismatch, unknown = {}, [], []
    for n in range(FIRST_ROW, ws.max_row + 1):
        name = _norm_name(ws.cell(n, 2).value)
        if not name:
            continue
        nums = {c: float(ws.cell(n, c).value) for c in range(5, 19)
                if isinstance(ws.cell(n, c).value, (int, float)) and ws.cell(n, c).value}
        if n not in names:
            if nums:
                unknown.append(f"рядок {n} «{ws.cell(n, 2).value}»: у переліку 2/Прод такого рядка немає")
            continue
        if _norm_name(names[n]) != name:
            mismatch.append(f"рядок {n}: у файлі «{ws.cell(n, 2).value}», у переліку «{names[n]}»")
            continue
        for c, v in nums.items():
            values[(n, c)] = v
    return {"year": int(m.group(1)), "values": values, "mismatch": mismatch, "unknown": unknown}


def store_submitted(con, year, values):
    con.execute("DELETE FROM form2_submitted WHERE year = ?", (year,))
    for (row, col), v in values.items():
        con.execute("INSERT INTO form2_submitted(year, row_no, col, value_milli) VALUES(?, ?, ?, ?)",
                    (year, row, col, int(round(v * MILLI))))
    return len(values)


def snapshot(con, year):
    """«Звіт здано»: графи 5–18 звіту року — у зданий звіт (з нього береться гр.8 наступного)."""
    rep = collect(con, year)
    values = {(r["row"], c): r["c"][c] for r in rep["rows"] if not r["header"]
              for c in range(5, 19) if abs(r["c"][c]) > 1e-9}
    return store_submitted(con, int(year), values)


def read_map(data, file_date):
    """Відповідність кодів рядкам 2/Прод (аркуш «Відповідність» звіту служби): прив'язки —
    перевіреними («звірено <рік файла>», дата файла), крім рядків «рядок — перевірити»; коди без
    рядка з приміткою — поза 2/прод з цією приміткою; рядки бланка після 1848 з одиницею —
    власні рядки."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "Відповідність" not in wb.sheetnames:
        raise ValueError("у файлі немає аркуша «Відповідність»")
    grid = list(wb["Відповідність"].iter_rows(values_only=True))
    head = [" ".join(str(h or "").split()) for h in grid[0]] if grid else []
    miss = [h for h in ("Код", "Рядок форми", "Множник", "Примітка") if h not in head]
    if miss:
        raise ValueError("на аркуші «Відповідність» немає колонок: " + ", ".join(miss))
    ix = {h: i for i, h in enumerate(head) if h}
    who_col = ix.get("Перевірено (хто, дата)")
    label = f"звірено {str(file_date)[:4]}"
    out, stats = {}, {"rows": 0, "checked": 0, "unsure": 0, "skip": 0, "unbound": 0}
    for r in grid[1:]:
        code = " ".join(str(r[ix["Код"]] or "").split())
        if not code:
            continue
        stats["rows"] += 1
        row = r[ix["Рядок форми"]]
        note = " ".join(str(r[ix["Примітка"]] or "").split())
        who = " ".join(str(r[who_col] or "").split()) if who_col is not None else ""
        try:
            factor = float(r[ix["Множник"]] or 1)
        except (TypeError, ValueError):
            factor = 1.0
        if row in (None, "") or not str(row).strip().isdigit():
            if not note:
                stats["unbound"] += 1
                continue
            out[code] = {"row": None, "factor": 1.0, "checked": None, "checkedOn": None, "skip": note}
            stats["skip"] += 1
            continue
        unsure = note == VERIFY
        out[code] = {"row": int(row), "factor": factor, "checked": None if unsure else (who or label),
                     "checkedOn": None if unsure else str(file_date), "skip": None}
        stats["unsure" if unsure else "checked"] += 1
    own = []
    if SHEET in wb.sheetnames:
        ws = wb[SHEET]
        for n in range(LAST_BLANK + 1, ws.max_row + 1):
            name = " ".join(str(ws.cell(n, 2).value or "").split())
            uom = " ".join(str(ws.cell(n, 3).value or "").split())
            if name and uom:                                       # під підписами одиниці немає
                own.append({"row": n, "name": name, "section": None, "uom": uom})
    return {"map": out, "own": own, "stats": stats}
