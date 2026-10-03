# -*- coding: utf-8 -*-
"""Заповнення бригадної «Відомості втрат та списання МТЗ продовольчої служби»
(додатки 1 і 2) з довідки про втрати, яку склала програма («Знищене майно» →
«Довідка про втрати», аркуш «Рядки»). Форму бригада надсилає свою; сюди — її файл.

    python reports/losses_brigade.py <довідка.xlsx> <форма.xlsx> <результат.xlsx> [звітна дата]

Аркуші результату: «2022» (додаток 1) і «2024», «2025», «2026» (додаток 2).
Суми — у млн грн, як вимагає форма, без округлення (формат клітинки з шістьма
знаками = точність до гривні). Техзасоби — за типами форми, столово-кухонний
посуд та інвентар — у рядок 5 «Інше майно продовольчої служби» сумою (тонн
програма не знає, кількість одиниць — у примітці).
"""
import copy
import datetime as dt
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
src, form, out = sys.argv[1:4]

# ---- рядки довідки ----------------------------------------------------------
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Рядки"]
rows = list(ws.iter_rows(values_only=True))
head_i = next(i for i, r in enumerate(rows) if r and r[0] == "Дата події")
head = [str(x or "") for x in rows[head_i]]
col = {h: i for i, h in enumerate(head)}
recs = []
for r in rows[head_i + 1:]:
    if not r or not r[0] or r[1] == "Разом":
        continue
    d = r[col["Дата події"]]
    if isinstance(d, dt.datetime):
        d = d.date().isoformat()
    a = r[col["Дата акта"]]
    if isinstance(a, dt.datetime):
        a = a.date().isoformat()
    recs.append({"date": str(d), "sub": r[col["Підрозділ"]] or "", "code": str(r[col["Код"]] or ""),
                 "name": str(r[col["Найменування"]] or ""), "qty": float(r[col["К-сть"]] or 0),
                 "sum": float(r[col["Сума, грн"]] or 0), "price": r[col["Ціна, грн"]],
                 "report": str(r[col["Рапорт"]] or ""), "act": str(r[col["Акт чи наказ"]] or ""),
                 "actDate": str(a or ""), "status": str(r[col["Стан"]] or "")})
print(f"довідка: {len(recs)} рядків, {sum(r['qty'] for r in recs):g} од., {sum(r['sum'] for r in recs):,.2f} грн")

# ---- класифікація за типами форми -------------------------------------------
# Посуд, інвентар, мийні засоби, немеханічне обладнання — рядок 5 форми
# (пояснення до додатків); засоби підвозу води — тип 3; засоби приготування їжі
# (кухні, плити, печі, пароконвектомати, кип'ятильники, фритюрниці, термоси) —
# тип 1; холодильне, технологічне (механічне) і ваговимірювальне — тип 4.
WATER = re.compile(r"цистерн|для води|водян|єврокуб|[iі][bв][cс]|контейнер полімерний|бочк", re.I)
FOOD_TRANSPORT = re.compile(r"підвоз|підвез|фургон|термофургон", re.I)
COOK = re.compile(r"кухн|плита|піч|пароконвектомат|сковород|кип.ятильник|термос|фр[иі]тюрниц|казан|котел|хлібопіч|жарочн", re.I)
UTENSIL = re.compile(r"гастроємн|кришк|каструл|кухоль|миск|тарілк|ложк|виделк|ніж|чайник|відро|стіл|мийк|стелаж|полиц|намет", re.I)
TECH = re.compile(r"холодильн|морозильн|м.ясорубк|слайсер|овочеріз|блендер|міксер|водонагрівач|насос|терез|ваги|шаф|дегідратор|пакувальник|тістоміс|хліборізк", re.I)


def kind_of(name):
    n = name.lower()
    if WATER.search(n):
        return 3
    if FOOD_TRANSPORT.search(n):
        return 2
    if COOK.search(n):
        return 1
    if UTENSIL.search(n):
        return 5
    if TECH.search(n):
        return 4
    return 4


def short(name):
    """Назва без хвоста з кодом ФЕС і без службових уточнень."""
    s = re.sub(r",\s*\d{7,}\s*$", "", name).strip()
    s = re.sub(r"\s+", " ", s)
    return s[:90]


# ---- зрізи за роками --------------------------------------------------------
def lost_in(r, y0, y1):
    return y0 <= r["date"] <= y1


def off_in(r, y0, y1):
    return r["status"] == "списано" and r["actDate"] and y0 <= r["actDate"] <= y1


def open_at(r, day):
    return r["date"] <= day and not (r["status"] == "списано" and r["actDate"] and r["actDate"] <= day)


def agg(items):
    return (sum(r["qty"] for r in items), sum(r["sum"] for r in items))


def by_name(items):
    out = OrderedDict()
    for r in items:
        out.setdefault((kind_of(r["name"]), short(r["name"])), []).append(r)
    return out


# ---- запис у форму ----------------------------------------------------------
fwb = openpyxl.load_workbook(form)
T1 = fwb["Додаток 1"]
T2 = fwb["Додаток 2"]
MLN = 1_000_000.0
FMT_Q = "0.###"
FMT_M = "0.000000"


def copy_style(src_cell, dst_cell):
    dst_cell.font = copy.copy(src_cell.font)
    dst_cell.border = copy.copy(src_cell.border)
    dst_cell.fill = copy.copy(src_cell.fill)
    dst_cell.alignment = copy.copy(src_cell.alignment)
    dst_cell.number_format = src_cell.number_format


def set_num(cell, v, fmt):
    cell.value = v
    cell.number_format = fmt
    cell.alignment = copy.copy(cell.alignment)


def fill_sheet(sheet, title, asof_text, slices, opening=None, prev_year_label=None):
    """slices: dict name -> (recs) for 'lost', 'off', 'open'; opening: recs (додаток 2)."""
    ws = sheet
    two = opening is not None
    # Об'єднані клітинки нижче шапки openpyxl при вставці рядків не зсуває:
    # знімаємо їх і повертаємо після заповнення за підписами рядків.
    for rng in [str(m) for m in ws.merged_cells.ranges]:
        if ws[rng.split(":")[0]].row >= 8:
            ws.unmerge_cells(rng)
    ws["A3"] = title
    if two:
        ws["J4"] = asof_text
        ws["D5"] = f"{prev_year_label} (не прийнято рішень)"
    # колонки: додаток 1 — D..I, додаток 2 — D..K (D,E — перехідне)
    q_cols = {"open0": ("D", "E"), "lost": ("F", "G"), "off": ("H", "I"), "open": ("J", "K")} if two \
        else {"lost": ("D", "E"), "off": ("F", "G"), "open": ("H", "I")}
    note_col = "L" if two else "J"
    type_rows = {1: 8, 2: 11, 3: 14, 4: 17}       # рядки типів у порожній формі
    sub_tpl = {1: (9, 10), 2: (12, 13), 3: (15, 16), 4: (18, 19)}
    total_row, other_row = 20, 21
    # Усі рядки за назвами по типах — щоб знати, скільки підрядків вставляти.
    every = []
    for key in ("lost", "off", "open"):
        every += slices[key]
    if opening:
        every += opening
    names = by_name(every)
    per_type = {k: sorted([n for (kk, n) in names if kk == k], key=lambda x: x.lower()) for k in (1, 2, 3, 4)}
    # Вставляємо рядки знизу вгору, щоб номери верхніх не зсувалися.
    for k in (4, 3, 2, 1):
        need = len(per_type[k])
        have = 2
        r0, r1 = sub_tpl[k]
        if need > have:
            ws.insert_rows(r1 + 1, need - have)
            for i in range(need - have):
                for c in range(1, ws.max_column + 1):
                    copy_style(ws.cell(r1, c), ws.cell(r1 + 1 + i, c))
                ws.row_dimensions[r1 + 1 + i].height = ws.row_dimensions[r1].height
    # Після вставок перераховуємо позиції.
    def find_row(label):
        for r in range(8, ws.max_row + 1):            # рядок 7 — нумерація граф, не тип
            v = ws.cell(r, 1).value
            if v == label or (isinstance(v, (int, float)) and isinstance(label, (int, float)) and v == label):
                return r
        raise KeyError(label)

    def put(row, keyname, items, unit="од."):
        q, m = agg(items)
        c1, c2 = q_cols[keyname]
        if q or m:
            set_num(ws[f"{c1}{row}"], q, FMT_Q)
            set_num(ws[f"{c2}{row}"], m / MLN, FMT_M)

    def put_all(row, kind_filter, extra_note=""):
        for keyname, items in (("lost", slices["lost"]), ("off", slices["off"]), ("open", slices["open"])):
            put(row, keyname, [r for r in items if kind_filter(r)])
        if two:
            put(row, "open0", [r for r in opening if kind_filter(r)])
        if extra_note:
            ws[f"{note_col}{row}"] = extra_note

    total_items = lambda r: kind_of(r["name"]) in (1, 2, 3, 4)  # noqa: E731
    for k in (1, 2, 3, 4):
        trow = find_row(k)
        put_all(trow, lambda r, k=k: kind_of(r["name"]) == k)
        # підрядки
        sub = per_type[k]
        for i, name in enumerate(sub):
            row = trow + 1 + i
            ws.cell(row, 1).value = f"{k}.{i + 1}."
            ws.cell(row, 1).number_format = "@"
            ws.cell(row, 2).value = name
            ws.cell(row, 3).value = "од."
            put_all(row, lambda r, k=k, name=name: kind_of(r["name"]) == k and short(r["name"]) == name)
        # зайві порожні підрядки шаблону — «…» лишаємо як у формі, лише якщо підрядків менше двох
        for i in range(len(sub), 2):
            row = trow + 1 + i
            ws.cell(row, 1).value = f"{k}.{i + 1}."
            ws.cell(row, 2).value = "…"
    trow = find_row("Всього")
    put_all(trow, total_items)
    orow = find_row(5)
    put_all(orow, lambda r: kind_of(r["name"]) == 5)
    # кількість посуду в тоннах програма не знає: у формі лишається сума, одиниці — у примітці
    for keyname in q_cols:
        c1, _ = q_cols[keyname]
        ws[f"{c1}{orow}"].value = None
    u = {}
    for keyname, items in (("втрачено", slices["lost"]), ("списано", slices["off"]), ("не списано", slices["open"])):
        n = sum(r["qty"] for r in items if kind_of(r["name"]) == 5)
        if n:
            u[keyname] = n
    if two:
        n = sum(r["qty"] for r in opening if kind_of(r["name"]) == 5)
        if n:
            u["на початок"] = n
    if u:
        ws[f"{note_col}{orow}"] = "посуд, інвентар: " + ", ".join(f"{k} {v:g} од." for k, v in u.items())
    # примітки до типів: рапорти
    reps = sorted({r["report"] for r in every if r["report"]}, key=lambda x: (len(x), x))
    ws[f"{note_col}{find_row('Всього')}"] = ("рапорти №" + ", ".join(reps)) if reps else "втрат не було"
    tr = find_row("Всього")
    ws.merge_cells(f"A{tr}:C{tr}")
    for r in range(tr, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith("Начальник продовольчої служби"):
            ws.merge_cells(f"A{r}:{note_col if two else 'I'}{r}")
            break
    return ws


def year_slices(y):
    y0, y1 = f"{y}-01-01", f"{y}-12-31"
    end = min(y1, ASOF)
    return {"lost": [r for r in recs if lost_in(r, y0, end)],
            "off": [r for r in recs if off_in(r, y0, end)],
            "open": [r for r in recs if open_at(r, end)]}


ASOF = sys.argv[4] if len(sys.argv) > 4 else dt.date.today().isoformat()
UA_MONTHS = ["січня", "лютого", "березня", "квітня", "травня", "червня", "липня", "серпня", "вересня", "жовтня", "листопада", "грудня"]


def ua_date(iso):
    d = dt.date.fromisoformat(iso)
    return f"{d.day:02d} {UA_MONTHS[d.month - 1]} {d.year} року"


# --- 2022: додаток 1
s2022 = year_slices(2022)
t1 = fwb.copy_worksheet(T1)
t1.title = "2022"
fill_sheet(t1, "втрат та списання матеріально-технічних засобів номенклатури продовольчої служби за 2022 рік",
           "", s2022)
# --- 2024–2026: додаток 2
for y in (2024, 2025, 2026):
    s = year_slices(y)
    day0 = f"{y - 1}-12-31"
    opening = [r for r in recs if open_at(r, day0)]
    end = min(f"{y}-12-31", ASOF)
    t = fwb.copy_worksheet(T2)
    t.title = str(y)
    fill_sheet(t, f"втрат та списання матеріально-технічних засобів номенклатури продовольчої служби за {y} рік",
               f"станом на {ua_date(end)}", s, opening=opening, prev_year_label=f"{y - 1} рік")
    print(f"{y}: на початок {agg(opening)}, втрачено {agg(s['lost'])}, списано {agg(s['off'])}, не списано {agg(s['open'])}")
print(f"2022: втрачено {agg(s2022['lost'])}, списано {agg(s2022['off'])}, не списано {agg(s2022['open'])}")

# порядок аркушів: заповнені спершу, шаблони й пояснення — далі
order = ["2022", "2024", "2025", "2026", "Додаток 1", "Додаток 2", "Пояснення по заповненню додаткі"]
fwb._sheets = [fwb[n] for n in order if n in fwb.sheetnames] + [w for w in fwb._sheets if w.title not in order]
fwb.active = 0
fwb.save(out)
print("збережено:", out)

# --- контроль класифікації
print("\nкласифікація найменувань:")
for (k, name), items in sorted(by_name(recs).items(), key=lambda x: (x[0][0], x[0][1])):
    print(f"  тип {k}: {name} — {sum(r['qty'] for r in items):g} од.")
