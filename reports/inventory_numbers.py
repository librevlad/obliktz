# -*- coding: utf-8 -*-
"""Відомість інвентарних номерів — те, з чим ідуть наносити номери.

Аркуш на підрозділ, у ньому рядок на позицію: що це, скільки одиниць і які саме
номери на них наносять. Номери всередині позиції йдуть підряд, тому в рядку
стоїть діапазон, а не перелік — «10201/001 — 10201/012».

Де заводський номер відомий (кухні, причепи-цистерни), він друкується поруч:
на цій одиниці буде і власний інвентарний номер, і паспортний заводський, і
сплутати їх не можна.

Номери власні, не з ФЕС. В 1С один інвентарний номер накриває кілька різних
активів — скажімо, і кришки для гастроємностей, і морський контейнер, — тому
для пооб'єктного обліку вони не годяться.
"""
import sys
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook                                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, Side        # noqa: E402
from openpyxl.utils import get_column_letter                     # noqa: E402

from desktop.excel_names import save_book                        # noqa: E402

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = Font(name="Times New Roman", size=10, bold=True)
BODY = Font(name="Times New Roman", size=10)
TITLE = Font(name="Times New Roman", size=12, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLUMNS = [("№", 5), ("Код", 9), ("Найменування", 52), ("Од.", 6),
           ("К-сть", 8), ("Інвентарні номери", 26), ("Заводські номери", 30)]


@dataclass
class Row:
    no: int
    code: str
    name: str
    uom: str
    qty: int
    first: str
    last: str
    serials: str


@dataclass
class Sheet:
    subdivision: str
    as_of: str
    rows: list = field(default_factory=list)

    @property
    def total(self) -> int:
        return sum(r.qty for r in self.rows)


def collect(con, as_of: str) -> list[Sheet]:
    """Аркуші відомості: по одному на підрозділ, де є пронумеровані одиниці."""
    rows = con.execute("""
        SELECT s.name AS sub, s.sort AS sub_sort, n.code, n.name, u.code AS uom,
               i.inv_no, i.serial_no, i.chassis_no
          FROM instance i
          JOIN nomen n ON n.id = i.nomen_id
          JOIN uom u ON u.id = n.uom_id
          JOIN instance_assignment a ON a.instance_id = i.id
          JOIN subdivision s ON s.id = a.subdivision_id
         WHERE i.inv_no IS NOT NULL AND a.on_date <= ?
         ORDER BY s.sort, s.name, n.code, i.inv_no""", (as_of,)).fetchall()

    sheets, cur, group = [], None, None
    for r in rows:
        if cur is None or cur.subdivision != r["sub"]:
            cur = Sheet(subdivision=r["sub"], as_of=as_of)
            sheets.append(cur)
            group = None
        if group is None or group.code != r["code"]:
            group = Row(no=len(cur.rows) + 1, code=r["code"], name=r["name"],
                        uom=r["uom"], qty=0, first=r["inv_no"], last=r["inv_no"],
                        serials="")
            cur.rows.append(group)
        group.qty += 1
        group.last = r["inv_no"]
        mark = r["serial_no"] or (f"шасі {r['chassis_no']}" if r["chassis_no"] else "")
        if mark:
            tail = f"{r['inv_no'].rsplit('/', 1)[-1]} — {mark}"
            group.serials = f"{group.serials}; {tail}" if group.serials else tail
    return sheets


def write(sheets: list[Sheet], path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        # Назва аркуша в Excel не може містити : \\ / ? * [ ] і довша за 31 знак.
        title = sh.subdivision.replace("·", "-")[:31]
        for bad in ':\\/?*[]':
            title = title.replace(bad, "-")
        ws = wb.create_sheet(title or "підрозділ")
        ws.append([f"ВІДОМІСТЬ інвентарних номерів — {sh.subdivision}"])
        ws["A1"].font = TITLE
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
        ws.append([f"станом на {sh.as_of}; номери наносяться на кожну одиницю"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
        ws.append([])
        ws.append([c for c, _ in COLUMNS])
        for i, (_, w) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
            cell = ws.cell(row=4, column=i)
            cell.font, cell.alignment, cell.border = HEAD, CENTER, BOX
        for r in sh.rows:
            span = r.first if r.qty == 1 else f"{r.first} — {r.last}"
            ws.append([r.no, r.code, r.name, r.uom, r.qty, span, r.serials])
            for i in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=ws.max_row, column=i)
                cell.font, cell.border = BODY, BOX
                cell.alignment = LEFT if i in (3, 6, 7) else CENTER
        ws.append(["", "", "Разом одиниць", "", sh.total, "", ""])
        for i in range(1, len(COLUMNS) + 1):
            ws.cell(row=ws.max_row, column=i).font = HEAD
        # Сім граф ширші за книжковий аркуш: альбомний А4, одна сторінка завширшки.
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0       # заввишки — скільки треба, а не один аркуш
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.freeze_panes = "A5"
        ws.print_title_rows = "4:4"
    save_book(wb, path)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    from db.connect import open_db, close_db
    from db.paths import db_path

    con = open_db(db_path(), migrate=False)
    try:
        as_of = con.execute("SELECT MAX(doc_date) FROM document").fetchone()[0]
        sheets = collect(con, as_of)
    finally:
        close_db(con)
    # Туди ж, куди й решта вивантажень програми, а не в корінь теки даних.
    out = Path(db_path()).parent / "вивантаження" / f"Відомість інвентарних номерів на {as_of}.xlsx"
    out.parent.mkdir(exist_ok=True)
    write(sheets, out)
    print(out)
    print(f"  аркушів: {len(sheets)}, одиниць: {sum(s.total for s in sheets)}")
    for s in sheets[:10]:
        print(f"    {s.subdivision[:30]:30} позицій {len(s.rows):>3}, одиниць {s.total:>5}")
    if len(sheets) > 10:
        print(f"    … ще {len(sheets) - 10}")


if __name__ == "__main__":
    main()
