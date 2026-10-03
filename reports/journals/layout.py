# -*- coding: utf-8 -*-
"""Спільне для друкованих книг обліку (форми 13/14 і 46/47): стилі, оцінка висоти рядків, сторінка A4
альбомна з дзеркальними полями, титул, зміст, засвідчувальний напис.

Розкладка — та сама, що погодив власник для книг сухпаю й продуктів (tools/kp_book/jr у iprod):
сторінка книги = аркуш Excel, ліва/права сторінки розвороту, номер сторінки у зовнішньому верхньому
куті, Times New Roman, таблиці без сітки Excel."""
import math

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

TNR = "Times New Roman"
MONTHS = ["січня", "лютого", "березня", "квітня", "травня", "червня", "липня", "серпня", "вересня", "жовтня",
          "листопада", "грудня"]
APPX = {"13": ("Додаток 13", "(пункт 12 розділу ІV)"), "14": ("Додаток 14", "(пункт 12 розділу ІV)"),
        "46": ("Додаток 46", "(пункт 4 розділу V)"), "47": ("Додаток 47", "(пункт 4 розділу V)")}
INSTR = "до Інструкції з обліку військового майна у Збройних Силах України"

# сторінка A4 альбомна, пт
PAGE_W, PAGE_H = 842.0, 595.0
M_IN, M_OUT, M_TOP, M_BOT, M_HDR, M_FTR = 0.8, 0.4, 0.55, 0.4, 0.25, 0.2      # дюйми
AVAIL_W = PAGE_W - (M_IN + M_OUT) * 72
AVAIL_H = PAGE_H - (M_TOP + M_BOT) * 72

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NOB = Border()
CC = Alignment(horizontal="center", vertical="center", wrap_text=True)
LC = Alignment(horizontal="left", vertical="center", wrap_text=True)
RC = Alignment(horizontal="right", vertical="center", wrap_text=True)
LB = Alignment(horizontal="left", vertical="bottom", wrap_text=False)
CT = Alignment(horizontal="center", vertical="top", wrap_text=True)
_fonts = {}


def font(sz, b=False, i=False, u=None):
    k = (sz, b, i, u)
    if k not in _fonts:
        _fonts[k] = Font(name=TNR, size=sz, bold=b, italic=i, underline=u)
    return _fonts[k]


def col_px(w):
    return int(w * 7 + 5)


def col_pt(w):
    return col_px(w) * 0.75


def wrap_pt(w):
    """найвужча ширина колонки на друці (96 dpi, без відступу) — для оцінки переносу тексту"""
    return w * 7 * 0.75


def _cw(ch):
    if ch.isdigit():
        return 0.5
    if ch == " ":
        return 0.25
    if ch in ".,:;/()-–—\"'«»№’":
        return 0.33
    if ch.isupper():
        return 0.7
    return 0.48


def text_w(s, pt):
    return sum(_cw(ch) for ch in s) * pt


def n_lines(s, width_pt, pt):
    """рядків тексту при переносі за словами в клітинці ширини width_pt"""
    if s in (None, ""):
        return 1
    avail = max(width_pt - 5, 5)
    total = 0
    for para in str(s).split("\n"):
        lines, cur = 1, 0.0
        for wd in para.split(" "):
            ww = text_w(wd, pt)
            sp = text_w(" ", pt) if cur else 0
            if cur + sp + ww <= avail:
                cur += sp + ww
            elif cur == 0:
                k = int(math.ceil(ww / avail))
                lines += k - 1
                cur = ww - (k - 1) * avail
            else:
                lines += 1
                if ww > avail:
                    k = int(math.ceil(ww / avail))
                    lines += k - 1
                    cur = ww - (k - 1) * avail
                else:
                    cur = ww
        total += lines
    return total


def line_h(pt, lines=1):
    return round(lines * pt * 1.18 + 2.5, 1)


def date_uk(d):
    return "%02d.%02d.%04d" % (d.day, d.month, d.year)


def date_words(d):
    return "“%02d” %s %d року" % (d.day, MONTHS[d.month - 1], d.year)


QTY_DEC = 3          # знаків після коми в кількостях: облік ведеться в тисячних


def qty(v):
    """кількість для клітинки: ціле — int, дробове — точно, до 3 знаків"""
    if v is None:
        return None
    v = round(float(v), QTY_DEC)
    if abs(v - round(v)) < 1e-9:
        return int(round(v))
    return v


def fix_qty_formats(wb):
    """дробові кількості з форматом «Загальний» — формат з рівно потрібною кількістю знаків (вузька колонка у форматі
    «Загальний» округлює відображення, а «0.###» на укр. локалі лишає кому в цілих)"""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, float) and c.number_format == "General" and abs(v - round(v)) > 1e-9:
                    k = len(("%.*f" % (QTY_DEC, v)).rstrip("0").split(".")[1])
                    c.number_format = "0." + "0" * max(1, k)


def money(v):
    """ціна: щонайменше 2 знаки після коми, більше — як у документі"""
    x = round(float(v or 0), 4)
    s = "{:,.4f}".format(x)
    whole, frac = s.split(".")
    frac = frac.rstrip("0")
    frac = frac + "0" * (2 - len(frac)) if len(frac) < 2 else frac
    return whole.replace(",", " ") + "," + frac


def setup_page(ws, side, page_no=None, fit=True):
    """A4 альбомна; side: 'L' — ліва сторінка розвороту (корінець праворуч), 'R' — права (корінець ліворуч),
    номер сторінки — у зовнішньому верхньому куті"""
    ps = ws.page_setup
    ps.paperSize = 9
    ps.orientation = "landscape"
    if fit:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ps.fitToWidth = 1
        ps.fitToHeight = 1
    m = ws.page_margins
    if side == "L":
        m.left, m.right = M_OUT, M_IN
    else:
        m.left, m.right = M_IN, M_OUT
    m.top, m.bottom, m.header, m.footer = M_TOP, M_BOT, M_HDR, M_FTR
    if page_no is not None:
        hf = ws.oddHeader.left if side == "L" else ws.oddHeader.right
        hf.text = str(page_no)
        hf.size = 10
        hf.font = TNR
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90


def box(ws, r1, c1, r2, c2, value=None, f=None, al=CC, merge=True, border=True, nf=None):
    if merge and (r1 != r2 or c1 != c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(r1, c1)
    if value is not None:
        cell.value = value
    cell.font = f or font(8)
    cell.alignment = al
    if nf:
        cell.number_format = nf
    if border:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = BOX
    return cell


def widths(ws, ws_list):
    for i, w in enumerate(ws_list, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def heights(ws, hs, start=1):
    for i, h in enumerate(hs, start):
        ws.row_dimensions[i].height = h


# ---------- титул, зміст, засвідчення ----------

def title_sheet(wb, appx, lines_mid, sub_caption, started, finished, unit, extra=None):
    """титул книги за формою: праворуч угорі «Додаток N …», по центру назва книги, служба/склад, в/ч, дати"""
    ws = wb.create_sheet("Титул")
    setup_page(ws, "R", None)
    widths(ws, [14.0] * 10)
    ncol = 10
    hs = [13] * 40
    heights(ws, hs)
    a, p = APPX[appx]
    box(ws, 1, 6, 1, 10, a, font(11), Alignment(horizontal="left", vertical="center"), border=False)
    box(ws, 2, 6, 3, 10, INSTR, font(11), Alignment(horizontal="left", vertical="top", wrap_text=True), border=False)
    box(ws, 4, 6, 4, 10, p, font(11), Alignment(horizontal="left", vertical="center"), border=False)
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 14
    r = 9
    for txt, f, h in lines_mid:
        box(ws, r, 1, r, ncol, txt, f, CC, border=False)
        ws.row_dimensions[r].height = h
        r += 1
    r += 2
    # служба / склад
    box(ws, r, 2, r, ncol - 1, sub_caption[0], font(14, b=True), CC, border=False)
    ws.cell(r, 2).border = NOB
    for c in range(2, ncol):
        ws.cell(r, c).border = Border(bottom=THIN)
    ws.row_dimensions[r].height = 20
    box(ws, r + 1, 2, r + 1, ncol - 1, sub_caption[1], font(9), CT, border=False)
    r += 3
    box(ws, r, 3, r, ncol - 2, unit, font(14, b=True), CC, border=False)
    for c in range(3, ncol - 1):
        ws.cell(r, c).border = Border(bottom=THIN)
    ws.row_dimensions[r].height = 20
    box(ws, r + 1, 3, r + 1, ncol - 2, "(військова частина)", font(9), CT, border=False)
    r += 5
    box(ws, r, 6, r, ncol, "Розпочато " + (date_words(started) if started else "“____” ______________ 20___ року"),
        font(12), Alignment(horizontal="left", vertical="center"), border=False)
    box(ws, r + 2, 6, r + 2, ncol, "Закінчено " + (date_words(finished) if finished else "“____” ______________ 20___ року"),
        font(12), Alignment(horizontal="left", vertical="center"), border=False)
    ws.print_area = "A1:%s%d" % (get_column_letter(ncol), r + 3)
    return ws


def blank_sheet(wb, name, side):
    ws = wb.create_sheet(name)
    setup_page(ws, side, None)
    widths(ws, [10.0] * 10)
    ws.cell(1, 1).value = "."                     # порожній аркуш Excel не друкує — біла крапка
    ws.cell(1, 1).font = Font(name=TNR, size=6, color="FFFFFF")
    ws.print_area = "A1:J2"
    return ws


CONTENTS_GRP = [29.0, 8.0, 8.0]


def _contents_pack(entries, confirm):
    """розкладка змісту по сторінках: на сторінці три групи «найменування — сторінки», заповнення зверху вниз по першій
    групі, далі друга, третя; під таблицею останньої сторінки — підтвердження первинних записів"""
    name_w = wrap_pt(CONTENTS_GRP[0])
    cap = AVAIL_H - 22 - 6 - 52 - (150 if confirm else 10)
    pages = [[[]]]
    used = 0.0
    for e in entries:
        h = line_h(9, n_lines(e[0], name_w, 9))
        grp = pages[-1][-1]
        if used + h > cap and grp:
            if len(pages[-1]) < 3:
                pages[-1].append([])
            else:
                pages.append([[]])
            used = 0.0
        pages[-1][-1].append((e, h))
        used += h
    return pages


def contents_pages(entries, confirm=None):
    """скільки сторінок займе зміст"""
    return len(_contents_pack(entries, confirm))


def contents_sheet(wb, page_no, entries, confirm=None):
    """ЗМІСТ за формою: три групи «найменування — сторінки (початкова, наступні)» на сторінці; заповнюється зверху вниз
    по першій групі, далі друга, третя, далі наступна сторінка. entries — [(назва, початкова, наступні)]. confirm — рядки
    підтвердження первинних записів під таблицею останньої сторінки. -> кількість сторінок"""
    pages = _contents_pack(entries, confirm)
    for k, placed in enumerate(pages):
        pno = page_no + k
        ws = wb.create_sheet("Зміст" if k == 0 else "Зміст (%d)" % (k + 1))
        setup_page(ws, "R" if pno % 2 else "L", pno)
        widths(ws, CONTENTS_GRP * 3)
        ws.row_dimensions[1].height = 22
        box(ws, 1, 1, 1, 9, "ЗМІСТ" if k == 0 else "ЗМІСТ (продовження)", font(13, b=True), CC, border=False)
        ws.row_dimensions[2].height = 6
        for g in range(3):
            c = 1 + g * 3
            box(ws, 3, c, 4, c, "Найменування військового майна (індекс, номер креслення)", font(9), CC)
            box(ws, 3, c + 1, 3, c + 2, "Сторінки книги", font(9), CC)
            box(ws, 4, c + 1, 4, c + 1, "початкова", font(8), CC)
            box(ws, 4, c + 2, 4, c + 2, "наступні", font(8), CC)
        ws.row_dimensions[3].height = 26
        ws.row_dimensions[4].height = 26
        while len(placed) < 3:
            placed.append([])
        nrows = max(len(p) for p in placed) if entries else 1
        for i in range(nrows):
            r = 5 + i
            h = 12.0
            for g in range(3):
                c = 1 + g * 3
                if i < len(placed[g]):
                    (name, p0, p1), hh = placed[g][i]
                    h = max(h, hh)
                    box(ws, r, c, r, c, name, font(9), LC)
                    box(ws, r, c + 1, r, c + 1, p0, font(9), CC)
                    box(ws, r, c + 2, r, c + 2, p1, font(9), CC)
                else:
                    for q in range(3):
                        box(ws, r, c + q, r, c + q, None, font(9), CC)
            ws.row_dimensions[r].height = h
        r = 5 + nrows + 1
        if confirm and k == len(pages) - 1:
            ws.row_dimensions[r].height = 10
            r += 1
            for txt, kind in confirm:
                if kind == "text":
                    box(ws, r, 1, r, 9, txt, font(10), Alignment(horizontal="left", vertical="bottom", wrap_text=True),
                        border=False)
                    ws.row_dimensions[r].height = 16
                elif kind == "line":
                    for c in range(1, 10):
                        ws.cell(r, c).border = Border(bottom=THIN)
                    ws.row_dimensions[r].height = 18
                elif kind == "cap":
                    box(ws, r, 1, r, 9, txt, font(7), CT, border=False)
                    ws.row_dimensions[r].height = 10
                r += 1
        ws.print_area = "A1:I%d" % max(r, 6)
    return len(pages)


def inscription_sheet(wb, side):
    """засвідчувальний напис — на звороті останнього аркуша, не нумерується"""
    ws = wb.create_sheet("Засвідчення")
    setup_page(ws, side, None)
    widths(ws, [14.0] * 10)
    heights(ws, [13] * 30)
    r = 8
    box(ws, r, 1, r + 1, 10, "У цій книзі пронумеровано, прошнуровано та скріплено мастиковою печаткою "
        "______ (________________________________________) аркушів.", font(12), LC, border=False)
    ws.row_dimensions[r].height = 18
    ws.row_dimensions[r + 1].height = 18
    r += 4
    for _ in range(2):
        for c in range(1, 11):
            ws.cell(r, c).border = Border(bottom=THIN)
        ws.row_dimensions[r].height = 20
        box(ws, r + 1, 1, r + 1, 10, "(посада, військове звання, підпис, власне ім’я, прізвище)", font(8), CT, border=False)
        r += 3
    box(ws, r, 1, r, 10, "“____” ______________ 20___ року", font(12), Alignment(horizontal="left", vertical="center"),
        border=False)
    ws.print_area = "A1:J%d" % (r + 1)
    return ws
