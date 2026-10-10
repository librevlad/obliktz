# -*- coding: utf-8 -*-
"""Паперові книги обліку томами за рік (друк, двобічний, розворотами): форма 46/47 — картки позицій з
місцями (служба), форма 13/14 — книга кожного підрозділу, що тримає майно. Сторінка = аркуш Excel:
ліва/права сторінки розвороту, вкладиші на місця понад п'ять, титул, зміст, засвідчувальний напис.
Номер форми — чинний у році (до 2023 — 13/46, з 2024 — 14/47). Том починається перенесенням залишків
з попереднього тому з посиланням на його сторінки (`refs`); `paper_year` віддає `refs` тому — для
наступного року. Перенесено з книг сухпаю й продуктів (tools/kp_book/jr у iprod)."""
import datetime
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

from .layout import (AVAIL_H, CC, LC, THIN, blank_sheet, box, contents_pages, contents_sheet, fix_qty_formats, font,
                     heights, inscription_sheet, line_h, money, n_lines, qty, setup_page, title_sheet, widths, wrap_pt)
from .model import EPS, ROMAN, Period, card47, docs_of, opening, rows13

F8 = font(8)
F8B = font(8, b=True)
F7 = font(7)
NF_DATE = "dd.mm.yyyy"
SIG = [("(посада, військове звання, підпис, власне ім’я, прізвище особи, яка відповідає за стан обліку "
        "військового майна)"), "(посада, військове звання, підпис, власне ім’я, прізвище особи, яка безпосередньо "
       "веде облік військового майна)"]
SIG_H = [8, 15, 9, 15, 9]
CFG = dict(sub="технічні засоби продовольчої служби", service="Продовольча служба", file="ТЗ",
           group="Технічні засоби продовольчої служби")
# Книга ОП — посуд одноразового використання, миючі засоби й серветки: ті самі форми,
# свій підзаголовок і свій префікс файлів.
CFG_OP = dict(CFG, sub="одноразовий посуд, миючі засоби, серветки", file="ОП",
              group="Одноразовий посуд, миючі засоби, серветки")


def forms(year):
    return ("13", "46") if year <= 2023 else ("14", "47")


def confirm_block(per):
    return [("Первинні записи (перенесення залишків на %s з книги обліку за %s) підтверджую:"
             % (per.d1.strftime("%d.%m.%Y"), per.prev().label), "text"),
            ("", "line"), ("(начальник продовольчої служби: військове звання, підпис, власне ім’я, прізвище)", "cap"),
            ("", "line"), ("(начальник фінансово-економічної служби: військове звання, підпис, власне ім’я, прізвище)",
                           "cap"),
            ("", "line"), ("(особа, яка веде облік військового майна: посада, військове звання, підпис, власне ім’я, "
                           "прізвище)", "cap")]


def title_lines(kind, cfg, per=None):
    t = [("КНИГА", font(18, b=True), 26), ("обліку наявності та руху", font(15, b=True), 21)]
    t.append(("військового майна (склад, підрозділ)" if kind == "sklad" else "військового майна (служба забезпечення)",
              font(15, b=True), 21))
    t.append((cfg["sub"], font(13, i=True), 19))
    if per is not None and per.volume:
        t.append((per.volume, font(14, b=True), 24))
    return t


def paginate(rows, hfun, cap, tail):
    """поділ рядків на сторінки за висотою; останній рядок («Разом») — разом із підписами (tail)"""
    blocks, cur, used = [], [], 0.0
    for i, r in enumerate(rows):
        h = hfun(r)
        need = h + (tail if i == len(rows) - 1 else 0)
        if cur and used + need > cap:
            blocks.append(cur)
            cur, used = [], 0.0
        cur.append(r)
        used += h
    if cur:
        blocks.append(cur)
    return blocks


def sig_rows(ws, r, ncols, left=True):
    """підписи під підсумком (на лівій сторінці); на правій — порожні рядки тієї ж висоти"""
    for k, h in enumerate(SIG_H):
        ws.row_dimensions[r + k].height = h
    if left:
        for j, k in enumerate((1, 3)):
            for c in range(1, ncols + 1):
                ws.cell(r + k, c).border = Border(bottom=THIN)
            box(ws, r + k + 1, 1, r + k + 1, ncols, SIG[j], F7, Alignment(horizontal="center", vertical="top"),
                border=False)
    return r + len(SIG_H)


# ------------------------------------------------------------------ форма 46/47

L47_W = [7.6, 10.5, 16.0, 7.6, 10.5, 6.9, 6.9, 7.0] + [2.4] * 5 + [7.0] + [2.4] * 5 + [7.0] + [2.4] * 5
R47_W = ([8.0] + [3.0] * 5) * 5
H47 = [15, 10, 4, 13, 4, 12, 20, 18, 10, 10]


def h47(r):
    if r["kind"] == "total":
        return line_h(8, n_lines(r["typ"], sum(wrap_pt(w) for w in L47_W[1:5]), 8))
    n8 = max(n_lines(r["typ"], wrap_pt(L47_W[1]), 8), n_lines(r["party"], wrap_pt(L47_W[4]), 8),
             n_lines(r["ddate"] if isinstance(r["ddate"], str) else "", wrap_pt(L47_W[3]), 8))
    return max(line_h(8, n8), line_h(7, n_lines(r["num"], wrap_pt(L47_W[2]), 7)))


def CT_():
    return Alignment(horizontal="center", vertical="top")


def head47_common(ws, item, ncols, cap_row4):
    box(ws, 1, 1, 1, ncols, item["name"], font(11, b=True), CC, border=False)
    for c in range(1, ncols + 1):
        ws.cell(1, c).border = Border(bottom=THIN)
    box(ws, 2, 1, 2, ncols, "(найменування військового майна, індекс, номер креслення)", F7, CT_(), border=False)
    box(ws, 4, 1, 4, ncols, cap_row4, font(9), CC, border=False)


def page47_left(wb, name, page_no, item, place_hdr, blk, gstart, last, head):
    ws = wb.create_sheet(name)
    setup_page(ws, "L", page_no)
    widths(ws, L47_W)
    heights(ws, H47)
    n = len(L47_W)
    head47_common(ws, item, n, "Нормативний запас: мінімальний ________________,     максимальний ________________")
    for c, t in ((1, "Дата запису"), (2, "Найменування документа"), (3, "Номер документа"), (4, "Дата документа"),
                 (5, "Постачальник (одержувач)"), (6, "Надійшло"), (7, "Вибуло")):
        box(ws, 6, c, 9, c, t, F8, CC)
    box(ws, 6, 8, 7, 13, "Перебуває згідно з документами", F8, CC)
    box(ws, 6, 14, 6, 25, "У тому числі на складі (у підрозділах,", F8, CC)
    box(ws, 7, 14, 7, 19, "на складі", F8B, CC)
    box(ws, 7, 20, 7, 25, place_hdr[0] if place_hdr else None, F8B, CC)
    for c0 in (8, 14, 20):
        box(ws, 8, c0, 9, c0, "усього", F7, CC)
        box(ws, 8, c0 + 1, 8, c0 + 5, "з них за категоріями (сортами)", F7, CC)
        for k in range(5):
            box(ws, 9, c0 + 1 + k, 9, c0 + 1 + k, k + 1, F7, CC)
    for c in range(1, n + 1):
        box(ws, 10, c, 10, c, gstart + c - 1, F7, CC)
    r = 11
    cols_place = {head: 14}
    if place_hdr and place_hdr[0]:
        cols_place[place_hdr[0]] = 20
    for row in blk:
        h = h47(row)
        ws.row_dimensions[r].height = h
        for c in range(1, n + 1):
            box(ws, r, c, r, c, None, F8, CC)
        if row["kind"] == "total":
            box(ws, r, 1, r, 1, None, F8, CC)
            box(ws, r, 2, r, 5, row["typ"], F8B, LC)
        else:
            box(ws, r, 1, r, 1, row["date"], F8, CC, nf=NF_DATE)
            box(ws, r, 2, r, 2, row["typ"], F8, CC)
            box(ws, r, 3, r, 3, row["num"], F7, CC)
            box(ws, r, 4, r, 4, row["ddate"], F8, CC, nf=NF_DATE if isinstance(row["ddate"], datetime.date) else None)
            box(ws, r, 5, r, 5, row["party"], F8, CC)
        fb = F8B if row["kind"] == "total" else F8
        for c, key in ((6, "inq"), (7, "outq")):
            if row[key] is not None:
                ws.cell(r, c).value = qty(row[key])
                ws.cell(r, c).font = fb
        ws.cell(r, 8).value = qty(row["total"])
        ws.cell(r, 8).font = fb
        for p, c in cols_place.items():
            if p in row["places"]:
                ws.cell(r, c).value = qty(row["places"][p])
                ws.cell(r, c).font = fb
        r += 1
    if last:
        r = sig_rows(ws, r, n, left=True)
    ws.print_area = "A1:%s%d" % (get_column_letter(n), r - 1)
    return ws


def page47_right(wb, name, page_no, item, subs5, blk, gstart, last):
    ws = wb.create_sheet(name)
    setup_page(ws, "R", page_no)
    widths(ws, R47_W)
    heights(ws, H47)
    n = len(R47_W)
    it = item
    cap4 = "Код номенклатури  %s       Одиниця виміру  %s       Вміст дорогоцінних металів ________       " \
           "Ціна за одиницю  %s грн" % (it["code"], it["unit"], money(it["price"]))
    head47_common(ws, item, n, cap4)
    box(ws, 6, 1, 6, n, "військових частинах)", F8, CC)
    for b in range(5):
        c0 = 1 + b * 6
        box(ws, 7, c0, 7, c0 + 5, subs5[b] if b < len(subs5) else None, F8B, CC)
        box(ws, 8, c0, 9, c0, "усього", F7, CC)
        box(ws, 8, c0 + 1, 8, c0 + 5, "з них за категоріями (сортами)", F7, CC)
        for k in range(5):
            box(ws, 9, c0 + 1 + k, 9, c0 + 1 + k, k + 1, F7, CC)
    for c in range(1, n + 1):
        box(ws, 10, c, 10, c, gstart + c - 1, F7, CC)
    r = 11
    for row in blk:
        ws.row_dimensions[r].height = h47(row)
        fb = F8B if row["kind"] == "total" else F8
        for c in range(1, n + 1):
            box(ws, r, c, r, c, None, F8, CC)
        for b, p in enumerate(subs5):
            if p and p in row["places"]:
                ws.cell(r, 1 + b * 6).value = qty(row["places"][p])
                ws.cell(r, 1 + b * 6).font = fb
        r += 1
    if last:
        r = sig_rows(ws, r, n, left=False)
    ws.print_area = "A1:%s%d" % (get_column_letter(n), r - 1)
    return ws


def book47(D, per, refs_prev, cfg):
    """том форми 46/47 за період; -> (книга, info) або (None, None), якщо ні руху, ні залишків"""
    f13, f46 = forms(per.year)
    docs = docs_of(D, per)
    ob = opening(D, per)
    codes = {c for (p, c), q in ob.items() if abs(q) > EPS}
    for g in docs:
        for c, pl in g["lines"].items():
            if any(abs(q) > EPS for q in pl.values()):
                codes.add(c)
    if not codes:
        return None, None
    codes = sorted(codes, key=lambda c: D.items[c]["order"])
    cap = AVAIL_H - sum(H47) - 4
    tail = sum(SIG_H)
    cards = []
    confirm = confirm_block(per) if ob else None
    n_cp = contents_pages([("%s, код %s" % (D.items[c]["name"], c), 0, "") for c in codes], confirm)
    pad = (3 + n_cp) % 2 == 1                      # картка починається на лівій (парній) сторінці
    nxt = 3 + n_cp + (1 if pad else 0)
    for code in codes:
        op = {p: q for (p, c), q in ob.items() if c == code}
        carry = (None, refs_prev.get("47", {}).get(str(code)))
        places, rows = card47(D, per, code, docs, op, carry)
        subs = places[1:]
        right_sets = [subs[1 + 5 * k:1 + 5 * (k + 1)] for k in range(max(1, int(math.ceil((len(subs) - 1) / 5.0))))] \
            if len(subs) > 1 else [[]]
        if (1 + len(right_sets)) % 2:
            right_sets.append([])
        blocks = paginate(rows, h47, cap, tail)
        pages = []
        for bi, blk in enumerate(blocks):
            pages.append(("L", nxt, bi, blk))
            nxt += 1
            for k, rs in enumerate(right_sets):
                pages.append(("R", nxt, bi, blk, rs, k))
                nxt += 1
        last_L = [p[1] for p in pages if p[0] == "L"][-1]
        cards.append(dict(code=code, item=D.items[code], places=places, rows=rows, pages=pages, first=pages[0][1],
                          last=pages[-1][1], last_L=last_L, sub1=subs[:1], nblocks=len(blocks)))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    started = per.d1 if ob else min(g["date"] for g in docs)
    title_sheet(wb, f46, title_lines("служба", cfg, per), (cfg["service"], "(служба забезпечення)"), started, per.d2,
                D.unit_of(D.head, per.year))
    blank_sheet(wb, "с.2", "L")
    entries = []
    for cd in cards:
        it = cd["item"]
        nx = "" if cd["last"] == cd["first"] else (str(cd["first"] + 1) if cd["last"] == cd["first"] + 1 else
                                                   "%d–%d" % (cd["first"] + 1, cd["last"]))
        entries.append(("%s, код %s" % (it["name"], it["code"]), cd["first"], nx))
    contents_sheet(wb, 3, entries, confirm)
    if pad:
        blank_sheet(wb, "с.%d" % (3 + n_cp), "R")
    for cd in cards:
        it = cd["item"]
        nb = cd["nblocks"]
        for pg in cd["pages"]:
            if pg[0] == "L":
                _, pno, bi, blk = pg
                page47_left(wb, "с.%d %s Л" % (pno, it["code"]), pno, it, cd["sub1"], blk, 1, bi == nb - 1, D.head)
            else:
                _, pno, bi, blk, rs, k = pg
                page47_right(wb, "с.%d %s %s" % (pno, it["code"], "П" if k == 0 else "В%d" % k), pno, it, rs, blk,
                             26 + 30 * k, bi == nb - 1)
    inscription_sheet(wb, "L")
    info = dict(pages=nxt - 1, refs={str(cd["code"]): cd["last_L"] for cd in cards},
                cards=[dict(code=cd["code"], first=cd["first"], last=cd["last"], places=cd["places"],
                            closing=cd["rows"][-1]["places"], total=cd["rows"][-1]["total"]) for cd in cards])
    return wb, info


# ------------------------------------------------------------------ форма 13/14 (підрозділ)

L13_W = [8.3, 11.0, 26.0, 11.5] + [6.4] * 12
R13_W = [6.4] * 18
H13_FIX = [15, 10, 4, 12, None, 11, 11, 11, 11, 11, 11, 11, 10]


def h13(r):
    if r["kind"] == "total":
        return line_h(8, 1)
    n = max(n_lines(r["typ"], wrap_pt(L13_W[1]), 8), n_lines(r["num"], wrap_pt(L13_W[2]), 8),
            n_lines(r["party"], wrap_pt(L13_W[3]), 8))
    return line_h(8, n)


def head13(ws, items, ncols, c0, gstart, name_h, left, group):
    box(ws, 1, 1, 1, ncols, group, font(11, b=True), CC, border=False)
    for c in range(1, ncols + 1):
        ws.cell(1, c).border = Border(bottom=THIN)
    box(ws, 2, 1, 2, ncols, "(найменування військового майна, індекс, номер креслення)", F7, CT_(), border=False)
    if left:
        for c, t in ((1, "Дата запису"), (2, "Найменування документа"), (3, "Номер і дата документа"),
                     (4, "Постачальник (одержувач)")):
            box(ws, 4, c, 12, c, t, F8, CC)
    box(ws, 4, c0, 4, ncols, "Найменування військового майна та категорія", F8, CC)
    box(ws, 6, c0, 6, ncols, "Код номенклатури", F8, CC)
    box(ws, 8, c0, 8, ncols, "Одиниця виміру", F8, CC)
    box(ws, 10, c0, 10, ncols, "Ціна за одиницю", F8, CC)
    for k in range((ncols - c0 + 1) // 3):
        c = c0 + 3 * k
        it = items[k] if k < len(items) else None
        box(ws, 5, c, 5, c + 2, it["name"] if it else None, F8, CC)
        box(ws, 7, c, 7, c + 2, it["code"] if it else None, F8, CC)
        box(ws, 9, c, 9, c + 2, it["unit"] if it else None, F8, CC)
        box(ws, 11, c, 11, c + 2, money(it["price"]) if it else None, F8, CC)
        for j, t in enumerate(("надійшло", "вибуло", "становить")):
            box(ws, 12, c + j, 12, c + j, t, F7, CC)
    for c in range(1, ncols + 1):
        box(ws, 13, c, 13, c, gstart + c - 1, F7, CC)
    ws.row_dimensions[5].height = name_h


def page13(wb, name, page_no, side, items, blk, gstart, last, name_h, group):
    ws = wb.create_sheet(name)
    setup_page(ws, side, page_no)
    W = L13_W if side == "L" else R13_W
    widths(ws, W)
    heights(ws, [h if h else name_h for h in H13_FIX])
    n = len(W)
    c0 = 5 if side == "L" else 1
    head13(ws, items, n, c0, gstart, name_h, side == "L", group)
    r = 14
    for row in blk:
        ws.row_dimensions[r].height = h13(row)
        fb = F8B if row["kind"] == "total" else F8
        for c in range(1, n + 1):
            box(ws, r, c, r, c, None, F8, CC)
        if side == "L":
            if row["kind"] == "total":
                box(ws, r, 2, r, 4, row["typ"], F8B, LC)
            else:
                box(ws, r, 1, r, 1, row["date"], F8, CC, nf=NF_DATE)
                box(ws, r, 2, r, 2, row["typ"], F8, CC)
                box(ws, r, 3, r, 3, row["num"], F8, CC)
                box(ws, r, 4, r, 4, row["party"], F8, CC)
        for k, it in enumerate(items):
            v = row["items"].get(it["code"])
            if v is None:
                continue
            c = c0 + 3 * k
            for j in range(3):
                if v[j] is not None:
                    ws.cell(r, c + j).value = qty(v[j])
                    ws.cell(r, c + j).font = fb
        r += 1
    if last:
        r = sig_rows(ws, r, n, left=(side == "L"))
    ws.print_area = "A1:%s%d" % (get_column_letter(n), r - 1)
    return ws


def book13(D, per, refs_prev, place, cfg):
    """том форми 13/14 підрозділу `place` за період; -> (книга, info)"""
    f13, f46 = forms(per.year)
    docs = [g for g in docs_of(D, per) if any(abs(pl.get(place, 0.0)) > EPS for pl in g["lines"].values())]
    ob = {c: q for (p, c), q in opening(D, per).items() if p == place}
    codes = {c for c, q in ob.items() if abs(q) > EPS}
    for g in docs:
        for c, pl in g["lines"].items():
            if abs(pl.get(place, 0.0)) > EPS:
                codes.add(c)
    codes = sorted(codes, key=lambda c: D.items[c]["order"])
    groups = [codes[i:i + 10] for i in range(0, len(codes), 10)]
    rp = refs_prev.get("13", {}).get(place, {})
    sub_name = "Склад продовольчої служби" if place == D.head else place
    cap_base = AVAIL_H - 4
    tail = sum(SIG_H)
    sections = []
    confirm = confirm_block(per) if ob else None
    n_cp = contents_pages([("%s, код %s" % (D.items[c]["name"], c), 0, "") for c in codes], confirm)
    pad = (3 + n_cp) % 2 == 1
    nxt = 3 + n_cp + (1 if pad else 0)
    for gi, grp in enumerate(groups):
        items = [D.items[c] for c in grp]
        pages_prev = sorted({rp.get(str(c)) for c in grp if abs(ob.get(c, 0.0)) > EPS and rp.get(str(c))})
        carry = ", ".join(str(p) for p in pages_prev) if pages_prev else None
        rows, closing = rows13(D, per, grp, docs, ob, carry, place)
        name_h = max(line_h(8, n_lines(it["name"], 3 * wrap_pt(6.4), 8)) for it in items)
        name_h = max(name_h, 12)
        cap = cap_base - sum(h for h in H13_FIX if h) - name_h
        blocks = paginate(rows, h13, cap, tail)
        pages = []
        for bi, blk in enumerate(blocks):
            pages.append(("L", nxt, bi, blk))
            pages.append(("R", nxt + 1, bi, blk))
            nxt += 2
        sections.append(dict(codes=grp, items=items, rows=rows, pages=pages, first=pages[0][1], last=pages[-1][1],
                             last_L=pages[-2][1], name_h=name_h, nblocks=len(blocks), closing=closing))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    started = per.d1 if ob else min(g["date"] for g in docs)
    title_sheet(wb, f13, title_lines("sklad", cfg, per),
                (sub_name, "(служба забезпечення, підрозділ, склад, сховище)"), started, per.d2,
                D.unit_of(place, per.year))
    blank_sheet(wb, "с.2", "L")
    entries = []
    for s in sections:
        nx = "%d–%d" % (s["first"] + 1, s["last"]) if s["last"] > s["first"] + 1 else str(s["last"])
        for it in s["items"]:
            entries.append(("%s, код %s" % (it["name"], it["code"]), s["first"], nx))
    contents_sheet(wb, 3, entries, confirm)
    if pad:
        blank_sheet(wb, "с.%d" % (3 + n_cp), "R")
    for si, s in enumerate(sections):
        nb = s["nblocks"]
        for pg in s["pages"]:
            side, pno, bi, blk = pg
            its = s["items"][:4] if side == "L" else s["items"][4:10]
            page13(wb, "с.%d гр.%d %s" % (pno, si + 1, "Л" if side == "L" else "П"), pno, side, its, blk,
                   1 if side == "L" else 17, bi == nb - 1, s["name_h"], cfg["group"])
    inscription_sheet(wb, "L")
    refs = {}
    for s in sections:
        for c in s["codes"]:
            refs[str(c)] = s["last_L"]
    info = dict(pages=nxt - 1, refs=refs, sections=[dict(codes=s["codes"], first=s["first"], last=s["last"],
                                                         closing=s["closing"]) for s in sections])
    return wb, info


# ------------------------------------------------------------------ PDF

def export_pdf(paths, timeout=3600):
    """PDF кожної книги через Excel (ExportAsFixedFormat); свій Excel закривається за PID."""
    import subprocess
    import threading
    import pythoncom
    import win32com.client
    import win32process
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    pid = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
    watch = threading.Timer(timeout, lambda: subprocess.call(["taskkill", "/F", "/PID", str(pid)]))
    watch.start()
    res = {}
    try:
        xl.Visible = False
        xl.DisplayAlerts = False
        for p in paths:
            wb = xl.Workbooks.Open(os.path.abspath(p), 0, True)
            try:
                multi = []
                for ws in wb.Worksheets:
                    n = ws.PageSetup.Pages.Count
                    if n != 1:
                        multi.append((ws.Name, n))
                pdf = os.path.splitext(p)[0] + ".pdf"
                wb.ExportAsFixedFormat(0, os.path.abspath(pdf), 0, True, False)
                res[p] = dict(sheets=wb.Worksheets.Count, multi=multi, pdf=pdf)
            finally:
                wb.Close(False)
    finally:
        try:
            xl.Quit()
        except Exception:                                              # noqa: BLE001
            pass
        watch.cancel()
    return res


BAD_CH = '\\/:*?"<>|'


def file_place(p):
    """назва місця для імені файла"""
    return "".join("_" if ch in BAD_CH else ch for ch in p.replace(" · ", " - "))


def paper_year(D, year, out_dir, refs_prev, cfg=CFG):
    """паперові книги року: 46/47 служби й 13/14 кожному підрозділу з майном; книга без руху й залишків
    не створюється. refs_prev — сторінки книг попереднього року.
    -> ([(шлях, вид, місце)], refs цього року)"""
    f13, f46 = forms(year)
    os.makedirs(out_dir, exist_ok=True)
    per = Period(year)
    made = []
    nxt = {"47": {}, "13": {}}
    wb47, i47 = book47(D, per, refs_prev, cfg)
    if wb47 is None:
        return made, refs_prev
    p47 = os.path.join(out_dir, "%s %s %d.xlsx" % (f46, cfg["file"], year))
    fix_qty_formats(wb47)
    wb47.save(p47)
    made.append((p47, "47", None))
    nxt["47"] = i47["refs"]
    for place in D.warehouses_of(per):
        wb13, i13 = book13(D, per, refs_prev, place, cfg)
        p13 = os.path.join(out_dir, "%s %s %s %d.xlsx" % (f13, cfg["file"], file_place(place), year))
        fix_qty_formats(wb13)
        wb13.save(p13)
        made.append((p13, "13", place))
        nxt["13"][place] = i13["refs"]
    return made, nxt
