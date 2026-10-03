# -*- coding: utf-8 -*-
"""Електронні книги обліку поточного року (форми 14 і 47) для подальшого ведення: залишки — формулами
(дописаний рядок рахується сам), перенесення залишків — з книги минулого року (сторінки з `refs`),
узагальнююча відомість (додаток 1) за період із клітинок «з/по». Перенесено з книг сухпаю й продуктів
(tools/kp_book/jr у iprod); перерахунок у Excel — `check`, окремо."""
import collections
import datetime
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter as L

from .layout import APPX, CC, CT, INSTR, THIN, box, date_words, font, money, qty, widths
from .model import EPS, Period, card47, docs_of, opening, rows13
from .paper import CFG, file_place

F8 = font(8)
F8B = font(8, b=True)
F7 = font(7)
F9 = font(9)
NF_DATE = "dd.mm.yyyy"
NF_HIDE0 = "General;-General;;@"      # нуль у графах місць не показувати
MAXR = 3000          # запас рядків у діапазонах формул відомості понад заповнені
SPARE14 = 3          # вільних груп граф для нових позицій у книзі підрозділу
SPARE_ROWS = 150     # порожніх рядків з формулами під дописування


def dtxt(ref):
    """дата клітинки текстом дд.мм.рррр без TEXT(…,"dd.mm.yyyy") (на укр. локалі коди формату інші)"""
    return 'TEXT(DAY(%s),"00")&"."&TEXT(MONTH(%s),"00")&"."&YEAR(%s)' % (ref, ref, ref)


def title(wb, appx, kind, sub, sub_cap, started, unit, cfg):
    ws = wb.create_sheet("Титул")
    widths(ws, [30.0, 8.0, 8.0, 30.0, 8.0, 8.0, 30.0, 8.0, 8.0])
    a, p = APPX[appx]
    box(ws, 1, 7, 1, 9, a, font(11), Alignment(horizontal="left"), border=False)
    box(ws, 2, 7, 3, 9, INSTR, font(11), Alignment(horizontal="left", vertical="top", wrap_text=True), border=False)
    box(ws, 4, 7, 4, 9, p, font(11), Alignment(horizontal="left"), border=False)
    r = 7
    for t, f in (("КНИГА", font(16, b=True)), ("обліку наявності та руху", font(14, b=True)),
                 ("військового майна (склад, підрозділ)" if kind == "sklad" else "військового майна (служба забезпечення)",
                  font(14, b=True)), (cfg["sub"], font(12, i=True))):
        box(ws, r, 1, r, 9, t, f, CC, border=False)
        ws.row_dimensions[r].height = 20
        r += 1
    r += 1
    box(ws, r, 2, r, 8, sub, font(13, b=True), CC, border=False)
    for c in range(2, 9):
        ws.cell(r, c).border = Border(bottom=THIN)
    box(ws, r + 1, 2, r + 1, 8, sub_cap, font(8), CT, border=False)
    r += 3
    box(ws, r, 3, r, 7, unit, font(13, b=True), CC, border=False)
    for c in range(3, 8):
        ws.cell(r, c).border = Border(bottom=THIN)
    box(ws, r + 1, 3, r + 1, 7, "(військова частина)", font(8), CT, border=False)
    r += 3
    box(ws, r, 6, r, 9, "Розпочато " + date_words(started), font(12), Alignment(horizontal="left"), border=False)
    box(ws, r + 1, 6, r + 1, 9, "Закінчено “____” ______________ 20___ року", font(12), Alignment(horizontal="left"),
        border=False)
    return ws, r + 3


def contents(ws, r, entries):
    """ЗМІСТ: найменування — аркуш книги (посилання)"""
    box(ws, r, 1, r, 9, "ЗМІСТ", font(12, b=True), CC, border=False)
    r += 1
    for g in range(3):
        c = 1 + g * 3
        box(ws, r, c, r + 1, c, "Найменування військового майна (індекс, номер креслення)", F8, CC)
        box(ws, r, c + 1, r, c + 2, "Сторінки книги", F8, CC)
        box(ws, r + 1, c + 1, r + 1, c + 1, "початкова", F7, CC)
        box(ws, r + 1, c + 2, r + 1, c + 2, "наступні", F7, CC)
    ws.row_dimensions[r].height = 24
    r += 2
    for name, sheet in entries:
        box(ws, r, 1, r, 1, name, F9, Alignment(horizontal="left", vertical="center", wrap_text=True))
        cell = box(ws, r, 2, r, 2, sheet, F9, CC)
        if sheet:
            cell.hyperlink = "#'%s'!A1" % sheet
            cell.font = Font(name="Times New Roman", size=9, color="0563C1", underline="single")
        box(ws, r, 3, r, 3, None, F9, CC)
        for c in range(4, 10):
            box(ws, r, c, r, c, None, F9, CC)
        r += 1
    return r


def sig_block(ws, r, ncols):
    for cap in ("(посада, військове звання, підпис, власне ім’я, прізвище посадової (службової) особи, яка відповідає "
                "за стан обліку військового майна, або проводить звірку)",
                "(посада, військове звання, підпис, власне ім’я, прізвище посадової (службової) особи, яка "
                "безпосередньо веде його облік, або з якою проводиться звірка)"):
        ws.row_dimensions[r].height = 22
        for c in range(1, ncols + 1):
            ws.cell(r, c).border = Border(bottom=THIN)
        box(ws, r + 1, 1, r + 1, ncols, cap, F7, CT, border=False)
        ws.row_dimensions[r + 1].height = 20
        r += 2
    return r


def summary_sheet(wb, book_name, lines, p_from, p_to):
    """узагальнююча відомість (додаток 1): lines — [(найменування, од., формула надійшло, вибуло, наявне)];
    період — клітинки C4/E4"""
    ws = wb.create_sheet("Узагальнююча відомість")
    widths(ws, [6.0, 46.0, 10.0, 16.0, 16.0, 16.0])
    box(ws, 1, 1, 1, 6, "Узагальнююча відомість №_________", font(12, b=True), CC, border=False)
    box(ws, 2, 1, 2, 6, "обліку військового майна в електронній формі", font(12, b=True), CC, border=False)
    box(ws, 4, 1, 4, 2, "Звітний період (для розрахунку):  з", F9, Alignment(horizontal="right"), border=False)
    box(ws, 4, 3, 4, 3, p_from, F9, CC, nf=NF_DATE)
    box(ws, 4, 4, 4, 4, "по", F9, CC, border=False)
    box(ws, 4, 5, 4, 5, p_to, F9, CC, nf=NF_DATE)
    box(ws, 6, 1, 6, 2, "Найменування книги обліку військового майна", F8, CC)
    box(ws, 6, 3, 6, 4, "Реквізити книги обліку військового майна", F8, CC)
    box(ws, 6, 5, 6, 5, "Звітний період", F8, CC)
    box(ws, 6, 6, 6, 6, "Дата складання узагальнюючої відомості", F8, CC)
    ws.row_dimensions[6].height = 26
    box(ws, 7, 1, 7, 2, book_name, F9, CC)
    box(ws, 7, 3, 7, 4, None, F9, CC)
    box(ws, 7, 5, 7, 5, '=%s&" – "&%s' % (dtxt("$C$4"), dtxt("$E$4")), F9, CC)
    box(ws, 7, 6, 7, 6, None, F9, CC)
    ws.row_dimensions[7].height = 40
    hdr = ["№ з/п", "Найменування військового майна", "Одиниця виміру", "Кількість військового майна, що надійшло",
           "Кількість військового майна, що вибуло", "Кількість наявного військового майна"]
    for i, t in enumerate(hdr, 1):
        box(ws, 9, i, 9, i, t, F8, CC)
        box(ws, 10, i, 10, i, i, F7, CC)
    ws.row_dimensions[9].height = 36
    r = 11
    for k, (name, unit, fn, fv, fs) in enumerate(lines, 1):
        box(ws, r, 1, r, 1, k, F9, CC)
        box(ws, r, 2, r, 2, name, F9, Alignment(horizontal="left", vertical="center", wrap_text=True))
        box(ws, r, 3, r, 3, unit, F9, CC)
        box(ws, r, 4, r, 4, fn, F9, CC)
        box(ws, r, 5, r, 5, fv, F9, CC)
        box(ws, r, 6, r, 6, fs, F9, CC)
        r += 1
    r += 2
    sig_block(ws, r, 6)
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


# ------------------------------------------------------------------ форма 14, підрозділ

def book14(D, place, sub_name, refs, per, p_from, p_to, cfg):
    docs = [g for g in docs_of(D, per) if any(abs(pl.get(place, 0.0)) > EPS for pl in g["lines"].values())]
    ob = {c: q for (p, c), q in opening(D, per).items() if p == place}
    codes = {c for c, q in ob.items() if abs(q) > EPS}
    for g in docs:
        for c, pl in g["lines"].items():
            if abs(pl.get(place, 0.0)) > EPS:
                codes.add(c)
    codes = sorted(codes, key=lambda c: D.items[c]["order"])
    rp = refs.get("13", {}).get(place, {})
    pages_prev = sorted({rp.get(str(c)) for c in codes if abs(ob.get(c, 0.0)) > EPS and rp.get(str(c))})
    carry = ", ".join(str(p) for p in pages_prev) if pages_prev else None
    rows, closing = rows13(D, per, codes, docs, ob, carry, place)
    rows = [r for r in rows if r["kind"] != "total"]
    items = [D.items[c] for c in codes]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_t, r_t = title(wb, "14", "sklad", sub_name, "(служба забезпечення, підрозділ, склад, сховище)",
                      per.d1, D.unit_of(place, per.year), cfg)
    ws = wb.create_sheet(cfg["sheet"])
    ngr = len(items) + SPARE14
    ncols = 4 + 3 * ngr
    widths(ws, [9.0, 13.0, 27.0, 17.0] + [7.0] * (3 * ngr))
    box(ws, 1, 1, 1, ncols, cfg["group"], font(11, b=True), Alignment(horizontal="left", vertical="center"),
        border=False)
    box(ws, 2, 1, 2, ncols, "(найменування військового майна, індекс, номер креслення)", F7,
        Alignment(horizontal="left", vertical="top"), border=False)
    for c, t in ((1, "Дата запису"), (2, "Найменування документа"), (3, "Номер і дата документа"),
                 (4, "Постачальник (одержувач)")):
        box(ws, 4, c, 12, c, t, F8, CC)
    box(ws, 4, 5, 4, ncols, "Найменування військового майна та категорія", F8, CC)
    box(ws, 6, 5, 6, ncols, "Код номенклатури", F8, CC)
    box(ws, 8, 5, 8, ncols, "Одиниця виміру", F8, CC)
    box(ws, 10, 5, 10, ncols, "Ціна за одиницю", F8, CC)
    for k in range(ngr):
        c = 5 + 3 * k
        it = items[k] if k < len(items) else None
        box(ws, 5, c, 5, c + 2, it["name"] if it else None, F8, CC)
        box(ws, 7, c, 7, c + 2, it["code"] if it else None, F8, CC)
        box(ws, 9, c, 9, c + 2, it["unit"] if it else None, F8, CC)
        box(ws, 11, c, 11, c + 2, it["price"] if it else None, F8, CC, nf="#,##0.00")
        for j, t in enumerate(("надійшло", "вибуло", "становить")):
            box(ws, 12, c + j, 12, c + j, t, F7, CC)
    for c in range(1, ncols + 1):
        box(ws, 13, c, 13, c, c, F7, CC)
    ws.row_dimensions[5].height = 64
    for rr in range(6, 13):
        ws.row_dimensions[rr].height = 13
    r0 = 14
    has_carry = bool(rows) and rows[0]["kind"] == "carry"
    first_op = r0 + 1 if has_carry else r0
    r = r0
    expect = {}
    for row in rows:
        ws.row_dimensions[r].height = 13
        for c in range(1, ncols + 1):
            box(ws, r, c, r, c, None, F8, CC)
        ws.cell(r, 1).value = row["date"]
        ws.cell(r, 1).number_format = NF_DATE
        ws.cell(r, 2).value = row["typ"]
        ws.cell(r, 3).value = row["num"]
        ws.cell(r, 4).value = row["party"]
        for k, it in enumerate(items):
            c = 5 + 3 * k
            v = row["items"].get(it["code"])
            if row["kind"] == "carry":
                if v is not None:
                    ws.cell(r, c + 2).value = qty(v[2])
                continue
            if v is None:
                continue
            if v[0] is not None:
                ws.cell(r, c).value = qty(v[0])
            if v[1] is not None:
                ws.cell(r, c + 1).value = qty(v[1])
            expect[(r, it["code"])] = v[2]
        r += 1
    last = r - 1
    for rr in range(last + 1, last + 1 + SPARE_ROWS):
        ws.row_dimensions[rr].height = 13
        for c in range(1, ncols + 1):
            box(ws, rr, c, rr, c, None, F8, CC)
        ws.cell(rr, 1).number_format = NF_DATE
    # «становить» формулою для всіх рядків операцій і запасу рядків під дописування
    for rr in range(first_op, last + 1 + SPARE_ROWS):
        for k in range(ngr):
            c = 5 + 3 * k
            n, v, s = L(c), L(c + 1), L(c + 2)
            base = "$%s$%d+" % (s, r0) if has_carry else ""
            ws.cell(rr, c + 2).value = '=IF(COUNT(%s%d,%s%d)=0,"",%sSUM(%s$%d:%s%d)-SUM(%s$%d:%s%d))' % (
                n, rr, v, rr, base, n, first_op, n, rr, v, first_op, v, rr)
    ws.freeze_panes = "E14"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:13"
    ws.print_title_cols = "A:D"
    entries = [("%s, код %s" % (it["name"], it["code"]), cfg["sheet"]) for it in items]
    contents(ws_t, r_t, entries)
    sname = "'%s'" % cfg["sheet"]
    maxr = last + SPARE_ROWS + MAXR
    lines = []
    for k, it in enumerate(items):
        c = 5 + 3 * k
        n, v, s = L(c), L(c + 1), L(c + 2)
        rng = lambda col: "%s!$%s$%d:$%s$%d" % (sname, col, first_op, col, maxr)       # noqa: E731
        a = "%s!$A$%d:$A$%d" % (sname, first_op, maxr)
        fn = '=SUMIFS(%s,%s,">="&$C$4,%s,"<="&$E$4)' % (rng(n), a, a)
        fv = '=SUMIFS(%s,%s,">="&$C$4,%s,"<="&$E$4)' % (rng(v), a, a)
        base = "%s!$%s$%d+" % (sname, s, r0) if has_carry else ""
        fs = '=%sSUMIFS(%s,%s,"<="&$E$4)-SUMIFS(%s,%s,"<="&$E$4)' % (base, rng(n), a, rng(v), a)
        lines.append(("%s, код %s" % (it["name"], it["code"]), it["unit"], fn, fv, fs))
    summary_sheet(wb, "Книга обліку наявності та руху військового майна (склад, підрозділ) — %s, %s" % (sub_name, cfg["sub"]),
                  lines, p_from, p_to)
    info = dict(codes=codes, closing={str(c): q for c, q in closing.items()}, rows=len(rows), carry=carry,
                expect={"%d|%s" % k: v for k, v in expect.items()}, first_op=first_op, r0=r0, has_carry=has_carry)
    return wb, info


# ------------------------------------------------------------------ форма 47

def book47e(D, refs, per, p_from, p_to, cfg):
    docs = docs_of(D, per)
    ob = opening(D, per)
    codes = {c for (p, c), q in ob.items() if abs(q) > EPS}
    act_places = {p for (p, c), q in ob.items() if abs(q) > EPS}
    for g in docs:
        for c, pl in g["lines"].items():
            for p, q in pl.items():
                if abs(q) > EPS:
                    codes.add(c)
                    act_places.add(p)
    codes = sorted(codes, key=lambda c: D.items[c]["order"])
    subs = sorted((p for p in act_places if p != D.head), key=D.place_key)
    places = [D.head] + subs
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_t, r_t = title(wb, "47", "служба", cfg["service"], "(служба забезпечення)", per.d1, D.unit_of(D.head, per.year), cfg)
    ncols = 13 + 6 * len(places)
    pc = {p: 14 + 6 * i for i, p in enumerate(places)}
    info = {}
    entries = []
    lines = []

    def card_sheet(name, it, rows):
        ws = wb.create_sheet(name)
        widths(ws, [9.5, 13.0, 17.0, 9.5, 17.0, 8.0, 8.0, 8.5] + [3.0] * 5 + ([8.5] + [3.0] * 5) * len(places))
        box(ws, 1, 1, 1, 12, it["name"] if it else None, font(12, b=True), Alignment(horizontal="left", vertical="center"),
            border=False)
        for c in range(1, 13):
            ws.cell(1, c).border = Border(bottom=THIN)
        box(ws, 2, 1, 2, 12, "(найменування військового майна, індекс, номер креслення)", F7,
            Alignment(horizontal="left", vertical="top"), border=False)
        box(ws, 4, 1, 4, 7, "Нормативний запас: мінімальний ____________, максимальний ____________", F9,
            Alignment(horizontal="left"), border=False)
        cap = ("Код номенклатури  %s     Одиниця виміру  %s     Вміст дорогоцінних металів ______     "
               "Ціна за одиницю  %s грн" % (it["code"], it["unit"], money(it["price"]))) if it else \
            "Код номенклатури ______   Одиниця виміру ______   Вміст дорогоцінних металів ______   Ціна за одиницю ______"
        box(ws, 4, 8, 4, max(ncols, 40), cap, F9, Alignment(horizontal="left"), border=False)
        for c, t in ((1, "Дата запису"), (2, "Найменування документа"), (3, "Номер документа"), (4, "Дата документа"),
                     (5, "Постачальник (одержувач)"), (6, "Надійшло"), (7, "Вибуло")):
            box(ws, 6, c, 9, c, t, F8, CC)
        box(ws, 6, 8, 7, 13, "Перебуває згідно з документами", F8, CC)
        box(ws, 6, 14, 6, ncols, "У тому числі на складі (у підрозділах, військових частинах)", F8, CC)
        for p in places:
            c0 = pc[p]
            box(ws, 7, c0, 7, c0 + 5, "на складі" if p == D.head else p, F8B, CC)
        for c0 in [8] + [pc[p] for p in places]:
            box(ws, 8, c0, 9, c0, "усього", F7, CC)
            box(ws, 8, c0 + 1, 8, c0 + 5, "з них за категоріями (сортами)", F7, CC)
            for k in range(5):
                box(ws, 9, c0 + 1 + k, 9, c0 + 1 + k, k + 1, F7, CC)
        for c in range(1, ncols + 1):
            box(ws, 10, c, 10, c, c, F7, CC)
        for rr, h in ((6, 14), (7, 20), (8, 20), (9, 11), (10, 11)):
            ws.row_dimensions[rr].height = h
        ws.freeze_panes = "F11"
        ws.sheet_view.zoomScale = 90
        ws.page_setup.paperSize = 9
        ws.page_setup.orientation = "landscape"
        ws.print_title_rows = "6:10"
        r = 11
        expect = {}
        prev = None
        bal = collections.defaultdict(float)
        for row in rows:
            if row["kind"] == "total":
                continue
            for c in range(1, ncols + 1):
                box(ws, r, c, r, c, None, F8, CC)
            ws.row_dimensions[r].height = 13
            ws.cell(r, 1).value = row["date"]
            ws.cell(r, 1).number_format = NF_DATE
            ws.cell(r, 2).value = row["typ"]
            ws.cell(r, 3).value = row["num"]
            ws.cell(r, 4).value = row["ddate"]
            if isinstance(row["ddate"], datetime.date):
                ws.cell(r, 4).number_format = NF_DATE
            ws.cell(r, 5).value = row["party"]
            if row["inq"] is not None:
                ws.cell(r, 6).value = qty(row["inq"])
            if row["outq"] is not None:
                ws.cell(r, 7).value = qty(row["outq"])
            ws.cell(r, 8).value = "=SUM(%s)" % ",".join("%s%d" % (L(pc[p]), r) for p in places)
            for p in places:
                col = L(pc[p])
                ws.cell(r, pc[p]).number_format = NF_HIDE0
                if prev is None:
                    v = row["places"].get(p)
                    ws.cell(r, pc[p]).value = qty(v) if v is not None else 0
                else:
                    if p in row["places"]:
                        d = round(row["places"][p] - bal[p], 6)
                        ws.cell(r, pc[p]).value = "=%s%d%s%s" % (col, prev, "+" if d >= 0 else "-", qty(abs(d)))
                    else:
                        ws.cell(r, pc[p]).value = "=%s%d" % (col, prev)
            if prev is None:
                bal = collections.defaultdict(float, {p: row["places"].get(p, 0.0) for p in places})
            else:
                for p, v in row["places"].items():
                    bal[p] = v
            expect[r] = (round(row["total"], 6), {p: round(bal[p], 6) for p in places})
            prev = r
            r += 1
        return ws, expect, r - 1

    for code in codes:
        it = D.items[code]
        op = {p: q for (p, c), q in ob.items() if c == code}
        carry = (None, refs.get("47", {}).get(str(code)))
        pls, rows = card47(D, per, code, docs, op, carry)
        name = str(code)
        ws, expect, last = card_sheet(name, it, rows)
        entries.append(("%s, код %s" % (it["name"], code), name))
        has_carry = rows[0]["kind"] == "carry"
        maxr = last + MAXR
        a = "'%s'!$A$%d:$A$%d" % (name, 11 if not has_carry else 12, maxr)
        f_ = "'%s'!$F$%d:$F$%d" % (name, 11 if not has_carry else 12, maxr)
        g_ = "'%s'!$G$%d:$G$%d" % (name, 11 if not has_carry else 12, maxr)
        fn = '=SUMIFS(%s,%s,">="&$C$4,%s,"<="&$E$4)' % (f_, a, a)
        fv = '=SUMIFS(%s,%s,">="&$C$4,%s,"<="&$E$4)' % (g_, a, a)
        base = "'%s'!$H$11+" % name if has_carry else ""
        fs = '=%sSUMIFS(%s,%s,"<="&$E$4)-SUMIFS(%s,%s,"<="&$E$4)' % (base, f_, a, g_, a)
        lines.append(("%s, код %s" % (it["name"], code), it["unit"], fn, fv, fs))
        info[str(code)] = dict(expect={str(k): v for k, v in expect.items()}, rows=len(rows) - 1,
                               closing=rows[-1]["places"], total=rows[-1]["total"], carry=has_carry)
    card_sheet("ШАБЛОН", None, [])
    contents(ws_t, r_t, entries)
    summary_sheet(wb, "Книга обліку наявності та руху військового майна (служба забезпечення) — %s, %s"
                  % (cfg["service"].lower(), cfg["sub"]), lines, p_from, p_to)
    return wb, dict(places=places, codes=codes, cards=info)


def electronic_year(D, year, d_to, out_dir, refs, cfg=None):
    """електронні книги поточного року (01.01 – d_to): 14 кожному підрозділу з майном і 47; refs — сторінки
    книг минулого року. -> [(шлях, вид, місце)]"""
    cfg = dict(CFG, sheet="Книга", **(cfg or {}))
    per = Period(year, None, d_to)
    os.makedirs(out_dir, exist_ok=True)
    q0 = datetime.date(d_to.year, 3 * ((d_to.month - 1) // 3) + 1, 1)
    made = []
    for place in D.warehouses_of(per):
        sub = "Склад продовольчої служби" if place == D.head else place
        wb, _ = book14(D, place, sub, refs, per, q0, d_to, cfg)
        p1 = os.path.join(out_dir, "14 %s %s %d.xlsx" % (cfg["file"], file_place(place), year))
        wb.save(p1)
        made.append((p1, "14", place))
    wb, _ = book47e(D, refs, per, q0, d_to, cfg)
    p3 = os.path.join(out_dir, "47 %s %d.xlsx" % (cfg["file"], year))
    wb.save(p3)
    made.append((p3, "47", None))
    return made
