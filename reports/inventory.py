# -*- coding: utf-8 -*-
"""Інвентаризаційний опис за формою наказу Мінфіну від 17.06.2015 №572.

Опис — це не окрема сутність, а залишок на дату, вкладений у бланк. Графи
«за даними бухгалтерського обліку» беруться з бази як є. Графи «фактична
наявність» — це той самий залишок мінус знищене, ще не списане: майно
числиться, доки не пройде акт списання, але в натурі його вже немає, і саме це
комісія й побачить. Решту розбіжностей комісія виправляє рукою, а коли їх
затвердять — проводить звичайним документом.

Необоротні активи й запаси друкуються окремими описами — це вимога форми, і
розділяє їх `nomen.is_fixed_asset`. Набір граф у них різний: шістнадцять проти
дванадцяти; шапка — як у бланку: назва групи граф («Номер», «Фактична наявність»,
«За даними бухгалтерського обліку») стоїть один раз над своїми графами, а графа
без групи займає обидва рядки шапки.

Сторінки розкладає програма, як у паперах інвентаризації (`Pager`), а не Excel:
кожному рядку задано висоту за текстом, шапку таблиці на кожній сторінці повторює
Excel, а підсумок і підписи комісії не лишаються самі під нею.
"""
import sys
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook                                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, Side        # noqa: E402
from openpyxl.utils import get_column_letter                     # noqa: E402

from db.uk_number import money_words, qty_words                  # noqa: E402
from desktop.excel_names import save_book                        # noqa: E402
from desktop.inventory_export import Pager, ROW_PT, _height      # noqa: E402
from desktop.mtz_export import LINE_PT, _lines as _lines12       # noqa: E402

UNIT = "Військова частина"              # якщо в налаштуваннях бази реквізитів немає
EDRPOU = ""
STAMP = ["ЗАТВЕРДЖЕНО", "Наказ Міністерства фінансів України", "від 17.06.2015 №572"]

# Графи бланка: (група, заголовок, ширина). Порядок і кількість — за формою; графи однієї
# групи стоять поруч і мають спільну верхню клітинку шапки, графа без групи займає обидва
# рядки шапки.
FIXED_COLUMNS = [
    ("", "№\nз/п", 5),
    ("", "Найменування, стисла характеристика\nта призначення об'єкта", 46),
    ("", "Рік випуску\n(будівництва)\nчи дата придбання", 11),
    ("Номер", "інвентарний/\nноменклатурний", 13),
    ("Номер", "заводський", 12),
    ("Номер", "паспорта", 10),
    ("", "Один.\nвимір.", 7),
    ("Фактична наявність", "кількість", 10),
    ("Фактична наявність", "первісна (переоцінена) вартість", 13),
    ("", "Відмітка\nпро\nвибуття", 9),
    ("За даними бухгалтерського обліку", "кількість", 10),
    ("За даними бухгалтерського обліку", "первісна (переоцінена) вартість", 13),
    ("За даними бухгалтерського обліку", "сума зносу", 11),
    ("За даними бухгалтерського обліку", "балансова вартість", 12),
    ("", "Строк корисного\nвикористання", 11),
    ("", "Інші\nвідомості", 12),
]

STOCK_COLUMNS = [
    ("", "№\nз/п", 5),
    ("", "Рахунок,\nсубрахунок", 10),
    ("Матеріальні цінності", "найменування, вид, сорт, група", 46),
    ("Матеріальні цінності", "номенклатурний номер", 14),
    ("", "Одиниця\nвиміру", 8),
    ("Фактична наявність", "кількість", 11),
    ("Фактична наявність", "вартість", 12),
    ("Фактична наявність", "сума", 13),
    ("За даними бухгалтерського обліку", "кількість", 11),
    ("За даними бухгалтерського обліку", "вартість", 12),
    ("За даними бухгалтерського обліку", "сума", 13),
    ("", "Інші\nвідомості", 12),
]

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = Font(name="Times New Roman", size=9, bold=True)
BODY = Font(name="Times New Roman", size=10)
SMALL = Font(name="Times New Roman", size=8, italic=True)
TITLE = Font(name="Times New Roman", size=12, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


@dataclass
class Row:
    no: int
    code: str
    name: str
    uom: str
    qty_milli: int
    price_kop: int
    sum_kop: int
    is_fixed: bool
    made_year: str = ""
    inv_no: str = ""
    serial_no: str = ""
    destroyed_milli: int = 0

    @property
    def actual_milli(self) -> int:
        """Фактична наявність: облікова мінус знищене, ще не списане."""
        return self.qty_milli - self.destroyed_milli

    @property
    def actual_sum_kop(self) -> int:
        """Вартість фактичної наявності — за ціною партії."""
        return round(self.actual_milli * self.price_kop / 1000)


@dataclass
class Description:
    subdivision: str
    as_of: str
    fixed: list = field(default_factory=list)
    stock: list = field(default_factory=list)
    responsible: str = ""
    responsible_position: str = ""
    commission_head: str = ""
    commission_members: list = field(default_factory=list)
    accountant: str = ""
    service_chief: str = ""
    unit_name: str = ""
    unit_edrpou: str = ""

    @property
    def total_qty_milli(self) -> int:
        return sum(r.qty_milli for r in self.fixed + self.stock)

    @property
    def total_sum_kop(self) -> int:
        return sum(r.sum_kop for r in self.fixed + self.stock)

    @property
    def total_actual_milli(self) -> int:
        """Фактично в наявності — без знищеного, ще не списаного."""
        return sum(r.actual_milli for r in self.fixed + self.stock)

    @property
    def total_actual_sum_kop(self) -> int:
        return sum(r.actual_sum_kop for r in self.fixed + self.stock)

    @property
    def words_positions(self) -> str:
        return qty_words((len(self.fixed) + len(self.stock)) * 1000)

    @property
    def words_qty(self) -> str:
        return qty_words(self.total_qty_milli)

    @property
    def words_value(self) -> str:
        return money_words(self.total_sum_kop)


_BALANCE_SQL = """
SELECT n.id AS nomen_id, n.code, n.name, u.code AS uom, n.is_fixed_asset,
       SUM(p.sign * p.qty_milli) AS qty_milli,
       b.price_kop,
       i.made_year, i.inv_no, i.serial_no
FROM posting p
JOIN nomen n ON n.id = p.nomen_id
JOIN uom u ON u.id = n.uom_id
LEFT JOIN batch b ON b.batch_line_id = p.batch_line_id
LEFT JOIN instance i ON i.id = p.instance_id
WHERE p.doc_date <= :as_of AND p.subdivision_id = :sub
GROUP BY n.id, b.price_kop, i.id
HAVING SUM(p.sign * p.qty_milli) <> 0
ORDER BY n.is_fixed_asset DESC, n.name, b.price_kop
"""


def _settings(con) -> dict:
    return {r["key"]: r["value"] for r in con.execute("SELECT key, value FROM settings")}


def collect(con, subdivision_id: int, as_of: str) -> Description:
    """Залишок підрозділу на дату, розкладений у два описи.

    Рядок опису — це партія, а не позиція: у бланку є графа ціни, а різні партії
    однієї позиції коштують по-різному. Звести їх в один рядок означало б
    вигадати середню ціну, якої в жодному документі немає.

    Мінусові залишки не відкидаються. Опис зобов'язаний дорівнювати обліку, а
    мінус — це саме те, що комісія має побачити й розібрати, а не те, що варто
    прибрати з паперу.
    """
    name = con.execute("SELECT name FROM subdivision WHERE id = ?",
                       (subdivision_id,)).fetchone()[0]
    d = Description(subdivision=name, as_of=as_of)
    gone = {r["nomen_id"]: r["qty"] for r in con.execute(
        "SELECT nomen_id, SUM(qty_milli) AS qty FROM destroyed_open "
        "WHERE doc_date <= ? AND subdivision_id = ? GROUP BY nomen_id",
        (as_of, subdivision_id))}
    left = dict(gone)
    n_fixed = n_stock = 0
    for r in con.execute(_BALANCE_SQL, dict(as_of=as_of, sub=subdivision_id)):
        price = r["price_kop"] or 0
        is_fixed = bool(r["is_fixed_asset"])
        if is_fixed:
            n_fixed += 1
            no = n_fixed
        else:
            n_stock += 1
            no = n_stock
        # Знищене списується на партії підряд: у якій саме партії лежало
        # втрачене, жоден документ не каже.
        take = min(left.get(r["nomen_id"], 0), r["qty_milli"])
        if take:
            left[r["nomen_id"]] -= take
        row = Row(destroyed_milli=take,
                  no=no, code=r["code"], name=r["name"], uom=r["uom"],
                  qty_milli=r["qty_milli"], price_kop=price,
                  sum_kop=round(r["qty_milli"] * price / 1000),
                  is_fixed=is_fixed,
                  made_year=str(r["made_year"] or ""),
                  inv_no=r["inv_no"] or "", serial_no=r["serial_no"] or "")
        (d.fixed if is_fixed else d.stock).append(row)

    person = con.execute("""
        SELECT p.full_name, p.position FROM responsible r
        JOIN person p ON p.id = r.person_id
        WHERE r.subdivision_id = ? AND r.valid_from <= ?
          AND (r.valid_to IS NULL OR r.valid_to > ?)""",
        (subdivision_id, as_of, as_of)).fetchone()
    if person:
        d.responsible, d.responsible_position = person["full_name"], person["position"]

    s = _settings(con)
    d.commission_head = s.get("commission_head", "")
    d.commission_members = [x for x in s.get("commission_members", "").split("|") if x]
    d.accountant = s.get("accountant", "")
    d.service_chief = s.get("service_chief", "")
    d.unit_name = s.get("unit_legal_name", "")
    d.unit_edrpou = s.get("unit_edrpou", "")
    return d


def _fmt_qty(qty_milli: int):
    return qty_milli // 1000 if qty_milli % 1000 == 0 else qty_milli / 1000


def _lines(text, width, size, bold=False):
    """Скільки рядків тексту такого кегля стане в графі завширшки width знаків. Рахує
    відомість МТЗ для кегля 12 — графа масштабується під кегль."""
    px = max(20.0, width * 7 - 5) * 12 / size
    return _lines12(text, (px + 5) / 7, bold)


def _pt(size, lines=1):
    """Висота рядка під стільки рядків тексту такого кегля: рядок — LINE_PT відомості МТЗ
    (кегль 12) у пропорції до кегля, кілька рядків — із запасом на поля клітинки."""
    return max(ROW_PT, LINE_PT * size / 12 * lines + (2 if lines > 1 else 0))


class _Pages(Pager):
    """Сторінки розкладає програма, як у паперах інвентаризації. Шапку таблиці на кожній
    наступній сторінці повторює Excel (print_title_rows), і її висота йде в рахунок."""

    def __init__(self, ws, widths, titles):
        super().__init__(ws, widths)
        self.titles = titles

    def page(self, r):
        super().page(r)
        self.y = self.titles


def _sheet(wb, title, columns, rows, d, kind_title):
    ws = wb.create_sheet(title)
    last = len(columns)
    letter = get_column_letter(last)
    widths = [width for _, _, width in columns]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    def fit(row, value, width, font):
        """Висота рядка під текст у графі такої ширини. Excel злитих клітинок не розтягує, а
        розкладка сторінок мусить знати висоту кожного рядка наперед."""
        need = _pt(font.sz, _lines(value, width, font.sz, font.b))
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0.0, need)

    def put(row, col, value, font=BODY, align=LEFT, span=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font, c.alignment = font, align
        if span:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=row, end_column=col + span - 1)
        fit(row, value, sum(widths[col - 1:col - 1 + (span or 1)]), font)
        return c

    for i, line in enumerate(STAMP, 1):
        put(i, max(1, last - 3), line, SMALL, RIGHT, span=4)
    put(4, 1, d.unit_name or UNIT, BODY, LEFT, span=6)
    put(5, 1, f"Ідентифікаційний код за ЄДРПОУ  {d.unit_edrpou or EDRPOU}", SMALL, LEFT, span=6)
    put(7, 1, "ІНВЕНТАРИЗАЦІЙНИЙ ОПИС", TITLE, CENTER, span=last)
    put(8, 1, kind_title, Font(name="Times New Roman", size=11), CENTER, span=last)
    put(9, 1, f"продовольчої служби, {d.subdivision}", BODY, CENTER, span=last)
    put(10, 1, f"станом на {_date_uk(d.as_of)}", BODY, CENTER, span=last)
    put(12, 1, f"Матеріально відповідальна особа: {d.responsible_position or '—'}, "
               f"{d.responsible or '—'}", BODY, LEFT, span=last)

    head = 14
    groups = {}                                   # назва групи → (перша графа, остання)
    for i, (group, label, width) in enumerate(columns, 1):
        for at in (head, head + 1):
            c = ws.cell(row=at, column=i)
            c.font, c.alignment, c.border = HEAD, CENTER, BOX
        if group:
            ws.cell(row=head + 1, column=i, value=label)
            fit(head + 1, label, width, HEAD)
            groups[group] = (groups.get(group, (i, i))[0], i)
        else:                                     # графа без групи — на обидва рядки шапки
            ws.cell(row=head, column=i, value=label)
            ws.merge_cells(start_row=head, start_column=i, end_row=head + 1, end_column=i)
        n = ws.cell(row=head + 2, column=i, value=i)
        n.font, n.alignment, n.border = SMALL, CENTER, BOX
        fit(head + 2, i, width, SMALL)
    for group, (a, b) in groups.items():
        ws.cell(row=head, column=a, value=group)
        ws.merge_cells(start_row=head, start_column=a, end_row=head, end_column=b)
        fit(head, group, sum(widths[a - 1:b]), HEAD)
    for i, (group, label, width) in enumerate(columns, 1):
        if not group:                             # обидва рядки разом мають умістити її текст
            need = _pt(HEAD.sz, _lines(label, width, HEAD.sz, HEAD.b))
            have = ws.row_dimensions[head].height + ws.row_dimensions[head + 1].height
            if need > have:
                ws.row_dimensions[head + 1].height += need - have

    first = head + 3
    r = first
    for row in rows:
        values = (_fixed_values(row) if row.is_fixed else _stock_values(row))
        for i, v in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border = BODY, BOX
            c.alignment = LEFT if i in (2, 3) else CENTER
            fit(r, v, widths[i - 1], BODY)
        r += 1

    total = r
    # Фактична наявність — без знищеного, ще не списаного; за даними обліку — усе,
    # що числиться. Рядки це розрізняють, тож підсумок і пропис — так само.
    total_qty = sum(x.qty_milli for x in rows)
    total_sum = sum(x.sum_kop for x in rows)
    fact_qty = sum(x.actual_milli for x in rows)
    fact_sum = sum(x.actual_sum_kop for x in rows)
    fact_qty_col, fact_sum_col, qty_col, sum_col = ((8, 9, 11, 12) if columns is FIXED_COLUMNS
                                                    else (6, 8, 9, 11))
    for i in range(1, last + 1):
        c = ws.cell(row=r, column=i, value="Разом" if i == 1 else "X")
        c.font, c.alignment, c.border = HEAD, CENTER, BOX
    for col, value in ((fact_qty_col, _fmt_qty(fact_qty)), (fact_sum_col, fact_sum / 100),
                       (qty_col, _fmt_qty(total_qty)), (sum_col, total_sum / 100)):
        cell = ws.cell(row=r, column=col, value=value)
        cell.font, cell.alignment, cell.border = HEAD, CENTER, BOX

    r += 2
    words = [
        ("а) кількість порядкових номерів", qty_words(len(rows) * 1000)),
        ("б) загальна кількість одиниць (фактично)", qty_words(fact_qty)),
        ("в) вартість фактична", money_words(fact_sum)),
        ("г) загальна кількість одиниць за даними бухгалтерського обліку",
         qty_words(total_qty)),
        ("ґ) вартість за даними бухгалтерського обліку", money_words(total_sum)),
    ]
    put(r, 1, "Разом за описом:", HEAD, LEFT, span=3)
    r += 1
    for label, value in words:
        put(r, 1, label, BODY, LEFT, span=4)
        put(r, 5, value, BODY, LEFT, span=max(1, last - 4))
        r += 1

    r += 1
    put(r, 1, "Голова комісії", BODY, LEFT, span=2)
    put(r, 3, d.commission_head, BODY, LEFT, span=3)
    r += 1
    for i, m in enumerate(d.commission_members):
        put(r, 1, "Члени комісії" if i == 0 else "", BODY, LEFT, span=2)
        put(r, 3, m, BODY, LEFT, span=3)
        r += 1
    r += 1
    put(r, 1, "Усі цінності, перелічені в цьому описі, комісією перевірено в натурі "
              "та внесено до опису, у зв'язку з чим претензій до комісії не маю.",
        SMALL, LEFT, span=last)
    r += 2
    put(r, 1, "Матеріально відповідальна особа", BODY, LEFT, span=3)
    put(r, 4, d.responsible or "", BODY, LEFT, span=3)
    r += 1
    put(r, 1, "Інформацію за даними бухгалтерського обліку внесено", BODY, LEFT, span=3)
    put(r, 4, d.accountant or "", BODY, LEFT, span=3)
    r += 1
    put(r, 1, "Дані перевірив", BODY, LEFT, span=3)
    put(r, 4, d.service_chief or "", BODY, LEFT, span=3)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0       # заввишки — скільки треба, а не один аркуш
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head}:{head + 2}"
    ws.print_area = f"A1:{letter}{r}"
    ws.freeze_panes = ws.cell(row=first, column=1)

    # Рядки таблиці йдуть сторінками як підуть; останні три разом із підсумком, «Разом за
    # описом» і підписами стоять на одній сторінці — підписи не лишаються самі під шапкою.
    pages = _Pages(ws, widths, titles=sum(_height(ws, x) for x in range(head, head + 3)))
    together = max(first, total - 3)
    pages.flow(together - 1)
    pages.keep(together, r)
    return ws


def _fixed_values(row: Row):
    # Вартість фактичної наявності — за тим, що є в натурі; за даними обліку — за тим,
    # що числиться: у рядку зі знищеним, ще не списаним, вони різні.
    fact, book, s, fs = (_fmt_qty(row.actual_milli), _fmt_qty(row.qty_milli),
                         row.sum_kop / 100, row.actual_sum_kop / 100)
    gone = "знищено, не списано" if row.destroyed_milli else ""
    return [row.no, row.name, row.made_year, row.inv_no or row.code, row.serial_no, "",
            row.uom, fact, fs, "", book, s, "", s, "", gone]


def _stock_values(row: Row):
    fact, book, s, fs = (_fmt_qty(row.actual_milli), _fmt_qty(row.qty_milli),
                         row.sum_kop / 100, row.actual_sum_kop / 100)
    gone = "знищено, не списано" if row.destroyed_milli else ""
    return [row.no, "", row.name, row.code, row.uom,
            fact, row.price_kop / 100, fs, book, row.price_kop / 100, s, gone]


MONTHS = ["січня", "лютого", "березня", "квітня", "травня", "червня", "липня",
          "серпня", "вересня", "жовтня", "листопада", "грудня"]


def _date_uk(iso: str) -> str:
    y, m, d = iso.split("-")
    return f"«{int(d)}» {MONTHS[int(m) - 1]} {y} р."


def build_inventory(con, subdivision_id: int, as_of: str, out_path) -> Description:
    d = collect(con, subdivision_id, as_of)
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Необоротні активи", FIXED_COLUMNS, d.fixed, d,
           "(основні засоби, інші необоротні матеріальні активи)")
    _sheet(wb, "Запаси", STOCK_COLUMNS, d.stock, d, "(запаси)")
    save_book(wb, out_path)
    return d
