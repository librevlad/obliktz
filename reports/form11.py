# -*- coding: utf-8 -*-
"""Звіт про продовольчу тару 11/прод — з книги СП програми, за календарний квартал.

Рядки — позиції групи «Тара» книги СП (категорія «Т»): чинні й ті архівні, що мали залишок чи рух у кварталі; назва й
одиниця — позиції. Гр.4 — залишок книги на кінець минулого кварталу по всій частині (майно інших служб — ні); рух —
документи кварталу через межу частини (переміщення всередині частини — ні). Графа — за правилом: прихід від
постачальника чи іншої частини — гр.5 «за нарядами від місцевих постачальників»; тара зі зміни якісного стану
(коробки → макулатура), лишки й початкові залишки — гр.6 «інший прибуток»; повернення постачальнику — гр.11; решта
вибуття (акт списання, зміна якісного стану, передача) — гр.12 «різний видаток». Гр.8–10 (здано тароремонтним
підприємствам) правило не дає. Гр.14 = 4 + 7 − 13 і має дорівнювати книзі на кінець кварталу; гр.15 («у тому числі
з продуктами») служба вписує в копії книги, збереженій без «(програма)» в назві.

Книга — як подані звіти частини: «ф.11» (затверджує командир; підписують начальник логістики й начальник служби —
«МВО й посадовці» на день після кварталу), «розшифровка гр 6» і «розшифровка гр 12» (завжди; інших граф — коли є
документи) і «Розрахунок» (графи кожного коду, книга на кінець, документи з графою й правилом, перевірки). Файл —
«<номер частини> 11прод <N> квартал <рік> - Звіт (програма).xlsx» у теці звітів «<рік>\\<N> квартал\\11прод»; ручного
файла програма не перезаписує.
"""
import datetime
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook                                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, Side        # noqa: E402
from openpyxl.utils import get_column_letter                     # noqa: E402

from reports.form2 import _commander, _own_subs, _save_wb       # noqa: E402

GROUP = "СП.3"                       # «Тара» книги СП (категорія «Т»)
MILLI = 1000
ROMAN = ("I", "II", "III", "IV")
MONTHS = ("січня", "лютого", "березня", "квітня", "травня", "червня", "липня", "серпня", "вересня", "жовтня",
          "листопада", "грудня")
INCOME = (5, 6)
OUTGO = (8, 9, 10, 11, 12)
TITLES = {5: "за нарядами від місцевих постачальників", 6: "інший прибуток", 8: "знижка на 1 категорію",
          9: "знижка на 2 та вищу категорії", 10: "інші види витрат", 11: "повернення виробнику",
          12: "різний видаток"}
ALWAYS = (6, 12)                     # розшифровки, які подані звіти мають завжди
HINT = "(посада, військове звання, підпис, прізвище та ініціали)"

TNR = "Times New Roman"
F14, F14B, F12, F12B, F8 = (Font(name=TNR, size=14), Font(name=TNR, size=14, bold=True), Font(name=TNR, size=12),
                            Font(name=TNR, size=12, bold=True), Font(name=TNR, size=8))
THIN, MED = Side(style="thin"), Side(style="medium")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
UP = Alignment(horizontal="center", vertical="center", wrap_text=True, text_rotation=90)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
WIDTHS = {"A": 6.55, "B": 33.22, "C": 7.44, "D": 12.44, "E": 7.44, "F": 11.66, "G": 10.78, "H": 8.22, "I": 8.22,
          "J": 8.22, "K": 8.22, "L": 8.22, "M": 8.22, "N": 14.11, "O": 20.66}
HEAD = {"A10": "№№ пп", "B10": "Найменування", "C10": "Одиниця виміру", "E10": "Прибуток", "H10": "Видаток",
        "E11": TITLES[5], "F11": TITLES[6], "G11": "усього", "H11": "здано тароремонтним підприємствам",
        "K11": TITLES[11], "L11": TITLES[12], "M11": "усього", "H12": TITLES[8], "I12": TITLES[9], "J12": TITLES[10],
        "N12": "усього", "O12": "у тому числі з продуктами"}
HEAD_MERGES = ("A10:A12", "B10:B12", "C10:C12", "D10:D12", "E10:G10", "H10:M10", "N10:O11", "E11:E12", "F11:F12",
               "G11:G12", "H11:J11", "K11:K12", "L11:L12", "M11:M12")
UPRIGHT = ("C10", "D10", "E11", "F11", "G11", "K11", "L11", "M11", "H12", "I12", "J12", "N12", "O12")
FIRST = 14                           # перший рядок позицій «ф.11»


# ---------------------------------------------------------------- рушій

def quarter_dates(year, q):
    """Календарний квартал: (2026, 3) -> (2026-07-01, 2026-09-30)."""
    d1 = datetime.date(year, 3 * q - 2, 1)
    nxt = datetime.date(year + 1, 1, 1) if q == 4 else datetime.date(year, 3 * q + 1, 1)
    return d1, nxt - datetime.timedelta(days=1)


def last_quarter(day):
    """Останній завершений квартал на день: 12.10.2026 -> (2026, 3); 15.01.2027 -> (2026, 4)."""
    q = (day.month - 1) // 3
    return (day.year, q) if q else (day.year - 1, 4)


def column_of(inward, kind, party, paper, number):
    """Графа документа й правило, що її дало."""
    words = f"{party} {paper} {number}".lower()
    change = "якісного стану" in words or "зяс" in words
    if inward:
        if kind == "opening":
            return 6, "початкові залишки"
        if change:
            return 6, "зміна якісного стану"
        if re.search(r"\b(?:над)?лишк", words):
            return 6, "лишки"
        return 5, "надходження з продуктами"
    if "постачальник" in words or "виробник" in words:
        return 11, "повернення постачальнику"
    if change:
        return 12, "зміна якісного стану"
    return 12, "передача" if party else "списання"


def _txt(v):
    v = round(float(v or 0), 3)
    return str(int(v)) if v.is_integer() else str(v)


def collect(con, year, q):
    """Звіт кварталу: рядки з графами 4–15 і залишком книги на кінець, документи з графою й правилом, перевірки."""
    year, q = int(year), int(q)
    if q not in (1, 2, 3, 4):
        raise ValueError("квартал — від 1 до 4")
    d1, d2 = quarter_dates(year, q)
    own = _own_subs(con)
    ids = ",".join(str(x) for x in sorted(own)) or "NULL"         # id — цілі з бази, не текст людини

    def stock(nid, day):
        return con.execute(f"""SELECT COALESCE(SUM(sign * qty_milli), 0) FROM posting
                                WHERE nomen_id = ? AND doc_date <= ? AND subdivision_id IN ({ids})""",
                           (nid, day.isoformat())).fetchone()[0] / MILLI

    items = con.execute("""SELECT n.id, n.code, n.name, COALESCE(u.code, '') AS uom, n.archived_at FROM nomen n
                             JOIN nomen_group g ON g.id = n.group_id LEFT JOIN uom u ON u.id = n.uom_id
                            WHERE g.code = ? ORDER BY n.code""", (GROUP,)).fetchall()
    rows = {}
    for it in items:
        c = {k: 0.0 for k in range(4, 16)}
        c[4] = stock(it["id"], d1 - datetime.timedelta(days=1))
        rows[it["id"]] = {"code": it["code"], "name": it["name"], "uom": it["uom"],
                          "archived": bool(it["archived_at"]), "c": c, "book": stock(it["id"], d2)}
    docs = []
    for r in con.execute("""
            SELECT d.id, d.number, d.doc_date, COALESCE(d.paper, k.name) AS paper, k.code AS kind,
                   d.from_subdivision_id AS fsub, d.to_subdivision_id AS tsub, COALESCE(c.name, '') AS party,
                   COALESCE(sf.name, '') AS fname, COALESCE(st.name, '') AS tname, l.nomen_id, SUM(l.qty_milli) AS q
              FROM document d JOIN doc_kind k ON k.id = d.kind_id
              JOIN document_line l ON l.document_id = d.id
              JOIN nomen n ON n.id = l.nomen_id JOIN nomen_group g ON g.id = n.group_id
              LEFT JOIN counterparty c ON c.id = d.counterparty_id
              LEFT JOIN subdivision sf ON sf.id = d.from_subdivision_id
              LEFT JOIN subdivision st ON st.id = d.to_subdivision_id
             WHERE k.affects_stock = 1 AND g.code = ? AND d.doc_date BETWEEN ? AND ?
             GROUP BY d.id, l.nomen_id
             ORDER BY d.doc_date, d.id""", (GROUP, d1.isoformat(), d2.isoformat())):
        inward, outward = r["tsub"] in own, r["fsub"] in own
        if inward == outward:
            continue                                            # усередині частини або поза нею
        party = r["party"] or (r["fname"] if inward else r["tname"])
        col, rule = column_of(inward, r["kind"], party, r["paper"] or "", r["number"] or "")
        row = rows[r["nomen_id"]]
        qty = r["q"] / MILLI
        row["c"][col] += qty
        docs.append({"id": r["id"], "date": r["doc_date"], "paper": r["paper"] or "", "no": r["number"] or "",
                     "party": party, "place": r["tname"] if inward else r["fname"], "code": row["code"],
                     "name": row["name"], "uom": row["uom"], "qty": round(qty, 6), "col": col, "rule": rule})
    out, checks = [], []
    for row in rows.values():
        c = row["c"]
        c[7] = c[5] + c[6]
        c[13] = c[8] + c[9] + c[10] + c[11] + c[12]
        c[14] = c[4] + c[7] - c[13]
        for k in c:
            c[k] = round(c[k], 6)
        row["book"] = round(row["book"], 6)
        if row["archived"] and not any(abs(c[k]) > 1e-9 for k in (4, 7, 13, 14)):
            continue                                            # архівна позиція без залишку й руху
        out.append(row)
        if abs(c[14] - row["book"]) > 1e-6:
            checks.append({"level": "✗", "text": f"код {row['code']}: гр.14 {_txt(c[14])} ≠ книга СП "
                                                 f"{_txt(row['book'])} на {d2:%d.%m.%Y}"})
        if c[14] < -1e-9:
            checks.append({"level": "✗", "text": f"код {row['code']}: від'ємний залишок на кінець кварталу"})
    if not items:
        checks.append({"level": "⚠", "text": "у книзі СП немає позицій тари (група «Тара», категорія «Т»)"})
    return {"year": year, "q": q, "from": d1.isoformat(), "to": d2.isoformat(), "rows": out, "docs": docs,
            "checks": checks}


# ---------------------------------------------------------------- книга звіту

def folder_of(root, year, q):
    return os.path.join(root, str(year), f"{q} квартал", "11прод")


def file_name(unit, year, q):
    """«А0001» -> «0001 11прод 3 квартал 2026 - Звіт (програма).xlsx» — як подані, з позначкою програми."""
    return f"{re.sub(r'[^0-9]', '', unit or '') or '0000'} 11прод {q} квартал {year} - Звіт (програма).xlsx"


def _num(v):
    v = round(float(v or 0), 3)
    return int(v) if v.is_integer() else v


def _day_text(day):
    return f' "{day.day:02d}" {MONTHS[day.month - 1]} {day.year} р.'


def _put(ws, ref, value, font=F14, align=None, border=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell


def _frame(ws, rng, side=MED):
    """Рамка навколо діапазону (шапка й рядок «Всього» — товста, як у поданих)."""
    rows = list(ws[rng])
    for i, line in enumerate(rows):
        for j, cell in enumerate(line):
            b = cell.border
            cell.border = Border(left=side if j == 0 else b.left, right=side if j == len(line) - 1 else b.right,
                                 top=side if i == 0 else b.top, bottom=side if i == len(rows) - 1 else b.bottom)


def _signers_block(ws, row, signers, first, last, name_col, date_ref=None):
    """Підписи: посада, звання й ім'я, підказка дрібно — для кожного підписанта."""
    for position, rank, name in signers:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        _put(ws, f"{get_column_letter(first)}{row}", position or "_" * 60, align=LEFT)
        ws.row_dimensions[row].height = 40.8
        if date_ref:
            _put(ws, f"B{row + 1}", f"={date_ref}")
        ws.merge_cells(start_row=row + 1, start_column=first, end_row=row + 1, end_column=max(first, name_col - 1))
        _put(ws, f"{get_column_letter(first)}{row + 1}", rank or "_" * 15, align=Alignment(horizontal="left"))
        _put(ws, f"{get_column_letter(name_col)}{row + 1}", name or "_" * 20, align=Alignment(horizontal="right"))
        ws.merge_cells(start_row=row + 2, start_column=first, end_row=row + 2, end_column=last)
        _put(ws, f"{get_column_letter(first)}{row + 2}", HINT, font=F8,
             align=Alignment(horizontal="center", vertical="top"))
        row += 3
    return row - 1


def _form(ws, rep, unit, commander, signers, cut):
    """Аркуш «ф.11» — як у поданих звітах частини."""
    year, q = rep["year"], rep["q"]
    d1 = datetime.date.fromisoformat(rep["from"])
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    pos, rank, name = commander
    _put(ws, "K1", "ЗАТВЕРДЖУЮ", align=Alignment(vertical="center"))
    ws.merge_cells("K2:O2")
    _put(ws, "K2", pos or f"Командир військової частини {unit or 'А____'}", align=LEFT)
    ws.merge_cells("K3:L3")
    _put(ws, "K3", rank or "_" * 15, align=LEFT)
    _put(ws, "O3", name or "_" * 20, align=Alignment(vertical="top"))
    _put(ws, "K4", _day_text(cut), align=Alignment(vertical="center"))
    for ref, text in (("A6", "ЗВІТ"), ("A7", f"про продовольчу тару за {ROMAN[q - 1]} квартал {year} року"),
                      ("A8", f"військова частина {unit or 'А____'}")):
        ws.merge_cells(f"{ref}:O{ref[1:]}")
        _put(ws, ref, text, font=F14B, align=Alignment(horizontal="center"))
    for rng in HEAD_MERGES:
        ws.merge_cells(rng)
    head = dict(HEAD, D10=f"Залишок на {d1:%d.%m.%Y} р.", N10=f"Залишок на {cut:%d.%m.%Y} р.")
    for r in range(10, 14):
        for c in range(1, 16):
            ws.cell(r, c).border = BOX
    for ref, text in head.items():
        _put(ws, ref, text, font=F14B, align=UP if ref in UPRIGHT else CENTER, border=BOX)
    for c in range(1, 16):
        _put(ws, f"{get_column_letter(c)}13", c, font=F14B, align=CENTER, border=BOX)
    _frame(ws, "A10:O13")
    for r, h in ((5, 7.95), (9, 12), (10, 20.25), (11, 54.75), (12, 118.2)):
        ws.row_dimensions[r].height = h
    n = FIRST
    for i, row in enumerate(rep["rows"], 1):
        c = row["c"]
        vals = [i, row["name"], row["uom"] or None, _num(c[4]), _num(c[5]), _num(c[6]), f"=E{n}+F{n}",
                _num(c[8]), _num(c[9]), _num(c[10]), _num(c[11]), _num(c[12]), f"=SUM(H{n}:L{n})",
                f"=D{n}+G{n}-M{n}", _num(c[15])]
        for col, v in enumerate(vals, 1):
            _put(ws, f"{get_column_letter(col)}{n}", v, align=CENTER, border=BOX)
        ws.row_dimensions[n].height = 19.05
        n += 1
    total = n
    ws.merge_cells(f"A{total}:B{total}")
    _put(ws, f"A{total}", "Всього", font=F14B, align=CENTER, border=BOX)
    ws[f"B{total}"].border = BOX
    _put(ws, f"C{total}", "Х", font=F14B, align=CENTER, border=BOX)
    for col in range(4, 16):
        letter = get_column_letter(col)
        _put(ws, f"{letter}{total}", f"=SUM({letter}{FIRST}:{letter}{total - 1})" if rep["rows"] else 0,
             font=F14B, align=CENTER, border=BOX)
    _frame(ws, f"A{total}:O{total}")
    end = _signers_block(ws, total + 1, signers, 9, 15, 15, date_ref="K4")
    ws.print_area = f"A1:O{end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _breakdown(ws, rep, col, signers):
    """Розшифровка графи — документи рядками, позиції колонками (як у поданих звітах)."""
    docs = [d for d in rep["docs"] if d["col"] == col]
    codes = sorted({d["code"] for d in docs})
    named = {d["code"]: (d["name"], d["uom"]) for d in docs}
    last = 5 + max(1, len(codes))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    _put(ws, "A1", f'РОЗШИФРОВКА до графи {col} "{TITLES[col][:1].upper()}{TITLES[col][1:]}"', font=F12B,
         align=Alignment(horizontal="center"))
    for c, title in enumerate(("№ з/п", "Назва документа", "Номер документа", "Дата документа",
                               "Від кого" if col in INCOME else "На кого"), 1):
        ws.merge_cells(start_row=3, start_column=c, end_row=7, end_column=c)
        _put(ws, f"{get_column_letter(c)}3", title, font=F12B, align=CENTER)
    for i in range(last - 5):
        letter = get_column_letter(6 + i)
        name, uom = named[codes[i]] if codes else ("", "")
        for r, v in ((3, "Найменування ТМЦ"), (4, name), (5, "Одиниця виміру"), (6, uom), (7, "Кількість")):
            _put(ws, f"{letter}{r}", v or None, font=F12B if r in (3, 5, 7) else F12, align=CENTER)
    for c in range(1, last + 1):
        _put(ws, f"{get_column_letter(c)}8", c, font=F12, align=CENTER)
    groups = {}
    for d in docs:
        groups.setdefault(d["id"], {"d": d, "q": {}})["q"][d["code"]] = d["qty"]
    r = 9
    for i, g in enumerate(groups.values(), 1):
        d = g["d"]
        day = datetime.date.fromisoformat(d["date"])
        for c, v in enumerate((i, d["paper"], d["no"], day, d["party"] or d["place"]), 1):
            _put(ws, f"{get_column_letter(c)}{r}", v, font=F12, align=CENTER)
        ws[f"D{r}"].number_format = "dd.mm.yyyy"
        for k, code in enumerate(codes):
            if code in g["q"]:
                _put(ws, f"{get_column_letter(6 + k)}{r}", _num(g["q"][code]), font=F12, align=CENTER)
        r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    _put(ws, f"A{r}", "ВСЬОГО", font=F12B, align=CENTER)
    for c in range(6, last + 1):
        letter = get_column_letter(c)
        _put(ws, f"{letter}{r}", f"=SUM({letter}9:{letter}{r - 1})" if r > 9 else 0, font=F12B, align=CENTER)
    for row in ws.iter_rows(min_row=3, max_row=r, max_col=last):
        for cell in row:
            cell.border = BOX
    end = _signers_block(ws, r + 2, signers, 1, last, last)
    for letter, width in zip("ABCDE", (7, 34, 22, 14, 24)):
        ws.column_dimensions[letter].width = width
    for c in range(6, last + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.print_area = f"A1:{get_column_letter(last)}{end}"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _calc(ws, rep):
    """«Розрахунок»: графи кожного коду, книга на кінець, документи з графою й правилом, перевірки."""
    d1, d2 = datetime.date.fromisoformat(rep["from"]), datetime.date.fromisoformat(rep["to"])
    ws["A1"] = f"Розрахунок до звіту 11/прод за {ROMAN[rep['q'] - 1]} квартал {rep['year']} року"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Період: {d1:%d.%m.%Y} – {d2:%d.%m.%Y}; книга СП програми, тара «Т»"
    head = ["Код", "Найменування", "Од.", "гр.4", "гр.5", "гр.6", "гр.8", "гр.9", "гр.10", "гр.11", "гр.12",
            "гр.14", f"Книга СП на {d2:%d.%m.%Y}", "Збіг"]
    for c, v in enumerate(head, 1):
        ws.cell(4, c, v).font = Font(bold=True)
    r = 5
    for row in rep["rows"]:
        c = row["c"]
        same = abs(c[14] - row["book"]) <= 1e-6
        for k, v in enumerate([row["code"], row["name"], row["uom"] or None] +
                              [_num(c[x]) for x in (4, 5, 6, 8, 9, 10, 11, 12, 14)] +
                              [_num(row["book"]), "✔" if same else "✗"], 1):
            ws.cell(r, k, v)
        r += 1
    r += 1
    ws.cell(r, 1, "Документи кварталу (книга СП, тара «Т», через межу частини)").font = Font(bold=True)
    r += 1
    for c, v in enumerate(["Дата", "Документ", "Номер", "Від кого / кому", "Місце", "Код", "Найменування",
                           "Кількість", "Графа ф.11", "Правило"], 1):
        ws.cell(r, c, v).font = Font(bold=True)
    for d in rep["docs"]:
        r += 1
        for c, v in enumerate([datetime.date.fromisoformat(d["date"]), d["paper"], d["no"], d["party"] or None,
                               d["place"] or None, d["code"], d["name"], _num(d["qty"]), d["col"], d["rule"]], 1):
            ws.cell(r, c, v)
        ws.cell(r, 1).number_format = "dd.mm.yyyy"
    r += 2
    ws.cell(r, 1, "Перевірки").font = Font(bold=True)
    for chk in rep["checks"] or [{"level": "✔", "text": "без зауважень"}]:
        r += 1
        ws.cell(r, 1, f"{chk['level']} {chk['text']}")
    for letter, width in zip("ABCDEFGHIJKLMN", (12, 34, 22, 26, 22, 9, 30, 11, 10, 26, 8, 8, 16, 7)):
        ws.column_dimensions[letter].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def save_book(con, spec, root):
    """Книга звіту кварталу в теку звітів; шлях файла. Файл, відкритий в Excel, не затирається — поруч лягає новий."""
    from reports.form21 import _signers, _unit_code               # noqa: PLC0415
    year, q = int(spec["year"]), int(spec["q"])
    rep = collect(con, year, q)
    unit = _unit_code(con)
    cut = datetime.date.fromisoformat(rep["to"]) + datetime.timedelta(days=1)
    signers = _signers(con, cut.isoformat())
    wb = Workbook()
    ws = wb.active
    ws.title = "ф.11"
    _form(ws, rep, unit, _commander(con, cut.isoformat()), signers, cut)
    for col in sorted(set(ALWAYS) | {d["col"] for d in rep["docs"]}):
        _breakdown(wb.create_sheet(f"розшифровка гр {col}"), rep, col, signers)
    _calc(wb.create_sheet("Розрахунок"), rep)
    folder = folder_of(root, year, q)
    os.makedirs(folder, exist_ok=True)
    return _save_wb(wb, os.path.join(folder, file_name(unit, year, q)))
