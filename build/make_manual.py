# -*- coding: utf-8 -*-
"""Інструкція користувача з Markdown у Word (.docx) — щоб її можна було
роздрукувати й передати разом із дистрибутивом.

Джерело — docs/інструкція-користувача.md. Підтримується те, що в ній
вживається: заголовки #…####, абзаци, переліки «- » і «1. », таблиці на «|»,
цитати «> », жирний **…** і код `…` усередині рядка.

    python build/make_manual.py [вихідний .docx]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "інструкція-користувача.md"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def _runs(par, text):
    """Жирний і код усередині рядка; решта — звичайним шрифтом."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            par.add_run(piece)


def _table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows
             if not re.match(r"^\s*\|?\s*:?-{2,}", r)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            _runs(par, row[j] if j < len(row) else "")
            if i == 0:
                for run in par.runs:
                    run.bold = True
    doc.add_paragraph()


def build_manual(source=SOURCE, out=None):
    """Зібрати .docx з Markdown; повертає шлях до файла."""
    source = Path(source)
    out = Path(out) if out else source.with_suffix(".docx")
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor(0x2B, 0x3A, 0x42)

    para, table, i = [], [], 0

    def flush_para():
        if para:
            p = doc.add_paragraph()
            _runs(p, " ".join(s.strip() for s in para))
            para.clear()

    def flush_table():
        if table:
            _table(doc, table)
            table.clear()

    while i < len(lines):
        line = lines[i]
        i += 1
        if line.strip().startswith("|"):
            flush_para()
            table.append(line)
            continue
        flush_table()
        if not line.strip():
            flush_para()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            flush_para()
            level = len(m.group(1))
            if level == 1:
                title = doc.add_paragraph(style="Title")
                _runs(title, m.group(2))
                title.alignment = WD_ALIGN_PARAGRAPH.LEFT
                doc.core_properties.title = m.group(2)
            else:
                doc.add_heading(m.group(2), level=level - 1)
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)", line)
        if m:
            flush_para()
            style = "List Bullet 2" if m.group(1) else "List Bullet"
            _runs(doc.add_paragraph(style=style), m.group(2))
            continue
        m = re.match(r"^\s*\d+\.\s+(.*)", line)
        if m:
            flush_para()
            _runs(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        if line.startswith("> "):
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(line[2:])
            run.italic = True
            continue
        para.append(line)
    flush_para()
    flush_table()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


if __name__ == "__main__":
    target = build_manual(out=sys.argv[1] if len(sys.argv) > 1 else None)
    sys.stdout.reconfigure(encoding="utf-8")
    print(f"{target}  ({target.stat().st_size / 1024:.0f} КБ)")
